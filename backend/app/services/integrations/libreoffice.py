"""
LibreOffice helper — headless document conversion and thumbnail rendering.

This module wraps the `soffice` wrapper script (at /usr/lib/libreoffice/program/soffice)
to provide PPTX → PNG thumbnail rendering, PPTX → PDF conversion, and
PPTX → per-slide JPEG rendering (for the in-app slide viewer).

The slide/page viewer pipeline is format-generic:
    PPTX --soffice --> PDF --fitz/pdftoppm      --> slide_NNN.jpg (+ _t thumb)
    DOCX --soffice --> PDF (into the cache dir) --> page JPEGs
    XLSX --soffice --> PDF (into the cache dir) --> page JPEGs
    PDF  (already a PDF — no soffice needed)    --> page JPEGs

XLSX workbooks honor the print setup stored in the workbook (the excel_gen
service writes landscape + fit-to-width + repeated header rows), so the
converted pages read like a properly paginated print preview.

LibreOffice is installed to its DEFAULT system location (/usr/lib/libreoffice)
via apt-get install, and that path is mounted as a Docker volume. This means:
  - All hardcoded paths in fundamentalrc are correct by default
  - The wrapper script handles URE_BOOTSTRAP, LD_LIBRARY_PATH, etc. itself
  - No env var hacks, no path patching, no overlay extraction needed

We just call the wrapper script. That's it.
"""

from __future__ import annotations

import asyncio
import logging
import os
import shutil
import tempfile
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def _log(msg: str, *args) -> None:
    try:
        formatted = msg % args if args else msg
    except (TypeError, ValueError):
        formatted = f"{msg} {args}"
    print(f"[libreoffice] {formatted}", flush=True)


# Serialize all soffice invocations — it can't handle concurrent runs
# with a single user profile, and even with separate profiles it's safer.
_so_lock = asyncio.Lock()

# Timeout for soffice conversions (seconds). LibreOffice can hang on
# corrupted files; we don't want to block the backend forever.
_SOFFICE_TIMEOUT = 60

# Path to the soffice wrapper script (on the mounted volume)
_SOFFICE_PATH = "/usr/lib/libreoffice/program/soffice"

# Path to soffice.bin (used for diagnostics)
_SOFFICE_BIN = "/usr/lib/libreoffice/program/soffice.bin"


def is_available() -> bool:
    """Check if LibreOffice (soffice) is installed."""
    return Path(_SOFFICE_BIN).exists() or shutil.which("soffice") is not None


def _find_soffice() -> Optional[str]:
    """Find the soffice wrapper script."""
    if Path(_SOFFICE_PATH).exists():
        return _SOFFICE_PATH
    return shutil.which("soffice")


async def _run_soffice(
    args: list[str],
    timeout: int = _SOFFICE_TIMEOUT,
) -> tuple[int, str, str]:
    """Run a soffice command. Returns (returncode, stdout, stderr).

    We call the wrapper script (not soffice.bin directly) because the wrapper
    handles all the environment setup: URE_BOOTSTRAP, LD_LIBRARY_PATH,
    fundamentalrc resolution, etc. Since LibreOffice is installed to its
    default location (/usr/lib/libreoffice), all paths are correct by default.
    """
    soffice = _find_soffice()
    if not soffice:
        return -1, "", "LibreOffice (soffice) not found"

    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")

    try:
        cmd = [
            soffice,
            "--headless",
            "--norestore",
            "--nodefault",
            "--nolockcheck",
            "--nologo",
            f"-env:UserInstallation=file://{profile_dir}",
        ] + args

        _log("running: %s ... %s", " ".join(cmd[:3]), " ".join(args))

        proc = await asyncio.create_subprocess_exec(
            *cmd,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
            env={**os.environ, "HOME": profile_dir},
        )

        try:
            stdout, stderr = await asyncio.wait_for(proc.communicate(), timeout=timeout)
        except asyncio.TimeoutError:
            proc.kill()
            await proc.wait()
            _log("soffice timed out after %ds", timeout)
            return -1, "", f"Timeout after {timeout}s"

        rc = proc.returncode
        out = stdout.decode(errors="replace") if stdout else ""
        err = stderr.decode(errors="replace") if stderr else ""
        return rc, out, err

    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)


async def _diagnose_soffice_failure() -> None:
    """Run diagnostic checks when soffice fails."""
    _log("=== running diagnostics ===")

    # Check 1: soffice --version
    rc, out, err = await _run_soffice(["--version"], timeout=15)
    if rc == 0:
        _log("soffice --version OK: %s", out.strip()[:100])
    else:
        _log("soffice --version FAILED (rc=%d): %s", rc, err.strip()[:500])

    # Check 2: fonts
    for d in [Path("/usr/share/fonts"), Path("/usr/local/share/fonts")]:
        if d.exists():
            fonts = list(d.rglob("*.ttf")) + list(d.rglob("*.otf"))
            _log("%s — %d font files", d, len(fonts))

    # Check 3: fontconfig
    if Path("/etc/fonts/fonts.conf").exists():
        _log("/etc/fonts/fonts.conf exists (fontconfig OK)")

    # Check 4: shared libraries
    if Path(_SOFFICE_BIN).exists():
        import subprocess

        try:
            result = subprocess.run(
                ["ldd", _SOFFICE_BIN],
                capture_output=True,
                text=True,
                timeout=10,
            )
            missing = [i for i in result.stdout.split("\n") if "not found" in i]
            if missing:
                _log("MISSING LIBRARIES:")
                for m in missing:
                    _log("  %s", m.strip())
            else:
                _log("all shared libraries resolved")
        except Exception as e:
            _log("could not run ldd: %s", e)

    _log("=== diagnostics complete ===")


async def render_pptx_thumbnail(
    pptx_path: Path,
    max_width: int = 400,
) -> Optional[bytes]:
    """Render the first slide of a PPTX as a JPEG thumbnail.

    Uses LibreOffice to convert PPTX to PNG, then Pillow to resize to JPEG.
    Returns JPEG bytes, or None if unavailable/failed.
    """
    if not is_available():
        _log("LibreOffice not available — cannot render thumbnail")
        return None

    if not pptx_path.exists():
        _log("PPTX file not found: %s", pptx_path)
        return None

    async with _so_lock:
        # Create a temp output dir for the PNG
        out_dir = Path(tempfile.mkdtemp(prefix="lo_thumb_"))
        try:
            # Convert PPTX to PNG (soffice renders the first slide as PNG)
            rc, stdout, stderr = await _run_soffice(
                ["--convert-to", "png", "--outdir", str(out_dir), str(pptx_path)]
            )

            if rc != 0:
                _log("soffice conversion failed (rc=%d)", rc)
                _log("stderr: %s", stderr.strip()[:1000])
                await _diagnose_soffice_failure()
                return None

            # Find the output PNG
            stem = pptx_path.stem
            png_path = out_dir / f"{stem}.png"
            if not png_path.exists():
                pngs = list(out_dir.glob("*.png"))
                if not pngs:
                    _log("no PNG output found in %s", out_dir)
                    return None
                png_path = pngs[0]

            # Resize with Pillow and convert to JPEG
            try:
                from PIL import Image  # noqa: F401
            except ImportError:
                _log("Pillow not installed — returning raw PNG")
                return png_path.read_bytes()

            return await asyncio.to_thread(_resize_to_jpeg, png_path, max_width)

        finally:
            shutil.rmtree(out_dir, ignore_errors=True)


def _resize_to_jpeg(png_path: Path, max_width: int) -> Optional[bytes]:
    """Resize a PNG to max_width (maintaining aspect ratio) and return JPEG bytes."""
    try:
        from PIL import Image
        import io

        img = Image.open(str(png_path))
        if img.mode in ("RGBA", "LA", "P"):
            # Flatten transparency onto white background for JPEG
            bg = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            bg.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
            img = bg
        elif img.mode != "RGB":
            img = img.convert("RGB")

        # Resize maintaining aspect ratio
        w, h = img.size
        if w > max_width:
            new_h = int(h * max_width / w)
            img = img.resize((max_width, new_h), Image.LANCZOS)

        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=85)
        data = buf.getvalue()
        _log(
            "thumbnail rendered: %dx%d → %dx%d (%d bytes JPEG)",
            w,
            h,
            img.size[0],
            img.size[1],
            len(data),
        )
        return data

    except Exception as e:
        _log("Pillow resize failed: %s", e)
        return None


async def convert_pptx_to_pdf(pptx_path: Path) -> Optional[Path]:
    """Convert a PPTX file to PDF using LibreOffice.

    The output PDF is saved next to the input file with a .pdf extension.

    Returns:
        Path to the generated PDF, or None if conversion failed.
    """
    if not is_available():
        _log("LibreOffice not available — cannot convert to PDF")
        return None

    if not pptx_path.exists():
        _log("PPTX file not found: %s", pptx_path)
        return None

    out_dir = pptx_path.parent

    async with _so_lock:
        rc, _, stderr = await _run_soffice(
            ["--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)]
        )

        if rc != 0:
            _log("soffice PDF conversion failed (rc=%d): %s", rc, stderr[:200])
            return None

        pdf_path = out_dir / f"{pptx_path.stem}.pdf"
        if not pdf_path.exists():
            _log("PDF output not found: %s", pdf_path)
            return None

        _log("PDF generated: %s (%d bytes)", pdf_path.name, pdf_path.stat().st_size)
        return pdf_path


async def generate_and_cache_thumbnail(
    pptx_path: Path,
    thumb_path: Path,
    max_width: int = 400,
) -> Optional[Path]:
    """Generate a thumbnail and cache it to disk.

    Args:
        pptx_path: Source PPTX file
        thumb_path: Where to save the JPEG thumbnail
        max_width: Maximum thumbnail width

    Returns:
        Path to the saved thumbnail, or None if generation failed.
    """
    jpeg_bytes = await render_pptx_thumbnail(pptx_path, max_width)
    if not jpeg_bytes:
        return None

    try:
        thumb_path.parent.mkdir(parents=True, exist_ok=True)
        thumb_path.write_bytes(jpeg_bytes)
        _log("thumbnail cached: %s (%d bytes)", thumb_path.name, len(jpeg_bytes))
        return thumb_path
    except Exception as e:
        _log("failed to cache thumbnail: %s", e)
        return None


async def clean_pptx_with_libreoffice(pptx_path: Path) -> bool:
    """Clean/repair a PPTX file by opening it in LibreOffice and re-saving.

    python-pptx generates valid OOXML, but when using slide duplication
    (deep-copying template slides), the result can have issues that
    PowerPoint flags as "needs repair": duplicate shape IDsd, dangling
    relationship references, malformed XML elements, etc.

    LibreOffice's OOXML writer validates everything and produces a clean,
    PowerPoint-compatible file. This function:
      1. Opens the PPTX in LibreOffice headless
      2. Saves it as PPTX to a temp directory
      3. Replaces the original file with the cleaned version

    Returns True if the file was cleaned successfully, False otherwise.
    The original file is preserved if cleaning fails.
    """
    if not is_available():
        _log("LibreOffice not available — skipping PPTX cleaning")
        return False

    if not pptx_path.exists():
        _log("PPTX file not found: %s", pptx_path)
        return False

    async with _so_lock:
        out_dir = Path(tempfile.mkdtemp(prefix="lo_clean_"))
        try:
            # Convert PPTX → PPTX (LibreOffice opens, validates, re-saves)
            rc, stdout, stderr = await _run_soffice(
                ["--convert-to", "pptx", "--outdir", str(out_dir), str(pptx_path)]
            )

            if rc != 0:
                _log("PPTX cleaning failed (rc=%d): %s", rc, stderr.strip()[:300])
                return False

            # Find the cleaned output
            cleaned_path = out_dir / f"{pptx_path.stem}.pptx"
            if not cleaned_path.exists():
                _log("cleaned PPTX not found in %s", out_dir)
                return False

            # Replace the original with the cleaned version
            original_size = pptx_path.stat().st_size
            cleaned_size = cleaned_path.stat().st_size
            shutil.move(str(cleaned_path), str(pptx_path))

            _log(
                "PPTX cleaned: %s (%d → %d bytes, %s diff)",
                pptx_path.name,
                original_size,
                cleaned_size,
                (
                    f"+{cleaned_size - original_size}"
                    if cleaned_size > original_size
                    else f"{cleaned_size - original_size}"
                ),
            )
            return True

        finally:
            shutil.rmtree(out_dir, ignore_errors=True)


# ══════════════════════════════════════════════════════════════════════
# PPTX → per-slide JPEG rendering (for the in-app slide viewer)
# ══════════════════════════════════════════════════════════════════════
#
# Pipeline:  PPTX --soffice--> PDF --fitz/pdftoppm--> slide_NNN.jpg
#                                                 └-> slide_NNN_t.jpg (thumb)
#
# Results are cached in a directory next to the PPTX:
#     data/reports/{report_id}_slides/
#         manifest.json          — count + dims + notes + source mtime/size
#         slide_001.jpg          — full-size render (≤ 1600px wide)
#         slide_001_t.jpg        — thumbnail (≤ 360px wide)
#         ...
#
# The cache is invalidated automatically whenever the source PPTX changes
# (mtime or size mismatch against the manifest).

# Manifest format version — bump when the layout/naming changes.
# v2: adds the per-slide speaker-notes list ("notes": [str, ...]).
SLIDES_MANIFEST_VERSION = 2

# Full render max width (px). 16:9 @1600px is crisp on retina yet ~150-300KB.
SLIDE_MAX_WIDTH = 1600

# Thumbnail max width (px) — used by the viewer's filmstrip.
SLIDE_THUMB_WIDTH = 360

# JPEG encoders quality.
_SLIDE_JPEG_QUALITY = 85
_THUMB_JPEG_QUALITY = 80


def slides_cache_dir(pptx_path: Path) -> Path:
    """Directory where per-slide renders for a PPTX are cached.

    E.g. data/reports/abc.pptx → data/reports/abc_slides/
    """
    return pptx_path.parent / f"{pptx_path.stem}_slides"


def slide_image_name(index: int, thumb: bool = False) -> str:
    """Filename for a slide render (1-based index)."""
    return f"slide_{index:03d}{'_t' if thumb else ''}.jpg"


def read_slides_manifest(cache_dir: Path) -> Optional[dict]:
    """Read and validate manifest.json from a slide cache dir.

    Returns the parsed manifest dict, or None if missing/corrupt/too old.
    """
    manifest_path = cache_dir / "manifest.json"
    if not manifest_path.exists():
        return None
    try:
        import json

        data = json.loads(manifest_path.read_text())
        if (
            not isinstance(data, dict)
            or data.get("version") != SLIDES_MANIFEST_VERSION
            or not isinstance(data.get("count"), int)
            or data["count"] < 1
        ):
            return None
        return data
    except Exception as e:
        _log("slides manifest unreadable (%s): %s", manifest_path, e)
        return None


def is_slide_cache_valid(pptx_path: Path, cache_dir: Optional[Path] = None) -> bool:
    """Check whether the on-disk slide render cache is fresh and complete.

    Pure filesystem check — no LibreOffice involvement. Valid when:
      - manifest.json exists, parses, and is the current version
      - the recorded source mtime+size match the PPTX on disk
      - every full-size AND thumbnail image file exists
    """
    if not pptx_path.exists():
        return False
    cache_dir = cache_dir or slides_cache_dir(pptx_path)

    manifest = read_slides_manifest(cache_dir)
    if manifest is None:
        return False

    try:
        src_stat = pptx_path.stat()
        if manifest.get("source_mtime") != src_stat.st_mtime:
            return False
        if manifest.get("source_size") != src_stat.st_size:
            return False
    except OSError:
        return False

    for i in range(1, manifest["count"] + 1):
        if not (cache_dir / slide_image_name(i)).is_file():
            return False
        if not (cache_dir / slide_image_name(i, thumb=True)).is_file():
            return False
    return True


def extract_pptx_speaker_notes(
    pptx_path: Path,
    expected_count: Optional[int] = None,
) -> list[str]:
    """Extract per-slide speaker-notes text from a PPTX.

    Returns one string per slide, in slide order — ``""`` for slides that
    carry no notes. Uses python-pptx (the same library the pptx_gen service
    already depends on) with a lazy import, and NEVER raises: on any failure
    (missing library, unreadable/corrupt file) it returns ``[]`` so the
    slide viewer simply degrades to "no notes available".

    If ``expected_count`` is given (e.g. the number of rendered PDF pages),
    the list is padded with "" / truncated so it always aligns 1:1 with the
    rendered slides.
    """
    try:
        from pptx import Presentation
    except ImportError as e:
        _log("python-pptx unavailable for notes extraction: %s", e)
        return []

    notes: list[str] = []
    try:
        prs = Presentation(str(pptx_path))
        for slide in prs.slides:
            text = ""
            if slide.has_notes_slide:
                # notes_text_frame.text joins paragraphs with \n already.
                text = (slide.notes_slide.notes_text_frame.text or "").strip()
            notes.append(text)
    except Exception as e:
        _log("speaker-notes extraction failed (%s): %s", pptx_path.name, e)
        return []

    if expected_count is not None:
        if len(notes) < expected_count:
            notes.extend([""] * (expected_count - len(notes)))
        elif len(notes) > expected_count:
            notes = notes[:expected_count]
    return notes


def _pil_to_jpeg_bytes(img, max_width: int, quality: int) -> tuple[bytes, int, int]:
    """Flatten to RGB, downscale to max_width, encode as JPEG.

    Returns (jpeg_bytes, width, height).
    """
    from PIL import Image

    if img.mode in ("RGBA", "LA", "P"):
        bg = Image.new("RGB", img.size, (255, 255, 255))
        if img.mode == "P":
            img = img.convert("RGBA")
        bg.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
        img = bg
    elif img.mode != "RGB":
        img = img.convert("RGB")

    w, h = img.size
    if w > max_width:
        img = img.resize((max_width, int(h * max_width / w)), Image.LANCZOS)

    import io

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=quality)
    return buf.getvalue(), img.size[0], img.size[1]


def _render_pdf_pages_with_fitz(
    pdf_path: Path,
    cache_dir: Path,
    max_width: int,
    thumb_width: int,
) -> Optional[list[dict]]:
    """Render every PDF page to slide_NNN.jpg (+ _t thumb) using PyMuPDF.

    Returns a list of {"index", "width", "height"} dicts, or None on failure.
    Runs synchronously — call via asyncio.to_thread().
    """
    try:
        import fitz  # PyMuPDF
        from PIL import Image
    except ImportError as e:
        _log("PyMuPDF/Pillow unavailable for rendering: %s", e)
        return None

    pages: list[dict] = []
    try:
        doc = fitz.open(str(pdf_path))
        with doc:
            for page in doc:
                # Scale so the page renders ≈ max_width pixels wide.
                zoom = max_width / max(page.rect.width, 1.0)
                pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)

                full, fw, fh = _pil_to_jpeg_bytes(img, max_width, _SLIDE_JPEG_QUALITY)
                thumb, _, _ = _pil_to_jpeg_bytes(img, thumb_width, _THUMB_JPEG_QUALITY)

                i = len(pages) + 1
                (cache_dir / slide_image_name(i)).write_bytes(full)
                (cache_dir / slide_image_name(i, thumb=True)).write_bytes(thumb)
                pages.append({"index": i, "width": fw, "height": fh})
        return pages or None
    except Exception as e:
        _log("fitz page rendering failed: %s", e)
        return None


def _render_pdf_pages_with_pdftoppm(
    pdf_path: Path,
    cache_dir: Path,
    max_width: int,
    thumb_width: int,
) -> Optional[list[dict]]:
    """Fallback renderer using poppler's pdftoppm CLI.

    Produces slide_NNN.jpg (+ _t thumb) like the fitz path. Returns the same
    shape, or None if pdftoppm is missing / fails.
    """
    import subprocess

    binary = shutil.which("pdftoppm")
    if not binary:
        _log("pdftoppm not found — cannot fall back to poppler rendering")
        return None

    out_dir = Path(tempfile.mkdtemp(prefix="pdftoppm_"))
    try:
        # --scale-to N scales the LONG side to N px; 16:9 decks are wider
        # than tall so this yields ≈ max_width-wide output.
        proc = subprocess.run(
            [
                binary,
                "-jpeg",
                "-jpegopt",
                "quality=90",
                "-scale-to",
                str(max_width),
                str(pdf_path),
                str(out_dir / "page"),
            ],
            capture_output=True,
            timeout=120,
        )
        if proc.returncode != 0:
            _log(
                "pdftoppm failed (rc=%d): %s",
                proc.returncode,
                proc.stderr.decode(errors="replace")[:300],
            )
            return None

        # pdftoppm names outputs page-1.jpg, page-01.jpg, page-001.jpg …
        # (zero-padded to the page count). Sort numerically.
        files = sorted(
            out_dir.glob("page-*.jpg"),
            key=lambda p: int(p.name.rsplit("-", 1)[1].split(".")[0]),
        )
        if not files:
            _log("pdftoppm produced no pages for %s", pdf_path)
            return None

        pages: list[dict] = []
        for i, f in enumerate(files, start=1):
            try:
                from PIL import Image

                img = Image.open(str(f))
                full, fw, fh = _pil_to_jpeg_bytes(img, max_width, _SLIDE_JPEG_QUALITY)
                thumb, _, _ = _pil_to_jpeg_bytes(img, thumb_width, _THUMB_JPEG_QUALITY)
            except ImportError:
                # No Pillow — serve pdftoppm's own JPEGs as-is, skip thumbs
                # by duplicating the full image as the "thumbnail".
                full = f.read_bytes()
                fw = fh = 0
                thumb = full

            (cache_dir / slide_image_name(i)).write_bytes(full)
            (cache_dir / slide_image_name(i, thumb=True)).write_bytes(thumb)
            pages.append({"index": i, "width": fw, "height": fh})

        return pages
    except Exception as e:
        _log("pdftoppm rendering failed: %s", e)
        return None
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


def _write_slides_manifest(
    cache_dir: Path,
    pptx_path: Path,
    pages: list[dict],
    notes: Optional[list[str]] = None,
) -> dict:
    """Persist manifest.json describing a fresh slide render cache.

    ``notes`` (per-slide speaker notes, aligned with ``pages``) is padded or
    truncated to ``len(pages)`` so it always aligns 1:1 with the renders.
    """
    import json
    import time

    src_stat = pptx_path.stat()
    first = pages[0]
    slide_notes = list(notes or [])
    if len(slide_notes) < len(pages):
        slide_notes.extend([""] * (len(pages) - len(slide_notes)))
    elif len(slide_notes) > len(pages):
        slide_notes = slide_notes[: len(pages)]
    manifest = {
        "version": SLIDES_MANIFEST_VERSION,
        "count": len(pages),
        "width": first.get("width") or 0,
        "height": first.get("height") or 0,
        "notes": slide_notes,
        "source_mtime": src_stat.st_mtime,
        "source_size": src_stat.st_size,
        "generated_at": time.time(),
    }
    (cache_dir / "manifest.json").write_text(json.dumps(manifest))
    return manifest


async def convert_pptx_to_slide_images(
    pptx_path: Path,
    cache_dir: Optional[Path] = None,
    max_width: int = SLIDE_MAX_WIDTH,
    thumb_width: int = SLIDE_THUMB_WIDTH,
    force: bool = False,
) -> Optional[dict]:
    """Render every slide of a PPTX as JPEG images (full + thumbnail).

    Pipeline: PPTX → PDF (LibreOffice, cached next to the PPTX) → per-page
    JPEGs (PyMuPDF primary, pdftoppm fallback) → manifest.json (which also
    embeds the per-slide speaker notes extracted from the PPTX).

    A fresh cache is written to {cache_dir}; any previous partial cache is
    wiped first so the directory never mixes generations.

    Returns the manifest dict on success, or None on failure (missing
    PPTX, LibreOffice unavailable, conversion error, zero pages).
    """
    if not pptx_path.exists():
        _log("PPTX file not found: %s", pptx_path)
        return None

    cache_dir = cache_dir or slides_cache_dir(pptx_path)

    # Serve from cache when still fresh (unless forced).
    if not force and is_slide_cache_valid(pptx_path, cache_dir):
        manifest = read_slides_manifest(cache_dir)
        if manifest is not None:
            return manifest

    # Fresh generation — start from an empty cache dir.
    if cache_dir.exists():
        shutil.rmtree(cache_dir, ignore_errors=True)
    cache_dir.mkdir(parents=True, exist_ok=True)

    # Step 1: PPTX → PDF. Drop a stale cached PDF first so a regenerated
    # PPTX never shows old slides.
    pdf_path = pptx_path.with_suffix(".pdf")
    try:
        if pdf_path.exists() and pdf_path.stat().st_mtime < pptx_path.stat().st_mtime:
            pdf_path.unlink()
            _log("dropped stale cached PDF (older than PPTX): %s", pdf_path.name)
    except OSError:
        pass

    if not pdf_path.exists():
        converted = await convert_pptx_to_pdf(pptx_path)
        if converted is None or not converted.exists():
            _log("PPTX→PDF failed — cannot render slides")
            return None
        pdf_path = converted

    # Step 2: PDF → per-page JPEGs (try fitz, fall back to pdftoppm).
    pages = await asyncio.to_thread(
        _render_pdf_pages_with_fitz, pdf_path, cache_dir, max_width, thumb_width
    )
    if pages is None:
        _log("fitz renderer failed — falling back to pdftoppm")
        pages = await asyncio.to_thread(
            _render_pdf_pages_with_pdftoppm,
            pdf_path,
            cache_dir,
            max_width,
            thumb_width,
        )
    if not pages:
        _log("no slide images rendered for %s", pptx_path.name)
        return None

    # Step 3: extract per-slide speaker notes from the PPTX itself (cheap,
    # python-pptx only) so the manifest carries the full viewing payload.
    notes = await asyncio.to_thread(extract_pptx_speaker_notes, pptx_path, len(pages))

    manifest = await asyncio.to_thread(
        _write_slides_manifest, cache_dir, pptx_path, pages, notes
    )
    _log(
        "slide images rendered: %s → %d slides in %s",
        pptx_path.name,
        manifest["count"],
        cache_dir.name,
    )
    return manifest


# ══════════════════════════════════════════════════════════════════════
# Generic document → per-page JPEG rendering (PDF / DOCX / XLSX viewer)
# ══════════════════════════════════════════════════════════════════════
#
# The PPTX pipeline above is specialized (speaker notes, cached PDF next
# to the source for the /pdf endpoint). The functions below reuse the
# same cache layout + manifest format for the other deliverable formats:
#
#   PDF  — the file IS a PDF, so pages are rasterized directly with
#          PyMuPDF (pdftoppm fallback). LibreOffice is NOT required.
#   DOCX — soffice converts DOCX → PDF *into the cache dir* (never next
#          to the source, so a stray {id}.pdf can never hijack the
#          download endpoint), then pages are rasterized like a PDF.
#          LibreOffice IS required.
#   XLSX — identical to DOCX: soffice converts the workbook to PDF into
#          the cache dir (honoring the workbook's print setup — the
#          excel_gen service writes landscape/fit-to-width/repeat-header
#          print settings), then pages are rasterized. LibreOffice IS
#          required.

# Formats the generic viewer pipeline understands.
VIEWER_SUPPORTED_FORMATS = ("pptx", "pdf", "docx", "xlsx")


def can_rasterize_pdf() -> bool:
    """True when at least one PDF rasterizer is available.

    PyMuPDF (fitz) is a bundled pip dependency (preferred); poppler's
    pdftoppm CLI is the fallback. Used to give a friendly 404 detail
    when neither is importable — no LibreOffice involvement.
    """
    try:
        import fitz  # noqa: F401

        return True
    except ImportError:
        return shutil.which("pdftoppm") is not None


def viewer_format_of(source_path: Path) -> Optional[str]:
    """Normalized viewer format for a source file ("pptx"/"pdf"/"docx"/"xlsx")."""
    suffix = source_path.suffix.lower().lstrip(".")
    return suffix if suffix in VIEWER_SUPPORTED_FORMATS else None


def viewer_cache_dir(source_path: Path) -> Path:
    """Directory where per-page renders for any viewer file are cached.

    Same convention as the PPTX pipeline: {stem}_slides next to the file
    (e.g. data/reports/abc.pdf → data/reports/abc_slides/).
    """
    return slides_cache_dir(source_path)


async def _convert_office_to_pdf_in_cache(
    source_path: Path, cache_dir: Path
) -> Optional[Path]:
    """Convert a DOCX/XLSX → PDF *inside the cache dir* using LibreOffice.

    The output lands at {cache_dir}/{stem}.pdf — deliberately NOT next
    to the source file. data/reports/ is probed by extension by the
    download endpoint, so writing {id}.pdf next to {id}.docx / {id}.xlsx
    would make downloads serve the PDF conversion instead of the original
    deliverable. (The cache dir is wiped whenever the source changes, so
    the embedded PDF can never go stale either.)

    For spreadsheets, LibreOffice paginates according to the workbook's
    print setup (page size / orientation / fit-to-width / repeated title
    rows) — the excel_gen service writes reader-friendly defaults on
    every generated sheet, so the preview reads like a print preview.
    """
    if not is_available():
        _log("LibreOffice not available — cannot convert %s to PDF", source_path.suffix)
        return None

    async with _so_lock:
        rc, _, stderr = await _run_soffice(
            ["--convert-to", "pdf", "--outdir", str(cache_dir), str(source_path)]
        )

        if rc != 0:
            _log(
                "soffice %s→PDF failed (rc=%d): %s",
                source_path.suffix,
                rc,
                stderr[:200],
            )
            return None

        pdf_path = cache_dir / f"{source_path.stem}.pdf"
        if not pdf_path.exists():
            _log("%s→PDF output not found: %s", source_path.suffix, pdf_path)
            return None

        _log(
            "%s→PDF converted into cache: %s (%d bytes)",
            source_path.suffix,
            pdf_path.name,
            pdf_path.stat().st_size,
        )
        return pdf_path


# Backwards-compatible alias (the DOCX path was the original caller).
_convert_docx_to_pdf_in_cache = _convert_office_to_pdf_in_cache


async def _rasterize_pdf_into_cache(
    pdf_path: Path,
    source_path: Path,
    cache_dir: Path,
    max_width: int,
    thumb_width: int,
) -> Optional[dict]:
    """Shared render core: PDF → per-page JPEGs + manifest in cache_dir.

    Used by both the PDF and DOCX paths of the generic viewer. Speaker
    notes are PPTX-only, so the manifest's notes list is all "".
    Returns the manifest dict, or None when no rasterizer worked.
    """
    pages = await asyncio.to_thread(
        _render_pdf_pages_with_fitz, pdf_path, cache_dir, max_width, thumb_width
    )
    if pages is None:
        _log("fitz renderer failed — falling back to pdftoppm")
        pages = await asyncio.to_thread(
            _render_pdf_pages_with_pdftoppm,
            pdf_path,
            cache_dir,
            max_width,
            thumb_width,
        )
    if not pages:
        _log("no page images rendered for %s", source_path.name)
        return None

    manifest = await asyncio.to_thread(
        _write_slides_manifest, cache_dir, source_path, pages, None
    )
    return manifest


async def convert_document_to_page_images(
    source_path: Path,
    cache_dir: Optional[Path] = None,
    max_width: int = SLIDE_MAX_WIDTH,
    thumb_width: int = SLIDE_THUMB_WIDTH,
    force: bool = False,
) -> Optional[dict]:
    """Render every page of a PDF/DOCX/XLSX/PPTX as JPEG images (full + thumb).

    Format-generic entry point for the in-app document viewer. Dispatches
    on the file extension:

      - .pptx → the dedicated PPTX pipeline (soffice → PDF next to the
        source, per-slide speaker notes in the manifest).
      - .pdf  → rasterize directly (PyMuPDF / pdftoppm). No LibreOffice.
      - .docx → soffice converts to PDF *inside the cache dir*, then
        rasterizes. Requires LibreOffice.
      - .xlsx → identical to .docx: soffice converts the workbook to PDF
        inside the cache dir (honoring the workbook's print setup — the
        excel_gen service writes landscape/fit-to-width/repeat-header
        print settings), then rasterizes. Requires LibreOffice.

    Cache layout + manifest format are identical to the PPTX pipeline
    ({stem}_slides/ with slide_NNN.jpg + slide_NNN_t.jpg + manifest.json),
    invalidated automatically when the source file's mtime/size changes.

    Returns the manifest dict on success, or None on failure (missing
    source, unsupported format, unavailable converter/rasterizer, or a
    conversion/render error).
    """
    if not source_path.exists():
        _log("source file not found: %s", source_path)
        return None

    fmt = viewer_format_of(source_path)
    if fmt is None:
        _log("unsupported viewer format: %s", source_path.name)
        return None

    # PPTX keeps its specialized pipeline (speaker notes + cached PDF
    # next to the source for the /pdf endpoint). Call it exactly like
    # the old endpoint did (max_width/thumb_width/force stay defaulted).
    if fmt == "pptx":
        return await convert_pptx_to_slide_images(source_path, cache_dir)

    cache_dir = cache_dir or viewer_cache_dir(source_path)

    # Serve from cache when still fresh (unless forced).
    if not force and is_slide_cache_valid(source_path, cache_dir):
        manifest = read_slides_manifest(cache_dir)
        if manifest is not None:
            return manifest

    # Fresh generation — start from an empty cache dir. (For DOCX this
    # also drops the previous embedded {id}.pdf conversion.)
    if cache_dir.exists():
        shutil.rmtree(cache_dir, ignore_errors=True)
    cache_dir.mkdir(parents=True, exist_ok=True)

    if fmt == "pdf":
        pdf_path = source_path
    else:  # docx / xlsx — same soffice → PDF-in-cache pipeline
        pdf_path = await _convert_office_to_pdf_in_cache(source_path, cache_dir)
        if pdf_path is None:
            _log("%s→PDF failed — cannot render pages", fmt.upper())
            return None

    manifest = await _rasterize_pdf_into_cache(
        pdf_path, source_path, cache_dir, max_width, thumb_width
    )
    if manifest is None:
        return None

    _log(
        "page images rendered: %s → %d pages in %s",
        source_path.name,
        manifest["count"],
        cache_dir.name,
    )
    return manifest


# ===========================================================================
# RAG document preview — thumbnails + page images for ANY document type
# ===========================================================================
#
# The Workspace > Documents section needs previews for everything a user can
# upload, not just the four report formats: txt, md, markdown, csv, tsv, pdf,
# docx, doc, xlsx, xls, pptx, ppt.
#
# Rendering strategy per family:
#   pdf          → PyMuPDF / pdftoppm raster directly (no LibreOffice needed)
#   office files → soffice → PDF inside the cache dir → raster (LibreOffice
#                  required; soffice happily converts legacy .doc/.xls/.ppt
#                  and modern .docx/.xlsx/.pptx alike)
#   text files   → soffice → PDF when available; PIL-rendered text pages as
#                  a universal fallback (works with zero external tools)
#
# Cache layout mirrors the slide viewer ({stem}_pages/ next to the source,
# slide_NNN.jpg + slide_NNN_t.jpg + manifest.json + thumb.jpg), invalidated
# by the source file's mtime + size. For RAG documents the cache lives at
# data/documents/{id}/{stem}_pages/ — inside the per-document directory —
# so delete_document's rmtree cleans the cache automatically.
# ---------------------------------------------------------------------------

# Every extension the documents module accepts can get a preview.
DOCUMENT_PREVIEW_FORMATS = {
    "pdf",
    "docx",
    "doc",
    "xlsx",
    "xls",
    "pptx",
    "ppt",
    "txt",
    "md",
    "markdown",
    "csv",
    "tsv",
}

_TEXT_PREVIEW_FORMATS = {"txt", "md", "markdown", "csv", "tsv"}

# Monospace-ish font for the PIL text fallback (probed in order).
_TEXT_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]


def document_preview_format(source_path: Path) -> Optional[str]:
    """Normalized preview format for a RAG document (any accepted ext)."""
    suffix = source_path.suffix.lower().lstrip(".")
    return suffix if suffix in DOCUMENT_PREVIEW_FORMATS else None


def document_preview_cache_dir(source_path: Path) -> Path:
    """Cache dir for a document's page renders — {stem}_pages next to it."""
    return source_path.parent / f"{source_path.stem}_pages"


def _load_text_font(size: int):
    """Load a TTF font for the PIL text renderer, with a default fallback."""
    from PIL import ImageFont

    for candidate in _TEXT_FONT_CANDIDATES:
        try:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size)
        except Exception:
            continue
    try:
        # Pillow ≥ 10 supports sized default fonts
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def _read_text_bytes(buf: bytes) -> str:
    """Decode text-ish bytes (utf-8 first, latin-1 fallback) with a size cap."""
    if len(buf) > 400_000:  # cap at ~400 KB for preview purposes
        buf = buf[:400_000]
    try:
        return buf.decode("utf-8")
    except UnicodeDecodeError:
        return buf.decode("latin-1", errors="replace")


def _paginate_text_lines(
    text: str, chars_per_line: int, lines_per_page: int, max_pages: int
) -> list[list[str]]:
    """Wrap + paginate raw text into screenfuls for the PIL renderer."""
    raw_lines = text.splitlines() or [""]
    # Long lines (CSV rows can be huge) get hard-wrapped.
    wrapped: list[str] = []
    for line in raw_lines:
        line = line.rstrip("\r")
        if len(line) <= chars_per_line:
            wrapped.append(line)
        else:
            for i in range(0, len(line), chars_per_line):
                wrapped.append(line[i : i + chars_per_line])
    pages: list[list[str]] = []
    for i in range(0, len(wrapped), lines_per_page):
        pages.append(wrapped[i : i + lines_per_page])
        if len(pages) >= max_pages:
            break
    return pages or [[""]]


def _render_text_page_image(
    filename: str,
    lines: list[str],
    page_number: int,
    total_pages: int,
    width: int,
) -> "object":
    """Render one 'text document' page as a PIL image (white bg, dark text,
    header bar with the filename + page indicator)."""
    from PIL import Image, ImageDraw

    line_height = 26
    margin_x, margin_top, margin_bottom = 56, 88, 56
    body_font = _load_text_font(17)
    header_font = _load_text_font(15)

    height = margin_top + line_height * len(lines) + margin_bottom
    height = max(height, 420)

    img = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    # Header band (subtle gray) with filename + page x/y
    draw.rectangle([0, 0, width, 64], fill=(245, 246, 248))
    draw.line([0, 64, width, 64], fill=(210, 214, 220), width=1)

    title = filename if len(filename) <= 70 else filename[:67] + "…"
    draw.text((margin_x, 22), title, font=header_font, fill=(55, 60, 70))
    page_label = f"page {page_number}/{total_pages}"
    pl_w = draw.textlength(page_label, font=header_font)
    draw.text(
        (width - margin_x - pl_w, 22),
        page_label,
        font=header_font,
        fill=(120, 126, 138),
    )

    y = margin_top
    for line in lines:
        # Render tabs as visible spacing, strip control chars
        safe = line.replace("\t", "    ")
        safe = "".join(ch for ch in safe if ch == " " or ch.isprintable())
        draw.text((margin_x, y), safe[:400], font=body_font, fill=(24, 26, 32))
        y += line_height

    return img


def _render_text_pages_sync(
    source_path: Path,
    cache_dir: Path,
    max_width: int,
    thumb_width: int,
) -> Optional[list[dict]]:
    """PIL-render a text-ish file (txt/md/csv/tsv) into page JPEGs.

    Returns the same page dict shape as the PDF rasterizers, or None on
    failure. Runs synchronously — call via asyncio.to_thread().
    """
    try:
        text = _read_text_bytes(source_path.read_bytes())
        pages = _paginate_text_lines(
            text, chars_per_line=96, lines_per_page=34, max_pages=20
        )
        rendered: list[dict] = []
        for i, lines in enumerate(pages, start=1):
            img = _render_text_page_image(source_path.name, lines, i, len(pages), 1000)
            full, fw, fh = _pil_to_jpeg_bytes(img, max_width, _SLIDE_JPEG_QUALITY)
            thumb, _, _ = _pil_to_jpeg_bytes(img, thumb_width, _THUMB_JPEG_QUALITY)
            (cache_dir / slide_image_name(i)).write_bytes(full)
            (cache_dir / slide_image_name(i, thumb=True)).write_bytes(thumb)
            rendered.append({"index": i, "width": fw, "height": fh})
        return rendered or None
    except Exception as e:
        _log("PIL text page rendering failed for %s: %s", source_path.name, e)
        return None


def _render_pdf_first_page_sync(pdf_path: Path, max_width: int):
    """Render ONLY page 1 of a PDF as a PIL image (cheap card thumbnails)."""
    try:
        import fitz  # PyMuPDF
        from PIL import Image
    except ImportError as e:
        _log("PyMuPDF/Pillow unavailable for first-page render: %s", e)
        return None
    try:
        doc = fitz.open(str(pdf_path))
        with doc:
            if doc.page_count < 1:
                return None
            page = doc[0]
            zoom = max_width / max(page.rect.width, 1.0)
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    except Exception as e:
        _log("fitz first-page render failed for %s: %s", pdf_path, e)
        return None


def _first_page_image_sync(pdf_path: Path, max_width: int):
    """Page-1 PIL image via fitz, falling back to pdftoppm."""
    img = _render_pdf_first_page_sync(pdf_path, max_width)
    if img is not None:
        return img
    # pdftoppm fallback (renders page 1 only with -f/-l)
    import subprocess
    import tempfile

    binary = shutil.which("pdftoppm")
    if not binary:
        return None
    out_dir = Path(tempfile.mkdtemp(prefix="doc_thumb_"))
    try:
        proc = subprocess.run(
            [
                binary,
                "-jpeg",
                "-f",
                "1",
                "-l",
                "1",
                "-scale-to",
                str(max_width),
                str(pdf_path),
                str(out_dir / "page"),
            ],
            capture_output=True,
            timeout=60,
        )
        files = sorted(out_dir.glob("page-*.jpg"))
        if proc.returncode != 0 or not files:
            return None
        from PIL import Image

        return Image.open(str(files[0])).convert("RGB")
    except Exception as e:
        _log("pdftoppm first-page fallback failed: %s", e)
        return None
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


def _thumb_is_fresh(source_path: Path, thumb_path: Path) -> bool:
    """A thumbnail is fresh when it's newer than the source file."""
    try:
        return (
            thumb_path.exists()
            and thumb_path.stat().st_mtime >= source_path.stat().st_mtime
        )
    except OSError:
        return False


async def render_document_thumbnail(
    source_path: Path,
    cache_dir: Optional[Path] = None,
    max_width: int = 480,
) -> Optional[bytes]:
    """Render a CARD thumbnail (first page) for any RAG document.

    Lightweight by design: renders only page 1 (never the whole document),
    caches it at {cache_dir}/thumb.jpg, and reuses an existing full-page
    render's slide_001_t.jpg when present. Returns JPEG bytes, or None when
    the format is unsupported / conversion fails (frontend falls back to a
    file-type icon).
    """
    fmt = document_preview_format(source_path)
    if fmt is None or not source_path.exists():
        return None

    cache_dir = cache_dir or document_preview_cache_dir(source_path)

    # Reuse a fresh full-render thumbnail when the pages cache is valid.
    if is_slide_cache_valid(source_path, cache_dir):
        first_thumb = cache_dir / slide_image_name(1, thumb=True)
        if first_thumb.is_file():
            return first_thumb.read_bytes()

    thumb_path = cache_dir / "thumb.jpg"
    if _thumb_is_fresh(source_path, thumb_path):
        return thumb_path.read_bytes()

    try:
        cache_dir.mkdir(parents=True, exist_ok=True)

        img = None
        if fmt == "pdf":
            img = await asyncio.to_thread(
                _first_page_image_sync, source_path, max_width
            )
        elif fmt in _TEXT_PREVIEW_FORMATS:
            # LibreOffice renders text files beautifully (Writer/Calc);
            # fall back to the PIL text renderer when soffice is missing
            # or refuses the format.
            if is_available():
                pdf_path = await _convert_office_to_pdf_in_cache(source_path, cache_dir)
                if pdf_path is not None:
                    img = await asyncio.to_thread(
                        _first_page_image_sync, pdf_path, max_width
                    )
            if img is None:
                text = _read_text_bytes(source_path.read_bytes())
                first_page = _paginate_text_lines(
                    text, chars_per_line=96, lines_per_page=22, max_pages=1
                )
                if first_page:
                    img = await asyncio.to_thread(
                        _render_text_page_image,
                        source_path.name,
                        first_page[0],
                        1,
                        1,
                        1000,
                    )
        else:  # office documents (docx/doc/xlsx/xls/pptx/ppt)
            if not is_available():
                _log("LibreOffice unavailable — no thumbnail for %s", source_path.name)
                return None
            pdf_path = await _convert_office_to_pdf_in_cache(source_path, cache_dir)
            if pdf_path is None:
                return None
            img = await asyncio.to_thread(_first_page_image_sync, pdf_path, max_width)

        if img is None:
            return None

        thumb_bytes, _, _ = await asyncio.to_thread(
            _pil_to_jpeg_bytes, img, max_width, _THUMB_JPEG_QUALITY
        )
        thumb_path.write_bytes(thumb_bytes)
        _log(
            "document thumbnail rendered: %s (%d bytes)",
            source_path.name,
            len(thumb_bytes),
        )
        return thumb_bytes
    except Exception as e:
        _log("document thumbnail failed for %s: %s", source_path.name, e)
        return None


async def render_document_pages(
    source_path: Path,
    cache_dir: Optional[Path] = None,
    max_width: int = SLIDE_MAX_WIDTH,
    thumb_width: int = SLIDE_THUMB_WIDTH,
    force: bool = False,
) -> Optional[dict]:
    """Render EVERY page of any RAG document as JPEGs + manifest.

    Same contract as convert_document_to_page_images (manifest dict on
    success, None on failure), but accepts the full documents-extension
    universe including legacy .doc/.xls/.ppt and plain text formats.
    Used by the document detail modal's preview pane.
    """
    fmt = document_preview_format(source_path)
    if fmt is None or not source_path.exists():
        _log("document pages: unsupported/missing source %s", source_path.name)
        return None

    cache_dir = cache_dir or document_preview_cache_dir(source_path)

    # Serve from cache when still fresh (unless forced).
    if not force and is_slide_cache_valid(source_path, cache_dir):
        manifest = read_slides_manifest(cache_dir)
        if manifest is not None:
            return manifest

    # Fresh generation — start from an empty cache dir (drops stale thumbs
    # and any previous soffice PDF conversion).
    if cache_dir.exists():
        shutil.rmtree(cache_dir, ignore_errors=True)
    cache_dir.mkdir(parents=True, exist_ok=True)

    if fmt == "pdf":
        pdf_path = source_path
    elif fmt in _TEXT_PREVIEW_FORMATS:
        pdf_path = None
        if is_available():
            pdf_path = await _convert_office_to_pdf_in_cache(source_path, cache_dir)
        if pdf_path is None:
            # PIL text pages — always works, no external tools.
            pages = await asyncio.to_thread(
                _render_text_pages_sync, source_path, cache_dir, max_width, thumb_width
            )
            if pages is None:
                return None
            manifest = await asyncio.to_thread(
                _write_slides_manifest, cache_dir, source_path, pages, None
            )
            _log(
                "text pages rendered: %s → %d pages (PIL)",
                source_path.name,
                manifest["count"],
            )
            return manifest
    else:  # office documents
        if not is_available():
            _log("LibreOffice unavailable — no page preview for %s", source_path.name)
            return None
        pdf_path = await _convert_office_to_pdf_in_cache(source_path, cache_dir)
        if pdf_path is None:
            _log("%s→PDF failed — cannot render document pages", fmt.upper())
            return None

    # PDF (native or soffice-produced) → per-page JPEGs.
    pages = await asyncio.to_thread(
        _render_pdf_pages_with_fitz, pdf_path, cache_dir, max_width, thumb_width
    )
    if pages is None:
        _log("fitz renderer failed — falling back to pdftoppm")
        pages = await asyncio.to_thread(
            _render_pdf_pages_with_pdftoppm, pdf_path, cache_dir, max_width, thumb_width
        )
    if not pages:
        _log("no page images rendered for %s", source_path.name)
        return None

    manifest = await asyncio.to_thread(
        _write_slides_manifest, cache_dir, source_path, pages, None
    )
    _log(
        "document pages rendered: %s → %d pages in %s",
        source_path.name,
        manifest["count"],
        cache_dir.name,
    )
    return manifest


async def convert_legacy_document(
    source_path: Path,
    target_format: str,
    out_dir: Optional[Path] = None,
) -> Optional[Path]:
    """Convert a legacy/odd office file to a modern format via soffice.

    Used by the RAG ingestion path: .doc → .docx, .xls → .xlsx,
    .ppt → .pptx so the standard Python extractors (python-docx /
    openpyxl / python-pptx) can read them. Returns the converted file's
    path (inside out_dir, which defaults to a temp dir), or None.
    """
    if not is_available():
        return None
    if target_format not in ("docx", "xlsx", "pptx", "pdf"):
        return None

    out_dir = out_dir or Path(tempfile.mkdtemp(prefix="lo_convert_"))
    out_dir.mkdir(parents=True, exist_ok=True)

    async with _so_lock:
        rc, _, stderr = await _run_soffice(
            ["--convert-to", target_format, "--outdir", str(out_dir), str(source_path)]
        )
        if rc != 0:
            _log(
                "soffice %s→%s failed (rc=%d): %s",
                source_path.suffix,
                target_format,
                rc,
                stderr[:200],
            )
            return None
        converted = out_dir / f"{source_path.stem}.{target_format}"
        if not converted.exists():
            # soffice sometimes names outputs differently on exotic inputs
            candidates = sorted(out_dir.glob(f"*.{target_format}"))
            if not candidates:
                _log(
                    "%s→%s output not found in %s",
                    source_path.suffix,
                    target_format,
                    out_dir,
                )
                return None
            converted = candidates[0]
        return converted
