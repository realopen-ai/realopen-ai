"""
Tests for the education & social pattern batch (app/services/patterns/):
gradebook, survey_results and expense_split — param coercion, builder
layout math (every formula pins to the row the converter actually
renders), converter round-trips, independently-computed math and the
routing integration in services/excel_gen.py.
"""

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _classify_response(payload):
    """Classifier stub returning a fixed routing JSON payload."""

    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


# ════════════════════════════════════════════════════════════════════
# Fixtures
# ════════════════════════════════════════════════════════════════════

GRADEBOOK_PARAMS = {
    "class_name": "Math 7B",
    "assessments": [
        {"name": "HW 1", "category": "Homework", "max_score": 100},
        {"name": "HW 2", "category": "Homework", "max_score": 50},
        {"name": "Quiz 1", "category": "Quiz", "max_score": 20},
        {"name": "Midterm", "category": "Exam", "max_score": 100},
    ],
    "students": [
        {"name": "Alice", "scores": [85, 40, 18, 92]},
        {"name": "Bob", "scores": [70, None, 15, 81]},
        {"name": "Cara", "scores": [99, 50, 20, 95]},
    ],
    "weights": None,
}

SURVEY_PARAMS = {
    "survey_name": "Customer Feedback 2026",
    "questions": [
        {
            "question": "How satisfied are you with our service?",
            "type": "rating",
            "options": None,
            "responses": {
                "Strongly Agree": 12,
                "Agree": 20,
                "Neutral": 5,
                "Disagree": 2,
            },
        },
        {
            "question": "Did our support team resolve your issue?",
            "type": "yesno",
            "options": None,
            "responses": {"Yes": 28, "No": 9},
        },
        {
            "question": "Which feature do you use most?",
            "type": "choice",
            "options": ["Dashboard", "Reports", "Exports", "API"],
            "responses": {"Dashboard": 14, "Reports": 9, "API": 3},
        },
        {
            "question": "What is your favorite color?",
            "type": "choice",
            "options": None,
            "responses": {"Blue": 7, "Green": 4, "Red": 2},
        },
        # a question without any counts given is OMITTED (stanza rule)
        {
            "question": "Anything else to add?",
            "type": "choice",
            "options": None,
            "responses": None,
        },
    ],
}

EXPENSE_EQUAL_PARAMS = {
    "group_name": "Flat 4B",
    "currency": "USD",
    "members": ["Alice", "Bob", "Claire", "Dave"],
    "expenses": [
        {"description": "Rent", "paid_by": "Alice", "amount": 1600},
        {"description": "Groceries", "paid_by": "Bob", "amount": 240.5},
        {"description": "Internet", "paid_by": "Claire", "amount": 60},
        {"description": "Cleaning", "paid_by": "Alice", "amount": 80},
    ],
}

EXPENSE_WEIGHTED_PARAMS = {
    "group_name": "Ski Trip",
    "members": ["Alice", "Bob", "Claire"],
    "expenses": [
        {
            "description": "Chalet",
            "paid_by": "Alice",
            "amount": 900,
            "shares": {"Alice": 2, "Bob": 1, "Claire": 1},
        },
        {"description": "Ski rental", "paid_by": "Bob", "amount": 300},
        {"description": "Dinner", "paid_by": "Claire", "amount": 150},
    ],
}


# ════════════════════════════════════════════════════════════════════
# Registry wiring
# ════════════════════════════════════════════════════════════════════


class TestRegistry:
    def test_patterns_registered(self):
        for name in ("gradebook", "survey_results", "expense_split"):
            assert name in ep.PATTERN_BUILDERS
            assert name in ep.PATTERN_DESCRIPTIONS
            assert ep.PATTERN_KEYWORDS.get(name)

    def test_module_constants(self):
        import app.services.patterns.gradebook as gb
        import app.services.patterns.expense_split as es
        import app.services.patterns.survey_results as sv

        assert gb.PATTERN_NAME == "gradebook"
        assert sv.PATTERN_NAME == "survey_results"
        assert es.PATTERN_NAME == "expense_split"
        for module in (gb, sv, es):
            assert callable(module.build_spec)
            assert callable(module.coerce_params)

    def test_shortlist_hits(self):
        assert "gradebook" in eg._shortlist_patterns(
            "gradebook for my class with student exam scores"
        )
        assert "survey_results" in eg._shortlist_patterns(
            "tabulate our customer survey poll responses"
        )
        assert "expense_split" in eg._shortlist_patterns(
            "split the trip costs between roommates"
        )
        assert eg._PATTERN_GATE_RE.search("split these costs between us")


# ════════════════════════════════════════════════════════════════════
# Gradebook builder
# ════════════════════════════════════════════════════════════════════


class TestGradebookBuilder:
    def setup_method(self):
        self.spec = ep.build_gradebook_spec(GRADEBOOK_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry(self):
        assert self.spec["filename"] == "math_7b_gradebook.xlsx"
        assert [s["name"] for s in self.norm["sheets"]] == [
            "Gradebook",
            "Assessments",
            "Calc",
        ]
        gb = self.norm["sheets"][0]
        table = gb["tables"][0]
        # start A4 without a table title → header row 4, data 5..7,
        # class-average total row 8
        assert table["start_cell"] == "A4"
        assert len(table["rows"]) == 3
        assert table["headers"] == [
            "Student",
            "HW 1",
            "HW 2",
            "Quiz 1",
            "Midterm",
            "Total Points",
            "Overall %",
            "Grade",
        ]
        assert gb["freeze_panes"] == "B5"

    def test_formula_lattice(self):
        rows = self.norm["sheets"][0]["tables"][0]["rows"]
        alice, bob, cara = rows
        # raw scores align positionally; null score → blank cell
        assert alice[:5] == ["Alice", 85.0, 40.0, 18.0, 92.0]
        assert bob[:5] == ["Bob", 70.0, None, 15.0, 81.0]
        assert cara[:5] == ["Cara", 99.0, 50.0, 20.0, 95.0]
        # every ref points at the row the converter actually renders
        assert alice[5] == "=SUM(B5:E5)"
        assert bob[5] == "=SUM(B6:E6)"
        assert alice[6] == (
            '=IF(SUMPRODUCT((Calc!G5:I5<>"")*Calc!$G$2:$I$2)=0,"",'
            "SUMPRODUCT(Calc!G5:I5,Calc!$G$2:$I$2)/"
            'SUMPRODUCT((Calc!G5:I5<>"")*Calc!$G$2:$I$2))'
        )
        assert bob[6] == (
            '=IF(SUMPRODUCT((Calc!G6:I6<>"")*Calc!$G$2:$I$2)=0,"",'
            "SUMPRODUCT(Calc!G6:I6,Calc!$G$2:$I$2)/"
            'SUMPRODUCT((Calc!G6:I6<>"")*Calc!$G$2:$I$2))'
        )
        # letter grade: nested IF on the overall % (G column)
        assert alice[7] == (
            '=IF(G5="","",IF(G5>=0.9,"A",IF(G5>=0.8,"B",'
            'IF(G5>=0.7,"C",IF(G5>=0.6,"D","F")))))'
        )
        # class-average total row: guarded AVERAGEs over rows 5..7
        total = self.norm["sheets"][0]["tables"][0]["total_row"]
        assert total[0] == "Class Average"
        assert total[1] == '=IF(COUNT(B5:B7)=0,"",AVERAGE(B5:B7))'
        assert total[5] == '=IF(COUNT(F5:F7)=0,"",AVERAGE(F5:F7))'
        assert total[6] == '=IF(COUNT(G5:G7)=0,"",AVERAGE(G5:G7))'
        assert total[7] == ""

    def test_calc_lattice(self):
        calc = self.norm["sheets"][2]
        assert calc["hidden"] is True
        helper, cats = calc["tables"]
        # helper table B4 → header row 4, data 5..7, one column per
        # assessment grouped by category (already grouped here)
        assert helper["start_cell"] == "B4"
        assert helper["headers"] == ["HW 1", "HW 2", "Quiz 1", "Midterm"]
        assert helper["rows"][0][0] == (
            '=IF(Gradebook!B5="","",Gradebook!B5/Assessments!$C$5)'
        )
        assert helper["rows"][0][1] == (
            '=IF(Gradebook!C5="","",Gradebook!C5/Assessments!$C$6)'
        )
        assert helper["rows"][1][1] == (
            '=IF(Gradebook!C6="","",Gradebook!C6/Assessments!$C$6)'
        )
        # category table G4 → one column per category; Homework averages
        # B:C, Quiz D, Exam E (AVERAGE over a range skips blanks)
        assert cats["start_cell"] == "G4"
        assert cats["headers"] == ["Homework", "Quiz", "Exam"]
        assert cats["rows"][0] == [
            '=IF(COUNT(B5:C5)=0,"",AVERAGE(B5:C5))',
            '=IF(COUNT(D5:D5)=0,"",AVERAGE(D5:D5))',
            '=IF(COUNT(E5:E5)=0,"",AVERAGE(E5:E5))',
        ]
        # row 2 mirrors the weights from the Assessments sheet live
        blocks = {b["cell"]: b for b in calc["text_blocks"]}
        assert blocks["G2"]["text"] == "=Assessments!$B$12"
        assert blocks["H2"]["text"] == "=Assessments!$B$13"
        assert blocks["I2"]["text"] == "=Assessments!$B$14"
        assert blocks["G2"]["number_format"] == "0.00%"

    def test_assessments_sheet(self):
        aw = self.norm["sheets"][1]
        ref, weights = aw["tables"]
        # reference block: title 3, header 4, data 5..8
        assert ref["start_cell"] == "A3"
        assert ref["title"] == "Assessments"
        assert ref["rows"] == [
            ["HW 1", "Homework", 100.0],
            ["HW 2", "Homework", 50.0],
            ["Quiz 1", "Quiz", 20.0],
            ["Midterm", "Exam", 100.0],
        ]
        # weights table: title 10, header 11, data 12..14, total 15
        assert weights["start_cell"] == "A10"
        assert weights["title"] == "Category Weights"
        assert weights["rows"][0] == [
            "Homework",
            pytest.approx(1 / 3),
            "=COUNTIF($B$5:$B$8,$A12)",
        ]
        assert weights["rows"][1] == [
            "Quiz",
            pytest.approx(1 / 3),
            "=COUNTIF($B$5:$B$8,$A13)",
        ]
        assert weights["rows"][2] == [
            "Exam",
            pytest.approx(1 / 3),
            "=COUNTIF($B$5:$B$8,$A14)",
        ]
        assert weights["total_row"] == ["Total", "=SUM(B12:B14)", "=SUM(C12:C14)"]
        # category dropdown sourced live from the weights table
        dv = aw["data_validation"][0]
        assert dv["range"] == "B5:B8"
        assert dv["source_range"] == "Assessments!$A$12:$A$14"

    def test_chart_and_conditional_formats(self):
        gb = self.norm["sheets"][0]
        chart = gb["charts"][0]
        assert chart["type"] == "bar"
        assert chart["categories_range"] == "Gradebook!A5:A7"
        assert chart["series"][0]["values_range"] == "Gradebook!G5:G7"
        assert chart["value_numfmt"] == "0%"
        cf = gb["conditional_formats"]
        assert cf[0]["range"] == "G5:G7"
        rules = cf[0]["rules"]
        assert rules[0]["operator"] == "greater_than_or_equal"
        assert rules[0]["value"] == 0.9
        assert rules[0]["fill"] == "C6EFCE"
        assert rules[1]["operator"] == "between"
        assert rules[1]["value"] == [0.6, 0.9]
        assert rules[1]["fill"] == "FFF2CC"
        assert rules[2]["operator"] == "less_than"
        assert rules[2]["value"] == 0.6
        assert rules[2]["fill"] == "FFC7CE"
        assert cf[1]["range"] == "H5:H7"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "gradebook.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Gradebook", "Assessments", "Calc"]
        assert wb["Calc"].sheet_state == "hidden"
        ws = wb["Gradebook"]
        assert ws["B5"].value == 85
        assert ws["C6"].value is None  # null score → blank cell
        assert ws["F5"].value == "=SUM(B5:E5)"
        assert ws["G5"].value.startswith("=IF(SUMPRODUCT((Calc!G5:I5")
        assert ws["H7"].value.startswith('=IF(G7="","",IF(G7>=0.9')
        assert ws["B8"].value == '=IF(COUNT(B5:B7)=0,"",AVERAGE(B5:B7))'
        assert len(ws._charts) == 1
        aw = wb["Assessments"]
        assert aw["C5"].value == 100
        assert aw["B12"].value == pytest.approx(1 / 3)
        assert aw["B15"].value == "=SUM(B12:B14)"
        calc = wb["Calc"]
        assert calc["G2"].value == "=Assessments!$B$12"
        assert (
            calc["B5"].value == '=IF(Gradebook!B5="","",Gradebook!B5/Assessments!$C$5)'
        )
        assert calc["G5"].value == '=IF(COUNT(B5:C5)=0,"",AVERAGE(B5:C5))'
        assert (
            calc["B6"].value == '=IF(Gradebook!B6="","",Gradebook!B6/Assessments!$C$5)'
        )
        assert (
            calc["C6"].value == '=IF(Gradebook!C6="","",Gradebook!C6/Assessments!$C$6)'
        )


class TestGradebookMath:
    """Independent Python verification of the weighted-grade math the
    formula lattice implements (equal-split weights, blanks skipped)."""

    @staticmethod
    def _reference(assessments, students, weights):
        categories = []
        for a in assessments:
            if a["category"] not in categories:
                categories.append(a["category"])
        w = weights or {c: 1 / len(categories) for c in categories}
        out = {}
        for st in students:
            by_cat = {}
            for j, a in enumerate(assessments):
                s = st["scores"][j]
                if s is None:
                    continue
                by_cat.setdefault(a["category"], []).append(s / a["max_score"])
            num = den = 0.0
            for c in categories:
                if c in by_cat:
                    avg = sum(by_cat[c]) / len(by_cat[c])
                    num += avg * w[c]
                    den += w[c]
            overall = num / den if den else None
            letter = (
                None
                if overall is None
                else (
                    "A"
                    if overall >= 0.9
                    else (
                        "B"
                        if overall >= 0.8
                        else "C" if overall >= 0.7 else "D" if overall >= 0.6 else "F"
                    )
                )
            )
            out[st["name"]] = (overall, letter)
        return out

    def test_equal_split_weights(self):
        p = ep.coerce_gradebook_params(GRADEBOOK_PARAMS)
        # weights null → equal split computed in Python as 1/#categories
        assert p["weights"] == {
            "Homework": pytest.approx(1 / 3),
            "Quiz": pytest.approx(1 / 3),
            "Exam": pytest.approx(1 / 3),
        }
        expected = self._reference(
            GRADEBOOK_PARAMS["assessments"],
            GRADEBOOK_PARAMS["students"],
            None,
        )
        # hand-computed: Alice 88.17% (B), Bob 75.33% (C — his missing
        # HW 2 is skipped by the Homework average), Cara 98.17% (A)
        assert expected["Alice"][0] == pytest.approx(0.8816666667)
        assert expected["Alice"][1] == "B"
        assert expected["Bob"][0] == pytest.approx(0.7533333333)
        assert expected["Bob"][1] == "C"
        assert expected["Cara"][0] == pytest.approx(0.9816666667)
        assert expected["Cara"][1] == "A"
        # the emitted weights table carries exactly these weights
        spec = ep.build_gradebook_spec(GRADEBOOK_PARAMS)
        wrows = spec["sheets"][1]["tables"][1]["rows"]
        assert [r[1] for r in wrows] == pytest.approx([1 / 3, 1 / 3, 1 / 3])

    def test_stated_weights(self):
        params = dict(GRADEBOOK_PARAMS, weights={"Homework": 0.3, "Exam": 0.7})
        p = ep.coerce_gradebook_params(params)
        assert p["weights"]["Homework"] == pytest.approx(0.3)
        assert p["weights"]["Exam"] == pytest.approx(0.7)
        assert p["weights"]["Quiz"] == 0.0  # unstated, fully allocated
        expected = self._reference(
            params["assessments"], params["students"], p["weights"]
        )
        assert expected["Alice"][0] == pytest.approx(0.3 * 0.825 + 0.7 * 0.92)
        assert expected["Bob"][0] == pytest.approx(0.3 * 0.7 + 0.7 * 0.81)
        assert expected["Alice"][1] == "B"

    def test_ratio_weights_normalized(self):
        p = ep.coerce_gradebook_params(
            dict(GRADEBOOK_PARAMS, weights={"Homework": 2, "Quiz": 1, "Exam": 1})
        )
        assert p["weights"] == {
            "Homework": pytest.approx(0.5),
            "Quiz": pytest.approx(0.25),
            "Exam": pytest.approx(0.25),
        }

    def test_unscored_student_gets_blank_grade(self):
        params = {
            "assessments": [{"name": "HW 1", "category": "Homework"}],
            "students": [{"name": "Ghost", "scores": [None]}],
        }
        spec = ep.build_gradebook_spec(params)
        norm = eg._normalize_spec(spec)
        row = norm["sheets"][0]["tables"][0]["rows"][0]
        assert row[1] is None
        assert row[2] == "=SUM(B5:B5)"
        # Overall % guard → "" when nothing is scored; Grade follows.
        # (n_a=1 → helper col B, spacer C, category col D; OV col D)
        assert 'Calc!D5:D5<>""' in row[3]
        assert row[4] == (
            '=IF(D5="","",IF(D5>=0.9,"A",IF(D5>=0.8,"B",IF(D5>=0.7,"C",'
            'IF(D5>=0.6,"D","F")))))'
        )


class TestGradebookCoercion:
    def test_quoted_numbers_and_aliases(self):
        p = ep.coerce_gradebook_params(
            {
                "class_name": "Physics",
                "assessments": [
                    {"name": "Lab 1", "category": "Labs", "max_score": "40"}
                ],
                "students": [{"name": "Zoe", "scores": ["38"]}],
                "weights": {"Labs": "1"},
            }
        )
        assert p["assessments"][0]["max_score"] == 40.0
        assert p["students"][0]["scores"] == [38.0]
        assert p["weights"] == {"Labs": 1.0}

    def test_bare_string_entries(self):
        p = ep.coerce_gradebook_params({"assessments": ["Essay"], "students": ["Ann"]})
        assert p["assessments"][0]["category"] == "General"
        assert p["assessments"][0]["max_score"] == 100.0
        assert p["students"][0]["scores"] == [None]

    def test_scores_padded_and_clamped(self):
        p = ep.coerce_gradebook_params(
            {
                "assessments": [{"name": "A"}, {"name": "B"}],
                "students": [{"name": "X", "scores": [5, 6, 7]}, {"name": "Y"}],
            }
        )
        assert p["students"][0]["scores"] == [5.0, 6.0]  # truncated
        assert p["students"][1]["scores"] == [None, None]  # padded
        neg = ep.coerce_gradebook_params(
            {
                "assessments": [{"name": "A"}],
                "students": [{"name": "Z", "scores": [-4]}],
            }
        )
        assert neg["students"][0]["scores"] == [None]

    def test_class_name_default(self):
        p = ep.coerce_gradebook_params({"assessments": ["A"], "students": ["S"]})
        assert p["class_name"] == "Class"

    def test_missing_required_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_gradebook_params({"assessments": [{"name": "A"}]})
        with pytest.raises(ValueError):
            ep.coerce_gradebook_params({"students": [{"name": "S"}]})
        with pytest.raises(ValueError):
            ep.coerce_gradebook_params({})
        with pytest.raises(ValueError):
            ep.coerce_gradebook_params("nope")

    def test_build_raises_on_insufficient(self):
        with pytest.raises(ValueError):
            ep.build_gradebook_spec({"class_name": "Empty"})
        with pytest.raises(ValueError):
            ep.build_gradebook_spec({"assessments": ["A"], "students": []})


class TestGradebookRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "gradebook", "params": GRADEBOOK_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "gradebook for my class with student exam and quiz scores"
        )
        assert result["pattern"] == "gradebook"
        assert result["sheet_names"] == ["Gradebook", "Assessments", "Calc"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
        assert "gradebook template" in result["summary"]

    @pytest.mark.asyncio
    async def test_insufficient_params_fall_back(self, tmp_path, monkeypatch):
        async def fake_ai(brief, requirements, model=None):
            return eg._normalize_spec(
                {
                    "sheets": [
                        {
                            "name": "S",
                            "tables": [
                                {"start_cell": "A1", "headers": ["A"], "rows": [[1]]}
                            ],
                        }
                    ]
                }
            )

        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "gradebook", "params": {}}),
        )
        monkeypatch.setattr(eg, "_generate_workbook_json", fake_ai)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("gradebook with no students listed")
        assert result["pattern"] is None


# ════════════════════════════════════════════════════════════════════
# Survey results builder
# ════════════════════════════════════════════════════════════════════


class TestSurveyBuilder:
    def setup_method(self):
        self.spec = ep.build_survey_results_spec(SURVEY_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry_and_defaults(self):
        assert self.spec["filename"] == "customer_feedback_2026_survey_results.xlsx"
        assert [s["name"] for s in self.norm["sheets"]] == [
            "Survey",
            "Summary",
            "Calc",
        ]
        sv = self.norm["sheets"][0]
        # four usable blocks (the count-less question is omitted)
        assert len(sv["tables"]) == 4
        q1, q2, q3, q4 = sv["tables"]
        # Q1: title 4, header 5, options 6..10, total 11
        assert q1["start_cell"] == "A4"
        assert q1["title"] == "Q1. How satisfied are you with our service? (rating)"
        assert [r[0] for r in q1["rows"]] == [
            "Strongly Agree",
            "Agree",
            "Neutral",
            "Disagree",
            "Strongly Disagree",
        ]
        assert [r[1] for r in q1["rows"]] == [12.0, 20.0, 5.0, 2.0, 0.0]
        # Q2: title 13, header 14, options 15..17, total 18 — default
        # yesno options Yes / No / No answer
        assert q2["start_cell"] == "A13"
        assert [r[0] for r in q2["rows"]] == ["Yes", "No", "No answer"]
        assert [r[1] for r in q2["rows"]] == [28.0, 9.0, 0.0]
        # Q3: request's own options, missing option counts as 0
        assert q3["start_cell"] == "A20"
        assert [r[1] for r in q3["rows"]] == [14.0, 9.0, 0.0, 3.0]
        # Q4: choice without options → the response keys are the options
        assert q4["start_cell"] == "A28"
        assert [r[0] for r in q4["rows"]] == ["Blue", "Green", "Red"]

    def test_formula_lattice(self):
        sv = self.norm["sheets"][0]
        q1, q2 = sv["tables"][0], sv["tables"][1]
        # share % guarded by the question total (row 11 for Q1)
        assert q1["rows"][0][2] == '=IF($B$11>0,B6/$B$11,"n/a")'
        assert q1["rows"][4][2] == '=IF($B$11>0,B10/$B$11,"n/a")'
        assert q1["total_row"] == ["Total", "=SUM(B6:B10)", '=IF($B$11>0,1,"n/a")']
        assert q2["total_row"] == ["Total", "=SUM(B15:B17)", '=IF($B$18>0,1,"n/a")']
        # winner-highlight CF per block, anchored at the range top-left
        # (normalized rules carry the expression under "value")
        cf = sv["conditional_formats"][0]
        assert cf["range"] == "B6:C10"
        assert cf["rules"][0]["value"] == "AND($B6=MAX($B$6:$B$10),$B6>0)"
        assert cf["rules"][0]["fill"] == "FFF2CC"
        assert sv["conditional_formats"][3]["range"] == "B30:C32"

    def test_summary_sheet(self):
        sm = self.norm["sheets"][1]
        table = sm["tables"][0]
        assert table["start_cell"] == "A3"
        assert table["headers"] == ["Question", "Type", "Respondents", "Average Rating"]
        rows = table["rows"]
        # respondents reference each block's Total cell (live)
        assert rows[0][2] == "=Survey!$B$11"
        assert rows[1][2] == "=Survey!$B$18"
        assert rows[2][2] == "=Survey!$B$26"
        assert rows[3][2] == "=Survey!$B$33"
        # rating average = SUMPRODUCT(scale, counts) / total — live
        assert rows[0][3] == (
            '=IF(Survey!$B$11=0,"n/a",'
            "SUMPRODUCT(Calc!$A$6:$A$10,Survey!$B$6:$B$10)/Survey!$B$11)"
        )
        assert rows[1][3] == "—"  # non-rating questions have no average
        assert table["total_row"] == ["Total responses", "", "=SUM(C5:C8)", ""]
        # ONE summary chart: the top question (Q1, 39 responses) is a
        # rating question → pie of its counts
        chart = sm["charts"][0]
        assert len(sm["charts"]) == 1
        assert chart["type"] == "pie"
        assert chart["categories_range"] == "Survey!A6:A10"
        assert chart["series"][0]["values_range"] == "Survey!B6:B10"

    def test_calc_sheet(self):
        calc = self.norm["sheets"][2]
        assert calc["hidden"] is True
        table = calc["tables"][0]
        assert table["start_cell"] == "A5"  # header 5, data 6.. aligned
        assert table["headers"] == ["Q1 scale"]
        assert [r[0] for r in table["rows"]] == [1.0, 2.0, 3.0, 4.0, 5.0]

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "survey.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Survey", "Summary", "Calc"]
        assert wb["Calc"].sheet_state == "hidden"
        ws = wb["Survey"]
        assert ws["A4"].value == "Q1. How satisfied are you with our service? (rating)"
        assert ws["B6"].value == 12
        assert ws["C6"].value == '=IF($B$11>0,B6/$B$11,"n/a")'
        assert ws["B11"].value == "=SUM(B6:B10)"
        assert ws["B15"].value == 28
        assert ws["A29"].value == "Option"  # Q4 header row
        sm = wb["Summary"]
        assert sm["C5"].value == "=Survey!$B$11"
        assert sm["D5"].value.startswith("=IF(Survey!$B$11=0")
        assert len(sm._charts) == 1
        calc = wb["Calc"]
        assert [calc.cell(row=r, column=1).value for r in range(6, 11)] == [
            1,
            2,
            3,
            4,
            5,
        ]

    def test_rating_average_math(self):
        # independent Python check of the emitted average-rating formula
        counts = [12, 20, 5, 2, 0]
        scale = [1, 2, 3, 4, 5]
        total = sum(counts)
        avg = sum(c * s for c, s in zip(counts, scale)) / total
        assert avg == pytest.approx(75 / 39)  # 1.92… — leaning agree
        # and the scale values the Calc sheet carries
        p = ep.coerce_survey_results_params(SURVEY_PARAMS)
        assert p["questions"][0]["scales"] == [1.0, 2.0, 3.0, 4.0, 5.0]
        assert p["questions"][0]["total"] == 39

    def test_pagination_over_six_questions(self):
        questions = []
        for i in range(8):
            questions.append(
                {
                    "question": f"Question number {i + 1}?",
                    "type": "yesno",
                    "responses": {"Yes": i + 1, "No": 1},
                }
            )
        spec = ep.build_survey_results_spec(
            {"survey_name": "Big", "questions": questions}
        )
        norm = eg._normalize_spec(spec)
        names = [s["name"] for s in norm["sheets"]]
        assert names == ["Survey", "Survey2", "Summary"]
        assert len(norm["sheets"][0]["tables"]) == 6
        assert len(norm["sheets"][1]["tables"]) == 2
        # Summary respondents for Q7/Q8 reference Survey2
        rows = norm["sheets"][2]["tables"][0]["rows"]
        assert rows[6][2].startswith("=Survey2!$B$")
        assert rows[7][2].startswith("=Survey2!$B$")


class TestSurveyCoercion:
    def test_defaults_per_type(self):
        p = ep.coerce_survey_results_params(
            {
                "questions": [
                    {
                        "question": "Q rating",
                        "type": "rating",
                        "responses": {"Agree": 3},
                    },
                    {"question": "Q yesno", "type": "yesno", "responses": {"yes": 3}},
                ]
            }
        )
        assert p["questions"][0]["options"] == [
            "Strongly Agree",
            "Agree",
            "Neutral",
            "Disagree",
            "Strongly Disagree",
        ]
        assert p["questions"][1]["options"] == ["Yes", "No", "No answer"]
        # case-insensitive response-key matching
        assert p["questions"][1]["counts"] == [3.0, 0.0, 0.0]

    def test_numeric_response_keys_become_scale(self):
        p = ep.coerce_survey_results_params(
            {
                "questions": [
                    {
                        "question": "Rate us 1-5",
                        "type": "rating",
                        "responses": {"5": 2, "1": 1},
                    }
                ]
            }
        )
        q = p["questions"][0]
        assert q["options"] == ["1", "5"]  # ascending
        assert q["counts"] == [1.0, 2.0]
        # numeric labels carry their own value as the scale
        assert q["scales"] == [1.0, 5.0]

    def test_quoted_counts(self):
        p = ep.coerce_survey_results_params(
            {
                "questions": [
                    {
                        "question": "Pick one",
                        "type": "choice",
                        "options": ["A", "B"],
                        "responses": {"A": "12", "B": "5"},
                    }
                ]
            }
        )
        assert p["questions"][0]["counts"] == [12.0, 5.0]

    def test_type_aliases(self):
        p = ep.coerce_survey_results_params(
            {
                "questions": [
                    {"question": "a", "type": "Yes/No", "responses": {"Yes": 1}},
                    {"question": "b", "type": "multiple choice", "responses": {"X": 1}},
                    {"question": "c", "type": "likert", "responses": {"Agree": 1}},
                    {"question": "d", "type": "weird", "responses": {"X": 1}},
                ]
            }
        )
        assert [q["type"] for q in p["questions"]] == [
            "yesno",
            "choice",
            "rating",
            "choice",
        ]

    def test_question_without_counts_omitted(self):
        p = ep.coerce_survey_results_params(
            {
                "questions": [
                    {"question": "no counts", "type": "choice"},
                    {"question": "garbage counts", "responses": {"A": "n/a"}},
                    {"question": "kept", "responses": {"A": 2}},
                ]
            }
        )
        assert [q["question"] for q in p["questions"]] == ["kept"]

    def test_missing_required_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_survey_results_params({})
        with pytest.raises(ValueError):
            ep.coerce_survey_results_params({"questions": []})
        with pytest.raises(ValueError):
            ep.coerce_survey_results_params(
                {"questions": [{"question": "only text, no counts"}]}
            )
        with pytest.raises(ValueError):
            ep.coerce_survey_results_params({"questions": "nope"})

    def test_question_cap(self):
        questions = [{"question": f"Q{i}", "responses": {"Yes": 1}} for i in range(20)]
        p = ep.coerce_survey_results_params({"questions": questions})
        assert len(p["questions"]) == 12


class TestSurveyRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        params = {
            "survey_name": "Team Pulse",
            "questions": [
                {
                    "question": "Are you happy?",
                    "type": "yesno",
                    "responses": {"Yes": 8, "No": 2},
                }
            ],
        }
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "survey_results", "params": params}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "tabulate the responses from our team survey poll"
        )
        assert result["pattern"] == "survey_results"
        assert result["sheet_names"] == ["Survey", "Summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ════════════════════════════════════════════════════════════════════
# Expense split builder — equal case
# ════════════════════════════════════════════════════════════════════


class TestExpenseSplitEqual:
    def setup_method(self):
        self.spec = ep.build_expense_split_spec(EXPENSE_EQUAL_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry_and_formulas(self):
        assert self.spec["filename"] == "flat_4b_expense_split.xlsx"
        # no unequal shares → no hidden Calc sheet
        assert [s["name"] for s in self.norm["sheets"]] == ["Expenses", "Summary"]
        ex = self.norm["sheets"][0]["tables"][0]
        assert ex["start_cell"] == "A4"
        assert ex["headers"] == ["Description", "Paid By", "Amount", "Shares"]
        assert ex["rows"][0] == ["Rent", "Alice", 1600.0, "equal"]
        assert ex["rows"][3] == ["Cleaning", "Alice", 80.0, "equal"]
        assert ex["total_row"] == ["Total", "", "=SUM(C5:C8)", ""]
        # Paid By dropdown sourced live from the Summary member column
        dv = self.norm["sheets"][0]["data_validation"][0]
        assert dv["range"] == "B5:B8"
        assert dv["source_range"] == "Summary!$A$5:$A$8"

    def test_summary_formula_lattice(self):
        sm = self.norm["sheets"][1]
        table = sm["tables"][0]
        assert table["start_cell"] == "A3"
        assert table["headers"] == [
            "Member",
            "Paid",
            "Share Owed",
            "Balance",
            "Status",
        ]
        alice = table["rows"][0]
        # paid = SUMIF by member over the expenses (rows 5..8)
        assert alice[1] == "=SUMIF(Expenses!$B$5:$B$8,$A5,Expenses!$C$5:$C$8)"
        # equal split → total / members, live
        assert alice[2] == "=Expenses!$C$9/4"
        assert alice[3] == "=B5-C5"
        assert alice[4] == '=IF(D5>0.005,"Is owed",IF(D5<-0.005,"Owes","Settled"))'
        # the balances MUST sum to zero — proven by the SUM total row
        assert table["total_row"] == [
            "Total",
            "=SUM(B5:B8)",
            "=SUM(C5:C8)",
            "=SUM(D5:D8)",
            "",
        ]

    def test_balances_and_settlement_math(self):
        p = ep.coerce_expense_split_params(EXPENSE_EQUAL_PARAMS)
        balances = ep.compute_balances(p["members"], p["expenses"])
        # total 1980.5 over 4 → 495.125 each
        assert balances["Alice"] == pytest.approx(1680.0 - 495.125)
        assert balances["Bob"] == pytest.approx(240.5 - 495.125)
        assert balances["Claire"] == pytest.approx(60.0 - 495.125)
        assert balances["Dave"] == pytest.approx(-495.125)
        assert sum(balances.values()) == pytest.approx(0.0, abs=1e-9)
        # greedy settlement: everyone pays Alice
        transfers = ep.compute_settlement(balances)
        assert [(t[0], t[1]) for t in transfers] == [
            ("Dave", "Alice"),
            ("Claire", "Alice"),
            ("Bob", "Alice"),
        ]
        assert [t[2] for t in transfers] == pytest.approx([495.125, 435.125, 254.625])
        assert sum(t[2] for t in transfers) == pytest.approx(balances["Alice"])
        # the emitted settlement table matches the Python computation
        settle = self.norm["sheets"][1]["tables"][1]
        assert settle["headers"] == ["From (owes)", "To (receives)", "Amount"]
        assert settle["rows"] == [
            ["Dave", "Alice", 495.125],
            ["Claire", "Alice", 435.125],
            ["Bob", "Alice", 254.625],
        ]
        # the build-time-computed exception is disclosed on the sheet
        blocks = {b["cell"]: b for b in self.norm["sheets"][1]["text_blocks"]}
        assert "Computed at build time" in blocks["A11"]["text"]

    def test_chart_and_cf(self):
        sm = self.norm["sheets"][1]
        chart = sm["charts"][0]
        assert chart["type"] == "bar"
        assert chart["categories_range"] == "Summary!A5:A8"
        assert chart["series"][0]["values_range"] == "Summary!D5:D8"
        assert chart["value_numfmt"] == '"$"#,##0.00'
        cf = sm["conditional_formats"][0]
        assert cf["range"] == "E5:E8"
        assert [r["value"] for r in cf["rules"]] == ["Is owed", "Owes", "Settled"]
        assert [r["fill"] for r in cf["rules"]] == ["C6EFCE", "FFC7CE", "FFF2CC"]

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "split.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Expenses", "Summary"]
        ws = wb["Expenses"]
        assert ws["C5"].value == 1600
        assert ws["C9"].value == "=SUM(C5:C8)"
        assert ws["C5"].number_format == '"$"#,##0.00'
        sm = wb["Summary"]
        assert sm["B5"].value == "=SUMIF(Expenses!$B$5:$B$8,$A5,Expenses!$C$5:$C$8)"
        assert sm["C5"].value == "=Expenses!$C$9/4"
        assert sm["D9"].value == "=SUM(D5:D8)"  # zero-sum proof
        assert sm["A15"].value == "Dave"  # first settlement row
        assert sm["C15"].value == 495.125
        assert len(sm._charts) == 1


# ════════════════════════════════════════════════════════════════════
# Expense split builder — unequal shares case
# ════════════════════════════════════════════════════════════════════


class TestExpenseSplitWeighted:
    def setup_method(self):
        self.spec = ep.build_expense_split_spec(EXPENSE_WEIGHTED_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_calc_weight_matrix(self):
        assert [s["name"] for s in self.norm["sheets"]] == [
            "Expenses",
            "Summary",
            "Calc",
        ]
        calc = self.norm["sheets"][2]
        assert calc["hidden"] is True
        weights, shares = calc["tables"]
        # weights table B4 → header 4, data 5..7 aligned with Expenses
        assert weights["start_cell"] == "B4"
        assert weights["headers"] == ["Alice", "Bob", "Claire", "Total"]
        assert weights["rows"][0] == [2.0, 1.0, 1.0, "=SUM(B5:D5)"]
        assert weights["rows"][1] == [1.0, 1.0, 1.0, "=SUM(B6:D6)"]
        # shares table G4 → renormalized fractions, live formulas
        assert shares["start_cell"] == "G4"
        assert shares["rows"][0] == [
            "=IF(SUM($B5:$D5)=0,0,B5/SUM($B5:$D5))",
            "=IF(SUM($B5:$D5)=0,0,C5/SUM($B5:$D5))",
            "=IF(SUM($B5:$D5)=0,0,D5/SUM($B5:$D5))",
        ]
        # expenses carry the shares note
        ex = self.norm["sheets"][0]["tables"][0]
        assert ex["rows"][0][3] == "Alice 2 · Bob 1 · Claire 1"
        assert ex["rows"][1][3] == "equal"

    def test_share_owed_sumproduct(self):
        table = self.norm["sheets"][1]["tables"][0]
        alice = table["rows"][0]
        assert alice[1] == "=SUMIF(Expenses!$B$5:$B$7,$A5,Expenses!$C$5:$C$7)"
        # weighted: SUMPRODUCT of amounts and the member's share column
        assert alice[2] == "=SUMPRODUCT(Expenses!$C$5:$C$7,Calc!G$5:G$7)"
        assert table["rows"][1][2] == "=SUMPRODUCT(Expenses!$C$5:$C$7,Calc!H$5:H$7)"
        assert table["rows"][2][2] == "=SUMPRODUCT(Expenses!$C$5:$C$7,Calc!I$5:I$7)"

    def test_weighted_math_and_settlement(self):
        p = ep.coerce_expense_split_params(EXPENSE_WEIGHTED_PARAMS)
        balances = ep.compute_balances(p["members"], p["expenses"])
        # Chalet 900 weights 2/1/1 → 450/225/225; Ski 300 and Dinner
        # 150 split equally → 100/50 each. Alice: 900 − 600 = +300.
        assert balances == {
            "Alice": pytest.approx(300.0),
            "Bob": pytest.approx(-75.0),
            "Claire": pytest.approx(-225.0),
        }
        assert sum(balances.values()) == pytest.approx(0.0, abs=1e-9)
        transfers = ep.compute_settlement(balances)
        assert transfers == [("Claire", "Alice", 225.0), ("Bob", "Alice", 75.0)]
        # settlement rows land on the sheet exactly as computed
        settle = self.norm["sheets"][1]["tables"][1]
        assert settle["rows"] == [
            ["Claire", "Alice", 225.0],
            ["Bob", "Alice", 75.0],
        ]

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "split_w.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Expenses", "Summary", "Calc"]
        calc = wb["Calc"]
        assert calc["B5"].value == 2
        assert calc["E5"].value == "=SUM(B5:D5)"
        assert calc["G5"].value == "=IF(SUM($B5:$D5)=0,0,B5/SUM($B5:$D5))"
        sm = wb["Summary"]
        assert sm["C5"].value == "=SUMPRODUCT(Expenses!$C$5:$C$7,Calc!G$5:G$7)"
        assert sm["D8"].value == "=SUM(D5:D7)"
        assert sm["A14"].value == "Claire"  # first settlement row
        assert sm["C14"].value == 225.0


class TestExpenseSplitCoercion:
    def test_currency_strings_and_aliases(self):
        p = ep.coerce_expense_split_params(
            {
                "group_name": "Trip",
                "members": ["Ann", "Ben"],
                "expenses": [
                    {
                        "description": "Hotel",
                        "payer": "Ann",
                        "cost": "$1,200",
                    }
                ],
            }
        )
        assert p["expenses"][0]["amount"] == 1200.0
        assert p["expenses"][0]["paid_by"] == "Ann"

    def test_unknown_payer_becomes_member(self):
        p = ep.coerce_expense_split_params(
            {
                "members": ["Ann", "Ben"],
                "expenses": [{"description": "Taxi", "paid_by": "Zoe", "amount": 30}],
            }
        )
        assert p["members"] == ["Ann", "Ben", "Zoe"]
        assert p["expenses"][0]["paid_by"] == "Zoe"

    def test_case_insensitive_member_match(self):
        p = ep.coerce_expense_split_params(
            {
                "members": ["Alice", "Bob"],
                "expenses": [{"description": "Food", "paid_by": "alice", "amount": 20}],
            }
        )
        assert p["members"] == ["Alice", "Bob"]
        assert p["expenses"][0]["paid_by"] == "Alice"

    def test_shares_quoted_and_unknown_keys(self):
        p = ep.coerce_expense_split_params(
            {
                "members": ["Alice", "Bob", "Claire"],
                "expenses": [
                    {
                        "description": "Chalet",
                        "paid_by": "Alice",
                        "amount": 900,
                        "shares": {"Alice": "2", "Bob": 1, "Nobody": 5},
                    }
                ],
            }
        )
        assert p["expenses"][0]["shares"] == {"Alice": 2.0, "Bob": 1.0}

    def test_unusable_shares_fall_back_to_equal(self):
        p = ep.coerce_expense_split_params(
            {
                "members": ["Ann", "Ben"],
                "expenses": [
                    {
                        "description": "Food",
                        "paid_by": "Ann",
                        "amount": 20,
                        "shares": {"Ann": "junk"},
                    }
                ],
            }
        )
        assert p["expenses"][0]["shares"] is None
        spec = ep.build_expense_split_spec(
            {
                "members": ["Ann", "Ben"],
                "expenses": [
                    {
                        "description": "Food",
                        "paid_by": "Ann",
                        "amount": 20,
                        "shares": {"Ann": "junk"},
                    }
                ],
            }
        )
        assert [s["name"] for s in spec["sheets"]] == ["Expenses", "Summary"]

    def test_unusable_expenses_skipped(self):
        p = ep.coerce_expense_split_params(
            {
                "members": ["Ann", "Ben"],
                "expenses": [
                    {"description": "ok", "paid_by": "Ann", "amount": 10},
                    {"description": "", "paid_by": "Ann", "amount": 10},
                    {"description": "no amount", "paid_by": "Ann"},
                    {"description": "negative", "paid_by": "Ann", "amount": -5},
                    {"amount": 10, "paid_by": "Ann"},
                ],
            }
        )
        assert len(p["expenses"]) == 1

    def test_settled_up_note(self):
        params = {
            "members": ["Ann", "Ben"],
            "expenses": [
                {"description": "Food", "paid_by": "Ann", "amount": 20},
                {"description": "Fuel", "paid_by": "Ben", "amount": 20},
            ],
        }
        spec = eg._normalize_spec(ep.build_expense_split_spec(params))
        sm = spec["sheets"][1]
        assert len(sm["tables"]) == 1  # no settlement table needed
        blocks = {b["cell"]: b for b in sm["text_blocks"]}
        # SM_RN=6, SM_TOT=7, SET_NOTE=SM_TOT+2=9
        assert "All settled up" in blocks["A9"]["text"]

    def test_missing_required_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_expense_split_params({"expenses": []})
        with pytest.raises(ValueError):
            ep.coerce_expense_split_params({"members": ["Ann"]})
        with pytest.raises(ValueError):
            ep.coerce_expense_split_params({})
        with pytest.raises(ValueError):
            ep.coerce_expense_split_params({"members": "nope"})
        with pytest.raises(ValueError):
            ep.coerce_expense_split_params({"members": ["Ann"], "expenses": "nope"})

    def test_build_raises_on_insufficient(self):
        with pytest.raises(ValueError):
            ep.build_expense_split_spec({"group_name": "X"})


class TestExpenseSplitRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {"pattern": "expense_split", "params": EXPENSE_EQUAL_PARAMS}
            ),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "split these trip costs between roommates and tell me who owes whom"
        )
        assert result["pattern"] == "expense_split"
        assert result["sheet_names"] == ["Expenses", "Summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
        assert "expense_split template" in result["summary"]
