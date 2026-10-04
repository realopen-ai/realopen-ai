"""Tests for the PPTX schematic thumbnail generator (app/services/thumbnail_gen.py).

Scope:
  - color/geometry helpers: _hex_to_rgb, _emu_to_px
  - shape rendering: _draw_shape dispatch (rect / rtTriangle rotations /
    ellipse / unknown-geometry fallback) with real PIL images + pixel checks
  - text overlay: _draw_text_overlay truncation, blank-text fallback, dark
    vs light background contrast, _get_font fallback chain (incl. the
    all-truetype-sources-missing → load_default branch) and the
    textbbox-measurement-failure fallback
  - _extract_layout_visuals: background color, decorative shapes (prst /
    rotation / fill), placeholder skipping, layout index clamping, and
    the corrupt-shape error branch (defaults kept, error logged) — using
    a real python-pptx Presentation with injected layout XML
  - generate_schematic_thumbnail: happy path (base64 JPEG output), broken
    PPTX file, missing python-pptx import, zero slide dimensions
  - auto_generate_thumbnail_for_template: async wrapper happy/failure path

All rendering is pure PIL (installed); no LibreOffice or subprocess is
spawned anywhere in this module, and no network is touched.
"""

import asyncio
import base64
import io
import sys
import uuid
from pathlib import Path

import pytest
from PIL import Image, ImageFont

BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.services import thumbnail_gen as tg  # noqa: E402

P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
NS = f"{P_NS} {A_NS}"


# ── helpers for building test PPTX files ──────────────────────────────


def _bg_xml(color: str) -> str:
    return (
        f'<p:bg {NS}><p:bgPr><a:solidFill><a:srgbClr val="{color}"/>'
        "</a:solidFill></p:bgPr></p:bg>"
    )


def _shape_xml(
    prst: str,
    left: int,
    top: int,
    width: int,
    height: int,
    fill: str | None,
    rot_deg: float | None = None,
    with_xfrm: bool = True,
) -> str:
    if with_xfrm:
        rot_attr = f' rot="{int(rot_deg * 60000)}"' if rot_deg is not None else ""
        xfrm = (
            f"<a:xfrm{rot_attr}>"
            f'<a:off x="{left}" y="{top}"/>'
            f'<a:ext cx="{width}" cy="{height}"/>'
            "</a:xfrm>"
        )
    else:
        xfrm = ""
    fill_xml = (
        f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else ""
    )
    return (
        f'<p:sp {NS}>'
        '<p:nvSpPr><p:cNvPr id="90" name="deco"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f"<p:spPr>{xfrm}"
        f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
        f"{fill_xml}"
        "</p:spPr>"
        "</p:sp>"
    )


def _build_presentation(tmp_path: Path, bg: str | None = "1F2937", shapes=()):
    """Create a real .pptx whose layout 0 has a background + decorations."""
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import qn

    prs = Presentation()
    layout = prs.slide_layouts[0]
    cSld = layout.element.find(qn("p:cSld"))
    if bg:
        cSld.insert(0, parse_xml(_bg_xml(bg)))
    spTree = cSld.find(qn("p:spTree"))
    for shape in shapes:
        spTree.append(parse_xml(shape))
    path = tmp_path / f"template-{uuid.uuid4().hex[:6]}.pptx"
    prs.save(str(path))
    return prs, path


# ── color + geometry helpers ──────────────────────────────────────────


class TestColorAndGeometryHelpers:
    def test_hex_with_hash(self):
        assert tg._hex_to_rgb("#FF8000") == (255, 128, 0)

    def test_hex_without_hash(self):
        assert tg._hex_to_rgb("00FF7F") == (0, 255, 127)

    def test_short_hex_falls_back_to_gray(self):
        assert tg._hex_to_rgb("#FFF") == (128, 128, 128)

    def test_empty_hex_falls_back_to_gray(self):
        assert tg._hex_to_rgb("") == (128, 128, 128)

    def test_invalid_hex_raises(self):
        with pytest.raises(ValueError):
            tg._hex_to_rgb("ZZZZZZ")

    def test_emu_to_px_scales(self):
        assert tg._emu_to_px(914400, 0.5) == 457200

    def test_emu_to_px_none_is_zero(self):
        assert tg._emu_to_px(None, 2.0) == 0


# ── shape rendering (real PIL canvas + pixel assertions) ──────────────


def _canvas(size=(100, 100)):
    img = Image.new("RGB", size, (255, 255, 255))
    return img


RED = "#FF0000"
BLUE = "#0000FF"


class TestShapeRendering:
    def test_rectangle_drawn_with_fill(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "rect", "left": 10, "top": 10, "width": 40, "height": 30,
             "fill_color": RED, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((20, 20)) == (255, 0, 0)
        assert img.getpixel((60, 45)) == (255, 255, 255)  # just outside

    def test_rectangle_without_fill_is_skipped(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "rect", "left": 0, "top": 0, "width": 100, "height": 100,
             "fill_color": None, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((50, 50)) == (255, 255, 255)

    def test_unknown_geometry_falls_back_to_rect(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "hexagon", "left": 10, "top": 10, "width": 40, "height": 40,
             "fill_color": BLUE, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((30, 30)) == (0, 0, 255)

    def test_ellipse_drawn(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "ellipse", "left": 10, "top": 10, "width": 40, "height": 40,
             "fill_color": RED, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((30, 30)) == (255, 0, 0)  # center
        assert img.getpixel((5, 5)) == (255, 255, 255)  # outside

    def test_ellipse_without_fill_skipped(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "ellipse", "left": 10, "top": 10, "width": 40, "height": 40,
             "fill_color": None, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((30, 30)) == (255, 255, 255)

    @pytest.mark.parametrize(
        "rot,inside,outside",
        [
            # 0°: right angle bottom-left → bottom-left half colored
            (0, (10, 45), (45, 10)),
            # 180°: right angle top-right
            (180, (45, 10), (10, 45)),
            # 270°: right angle top-left
            (270, (20, 12), (40, 45)),
            # 90°: right angle bottom-right
            (90, (40, 45), (12, 12)),
            # near-180 still treated as 180 (±5° tolerance)
            (184, (45, 10), (10, 45)),
        ],
    )
    def test_right_triangle_rotations(self, rot, inside, outside):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "rtTriangle", "left": 10, "top": 10, "width": 40,
             "height": 40, "fill_color": RED, "rotation": rot},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel(inside) == (255, 0, 0), f"rot={rot} inside {inside}"
        assert img.getpixel(outside) == (255, 255, 255), f"rot={rot} outside {outside}"

    def test_right_triangle_without_fill_skipped(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        tg._draw_shape(
            draw,
            {"prst": "rtTriangle", "left": 10, "top": 10, "width": 40,
             "height": 40, "fill_color": None, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((15, 40)) == (255, 255, 255)

    def test_zero_sized_triangle_clamped_to_one_pixel(self):
        img = _canvas()
        from PIL import ImageDraw

        draw = ImageDraw.Draw(img, "RGBA")
        # w/h of 0 are clamped to >= 1 so PIL polygon doesn't fail
        tg._draw_shape(
            draw,
            {"prst": "rtTriangle", "left": 10, "top": 10, "width": 0,
             "height": 0, "fill_color": RED, "rotation": 0},
            scale=1.0,
            img_w=100,
            img_h=100,
        )
        assert img.getpixel((10, 10)) == (255, 0, 0)


# ── text overlay ──────────────────────────────────────────────────────


class TestTextOverlay:
    def _bottom_region_colors(self, img, y0, y1):
        return {
            img.getpixel((x, y))
            for y in range(y0, y1)
            for x in range(10, img.width - 10, 5)
        }

    WHITE = (255, 255, 255)

    def test_overlay_returns_image_and_draws_bottom_bar(self):
        img = _canvas((200, 100))
        result = tg._draw_text_overlay(img, "Template", None)
        assert result is img
        # The bottom strip must contain bar/text pixels, not just background
        colors = self._bottom_region_colors(result, 60, 99)
        assert any(c != self.WHITE for c in colors)

    def test_long_text_is_truncated_without_crash(self):
        img = _canvas((120, 80))
        tg._draw_text_overlay(img, "x" * 60, "#FFFFFF")
        # 25-char truncation + ellipsis happens internally; rendering succeeds
        colors = self._bottom_region_colors(img, 40, 78)
        assert any(c != self.WHITE for c in colors)

    def test_blank_text_falls_back_to_lorem_ipsum(self):
        img = _canvas((120, 80))
        tg._draw_text_overlay(img, "   ", None)
        colors = self._bottom_region_colors(img, 40, 78)
        assert any(c != self.WHITE for c in colors)

    def test_dark_background_uses_light_text(self):
        img = Image.new("RGB", (200, 100), (10, 10, 10))
        tg._draw_text_overlay(img, "Dark", "#0A0A0A")
        # Some light pixels (white text/bar) must appear near the bottom
        found_light = any(
            img.getpixel((x, y))[0] > 100
            for y in range(60, 99)
            for x in range(20, 180, 5)
        )
        assert found_light

    def test_font_loader_returns_a_font(self):
        font = tg._get_font(16)
        assert font is not None

    def test_font_loader_falls_back_to_default_when_truetype_fails(self, monkeypatch):
        # Every bundled TTF path fails to load → PIL default font, no raise.
        # (load_default() itself may call truetype() with an in-memory file —
        # those calls delegate to the real implementation.)
        real_truetype = ImageFont.truetype

        def broken(path, size, *args, **kwargs):
            if isinstance(path, str):
                raise OSError(f"font not readable: {path}")
            return real_truetype(path, size, *args, **kwargs)

        monkeypatch.setattr(tg.ImageFont, "truetype", broken)
        font = tg._get_font(14)
        assert font is not None  # ImageFont.load_default() result

    def test_overlay_survives_text_measurement_failure(self, monkeypatch):
        from PIL import ImageDraw

        def broken_bbox(self, xy, text, font=None, **kwargs):
            raise RuntimeError("textbbox unavailable")

        monkeypatch.setattr(ImageDraw.ImageDraw, "textbbox", broken_bbox)
        img = _canvas((160, 90))
        result = tg._draw_text_overlay(img, "Fallback", None)
        assert result is img
        # Fallback metrics (len*8 x 16) were used — bar + text still drawn.
        colors = self._bottom_region_colors(img, 50, 88)
        assert any(c != self.WHITE for c in colors)


# ── _extract_layout_visuals (real python-pptx XML) ────────────────────


class TestExtractLayoutVisuals:
    def test_extracts_background_and_shapes(self, tmp_path):
        prs, _ = _build_presentation(
            tmp_path,
            bg="1F2937",
            shapes=[
                _shape_xml("rect", 100000, 50000, 200000, 100000, "3B82F6"),
                _shape_xml(
                    "rtTriangle", 0, 0, 300000, 300000, "F59E0B", rot_deg=180
                ),
                _shape_xml("ellipse", 400000, 400000, 100000, 100000, None),
                # spPr without xfrm → left/top/width/height read as None → 0
                _shape_xml("rect", 0, 0, 10, 10, None, with_xfrm=False),
            ],
        )
        visuals = tg._extract_layout_visuals(prs, 0)
        assert visuals["bg_color"] == "1F2937"
        assert visuals["slide_w"] == prs.slide_width
        assert visuals["slide_h"] == prs.slide_height
        assert len(visuals["shapes"]) == 4

        rect = visuals["shapes"][0]
        assert rect["prst"] == "rect"
        assert rect["fill_color"] == "3B82F6"
        assert rect["rotation"] == 0

        triangle = visuals["shapes"][1]
        assert triangle["prst"] == "rtTriangle"
        assert triangle["rotation"] == 180.0
        assert triangle["fill_color"] == "F59E0B"

        ellipse = visuals["shapes"][2]
        assert ellipse["prst"] == "ellipse"
        assert ellipse["fill_color"] is None

        bare = visuals["shapes"][3]
        assert bare["prst"] == "rect"
        assert bare["fill_color"] is None
        assert bare["rotation"] == 0

    def test_placeholders_are_ignored(self, tmp_path):
        # Layout 0 of the default template contains title/subtitle
        # placeholders — none of them may leak into the shape list.
        prs, _ = _build_presentation(tmp_path, bg=None)
        visuals = tg._extract_layout_visuals(prs, 0)
        assert visuals["shapes"] == []

    def test_layout_index_out_of_range_clamps_to_zero(self, tmp_path):
        prs, path = _build_presentation(
            tmp_path, bg="112233", shapes=[_shape_xml("rect", 1, 2, 3, 4, "AABBCC")]
        )
        visuals = tg._extract_layout_visuals(prs, 999)
        assert visuals["bg_color"] == "112233"
        assert len(visuals["shapes"]) == 1

    def test_missing_background_yields_none(self, tmp_path):
        prs, _ = _build_presentation(tmp_path, bg=None)
        visuals = tg._extract_layout_visuals(prs, 0)
        assert visuals["bg_color"] is None

    def test_rotation_270_parsed(self, tmp_path):
        prs, _ = _build_presentation(
            tmp_path,
            bg=None,
            shapes=[_shape_xml("rtTriangle", 5, 5, 100, 100, "FF0000", rot_deg=270)],
        )
        visuals = tg._extract_layout_visuals(prs, 0)
        assert visuals["shapes"][0]["rotation"] == 270.0

    def test_corrupt_shape_is_logged_and_kept_with_defaults(self, tmp_path):
        # A shape whose XML cannot be read must not abort extraction: the
        # error is logged and the shape is kept with its default fields.
        prs, _ = _build_presentation(tmp_path, bg=None)
        real_layout = prs.slide_layouts[0]

        class _BrokenShape:
            is_placeholder = False
            left = top = width = height = 0

            @property
            def _element(self):
                raise RuntimeError("corrupt shape xml")

        class _LayoutProxy:
            def __init__(self, real):
                self.element = real.element
                self.shapes = list(real.shapes) + [_BrokenShape()]

        class _PrsProxy:
            slide_width = prs.slide_width
            slide_height = prs.slide_height
            slide_layouts = [_LayoutProxy(real_layout)]

        visuals = tg._extract_layout_visuals(_PrsProxy(), 0)
        # Only the broken shape is non-placeholder → exactly one entry,
        # carrying the defaults (prst rect, no fill, no rotation).
        assert len(visuals["shapes"]) == 1
        broken = visuals["shapes"][0]
        assert broken == {
            "prst": "rect",
            "left": 0,
            "top": 0,
            "width": 0,
            "height": 0,
            "fill_color": None,
            "rotation": 0,
        }


# ── generate_schematic_thumbnail ──────────────────────────────────────


class TestGenerateSchematicThumbnail:
    def test_happy_path_returns_base64_jpeg(self, tmp_path):
        _, path = _build_presentation(
            tmp_path,
            bg="1F2937",
            shapes=[
                _shape_xml("rect", 0, 0, 9144000, 6858000, "111827"),
                _shape_xml("rtTriangle", 5000000, 0, 4144000, 6858000, "3B82F6"),
            ],
        )
        b64 = tg.generate_schematic_thumbnail(path, "Corporate Dark")
        assert isinstance(b64, str) and len(b64) > 100

        img = Image.open(io.BytesIO(base64.b64decode(b64)))
        assert img.format == "JPEG"
        # Default template is 4:3 → 300x225 within the 400x225 box
        assert img.size == (300, 225)

    def test_custom_size_respected(self, tmp_path):
        _, path = _build_presentation(tmp_path, bg="FFFFFF")
        b64 = tg.generate_schematic_thumbnail(path, "T", width=200, height=150)
        img = Image.open(io.BytesIO(base64.b64decode(b64)))
        assert img.size == (200, 150)

    def test_broken_pptx_returns_none(self, tmp_path):
        bad = tmp_path / "broken.pptx"
        bad.write_bytes(b"this is not a zip file")
        assert tg.generate_schematic_thumbnail(bad, "Broken") is None

    def test_missing_file_returns_none(self, tmp_path):
        assert (
            tg.generate_schematic_thumbnail(tmp_path / "nope.pptx", "X") is None
        )

    def test_zero_slide_dimensions_returns_none(self, tmp_path, monkeypatch):
        _, path = _build_presentation(tmp_path)
        monkeypatch.setattr(
            tg,
            "_extract_layout_visuals",
            lambda prs, layout_index=0: {
                "bg_color": None,
                "slide_w": 0,
                "slide_h": 0,
                "shapes": [],
            },
        )
        assert tg.generate_schematic_thumbnail(path, "Zero") is None

    def test_missing_pptx_library_returns_none(self, tmp_path, monkeypatch):
        _, path = _build_presentation(tmp_path)
        # None in sys.modules makes `from pptx import Presentation` raise ImportError
        monkeypatch.setitem(sys.modules, "pptx", None)
        assert tg.generate_schematic_thumbnail(path, "NoLib") is None

    def test_dark_background_renders_dark_pixels(self, tmp_path):
        _, path = _build_presentation(tmp_path, bg="05070D")
        b64 = tg.generate_schematic_thumbnail(path, "Very Dark", width=80, height=60)
        img = Image.open(io.BytesIO(base64.b64decode(b64)))
        # Top-left corner is pure background (text bar is at the bottom only);
        # JPEG is lossy so allow a small per-channel drift.
        r, g, b = img.getpixel((2, 2))
        assert abs(r - 5) <= 4 and abs(g - 7) <= 4 and abs(b - 13) <= 4


# ── async wrapper ─────────────────────────────────────────────────────


class TestAutoGenerateThumbnail:
    @pytest.mark.asyncio
    async def test_returns_base64_for_valid_template(self, tmp_path):
        _, path = _build_presentation(tmp_path, bg="1F2937")
        b64 = await tg.auto_generate_thumbnail_for_template(path, "Auto")
        assert isinstance(b64, str) and len(b64) > 100
        img = Image.open(io.BytesIO(base64.b64decode(b64)))
        assert img.format == "JPEG"

    @pytest.mark.asyncio
    async def test_returns_none_for_invalid_path(self, tmp_path):
        result = await tg.auto_generate_thumbnail_for_template(
            tmp_path / "missing.pptx", "Auto"
        )
        assert result is None


# ── module logging helper ─────────────────────────────────────────────


def test_log_helper_handles_bad_format_args():
    assert tg._log("count=%d", "oops") is None  # falls back, no raise
    assert tg._log("plain") is None
    assert tg._log("two %s %s", "a") is None  # too few args → fallback


def test_decode_round_trip_of_generated_thumbnail(tmp_path):
    """The base64 payload must decode to the same bytes PIL saved."""
    _, path = _build_presentation(tmp_path, bg="123456")
    b64 = tg.generate_schematic_thumbnail(path, "RoundTrip")
    raw = base64.b64decode(b64, validate=True)
    assert raw.startswith(b"\xff\xd8")  # JPEG magic
    assert asyncio.iscoroutinefunction(tg.auto_generate_thumbnail_for_template)
