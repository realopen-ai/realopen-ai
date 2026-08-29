"""
Tests for the PPTX slide-image viewer pipeline.

Covers three layers:

1. Pure-filesystem units — manifest read/write, cache validity, naming
   (no LibreOffice, no PDF tooling required).
2. API endpoints — /reports/{id}/slides manifest + /reports/{id}/slides/{n}
   image serving, via httpx ASGI transport against an isolated FastAPI app
   with the reports router. LibreOffice calls are mocked — no soffice needed.
3. Real pipeline — PPTX → PDF (LibreOffice) → per-slide JPEGs (PyMuPDF /
   pdftoppm). These tests SKIP automatically when soffice / PyMuPDF are not
   installed on the machine running the tests.
"""

import asyncio
import json
import os
import shutil
import sys
import time
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI

BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.api import reports  # noqa: E402
from app.services.integrations import libreoffice  # noqa: E402

# ══════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════


def _make_pptx(path: Path, slide_titles=("One", "Two", "Three")) -> Path:
    """Build a small real PPTX with python-pptx (title-only slides)."""
    from pptx import Presentation

    prs = Presentation()
    for title in slide_titles:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _write_fake_cache(pptx: Path, count: int = 2) -> Path:
    """Create a COMPLETE, fresh slide cache (manifest + all images)."""
    import io

    from PIL import Image

    cache_dir = libreoffice.slides_cache_dir(pptx)
    cache_dir.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    Image.new("RGB", (1600, 900), (200, 60, 60)).save(buf, format="JPEG")
    full_jpeg = buf.getvalue()
    buf2 = io.BytesIO()
    Image.new("RGB", (360, 202), (200, 60, 60)).save(buf2, format="JPEG")
    thumb_jpeg = buf2.getvalue()

    for i in range(1, count + 1):
        (cache_dir / libreoffice.slide_image_name(i)).write_bytes(full_jpeg)
        (cache_dir / libreoffice.slide_image_name(i, thumb=True)).write_bytes(
            thumb_jpeg
        )

    st = pptx.stat()
    manifest = {
        "version": libreoffice.SLIDES_MANIFEST_VERSION,
        "count": count,
        "width": 1600,
        "height": 900,
        "source_mtime": st.st_mtime,
        "source_size": st.st_size,
        "generated_at": time.time(),
    }
    (cache_dir / "manifest.json").write_text(json.dumps(manifest))
    return cache_dir


def _make_client(tmp_data: Path, monkeypatch) -> httpx.AsyncClient:
    """httpx client wired to a FastAPI app whose reports data dir is tmp_data."""
    app = FastAPI()
    # Mount with the same /api prefix as app.main uses in production.
    app.include_router(reports.router, prefix="/api")
    monkeypatch.setattr(reports, "_get_data_dir", lambda: tmp_data)
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app),
        base_url="http://testserver",
    )


# ══════════════════════════════════════════════════════════════════════
# 1. Pure filesystem units
# ══════════════════════════════════════════════════════════════════════


def test_slide_image_naming():
    assert libreoffice.slide_image_name(1) == "slide_001.jpg"
    assert libreoffice.slide_image_name(12) == "slide_012.jpg"
    assert libreoffice.slide_image_name(1234) == "slide_1234.jpg"
    assert libreoffice.slide_image_name(3, thumb=True) == "slide_003_t.jpg"


def test_slides_cache_dir_derivation(tmp_path):
    pptx = tmp_path / "reports" / "abc.pptx"
    assert libreoffice.slides_cache_dir(pptx) == tmp_path / "reports" / "abc_slides"


def test_read_slides_manifest_missing(tmp_path):
    assert libreoffice.read_slides_manifest(tmp_path) is None


def test_read_slides_manifest_corrupt(tmp_path):
    (tmp_path / "manifest.json").write_text("{not json")
    assert libreoffice.read_slides_manifest(tmp_path) is None


def test_read_slides_manifest_wrong_version(tmp_path):
    (tmp_path / "manifest.json").write_text(json.dumps({"version": 999, "count": 3}))
    assert libreoffice.read_slides_manifest(tmp_path) is None


def test_read_slides_manifest_valid(tmp_path):
    (tmp_path / "manifest.json").write_text(
        json.dumps({"version": libreoffice.SLIDES_MANIFEST_VERSION, "count": 3})
    )
    m = libreoffice.read_slides_manifest(tmp_path)
    assert m is not None and m["count"] == 3


def test_cache_invalid_when_pptx_missing(tmp_path):
    pptx = tmp_path / "nope.pptx"
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_cache_invalid_without_manifest(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    cache = libreoffice.slides_cache_dir(pptx)
    cache.mkdir()
    (cache / "slide_001.jpg").write_bytes(b"x")
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_cache_valid_complete(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    _write_fake_cache(pptx, count=3)
    assert libreoffice.is_slide_cache_valid(pptx) is True


def test_cache_invalid_on_mtime_change(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    _write_fake_cache(pptx, count=2)
    size_before = pptx.stat().st_size
    # Bump mtime only (size unchanged) → must invalidate.
    future = time.time() + 10
    os.utime(pptx, (future, future))
    assert pptx.stat().st_size == size_before
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_cache_invalid_on_size_change(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    _write_fake_cache(pptx, count=2)
    mtime = pptx.stat().st_mtime
    # Append padding so the size changes, then restore the exact old mtime.
    with open(pptx, "ab") as f:
        f.write(b"\0" * 64)
    os.utime(pptx, (mtime, mtime))
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_cache_invalid_when_slide_file_deleted(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    cache = _write_fake_cache(pptx, count=3)
    (cache / "slide_002.jpg").unlink()
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_cache_invalid_when_thumb_deleted(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    cache = _write_fake_cache(pptx, count=3)
    (cache / "slide_003_t.jpg").unlink()
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_write_slides_manifest_roundtrip(tmp_path):
    pptx = _make_pptx(tmp_path / "a.pptx")
    cache = tmp_path / "cache"
    cache.mkdir()
    pages = [{"index": 1, "width": 1600, "height": 900}]
    m = libreoffice._write_slides_manifest(cache, pptx, pages)
    assert m["count"] == 1
    on_disk = libreoffice.read_slides_manifest(cache)
    assert on_disk == m
    assert on_disk["source_size"] == pptx.stat().st_size


# ══════════════════════════════════════════════════════════════════════
# 2. API endpoints (LibreOffice mocked out)
# ══════════════════════════════════════════════════════════════════════


def test_slides_endpoint_404_without_pptx(tmp_path, monkeypatch):
    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/doesnotexist/slides")
            assert r.status_code == 404
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_404_when_libreoffice_missing(tmp_path, monkeypatch):
    _make_pptx(tmp_path / "reports" / "abc.pptx")
    monkeypatch.setattr(libreoffice, "is_available", lambda: False)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 404
            assert "LibreOffice" in r.json()["detail"]
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_generates_and_returns_manifest(tmp_path, monkeypatch):
    _make_pptx(tmp_path / "reports" / "abc.pptx")
    calls = []

    async def fake_convert(p, cache_dir=None, **kw):
        calls.append(p)
        return json.loads((_write_fake_cache(p, count=3) / "manifest.json").read_text())

    monkeypatch.setattr(libreoffice, "is_available", lambda: True)
    monkeypatch.setattr(libreoffice, "convert_pptx_to_slide_images", fake_convert)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 200
            data = r.json()
            assert data["count"] == 3
            assert len(data["slides"]) == 3
            assert data["slides"][0]["index"] == 1
            assert data["slides"][2]["url"].startswith("/api/reports/abc/slides/3?")
            assert "variant=thumb" in data["slides"][0]["thumb_url"]
            assert calls, "convert should have been invoked"
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_serves_from_cache_without_conversion(tmp_path, monkeypatch):
    pptx = _make_pptx(tmp_path / "reports" / "abc.pptx")
    _write_fake_cache(pptx, count=2)

    async def boom(*a, **kw):  # pragma: no cover - must not run
        raise AssertionError("convert must not be called for a fresh cache")

    monkeypatch.setattr(libreoffice, "convert_pptx_to_slide_images", boom)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 200
            assert r.json()["count"] == 2
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_404_when_conversion_fails(tmp_path, monkeypatch):
    _make_pptx(tmp_path / "reports" / "abc.pptx")
    monkeypatch.setattr(libreoffice, "is_available", lambda: True)

    async def fail(*a, **kw):
        return None

    monkeypatch.setattr(libreoffice, "convert_pptx_to_slide_images", fail)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 404
            assert "download" in r.json()["detail"].lower()
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slide_image_endpoint_serves_jpeg(tmp_path, monkeypatch):
    pptx = _make_pptx(tmp_path / "reports" / "abc.pptx")
    _write_fake_cache(pptx, count=2)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides/1")
            assert r.status_code == 200
            assert r.headers["content-type"] == "image/jpeg"
            assert r.content[:2] == b"\xff\xd8"  # JPEG magic
            assert "max-age" in r.headers.get("cache-control", "")

            # Thumbnail variant
            rt = await client.get("/api/reports/abc/slides/2?variant=thumb")
            assert rt.status_code == 200
            assert rt.headers["content-type"] == "image/jpeg"
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slide_image_endpoint_out_of_range(tmp_path, monkeypatch):
    pptx = _make_pptx(tmp_path / "reports" / "abc.pptx")
    _write_fake_cache(pptx, count=2)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            for bad in (0, 3, 99, -1):
                r = await client.get(f"/api/reports/abc/slides/{bad}")
                assert r.status_code == 404, bad
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slide_image_endpoint_404_without_manifest(tmp_path, monkeypatch):
    _make_pptx(tmp_path / "reports" / "abc.pptx")  # pptx exists, no cache

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides/1")
            assert r.status_code == 404
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slide_image_endpoint_rejects_bad_variant(tmp_path, monkeypatch):
    pptx = _make_pptx(tmp_path / "reports" / "abc.pptx")
    _write_fake_cache(pptx, count=1)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides/1?variant=huge")
            assert r.status_code == 422  # Query pattern validation
        finally:
            await client.aclose()

    asyncio.run(run())


def test_report_id_path_traversal_is_neutralized(tmp_path, monkeypatch):
    # A traversal-ish id is treated as a literal basename → no pptx → 404,
    # and nothing outside the reports dir is touched.
    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/..%2F..%2Fetc%2Fpasswd/slides")
            assert r.status_code == 404
        finally:
            await client.aclose()

    asyncio.run(run())


# ══════════════════════════════════════════════════════════════════════
# 3. Real rendering pipeline (skips when tools are absent)
# ══════════════════════════════════════════════════════════════════════


def _fitz_available() -> bool:
    try:
        import fitz  # noqa: F401

        return True
    except ImportError:
        return False


def _make_synthetic_pdf(path: Path, pages: int = 3) -> Path:
    """Build a simple PDF with fitz (colored pages, numbered text)."""
    import fitz

    doc = fitz.open()
    colors = [(0.85, 0.2, 0.2), (0.2, 0.6, 0.85), (0.2, 0.75, 0.35)]
    for i in range(pages):
        page = doc.new_page(width=960, height=540)  # 16:9, like a slide
        page.draw_rect(
            fitz.Rect(0, 0, 960, 540),
            color=None,
            fill=colors[i % len(colors)],
        )
        page.insert_text((60, 120), f"Slide {i + 1}", fontsize=48)
    doc.save(str(path))
    doc.close()
    return path


@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_fitz_renderer_produces_slides_and_thumbs(tmp_path):
    pdf = _make_synthetic_pdf(tmp_path / "deck.pdf", pages=3)
    cache = tmp_path / "cache"
    cache.mkdir()

    pages = libreoffice._render_pdf_pages_with_fitz(
        pdf, cache, libreoffice.SLIDE_MAX_WIDTH, libreoffice.SLIDE_THUMB_WIDTH
    )

    assert pages is not None and len(pages) == 3
    for i in range(1, 4):
        full = cache / libreoffice.slide_image_name(i)
        thumb = cache / libreoffice.slide_image_name(i, thumb=True)
        assert full.is_file() and full.read_bytes()[:2] == b"\xff\xd8"
        assert thumb.is_file() and thumb.read_bytes()[:2] == b"\xff\xd8"

    # Dimensions recorded and bounded.
    assert pages[0]["width"] <= libreoffice.SLIDE_MAX_WIDTH
    assert pages[0]["width"] > 0 and pages[0]["height"] > 0


@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_pdftoppm_fallback_renderer(tmp_path):
    if shutil.which("pdftoppm") is None:
        pytest.skip("pdftoppm not installed")

    pdf = _make_synthetic_pdf(tmp_path / "deck.pdf", pages=2)
    cache = tmp_path / "cache"
    cache.mkdir()

    pages = libreoffice._render_pdf_pages_with_pdftoppm(
        pdf, cache, libreoffice.SLIDE_MAX_WIDTH, libreoffice.SLIDE_THUMB_WIDTH
    )

    assert pages is not None and len(pages) == 2
    for i in range(1, 3):
        full = cache / libreoffice.slide_image_name(i)
        thumb = cache / libreoffice.slide_image_name(i, thumb=True)
        assert full.is_file() and full.read_bytes()[:2] == b"\xff\xd8"
        assert thumb.is_file()


def _soffice_available() -> bool:
    return libreoffice.is_available()


@pytest.mark.skipif(not _soffice_available(), reason="LibreOffice not installed")
@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_end_to_end_pptx_to_slide_images(tmp_path):
    """Real pipeline: python-pptx deck → soffice PDF → per-slide JPEGs."""
    pptx = _make_pptx(tmp_path / "reports" / "e2e.pptx", ("A", "B", "C"))

    manifest = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))

    assert manifest is not None
    assert manifest["count"] == 3
    cache = libreoffice.slides_cache_dir(pptx)
    for i in range(1, 4):
        full = cache / libreoffice.slide_image_name(i)
        assert full.is_file()
        assert full.read_bytes()[:2] == b"\xff\xd8"
        assert full.stat().st_size > 5_000, "slide render suspiciously small"

    # PDF got cached alongside
    assert pptx.with_suffix(".pdf").is_file()

    # Cache is now valid — second call short-circuits (convert_pptx_to_pdf
    # must NOT be invoked again).
    import unittest.mock as mock

    with mock.patch.object(
        libreoffice,
        "convert_pptx_to_pdf",
        side_effect=AssertionError("cached path must not reconvert"),
    ):
        m2 = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))
    assert m2 == manifest

    assert libreoffice.is_slide_cache_valid(pptx) is True


@pytest.mark.skipif(not _soffice_available(), reason="LibreOffice not installed")
@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_stale_cache_regenerates_after_pptx_change(tmp_path):
    pptx = _make_pptx(tmp_path / "reports" / "stale.pptx", ("A", "B"))

    m1 = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))
    assert m1["count"] == 2

    # Simulate regeneration: 3 slides + new mtime.
    time.sleep(0.05)  # ensure mtime differs
    _make_pptx(pptx, ("A", "B", "C"))
    os.utime(pptx)  # refresh mtime to now

    assert libreoffice.is_slide_cache_valid(pptx) is False

    m2 = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))
    assert m2["count"] == 3
    assert libreoffice.is_slide_cache_valid(pptx) is True


@pytest.mark.skipif(not _soffice_available(), reason="LibreOffice not installed")
@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_stale_cached_pdf_is_dropped(tmp_path):
    """A cached PDF older than the PPTX must be deleted before converting."""
    pptx = _make_pptx(tmp_path / "reports" / "pdfstale.pptx", ("A",))
    stale_pdf = pptx.with_suffix(".pdf")
    stale_pdf.write_bytes(b"not a real pdf")
    # Make the PDF look OLD.
    old = time.time() - 3600
    os.utime(stale_pdf, (old, old))

    manifest = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))
    assert manifest is not None
    assert stale_pdf.stat().st_size > 100, "stale pdf should have been replaced"


@pytest.mark.skipif(not _soffice_available(), reason="LibreOffice not installed")
@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_end_to_end_via_api_endpoints(tmp_path, monkeypatch):
    """Full HTTP round trip against the real router + real converters."""
    _make_pptx(tmp_path / "reports" / "via-api.pptx", ("One", "Two"))

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/via-api/slides")
            assert r.status_code == 200
            data = r.json()
            assert data["count"] == 2

            img = await client.get(data["slides"][0]["url"])
            assert img.status_code == 200
            assert img.headers["content-type"] == "image/jpeg"
            assert img.content[:2] == b"\xff\xd8"

            thumb = await client.get(data["slides"][1]["thumb_url"])
            assert thumb.status_code == 200
            assert thumb.headers["content-type"] == "image/jpeg"
        finally:
            await client.aclose()

    asyncio.run(run())
