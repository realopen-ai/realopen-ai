"""
Tests for the investment-portfolio pattern (app/services/patterns/portfolio.py).

Same contract as tests/test_patterns_personal_finance.py:
  • param coercion (holding/transaction normalization, alias keys,
    quoted/currency strings, caps, missing params → ValueError)
  • transaction mode: average-cost position math, chronological trade
    ordering, clamped/fully-closed sells, last-price fallback
  • builder layout math — every formula references the row the
    converter will actually render (title rows included)
  • _build_xlsx round-trip via openpyxl (real in-memory workbooks)
  • independent Python verification of the math behind the formulas
    (invested/value/gain/return/allocation)
  • routing: classifier JSON → template with the AI path skipped

Nothing external is touched: no Ollama (the classifier LLM is stubbed),
no DB, no filesystem outside tmp_path.
"""

import json
import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from app.services.patterns import portfolio as pf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

# ── helpers ───────────────────────────────────────────────────────────


def _norm_txns(raw):
    """Raw transaction dicts → the normalized form the builder uses."""
    return pf.coerce_portfolio_params({"transactions": raw})["transactions"]


def _classify_response(payload):
    """Classifier LLM stub returning a fixed routing JSON object."""

    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


HOLDINGS = [
    {"symbol": "IAM", "name": "IAM SA", "quantity": 100, "average_cost": 95.5,
     "current_price": 102.0},
    {"symbol": "TQM", "name": "TQM SA", "quantity": 50, "average_cost": 70.0,
     "current_price": 63.0},
    {"symbol": "BTC", "name": "Bitcoin", "quantity": 0.5, "average_cost": 50000,
     "current_price": 64000},
]

PORTFOLIO_PARAMS = dict(
    portfolio_name="Retirement Portfolio",
    currency="USD",
    holdings=HOLDINGS,
)

TRANSACTIONS = [
    {"date": "2025-01-10", "symbol": "IAM", "name": "IAM SA", "quantity": 100,
     "price": 95.5, "type": "buy"},
    {"date": "2025-02-01", "symbol": "IAM", "quantity": 50, "price": 98.0,
     "type": "buy"},
    {"date": "2025-03-01", "symbol": "TQM", "quantity": 20, "price": 70.0},
    {"date": "2025-03-05", "symbol": "IAM", "quantity": 40, "price": 101.0,
     "type": "SELL"},
]


# ── registry / stanza contract ────────────────────────────────────────


class TestRegistry:
    def test_pattern_registered(self):
        assert "portfolio" in ep.PATTERN_BUILDERS
        assert ep.PATTERN_BUILDERS["portfolio"] is pf.build_spec
        assert callable(ep.coerce_portfolio_params)
        assert "portfolio" in ep.PATTERN_DESCRIPTIONS

    def test_stanza_names_match_registry_keys(self):
        prompt = (
            Path(__file__).resolve().parent.parent
            / "app"
            / "prompts"
            / "pattern_classifier.md"
        ).read_text(encoding="utf-8")
        assert "<!-- stanza: portfolio -->" in prompt

    def test_keywords_registered(self):
        kws = ep.PATTERN_KEYWORDS["portfolio"]
        for kw in ("portfolio", "invest", "stock", "holding", "crypto", "etf"):
            assert kw in kws

    def test_gate_and_shortlist(self):
        assert eg._PATTERN_GATE_RE.search("track my investment portfolio")
        assert "portfolio" in eg._shortlist_patterns(
            "portfolio of stocks and crypto holdings"
        )


# ── holding normalization ─────────────────────────────────────────────


class TestNormalizeHolding:
    def test_documented_object_form(self):
        h = pf._normalize_holding(
            {"symbol": "IAM", "name": "IAM SA", "quantity": 100,
             "average_cost": 95.5, "current_price": 102.0}
        )
        assert h == {
            "symbol": "IAM",
            "name": "IAM SA",
            "quantity": 100.0,
            "average_cost": 95.5,
            "current_price": 102.0,
        }

    def test_alias_keys_and_quoted_numbers(self):
        h = pf._normalize_holding(
            {"ticker": " X ", "asset": "X Corp", "qty": "1,000",
             "cost_basis": "$5.50", "market_price": "6"}
        )
        assert h["symbol"] == "X"
        assert h["name"] == "X Corp"
        assert h["quantity"] == 1000.0
        assert h["average_cost"] == 5.5
        assert h["current_price"] == 6.0

    def test_compact_list_form(self):
        h = pf._normalize_holding(["AAPL", 100, 150.0])
        assert h["symbol"] == "AAPL"
        assert h["quantity"] == 100.0
        assert h["average_cost"] == 150.0
        # no market price given → flat position (current = avg cost)
        assert h["current_price"] == 150.0
        assert h["name"] == "AAPL"  # name falls back to the symbol

    def test_compact_list_form_full(self):
        h = pf._normalize_holding(["AAPL", 100, 150.0, 175.0, "Apple Inc."])
        assert h["current_price"] == 175.0
        assert h["name"] == "Apple Inc."

    def test_unusable_entries_return_none(self):
        assert pf._normalize_holding(None) is None
        assert pf._normalize_holding(42) is None
        assert pf._normalize_holding("IAM") is None
        assert pf._normalize_holding([]) is None  # < 3 elements
        assert pf._normalize_holding(["AAPL", 100]) is None
        # missing / empty symbol
        assert pf._normalize_holding({"symbol": "  ", "quantity": 1,
                                      "average_cost": 2}) is None
        assert pf._normalize_holding({"quantity": 1, "average_cost": 2}) is None
        # quantity required and must be > 0
        assert pf._normalize_holding({"symbol": "A", "average_cost": 2}) is None
        assert pf._normalize_holding({"symbol": "A", "quantity": 0,
                                      "average_cost": 2}) is None
        assert pf._normalize_holding({"symbol": "A", "quantity": -5,
                                      "average_cost": 2}) is None
        assert pf._normalize_holding({"symbol": "A", "quantity": "n/a",
                                      "average_cost": 2}) is None
        # invested cost is required — cannot compute returns without it
        assert pf._normalize_holding({"symbol": "A", "quantity": 1}) is None
        assert pf._normalize_holding({"symbol": "A", "quantity": 1,
                                      "average_cost": -1}) is None

    def test_missing_or_negative_price_defaults_to_cost(self):
        h = pf._normalize_holding({"symbol": "A", "quantity": 10, "average_cost": 20})
        assert h["current_price"] == 20.0
        h = pf._normalize_holding({"symbol": "A", "quantity": 10,
                                   "average_cost": 20, "current_price": -1})
        assert h["current_price"] == 20.0

    def test_zero_price_is_kept(self):
        h = pf._normalize_holding({"symbol": "A", "quantity": 10,
                                   "average_cost": 20, "current_price": 0})
        assert h["current_price"] == 0.0

    def test_blank_name_falls_back_to_symbol(self):
        h = pf._normalize_holding({"symbol": "A", "quantity": 1,
                                   "average_cost": 2, "name": "   "})
        assert h["name"] == "A"


# ── transaction normalization ─────────────────────────────────────────


class TestNormalizeTransaction:
    def test_documented_object_form(self):
        t = pf._normalize_transaction(
            {"date": "2025-01-10", "symbol": "IAM", "type": "buy",
             "name": "IAM SA", "quantity": 100, "price": 95.5}
        )
        assert t == {"date": "2025-01-10", "symbol": "IAM", "type": "buy",
                     "name": "IAM SA", "quantity": 100.0, "price": 95.5}

    def test_sell_aliases(self):
        for side in ("sell", "Sold", "SALE", "s", "out", "divest"):
            t = pf._normalize_transaction(
                {"symbol": "A", "quantity": 1, "price": 2, "side": side}
            )
            assert t["type"] == "sell", side

    def test_missing_or_unknown_type_becomes_buy(self):
        for side in (None, "", "purchase", "acquire", "weird"):
            t = pf._normalize_transaction(
                {"symbol": "A", "quantity": 1, "price": 2, "type": side}
            )
            assert t["type"] == "buy", side

    def test_compact_list_form(self):
        t = pf._normalize_transaction(["2025-01-10", "IAM", "buy", 100, 95.5])
        assert t["symbol"] == "IAM"
        assert t["type"] == "buy"
        assert t["quantity"] == 100.0
        assert t["price"] == 95.5
        assert t["name"] is None
        t = pf._normalize_transaction(["2025-01-10", "IAM", "sell", 10, 99, "IAM SA"])
        assert t["type"] == "sell"
        assert t["name"] == "IAM SA"

    def test_date_parsing(self):
        t = pf._normalize_transaction({"symbol": "A", "quantity": 1, "price": 2,
                                       "date": "2025-1-5"})
        assert t["date"] == "2025-01-05"
        # invalid calendar date / non-ISO → date kept as None
        for bad in ("2025-13-01", "01/02/2025", 20250101):
            t = pf._normalize_transaction({"symbol": "A", "quantity": 1,
                                           "price": 2, "date": bad})
            assert t["date"] is None, bad
        # non-string first element in list form → no date
        t = pf._normalize_transaction([20250101, "A", "buy", 1, 2])
        assert t["date"] is None

    def test_unusable_entries_return_none(self):
        assert pf._normalize_transaction(None) is None
        assert pf._normalize_transaction("x") is None
        assert pf._normalize_transaction(["A", 1, 2]) is None  # < 5 elements
        assert pf._normalize_transaction(
            {"quantity": 1, "price": 2}) is None  # no symbol
        assert pf._normalize_transaction(
            {"symbol": "A", "price": 2}) is None  # no quantity
        assert pf._normalize_transaction(
            {"symbol": "A", "quantity": 0, "price": 2}) is None
        assert pf._normalize_transaction(
            {"symbol": "A", "quantity": 1}) is None  # no price
        assert pf._normalize_transaction(
            {"symbol": "A", "quantity": 1, "price": -5}) is None

    def test_blank_name_stays_none(self):
        t = pf._normalize_transaction({"symbol": "A", "quantity": 1, "price": 2,
                                       "name": "  "})
        assert t["name"] is None


# ── coerce_portfolio_params ───────────────────────────────────────────


class TestCoercePortfolioParams:
    def test_holdings_key_aliases(self):
        for key in ("holdings", "positions", "assets"):
            p = pf.coerce_portfolio_params({key: list(HOLDINGS)})
            assert len(p["holdings"]) == 3, key

    def test_transactions_key_aliases(self):
        for key in ("transactions", "trades", "history"):
            p = pf.coerce_portfolio_params({key: list(TRANSACTIONS)})
            assert len(p["transactions"]) == 4, key

    def test_invalid_container_types_raise(self):
        with pytest.raises(ValueError):
            pf.coerce_portfolio_params("nope")
        with pytest.raises(ValueError):
            pf.coerce_portfolio_params({"holdings": {"symbol": "A"}})
        with pytest.raises(ValueError):
            pf.coerce_portfolio_params({"transactions": "buy 100 IAM"})

    def test_empty_portfolio_raises(self):
        with pytest.raises(ValueError, match="no holdings or transactions"):
            pf.coerce_portfolio_params({})
        with pytest.raises(ValueError):
            pf.coerce_portfolio_params({"holdings": [], "transactions": []})

    def test_all_holdings_unusable_raises(self):
        with pytest.raises(ValueError):
            pf.coerce_portfolio_params({"holdings": [
                {"symbol": "A"},  # no quantity/cost
                "garbage",
                17,
            ]})

    def test_unusable_entries_are_dropped_not_fatal(self):
        p = pf.coerce_portfolio_params({"holdings": [
            *HOLDINGS,
            {"symbol": "BAD", "quantity": 0, "average_cost": 1},  # dropped
            "junk",  # dropped
        ]})
        assert len(p["holdings"]) == 3

    def test_defaults(self):
        p = pf.coerce_portfolio_params({"holdings": HOLDINGS})
        assert p["portfolio_name"] == "Investment Portfolio"
        assert p["currency"] is None
        assert p["notes"] is None
        assert p["transactions"] == []

    def test_name_notes_and_currency(self):
        p = pf.coerce_portfolio_params({
            "name": "  My Stocks  ",
            "notes": "  long-term horizon  ",
            "currency": "usd",
            "holdings": HOLDINGS,
        })
        assert p["portfolio_name"] == "My Stocks"
        assert p["notes"] == "long-term horizon"
        assert p["currency"] == '"$"#,##0.00'
        # unknown currency → None (builder falls back to plain money fmt)
        p = pf.coerce_portfolio_params({"currency": "XYZ", "holdings": HOLDINGS})
        assert p["currency"] is None
        # notes capped at 1000 chars, blank → None, non-string → None
        p = pf.coerce_portfolio_params({"notes": "x" * 1500, "holdings": HOLDINGS})
        assert len(p["notes"]) == 1000
        p = pf.coerce_portfolio_params({"notes": "   ", "holdings": HOLDINGS})
        assert p["notes"] is None
        p = pf.coerce_portfolio_params({"notes": 42, "holdings": HOLDINGS})
        assert p["notes"] is None

    def test_holdings_capped_at_200(self):
        many = [{"symbol": f"S{i}", "quantity": 1, "average_cost": 1}
                for i in range(205)]
        p = pf.coerce_portfolio_params({"holdings": many})
        assert len(p["holdings"]) == pf.MAX_PORTFOLIO_HOLDINGS

    def test_transactions_capped_at_500(self):
        many = [{"symbol": "S", "quantity": 1, "price": 1, "date": "2025-01-01"}
                for _ in range(505)]
        p = pf.coerce_portfolio_params({"transactions": many})
        assert len(p["transactions"]) == pf.MAX_PORTFOLIO_TRANSACTIONS


# ── transaction mode: positions from the trade history ───────────────


class TestHoldingsFromTransactions:
    def test_average_cost_method(self):
        holdings = pf._holdings_from_transactions(_norm_txns(TRANSACTIONS))
        by_symbol = {h["symbol"]: h for h in holdings}
        # IAM: buys 100@95.5 + 50@98 → cost 14450 / 150 = 96.33̄;
        # sell 40 at the current average cost → qty 110, cost 10596.67
        iam = by_symbol["IAM"]
        assert iam["quantity"] == pytest.approx(110.0)
        assert iam["average_cost"] == pytest.approx(14450.0 / 150.0)
        assert iam["current_price"] == 101.0  # most recent trade price
        assert iam["name"] == "IAM SA"  # from the named trade
        tqm = by_symbol["TQM"]
        assert tqm["quantity"] == 20.0
        assert tqm["average_cost"] == 70.0
        assert tqm["current_price"] == 70.0
        assert tqm["name"] == "TQM"  # no name given → symbol

    def test_dated_trades_sorted_chronologically(self):
        # sell listed before its buy — must be applied in DATE order
        holdings = pf._holdings_from_transactions(_norm_txns([
            {"date": "2025-02-01", "symbol": "A", "quantity": 5, "price": 99,
             "type": "sell"},
            {"date": "2025-01-01", "symbol": "A", "quantity": 10, "price": 100,
             "type": "buy"},
        ]))
        assert len(holdings) == 1
        assert holdings[0]["quantity"] == 5.0
        assert holdings[0]["average_cost"] == 100.0

    def test_undated_trades_applied_after_dated(self):
        holdings = pf._holdings_from_transactions(_norm_txns([
            {"symbol": "A", "quantity": 5, "price": 50, "type": "sell"},  # undated
            {"date": "2025-01-01", "symbol": "A", "quantity": 10, "price": 100,
             "type": "buy"},
        ]))
        # ordered buy → sell leaves 5 shares at cost 500
        assert holdings[0]["quantity"] == 5.0
        assert holdings[0]["average_cost"] == 100.0

    def test_sell_more_than_held_is_clamped(self):
        holdings = pf._holdings_from_transactions(_norm_txns([
            {"date": "2025-01-01", "symbol": "A", "quantity": 10, "price": 100,
             "type": "buy"},
            {"date": "2025-01-02", "symbol": "A", "quantity": 15, "price": 120,
             "type": "sell"},  # oversell — clamped to the 10 held
            {"date": "2025-01-03", "symbol": "B", "quantity": 5, "price": 10,
             "type": "buy"},
        ]))
        # the oversold position fully closed → dropped; B survives
        assert [h["symbol"] for h in holdings] == ["B"]
        assert holdings[0]["quantity"] == 5.0

    def test_oversold_only_position_raises(self):
        with pytest.raises(ValueError, match="closed all positions"):
            pf._holdings_from_transactions(_norm_txns([
                {"date": "2025-01-01", "symbol": "A", "quantity": 10, "price": 100,
                 "type": "buy"},
                {"date": "2025-01-02", "symbol": "A", "quantity": 15, "price": 120,
                 "type": "sell"},
            ]))

    def test_all_positions_closed_raises(self):
        with pytest.raises(ValueError, match="closed all positions"):
            pf._holdings_from_transactions(_norm_txns([
                {"date": "2025-01-01", "symbol": "A", "quantity": 10, "price": 100,
                 "type": "buy"},
                {"date": "2025-01-02", "symbol": "A", "quantity": 10, "price": 110,
                 "type": "sell"},
            ]))

    def test_last_txn_name_wins(self):
        holdings = pf._holdings_from_transactions(_norm_txns([
            {"date": "2025-01-01", "symbol": "A", "quantity": 10, "price": 100,
             "name": "First Name"},
            {"date": "2025-01-02", "symbol": "A", "quantity": 1, "price": 100,
             "name": "Final Name"},
        ]))
        assert holdings[0]["name"] == "Final Name"


# ── builder: holdings mode ────────────────────────────────────────────


class TestPortfolioBuilder:
    def setup_method(self):
        self.spec = pf.build_portfolio_spec(PORTFOLIO_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        self.sheets = {s["name"]: s for s in self.norm["sheets"]}
        # first_data=4, 3 holdings → last_data=6, total_row=7
        self.first, self.last, self.total = 4, 6, 7

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert warnings == []

    def test_sheet_order_and_names(self):
        assert [s["name"] for s in self.norm["sheets"]] == [
            "Overview", "Holdings", "Performance",
        ]
        assert self.spec["filename"] == "retirement_portfolio.xlsx"

    def test_holdings_geometry_and_formulas(self):
        sheet = self.sheets["Holdings"]
        assert sheet["freeze_panes"] == "A4"
        assert sheet["tab_color"] == "16304F"
        table = sheet["tables"][0]
        assert table["start_cell"] == "A3"
        assert table["headers"] == [
            "Symbol", "Asset", "Quantity", "Avg. Cost", "Invested",
            "Current Price", "Current Value", "Gain/Loss", "Return %",
            "Allocation",
        ]
        assert len(table["rows"]) == 3
        row0 = table["rows"][0]
        # static inputs first, then the live formula lattice
        assert row0[0] == "IAM"
        assert row0[1] == "IAM SA"
        assert row0[2] == 100
        assert row0[3] == 95.5
        assert row0[4] == "=C4*D4"  # Invested = Qty × Avg. Cost
        assert row0[5] == 102.0
        assert row0[6] == "=C4*F4"  # Current Value = Qty × Price
        assert row0[7] == "=G4-E4"  # Gain/Loss
        assert row0[8] == "=IF(E4=0,0,H4/E4)"  # Return %
        assert row0[9] == "=IF($G$7=0,0,G4/$G$7)"  # Allocation
        # middle + last rows reference THEIR OWN row numbers
        assert table["rows"][1][9] == "=IF($G$7=0,0,G5/$G$7)"
        assert table["rows"][2][9] == "=IF($G$7=0,0,G6/$G$7)"
        # totals SUM over exactly the data rows
        assert table["total_row"] == [
            "Total", "", "=SUM(C4:C6)", "", "=SUM(E4:E6)", "",
            "=SUM(G4:G6)", "=SUM(H4:H6)", "", "=SUM(J4:J6)",
        ]

    def test_no_round_inside_formulas(self):
        """Display rounding is the number format's job (docstring contract)."""
        for sheet in self.norm["sheets"]:
            for table in sheet.get("tables", []):
                for row in table["rows"]:
                    for cell in row:
                        if isinstance(cell, str):
                            assert "ROUND(" not in cell

    def test_holdings_number_formats(self):
        table = self.sheets["Holdings"]["tables"][0]
        fmts = table["number_formats"]
        assert fmts["C"] == "#,##0.##"  # QTY_FMT
        assert fmts["D"] == '"$"#,##0.00'  # currency USD
        for col in "EFGH":
            assert fmts[col] == '"$"#,##0.00'
        for col in "IJ":
            assert fmts[col] == "0.00%"  # PCT_FMT

    def test_plain_money_format_without_currency(self):
        spec = pf.build_portfolio_spec({"holdings": HOLDINGS})
        fmts = eg._normalize_spec(spec)["sheets"][
            [s["name"] for s in spec["sheets"]].index("Holdings")
        ]["tables"][0]["number_formats"]
        assert fmts["D"] == "#,##0.00"  # MONEY_FMT fallback

    def test_overview_blocks_are_cross_sheet_formulas(self):
        blocks = {b["cell"]: b for b in self.sheets["Overview"]["text_blocks"]}
        assert blocks["A1"]["text"] == "Retirement Portfolio"
        assert blocks["B3"]["text"] == "=COUNTA(Holdings!A4:A6)"
        assert blocks["B4"]["text"] == "=Holdings!E7"
        assert blocks["B5"]["text"] == "=Holdings!G7"
        assert blocks["B6"]["text"] == "=Holdings!H7"
        assert blocks["B7"]["text"] == "=IF(Holdings!E7=0,0,Holdings!H7/Holdings!E7)"
        assert blocks["B8"]["text"] == (
            "=INDEX(Holdings!A4:A6,"
            "MATCH(MAX(Holdings!I4:I6),Holdings!I4:I6,0))"
        )
        assert blocks["B9"]["text"] == (
            "=INDEX(Holdings!A4:A6,"
            "MATCH(MIN(Holdings!I4:I6),Holdings!I4:I6,0))"
        )
        # money vs percent formats on their own blocks
        assert blocks["B4"]["number_format"] == '"$"#,##0.00'
        assert blocks["B7"]["number_format"] == "0.00%"

    def test_overview_pie_chart_ranges_match_layout(self):
        chart = self.sheets["Overview"]["charts"][0]
        assert chart["type"] == "pie"
        assert chart["title"] == "Portfolio Allocation"
        assert chart["categories_range"] == "Holdings!A4:A6"
        assert chart["series"][0]["values_range"] == "Holdings!G4:G6"

    def test_performance_blocks_and_bar_chart(self):
        sheet = self.sheets["Performance"]
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["B3"]["text"] == '=COUNTIF(Holdings!I4:I6,">0")'
        assert blocks["B4"]["text"] == '=COUNTIF(Holdings!I4:I6,"<0")'
        assert blocks["B5"]["text"] == "=MAX(Holdings!I4:I6)"
        assert blocks["B6"]["text"] == "=MIN(Holdings!I4:I6)"
        assert blocks["B7"]["text"] == "=AVERAGE(Holdings!I4:I6)"
        chart = sheet["charts"][0]
        assert chart["type"] == "bar"
        assert chart["show_values"] is True
        assert chart["categories_range"] == "Holdings!A4:A6"
        assert chart["series"][0]["values_range"] == "Holdings!H4:H6"

    def test_notes_default(self):
        assert "Invested = Quantity" in self.sheets["Holdings"]["notes"]
        assert "this overview updates" in self.sheets["Overview"]["notes"]
        # no transaction history → the "computed from transactions" note
        # must NOT appear
        assert "transaction" not in self.sheets["Holdings"]["notes"]

    def test_python_math_matches_formula_lattice(self):
        """Independent re-computation of every emitted formula."""
        invested = [h["quantity"] * h["average_cost"] for h in HOLDINGS]
        value = [h["quantity"] * h["current_price"] for h in HOLDINGS]
        gain = [v - i for i, v in zip(invested, value)]
        ret = [g / i for i, g in zip(invested, gain)]
        alloc = [v / sum(value) for v in value]
        # the fixture is deliberately mixed: one gainer, one loser, one big
        # winner — Best/Worst and COUNTIF stats have real signal
        assert gain[0] > 0 and gain[1] < 0 and gain[2] > 0
        assert sum(alloc) == pytest.approx(1.0)
        assert ret.index(max(ret)) == 2  # BTC best performer
        assert ret.index(min(ret)) == 1  # TQM worst performer
        assert sum(g > 0 for g in gain) == 2  # Holdings Gaining
        assert sum(g < 0 for g in gain) == 1  # Holdings Losing

    def test_single_position_edge_case(self):
        spec = pf.build_portfolio_spec({
            "holdings": [{"symbol": "SOLO", "quantity": 10, "average_cost": 5,
                          "current_price": 7}],
        })
        norm = eg._normalize_spec(spec)
        holdings = [s for s in norm["sheets"] if s["name"] == "Holdings"][0]
        table = holdings["tables"][0]
        assert len(table["rows"]) == 1
        assert table["total_row"][2] == "=SUM(C4:C4)"
        assert table["rows"][0][9] == "=IF($G$5=0,0,G4/$G$5)"
        blocks = {b["cell"]: b for b in [
            s for s in norm["sheets"] if s["name"] == "Overview"
        ][0]["text_blocks"]}
        assert blocks["B3"]["text"] == "=COUNTA(Holdings!A4:A4)"

    def test_flat_position_when_price_missing(self):
        spec = pf.build_portfolio_spec({
            "holdings": [{"symbol": "A", "quantity": 10, "average_cost": 5}],
        })
        rows = eg._normalize_spec(spec)["sheets"][1]["tables"][0]["rows"]
        # current price defaulted to the average cost → zero gain/loss
        assert rows[0][5] == 5.0

    def test_user_notes_flow_to_overview(self):
        spec = pf.build_portfolio_spec(dict(PORTFOLIO_PARAMS,
                                            notes="My custom note"))
        norm = eg._normalize_spec(spec)
        overview = [s for s in norm["sheets"] if s["name"] == "Overview"][0]
        assert overview["notes"] == "My custom note"

    def test_filename_sanitization(self):
        spec = pf.build_portfolio_spec(dict(PORTFOLIO_PARAMS,
                                            portfolio_name="My *Retirement* 2026!"))
        assert spec["filename"] == "my_retirement_2026.xlsx"
        spec = pf.build_portfolio_spec(dict(PORTFOLIO_PARAMS,
                                            portfolio_name="###"))
        assert spec["filename"] == "portfolio.xlsx"
        long_name = "Very Long Portfolio Name " * 5
        spec = pf.build_portfolio_spec(dict(PORTFOLIO_PARAMS,
                                            portfolio_name=long_name))
        base = spec["filename"][: -len(".xlsx")]
        assert len(base) <= 40  # the 40-char name cap

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "portfolio.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Overview", "Holdings", "Performance"]
        ws = wb["Holdings"]
        assert ws["A3"].value == "Symbol"
        assert ws["A4"].value == "IAM"
        assert ws["E4"].value == "=C4*D4"
        assert ws["G4"].value == "=C4*F4"
        assert ws["H4"].value == "=G4-E4"
        assert ws["I4"].value == "=IF(E4=0,0,H4/E4)"
        assert ws["J4"].value == "=IF($G$7=0,0,G4/$G$7)"
        assert ws["A7"].value == "Total"
        assert ws["E7"].value == "=SUM(E4:E6)"
        assert ws["J7"].value == "=SUM(J4:J6)"
        assert ws["E4"].number_format == '"$"#,##0.00'
        assert ws["I4"].number_format == "0.00%"
        assert ws.freeze_panes == "A4"
        assert len(ws._charts) == 0  # charts live on Overview/Performance
        ov = wb["Overview"]
        assert ov["A1"].value == "Retirement Portfolio"
        assert ov["B3"].value == "=COUNTA(Holdings!A4:A6)"
        assert ov["B4"].value == "=Holdings!E7"
        assert len(ov._charts) == 1
        perf = wb["Performance"]
        assert perf["B3"].value == '=COUNTIF(Holdings!I4:I6,">0")'
        assert len(perf._charts) == 1


# ── builder: transaction mode + trade log ─────────────────────────────


class TestPortfolioBuilderTransactions:
    def setup_method(self):
        self.spec = pf.build_portfolio_spec({"transactions": TRANSACTIONS})
        self.norm = eg._normalize_spec(self.spec)
        self.sheets = {s["name"]: s for s in self.norm["sheets"]}

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert warnings == []

    def test_positions_derived_from_trades(self):
        table = self.sheets["Holdings"]["tables"][0]
        rows = {r[0]: r for r in table["rows"]}
        assert set(rows) == {"IAM", "TQM"}
        assert rows["IAM"][2] == pytest.approx(110.0)
        assert rows["IAM"][3] == pytest.approx(14450.0 / 150.0)
        assert rows["IAM"][5] == 101.0
        assert rows["TQM"][2] == 20.0

    def test_holdings_note_mentions_computed_positions(self):
        notes = self.sheets["Holdings"]["notes"]
        assert "average-cost method" in notes
        assert "most recent trade price" in notes

    def test_transactions_sheet_layout(self):
        assert "Transactions" in self.sheets
        sheet = self.sheets["Transactions"]
        table = sheet["tables"][0]
        assert table["headers"] == [
            "Date", "Symbol", "Type", "Quantity", "Price", "Cash Flow",
        ]
        assert len(table["rows"]) == 4
        first = table["rows"][0]
        # the normalizer coerces ISO date strings → datetime.date
        assert first == [date(2025, 1, 10), "IAM", "Buy", 100, 95.5, "=-D4*E4"]
        # sells are cash IN (positive) — row 4 is the SELL at row 7
        sell = table["rows"][3]
        assert sell[2] == "Sell"
        assert sell[5] == "=D7*E7"
        # Net Cash Flow total over the exact trade rows
        assert table["total_row"] == ["", "", "", "", "Net Cash Flow",
                                      "=SUM(F4:F7)"]
        fmts = table["number_formats"]
        assert fmts["A"] == "yyyy-mm-dd"
        assert fmts["F"] == "#,##0.00"  # no currency given → plain money

    def test_undated_trades_render_blank_date(self):
        spec = pf.build_portfolio_spec({"transactions": [
            {"symbol": "A", "quantity": 10, "price": 100},
        ]})
        norm = eg._normalize_spec(spec)
        sheet = {s["name"]: s for s in norm["sheets"]}["Transactions"]
        assert sheet["tables"][0]["rows"][0][0] == ""

    def test_both_modes_holdings_win(self):
        """Explicit holdings win; transactions still render as the log."""
        params = dict(holdings=list(HOLDINGS), transactions=list(TRANSACTIONS))
        spec = pf.build_portfolio_spec(params)
        norm = eg._normalize_spec(spec)
        sheets = {s["name"]: s for s in norm["sheets"]}
        table = sheets["Holdings"]["tables"][0]
        # explicit holdings are the positions (IAM qty 100, not 110)
        rows = {r[0]: r for r in table["rows"]}
        assert rows["IAM"][2] == 100
        assert "BTC" in rows
        # ...and the trade log is still rendered
        assert "Transactions" in sheets
        assert len(sheets["Transactions"]["tables"][0]["rows"]) == 4
        # no "computed from transactions" note (holdings were explicit)
        assert "average-cost method" not in sheets["Holdings"]["notes"]

    def test_built_workbook_transactions(self, tmp_path):
        out = tmp_path / "portfolio_txn.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Overview", "Holdings", "Performance",
                                 "Transactions"]
        ws = wb["Transactions"]
        assert ws["C4"].value == "Buy"
        assert ws["F4"].value == "=-D4*E4"
        assert ws["C7"].value == "Sell"
        assert ws["F7"].value == "=D7*E7"
        assert ws["E8"].value == "Net Cash Flow"
        assert ws["F8"].value == "=SUM(F4:F7)"


# ── routing ───────────────────────────────────────────────────────────


class TestPortfolioRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg, "_call_llm",
            _classify_response({
                "pattern": "portfolio",
                "params": {
                    "portfolio_name": "My Stocks",
                    "currency": "USD",
                    "holdings": [
                        {"symbol": "IAM", "quantity": 100, "average_cost": 95.5,
                         "current_price": 102},
                    ],
                },
            }),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "portfolio tracker: 100 IAM shares at 95.50 average cost, "
            "now trading at 102"
        )
        assert result["pattern"] == "portfolio"
        assert result["sheet_names"] == ["Overview", "Holdings", "Performance"]
        assert result["filename"] == "my_stocks.xlsx"
        assert result["chart_count"] == 2
        assert result["formula_count"] > 10
        assert "portfolio template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    @pytest.mark.asyncio
    async def test_transaction_mode_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg, "_call_llm",
            _classify_response({
                "pattern": "portfolio",
                "params": {"transactions": TRANSACTIONS},
            }),
        )
        monkeypatch.setattr(eg, "_generate_workbook_json",
                            AssertionError("AI path must not run"))
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "investment tracker: bought 100 IAM at 95.5, another 50 at 98, "
            "sold 40 at 101"
        )
        assert result["pattern"] == "portfolio"
        assert result["sheet_names"] == ["Overview", "Holdings", "Performance",
                                         "Transactions"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
