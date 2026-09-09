"""
Tests for the deterministic template layer (the app/services/patterns/
package: one module per pattern, discovered dynamically) and its
routing integration in services/excel_gen.py.

The templates exist because models hand-writing formula lattices emit
refs that don't match the rendered layout (off by 1..N rows) — see the
round-3 user report (interest reading the balance three rows back,
starting balance on an empty cell → negative-balance spiral). These
tests pin down:
  • param coercion (quoted/currency/percent strings from classifiers)
  • builder layout math — every formula references the row the
    converter will actually render (title rows included!)
  • routing: classifier → template with the AI path skipped; every
    failure mode falls back to the AI path without raising
  • the new validation warning that flags the off-by-N class
  • the heal no longer false-positives on cross-sheet qualified refs
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

# ── Param coercion ───────────────────────────────────────────────────


class TestCoercion:
    def test_to_number_handles_strings(self):
        assert ep.to_number("25000") == 25000.0
        assert ep.to_number("$25,000") == 25000.0
        assert ep.to_number("25 000 €") == 25000.0
        assert ep.to_number("6.5") == 6.5
        assert ep.to_number(7) == 7.0
        assert ep.to_number("abc") is None
        assert ep.to_number(None) is None
        assert ep.to_number(True) is None

    def test_to_rate_percent_semantics(self):
        assert ep.to_rate("6.5%") == 0.065
        assert ep.to_rate(6.5) == 0.065  # bare > 1 → percent
        assert ep.to_rate(0.065) == 0.065  # decimal stays
        assert ep.to_rate("0.065") == 0.065
        assert ep.to_rate(0) == 0.0
        assert ep.to_rate(None) is None

    def test_to_int(self):
        assert ep.to_int("36") == 36
        assert ep.to_int(3.0) == 3
        assert ep.to_int(3.5) is None

    def test_to_iso_date(self):
        assert ep.to_iso_date("2026-09-01") == "2026-09-01"
        assert ep.to_iso_date("2026-9-1") == "2026-09-01"
        assert ep.to_iso_date("2026-13-01") is None
        assert ep.to_iso_date("sept 2026") is None

    def test_amortization_params_aliases_and_defaults(self):
        p = ep.coerce_amortization_params(
            {"principal": "$25,000", "interest_rate": "6.5%", "term_years": 3}
        )
        assert p["loan_amount"] == 25000.0
        assert p["annual_rate"] == 0.065
        assert p["term_months"] == 36
        assert p["payment"] is None

    def test_amortization_params_missing_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_amortization_params({"loan_amount": 25000})
        with pytest.raises(ValueError):
            ep.coerce_amortization_params(
                {"loan_amount": -5, "annual_rate": 0.05, "term_months": 12}
            )


# ── Amortization builder: layout math is the contract ────────────────


class TestAmortizationBuilder:
    def setup_method(self):
        self.spec = ep.build_amortization_spec(
            dict(
                loan_amount=25000,
                annual_rate=0.065,
                term_months=36,
                start_date="2026-09-01",
            )
        )
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry_and_formula_lattice(self):
        sched = self.norm["sheets"][0]
        assert sched["name"] == "Amortization"
        table = sched["tables"][0]
        # start A7, no table title → header row 7, data 8..43, total 44
        assert table["start_cell"] == "A7"
        assert len(table["rows"]) == 36
        first, second = table["rows"][0], table["rows"][1]
        # every ref points at the row the converter actually renders.
        # NO ROUND() — display rounding is the number format's job
        # (formula-level rounding accumulates cent drift: a final
        # balance of 0.14 instead of 0).
        assert first[1] == "=$B$5"
        assert first[2] == "=B8-D8"
        assert first[3] == "=$B$2*$B$3/12"
        assert first[4] == "=$B$2-C8"
        assert second[3] == "=E8*$B$3/12"
        assert second[4] == "=E8-C9"
        # params blocks populate B2..B5 above the table
        blocks = {b["cell"]: b for b in sched["text_blocks"]}
        assert blocks["B2"]["text"] == 25000.0
        assert blocks["B3"]["text"] == 0.065
        assert blocks["B4"]["text"] == 36
        assert blocks["B5"]["text"] == "=-PMT(B3/12,B4,B2)"
        # B5 (Monthly Payment) carries the currency format so it reads
        # as money like the schedule columns, not a raw float
        assert blocks["B5"]["number_format"] == "#,##0.00"

    def test_summary_sheet_cross_refs(self):
        summary = self.norm["sheets"][1]
        assert summary["name"] == "Summary"
        rows = summary["tables"][0]["rows"]
        refs = {r[0]: r[1] for r in rows}
        assert refs["Loan Amount"] == "=Amortization!$B$2"
        # 36 data rows: 8..43 → totals row 44
        assert refs["Total Paid"] == "=Amortization!$B$44"
        assert refs["Total Interest"] == "=Amortization!$D$44"
        assert refs["Final Balance"] == "=Amortization!$E$43"

    def test_chart_ranges_match_layout(self):
        chart = self.norm["sheets"][0]["charts"][0]
        assert chart["categories_range"] == "Amortization!A8:A43"
        assert chart["series"][0]["values_range"] == "Amortization!E8:E43"

    def test_built_workbook_formulas(self, tmp_path):
        out = tmp_path / "amort.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        ws = wb["Amortization"]
        assert ws["B5"].value == "=-PMT(B3/12,B4,B2)"
        assert ws["B2"].value == 25000
        assert ws["B3"].value == 0.065
        assert ws["D9"].value == "=E8*$B$3/12"
        assert ws["E9"].value == "=E8-C9"
        assert ws["D43"].value == "=E42*$B$3/12"
        # total row: placeholders expanded to the real data rows
        assert ws["B44"].value == "=SUM(B8:B43)"
        assert ws["D44"].value == "=SUM(D8:D43)"

    def test_b5_monthly_payment_uses_currency_format(self, tmp_path):
        # B5 must render with the requested currency's format, exactly
        # like the schedule's money columns
        spec = ep.build_amortization_spec(
            dict(
                loan_amount=25000,
                annual_rate=0.065,
                term_months=36,
                currency="USD",
            )
        )
        norm = eg._normalize_spec(spec)
        out = tmp_path / "amort.xlsx"
        eg._build_xlsx(norm, out)
        ws = load_workbook(out)["Amortization"]
        assert ws["B5"].number_format == '"$"#,##0.00'
        # and the default (no currency param) still gets money format
        out2 = tmp_path / "a2.xlsx"
        eg._build_xlsx(
            eg._normalize_spec(
                ep.build_amortization_spec(
                    dict(loan_amount=1000, annual_rate=0.05, term_months=12)
                )
            ),
            out2,
        )
        assert load_workbook(out2)["Amortization"]["B5"].number_format == "#,##0.00"

    def test_heal_never_fires_on_template(self):
        # the normalized spec must be unchanged by a second heal pass —
        # formulas never reference the header row (row 7)
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after


# ── Invoice builder ──────────────────────────────────────────────────


class TestInvoiceBuilder:
    @staticmethod
    def _blocks_by_text(sheet):
        return {b["text"]: b["cell"] for b in sheet["text_blocks"]}

    def test_layout_and_totals(self):
        spec = ep.build_invoice_spec(
            dict(
                seller="Acme\n12 Rue",
                client="Client SARL",
                invoice_number="INV-1",
                date="2026-02-10",
                due_date="2026-03-12",
                items=[
                    {"description": "Logo", "quantity": 1, "unit_price": 1200},
                    {"description": "Consulting", "quantity": 7.5, "unit_price": 90},
                ],
                tax_rate=0.2,
            )
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        sheet = eg._normalize_spec(spec)["sheets"][0]
        table = sheet["tables"][0]
        # Meta rows 3-5 (inv #, date, due), blank 6, From heading 7 +
        # name 8 + address 9, blank 10, Bill To 11 + name 12, blank 13
        # → items table starts at row 14
        items_start = 14
        first_data = items_start + 1
        assert table["start_cell"] == f"A{items_start}"
        assert table["rows"][0][3] == f"=B{first_data}*C{first_data}"
        assert table["rows"][1][3] == f"=B{first_data + 1}*C{first_data + 1}"
        assert table["total_row"][3] == "=SUM(D{first_row}:D{last_row})"
        blocks = {b["text"]: b["cell"] for b in sheet["text_blocks"]}
        subtotal_row = first_data + 2  # data 15-16 → total row 17
        # professional party blocks
        assert blocks["From"] == "A7"
        assert blocks["Acme"] == "A8"
        assert blocks["12 Rue"] == "A9"
        assert blocks["Bill To"] == "A11"
        assert blocks["Client SARL"] == "A12"
        # totals: tax on the full subtotal when there is no discount
        assert blocks["Tax (20%)"] == f"C{subtotal_row + 1}"  # label col C
        assert blocks["TOTAL DUE"] == f"C{subtotal_row + 2}"
        by_cell = {b["cell"]: b for b in sheet["text_blocks"]}
        assert by_cell[f"D{subtotal_row + 1}"]["text"] == f"=D{subtotal_row}*0.2"
        assert (
            by_cell[f"D{subtotal_row + 2}"]["text"]
            == f"=D{subtotal_row}+D{subtotal_row + 1}"
        )

    def test_discount_percentage_applied_before_tax(self):
        # "10% discount then 20% tax": discount = -subtotal*10%,
        # tax computed on the DISCOUNTED subtotal, total sums all three
        spec = ep.build_invoice_spec(
            dict(
                items=[
                    {"description": "A", "quantity": 1, "unit_price": 100},
                    {"description": "B", "quantity": 2, "unit_price": 50},
                ],
                tax_rate=0.2,
                discount=10,
                discount_type="percent",
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        by_cell = {b["cell"]: b for b in sheet["text_blocks"]}
        # no meta/parties → items at row 4, data 5-6, subtotal 7,
        # discount 8, tax 9, TOTAL DUE 10
        assert by_cell["C8"]["text"] == "Discount (10%)"
        assert by_cell["D8"]["text"] == "=-D7*0.1"
        assert by_cell["C9"]["text"] == "Tax (20%)"
        # tax base = subtotal + (negative) discount
        assert by_cell["D9"]["text"] == "=(D7+D8)*0.2"
        assert by_cell["D10"]["text"] == "=D7+D8+D9"

    def test_discount_percentage_detected_from_percent_string(self):
        spec = ep.build_invoice_spec(
            dict(
                items=[{"description": "A", "quantity": 1, "unit_price": 200}],
                tax_rate=0.2,
                discount="10%",
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        by_cell = {b["cell"]: b for b in sheet["text_blocks"]}
        # items 4, data 5, subtotal 6, discount 7, tax 8, total 9
        assert by_cell["D7"]["text"] == "=-D6*0.1"
        assert by_cell["D8"]["text"] == "=(D6+D7)*0.2"

    def test_discount_fixed_amount_tax_on_discounted_subtotal(self):
        spec = ep.build_invoice_spec(
            dict(
                items=[{"description": "A", "quantity": 1, "unit_price": 100}],
                tax_rate=0.1,
                discount=30,
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        by_cell = {b["cell"]: b for b in sheet["text_blocks"]}
        # items 4, data 5, subtotal 6, discount 7, tax 8, total 9
        assert by_cell["C7"]["text"] == "Discount"
        assert by_cell["D7"]["text"] == "=-30"
        assert by_cell["D8"]["text"] == "=(D6+D7)*0.1"
        assert by_cell["D9"]["text"] == "=D6+D7+D8"

    def test_structured_seller_and_client_blocks(self):
        spec = ep.build_invoice_spec(
            dict(
                seller={
                    "name": "Acme Corp",
                    "address": "12 Rue des Fleurs, Casablanca 20250",
                    "phone": "+212 600 000 000",
                    "fax": "+212 500 000 000",
                    "email": "billing@acme.com",
                },
                client={
                    "name": "Client SARL",
                    "id": "ICE-12345",
                    "address": "45 Ave Hassan II, Rabat 10100",
                    "phone": "+212 611 111 111",
                    "email": "ap@client.ma",
                },
                items=[{"description": "Work", "quantity": 1, "unit_price": 10}],
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        texts = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        # no meta rows → seller block starts right after the title:
        # heading 4, bold name 5, then the optional detail lines
        assert texts["A4"] == "From"
        assert texts["A5"] == "Acme Corp"
        assert texts["A6"] == "12 Rue des Fleurs, Casablanca 20250"
        assert texts["A7"] == "Phone: +212 600 000 000"
        assert texts["A8"] == "Fax: +212 500 000 000"
        assert texts["A9"] == "Email: billing@acme.com"
        # client block with ID line
        assert texts["A11"] == "Bill To"
        assert texts["A12"] == "Client SARL"
        assert texts["A13"] == "ID: ICE-12345"
        assert texts["A14"] == "45 Ave Hassan II, Rabat 10100"
        assert texts["A15"] == "Phone: +212 611 111 111"
        assert texts["A16"] == "Email: ap@client.ma"

    def test_notes_terms_section_rendered(self):
        spec = ep.build_invoice_spec(
            dict(
                items=[{"description": "A", "quantity": 1, "unit_price": 100}],
                notes="Payment due within 30 days.\nLate payments accrue 2% monthly interest.",
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        texts = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        # items 4, data 5, subtotal 6, TOTAL DUE 7 (no tax/discount)
        # → notes section starts 2 rows below the total
        assert texts["A9"] == "Notes / Terms & Conditions"
        assert texts["A10"] == "Payment due within 30 days."
        assert texts["A11"] == "Late payments accrue 2% monthly interest."

    def test_tax_zero_discount_positive_no_row_collision(self):
        # regression: tax=0 + discount>0 must not stack Discount and
        # TOTAL DUE on the same row
        spec = ep.build_invoice_spec(
            dict(
                items=[{"description": "A", "quantity": 1, "unit_price": 100}],
                discount=10,
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        by_cell = {b["cell"]: b for b in sheet["text_blocks"]}
        disc_row = int(
            next(c for c, b in by_cell.items() if b["text"] == "Discount")[1:]
        )  # label C7
        total_row = int(
            next(c for c, b in by_cell.items() if b["text"] == "TOTAL DUE")[1:]
        )
        assert total_row == disc_row + 1
        assert by_cell[f"D{total_row}"]["text"] == f"=D{disc_row - 1}+D{disc_row}"

    def test_missing_items_raise(self):
        with pytest.raises(ValueError):
            ep.build_invoice_spec({"seller": "x"})


# ── Budget builder ───────────────────────────────────────────────────


class TestBudgetBuilder:
    def test_title_row_math(self):
        # regression: tables have title fields — data starts TWO rows
        # below the anchor, the total row right after the data
        spec = ep.build_budget_spec(
            dict(
                period="monthly",
                income=[
                    {"source": "Salary", "amount": 6500},
                    {"source": "Freelance", "amount": 800},
                ],
                expenses=[
                    {"category": "Rent", "amount": 1800},
                    {"category": "Food", "amount": 550},
                    {"category": "Bus", "amount": 220},
                ],
            )
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        sheet = eg._normalize_spec(spec)["sheets"][0]
        income_t, expenses_t, summary_t = sheet["tables"]
        # Income anchored at 3 with title → data 5-6, total 7
        assert income_t["start_cell"] == "A3"
        assert income_t["total_row"][0] == "Total Income"
        # Expenses anchored at 9 → data 11-13, total 14
        assert expenses_t["start_cell"] == "A9"
        # Summary anchored at 16 → data 18-21 referencing B7 / B14
        rows = {r[0]: r[1] for r in summary_t["rows"]}
        assert rows["Total Income"] == "=B7"
        assert rows["Total Expenses"] == "=B14"
        assert rows["Net (Income - Expenses)"] == "=B7-B14"
        assert 'TEXT((B7-B14)/B7,"0.0%")' in rows["Savings Rate"]

    def test_no_freeze_panes_budget_scrolls_freely(self):
        # budgets are meant to be scrolled — nothing stays pinned
        spec = ep.build_budget_spec(
            dict(
                income=[{"source": "Salary", "amount": 6500}],
                expenses=[{"category": "Rent", "amount": 1800}],
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        assert "freeze_panes" not in sheet or sheet.get("freeze_panes") is None

    def test_pie_chart_next_to_tables_visualizes_expenses(self, tmp_path):
        spec = ep.build_budget_spec(
            dict(
                period="monthly",
                income=[{"source": "Salary", "amount": 6500}],
                expenses=[
                    {"category": "Rent", "amount": 1800},
                    {"category": "Food", "amount": 550},
                    {"category": "Bus", "amount": 220},
                ],
            )
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        charts = sheet.get("charts") or []
        assert len(charts) == 1
        pie = charts[0]
        assert pie["type"] == "pie"
        assert pie["title"] == "Monthly Expenses"
        # anchored in column D — next to the A/B table band
        assert pie["anchor"] == "D3"
        # one slice per expense row: 1 income line → Income table 3-6,
        # Expenses anchored at row 8 (title) → data rows 10-12
        assert pie["categories_range"] == "Budget!A10:A12"
        assert pie["series"][0]["values_range"] == "Budget!B10:B12"
        # real workbook: the pie chart object must exist on the sheet
        out = tmp_path / "budget.xlsx"
        eg._build_xlsx(eg._normalize_spec(spec), out)
        ws = load_workbook(out)["Budget"]
        assert len(ws._charts) == 1
        assert ws.freeze_panes is None

    def test_income_only_budget_has_no_pie(self):
        # nothing to slice — no expenses, no pie
        spec = ep.build_budget_spec(dict(income=[{"source": "Salary", "amount": 6500}]))
        sheet = eg._normalize_spec(spec)["sheets"][0]
        assert not sheet.get("charts")

    def test_expenses_only_budget(self):
        spec = ep.build_budget_spec(
            dict(expenses=[{"category": "Rent", "amount": 1000}])
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        rows = {r[0]: r[1] for r in sheet["tables"][-1]["rows"]}
        assert rows["Total Income"] == "=0"
        assert rows["Net (Income - Expenses)"].startswith("=0-")

    def test_no_lines_raise(self):
        with pytest.raises(ValueError):
            ep.build_budget_spec({"period": "monthly"})


# ── Routing integration ──────────────────────────────────────────────


def _classify_response(payload):
    import json

    async def fake_llm(messages):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


class TestPatternRouting:
    @pytest.mark.asyncio
    async def test_amortization_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {
                    "pattern": "amortization",
                    "params": {
                        "loan_amount": 25000,
                        "annual_rate": 0.065,
                        "term_months": 36,
                        "start_date": "2026-09-01",
                    },
                }
            ),
        )

        async def must_not_run(brief, requirements):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("loan amortization $25,000 6.5% 3 years")
        assert result["pattern"] == "amortization"
        assert result["sheet_count"] == 2
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
        # workbook summary returned to the agent (report/pptx contract)
        assert result["sheet_names"] == ["Amortization", "Summary"]
        assert result["table_count"] == 2
        assert result["chart_count"] == 1
        assert result["formula_count"] > 100  # 36 rows x 4 + totals + PMT
        assert "2 sheets" in result["summary"]
        assert "amortization template" in result["summary"]

    @pytest.mark.asyncio
    async def test_none_falls_back(self, tmp_path, monkeypatch):
        import json

        async def fake_llm(messages):
            return json.dumps({"pattern": "none", "params": {}})

        async def fake_ai(brief, requirements):
            return eg._normalize_spec(
                {
                    "sheets": [
                        {
                            "name": "Data",
                            "tables": [
                                {"start_cell": "A1", "headers": ["A"], "rows": [[1]]}
                            ],
                        }
                    ]
                }
            )

        monkeypatch.setattr(eg, "_call_llm", fake_llm)
        monkeypatch.setattr(eg, "_generate_workbook_json", fake_ai)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("compile budget data")
        assert result["pattern"] is None

    @pytest.mark.asyncio
    async def test_gate_skips_classifier(self, tmp_path, monkeypatch):
        async def must_not_run_llm(messages):
            raise AssertionError("classifier must be skipped without keywords/digits")

        async def fake_ai(brief, requirements):
            return eg._normalize_spec(
                {
                    "sheets": [
                        {
                            "name": "S",
                            "tables": [
                                {"start_cell": "A1", "headers": ["A"], "rows": [[1]]}
                            ],
                        }
                    ]
                }
            )

        monkeypatch.setattr(eg, "_call_llm", must_not_run_llm)
        monkeypatch.setattr(eg, "_generate_workbook_json", fake_ai)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("grocery list")
        assert result["pattern"] is None

    @pytest.mark.asyncio
    async def test_insufficient_params_fall_back(self, tmp_path, monkeypatch):
        async def fake_ai(brief, requirements):
            return eg._normalize_spec(
                {
                    "sheets": [
                        {
                            "name": "S",
                            "tables": [
                                {"start_cell": "A1", "headers": ["A"], "rows": [[1]]}
                            ],
                        }
                    ]
                }
            )

        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {"pattern": "amortization", "params": {"loan_amount": 25000}}
            ),
        )
        monkeypatch.setattr(eg, "_generate_workbook_json", fake_ai)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("loan request")
        assert result["pattern"] is None

    @pytest.mark.asyncio
    async def test_garbage_classifier_falls_back(self, tmp_path, monkeypatch):
        async def fake_llm(messages):
            return "I cannot answer that"

        async def fake_ai(brief, requirements):
            return eg._normalize_spec(
                {
                    "sheets": [
                        {
                            "name": "S",
                            "tables": [
                                {"start_cell": "A1", "headers": ["A"], "rows": [[1]]}
                            ],
                        }
                    ]
                }
            )

        monkeypatch.setattr(eg, "_call_llm", fake_llm)
        monkeypatch.setattr(eg, "_generate_workbook_json", fake_ai)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("budget for my salary")
        assert result["pattern"] is None


# ── New off-by-N warning + heal cross-sheet fix ──────────────────────


class TestOffByNDetection:
    def _sheet_with_formula(self, formula, blocks=None):
        return {
            "name": "S",
            "text_blocks": blocks or [],
            "tables": [
                {
                    "start_cell": "A2",
                    "title": "T",
                    "headers": ["A", "B"],
                    "rows": [[1, formula]],
                }
            ],
        }

    def test_unpopulated_above_header_ref_warns(self):
        # column B is inside the table band (headers A,B); B2 sits on
        # the title row — only the anchor A2 holds text, so B2 renders
        # empty and the formula silently reads 0
        spec = {"sheets": [self._sheet_with_formula("=B2*0.5")]}
        _errors, warnings = eg.validate_workbook_spec(spec)
        assert any("above the table" in w for w in warnings)

    def test_populated_ref_does_not_warn(self):
        spec = {
            "sheets": [
                self._sheet_with_formula(
                    "=$B$1*0.5", blocks=[{"cell": "B1", "text": 25000}]
                )
            ]
        }
        _errors, warnings = eg.validate_workbook_spec(spec)
        assert not [w for w in warnings if "above the table" in w]

    def test_cross_sheet_qualified_ref_does_not_warn(self):
        # a qualified ref targets another sheet — not our header row
        spec = {"sheets": [self._sheet_with_formula("=Other!$B$2*0.5")]}
        _errors, warnings = eg.validate_workbook_spec(spec)
        assert not [w for w in warnings if "above the table" in w]

    def test_heal_ignores_cross_sheet_refs_to_own_header_row(self):
        # regression (found via the template Summary sheet): a table on
        # "Summary" whose formulas reference Amortization!$B$4 must NOT
        # be detected as "references own header row 4"
        spec = {
            "sheets": [
                {
                    "name": "Amortization",
                    "tables": [
                        {
                            "start_cell": "A7",
                            "headers": ["M", "P"],
                            "rows": [[1, 2]],
                        }
                    ],
                },
                {
                    "name": "Summary",
                    "tables": [
                        {
                            "start_cell": "A3",
                            "title": "Loan Summary",
                            "headers": ["Item", "Value"],
                            "rows": [
                                ["Term", "=Amortization!$B$7"],
                                [
                                    "X",
                                    "=Amortization!$B$4",
                                ],  # row 4 = Summary's header row
                            ],
                        }
                    ],
                },
            ]
        }
        norm = eg._normalize_spec(spec)
        summary_rows = norm["sheets"][1]["tables"][0]["rows"]
        # heal must not shift the cross-sheet refs (no false positive)
        assert summary_rows[1][1] == "=Amortization!$B$4"
