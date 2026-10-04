"""
Tests for the inventory-tracker pattern (app/services/patterns/inventory.py).

Same contract as tests/test_patterns_personal_finance.py:
  • param coercion (item/status/condition/category/location lists,
    alias keys, quoted numbers, caps, missing params → ValueError)
  • builder layout math — 50-row Inventory grid with the ID / total /
    days-owned formulas pinned to rows 5..54; live Categories /
    Locations dropdown sources; hidden Calc helper lattice; Dashboard
    KPIs + Top-5 / recency / lending tables
  • _build_xlsx round-trip via openpyxl (real in-memory workbooks)
  • independent Python verification of the math behind the formulas
    (valuation, low-stock, units-in-stock)
  • routing: classifier JSON → template with the AI path skipped

Nothing external is touched: no Ollama (the classifier LLM is stubbed),
no DB, no filesystem outside tmp_path.
"""

import json
import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from app.services.patterns import inventory as inv  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

# ── helpers ───────────────────────────────────────────────────────────


def _classify_response(payload):
    """Classifier LLM stub returning a fixed routing JSON object."""

    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


ITEMS = [
    {
        "name": "MacBook Pro 14",
        "category": "Electronics",
        "brand_model": "Apple M3 Pro",
        "quantity": 1,
        "min_stock": None,
        "location": "Home Office",
        "purchase_date": "2024-06-15",
        "purchase_price": 2400,
        "current_value": 1800,
        "condition": "Excellent",
        "serial_number": "C02X1234",
        "warranty_expiry": "2027-06-15",
        "status": "In Stock",
        "notes": "main laptop",
    },
    {
        "name": "Phone chargers",
        "quantity": 2,
        "min_stock": 3,  # 2 <= 3 → low stock
        "category": "Electronics",
        "location": "Living Room",
        "purchase_price": 20,
        "current_value": 15,
        "status": "In Stock",
    },
    {
        "name": "Cordless drill",
        "category": "Tools",
        "quantity": 1,
        "location": "Garage",
        "purchase_price": 150,
        "current_value": 90,
        "condition": "Good",
        "status": "Lent Out",
        "lent_to": "Sam",
        "date_lent": "2025-11-01",
    },
    {
        "name": "Vintage camera",
        "category": "Other",
        "quantity": 1,
        "location": "Storage Unit",
        "purchase_price": 300,
        "current_value": 450,
        "status": "In Stock",
    },
    # a bare item-name string — quantity defaults to 1, status In Stock
    "Coffee grinder",
]

INVENTORY_PARAMS = dict(
    inventory_name="Home Inventory",
    currency="USD",
    categories=["Electronics", "Tools", "Kitchen", "Books", "Other"],
    locations=["Home Office", "Living Room", "Garage", "Storage Unit"],
    items=ITEMS,
)


# ── registry / stanza contract ────────────────────────────────────────


class TestRegistry:
    def test_pattern_registered(self):
        assert "inventory" in ep.PATTERN_BUILDERS
        assert ep.PATTERN_BUILDERS["inventory"] is inv.build_spec
        assert callable(ep.coerce_inventory_params)
        assert "inventory" in ep.PATTERN_DESCRIPTIONS

    def test_stanza_names_match_registry_keys(self):
        prompt = (
            Path(__file__).resolve().parent.parent
            / "app"
            / "prompts"
            / "pattern_classifier.md"
        ).read_text(encoding="utf-8")
        assert "<!-- stanza: inventory -->" in prompt

    def test_keywords_registered(self):
        kws = ep.PATTERN_KEYWORDS["inventory"]
        for kw in ("inventory", "belonging", "collection", "warranty", "serial"):
            assert kw in kws

    def test_gate_and_shortlist(self):
        assert eg._PATTERN_GATE_RE.search("home inventory of my belongings")
        assert "inventory" in eg._shortlist_patterns(
            "inventory of my belongings with warranty dates and serial numbers"
        )


# ── list cleaning + status/condition normalization ────────────────────


class TestCleanHelpers:
    def test_clean_list_strips_dedupes_and_caps(self):
        out = inv._clean_list(
            ["  Electronics ", "electronics", "ELECTRONICS ", 42, "", "Tools"],
            limit=20,
        )
        assert out == ["Electronics", "Tools"]
        assert inv._clean_list(None, 20) == []
        assert inv._clean_list("Electronics", 20) == []
        assert inv._clean_list([1, 2, 3], 20) == []
        # entries beyond the limit never appear
        out = inv._clean_list([f"Cat{i}" for i in range(30)], limit=5)
        assert out == [f"Cat{i}" for i in range(5)]
        # 40-char per-entry cap
        out = inv._clean_list(["x" * 80], limit=5)
        assert out == ["x" * 40]

    def test_normalize_status_aliases(self):
        cases = {
            "in stock": "In Stock",
            "stock": "In Stock",
            "available": "In Stock",
            "lent": "Lent Out",
            "on loan": "Lent Out",
            "borrowed": "Lent Out",
            "sold": "Sold",
            "lost": "Lost",
            "missing": "Lost",
            "retired": "Retired",
            "disposed": "Retired",
            "donated": "Retired",
        }
        for raw, want in cases.items():
            assert inv._normalize_status(raw) == want, raw
        # unknown or non-string statuses → None
        assert inv._normalize_status("exploded") is None
        assert inv._normalize_status(None) is None
        assert inv._normalize_status(7) is None
        # title-cased canonical spellings pass through
        assert inv._normalize_status("Lent Out") == "Lent Out"
        assert inv._normalize_status("In Stock") == "In Stock"

    def test_normalize_condition(self):
        for cond in inv.CONDITIONS:
            assert inv._normalize_condition(cond) == cond
            assert inv._normalize_condition(cond.lower()) == cond
        assert inv._normalize_condition("in stock") is None  # not a condition
        assert inv._normalize_condition(None) is None
        assert inv._normalize_condition("mint") is None


# ── item normalization + coerce_inventory_params ──────────────────────


class TestCoerceInventoryParams:
    def test_items_key_aliases(self):
        for key in ("items", "item_list", "inventory"):
            p = inv.coerce_inventory_params({key: list(ITEMS)})
            assert len(p["items"]) == 5, key

    def test_invalid_container_types_raise(self):
        with pytest.raises(ValueError, match="params must be an object"):
            inv.coerce_inventory_params("nope")
        with pytest.raises(ValueError, match="items must be an array"):
            inv.coerce_inventory_params({"items": "just one thing"})

    def test_bare_string_item_gets_defaults(self):
        p = inv.coerce_inventory_params({"items": ["Coffee grinder", "  "]})
        item = p["items"][0]
        assert item == {
            "name": "Coffee grinder",
            "category": None,
            "brand_model": None,
            "quantity": 1,
            "min_stock": None,
            "location": None,
            "purchase_date": None,
            "purchase_price": None,
            "current_value": None,
            "condition": None,
            "serial_number": None,
            "warranty_expiry": None,
            "status": "In Stock",
            "lent_to": None,
            "date_lent": None,
            "notes": None,
        }
        # a blank string is dropped, not turned into an empty-name row
        assert len(p["items"]) == 1

    def test_dict_item_full_aliases(self):
        p = inv.coerce_inventory_params({"items": [{
            "title": "Drill",  # name alias
            "room": "Garage",  # location alias
            "qty": "2",  # quantity alias (quoted number)
            "count": None,  # ignored — quantity wins first
            "min": 1,  # min_stock alias
            "bought": "2024-02-01",  # purchase_date alias
            "cost": "$150",  # purchase_price alias
            "worth": "90",  # current_value alias
            "sn": "DR-77",  # serial alias
            "warranty": "2026-02-01",  # warranty alias
            "state": "on loan",  # status alias
            "borrower": "Sam",  # lent_to alias
            "lent_date": "2025-11-01",  # date_lent alias
            "comment": "bring back",  # notes alias
            "brand": "Bosch",  # brand_model alias
            "state2": None,
        }]})
        item = p["items"][0]
        assert item["name"] == "Drill"
        assert item["location"] == "Garage"
        assert item["quantity"] == 2.0
        assert item["min_stock"] == 1.0
        assert item["purchase_date"] == "2024-02-01"
        assert item["purchase_price"] == 150.0
        assert item["current_value"] == 90.0
        assert item["serial_number"] == "DR-77"
        assert item["warranty_expiry"] == "2026-02-01"
        assert item["status"] == "Lent Out"
        assert item["lent_to"] == "Sam"
        assert item["date_lent"] == "2025-11-01"
        assert item["notes"] == "bring back"
        assert item["brand_model"] == "Bosch"
        assert item["condition"] is None  # not given

    def test_unusable_items_dropped(self):
        p = inv.coerce_inventory_params({"items": [
            ITEMS[0], 42, None, {"name": "   "}, {"category": "No name"},
        ]})
        assert len(p["items"]) == 1

    def test_quantity_fallbacks(self):
        p = inv.coerce_inventory_params({"items": [
            {"name": "A", "quantity": -3},  # negative → 1
            {"name": "B", "quantity": "n/a"},  # non-numeric → 1
            {"name": "C"},  # missing → 1
        ]})
        assert [i["quantity"] for i in p["items"]] == [1, 1, 1]
        # zero is a legitimate count (all-out item) — kept as-is
        p = inv.coerce_inventory_params({"items": [{"name": "D", "quantity": 0}]})
        assert p["items"][0]["quantity"] == 0.0

    def test_string_caps_and_bad_values(self):
        p = inv.coerce_inventory_params({"items": [{
            "name": "N" * 100,  # capped at 60
            "category": "C" * 100,  # capped at 40
            "brand_model": "B" * 100,  # capped at 60
            "serial_number": "S" * 100,  # capped at 40
            "notes": "N" * 300,  # capped at 200
            "location": "L" * 100,  # capped at 40
            "lent_to": "P" * 100,  # capped at 40
            "purchase_date": "2024-13-45",  # invalid → None
            "warranty_expiry": "not a date",  # → None
            "date_lent": "2025-11-01",
        }]})
        item = p["items"][0]
        assert len(item["name"]) == 60
        assert len(item["category"]) == 40
        assert len(item["brand_model"]) == 60
        assert len(item["serial_number"]) == 40
        assert len(item["notes"]) == 200
        assert len(item["location"]) == 40
        assert len(item["lent_to"]) == 40
        assert item["purchase_date"] is None
        assert item["warranty_expiry"] is None
        assert item["date_lent"] == "2025-11-01"

    def test_title_currency_categories_locations_defaults(self):
        p = inv.coerce_inventory_params({})
        assert p["title"] == "Inventory Tracker"
        assert p["currency"] is None
        assert p["categories"] == list(inv.DEFAULT_CATEGORIES)
        assert p["locations"] == list(inv.DEFAULT_LOCATIONS)
        assert p["items"] == []
        assert p["notes"] is None

    def test_title_aliases_and_caps(self):
        p = inv.coerce_inventory_params({"title": "  My Stuff "})
        assert p["title"] == "My Stuff"
        p = inv.coerce_inventory_params({"name": "N" * 100})
        assert len(p["title"]) == 60
        p = inv.coerce_inventory_params({"currency": "mad"})
        assert p["currency"] == "MAD"
        p = inv.coerce_inventory_params({"currency": "  "})
        assert p["currency"] is None

    def test_custom_categories_and_locations(self):
        p = inv.coerce_inventory_params(INVENTORY_PARAMS)
        assert p["categories"] == ["Electronics", "Tools", "Kitchen", "Books",
                                   "Other"]
        assert p["locations"] == ["Home Office", "Living Room", "Garage",
                                  "Storage Unit"]

    def test_category_list_aliases(self):
        for key in ("categories", "category_list"):
            p = inv.coerce_inventory_params({key: ["A", "B"]})
            assert p["categories"] == ["A", "B"], key
        for key in ("locations", "location_list"):
            p = inv.coerce_inventory_params({key: ["A", "B"]})
            assert p["locations"] == ["A", "B"], key

    def test_categories_capped_at_20_locations_at_15(self):
        p = inv.coerce_inventory_params(
            {"categories": [f"C{i}" for i in range(30)]}
        )
        assert len(p["categories"]) == inv.MAX_CATEGORY_ROWS
        p = inv.coerce_inventory_params(
            {"locations": [f"L{i}" for i in range(30)]}
        )
        assert len(p["locations"]) == inv.MAX_LOCATION_ROWS

    def test_items_capped_at_50(self):
        many = [{"name": f"Item {i}"} for i in range(60)]
        p = inv.coerce_inventory_params({"items": many})
        assert len(p["items"]) == inv.MAX_ITEM_ROWS

    def test_notes(self):
        p = inv.coerce_inventory_params({"notes": "  insured list  "})
        assert p["notes"] == "insured list"
        p = inv.coerce_inventory_params({"notes": "   "})
        assert p["notes"] is None
        p = inv.coerce_inventory_params({"notes": "x" * 2000})
        assert len(p["notes"]) == 1000


# ── builder: layout + formulas ────────────────────────────────────────


class TestInventoryBuilder:
    def setup_method(self):
        self.spec = inv.build_inventory_spec(INVENTORY_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        self.sheets = {s["name"]: s for s in self.norm["sheets"]}

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        # the only warning is the Instructions tab's 36 doc lines being
        # capped at 30 by the converter — expected for this pattern
        assert [w for w in warnings if "text_blocks" not in w] == []
        assert any("36 text_blocks" in w for w in warnings)

    def test_sheet_order_and_filename(self):
        assert self.spec["filename"] == "inventory_tracker.xlsx"
        assert [s["name"] for s in self.norm["sheets"]] == [
            "Inventory", "Dashboard", "Categories", "Locations",
            "Instructions", "Calc",
        ]

    # ── Inventory sheet ──────────────────────────────────────────────

    def test_inventory_geometry(self):
        sheet = self.sheets["Inventory"]
        assert sheet["freeze_panes"] == "C5"
        assert sheet["tab_color"] == inv.NAVY
        table = sheet["tables"][0]
        assert table["start_cell"] == "A4"
        assert len(table["headers"]) == 21
        assert table["headers"][0] == "Item ID"
        assert table["headers"][17] == "Last Updated"
        # 50 pre-wired rows regardless of the 5 seeded items
        assert len(table["rows"]) == inv.MAX_ITEM_ROWS

    def test_inventory_row_formulas_pin_to_layout(self):
        rows = self.sheets["Inventory"]["tables"][0]["rows"]
        first, last = rows[0], rows[-1]
        # row 5 — the first data row
        assert first[0] == '=IF($B5="","","INV-"&TEXT(COUNTA($B$5:$B5),"000"))'
        assert first[18] == '=IF(OR($B5="",$E5="",$I5=""),"",$E5*$I5)'  # Total P.
        assert first[19] == '=IF(OR($B5="",$E5="",$J5=""),"",$E5*$J5)'  # Total C.
        assert first[20] == '=IF(OR($B5="",$H5=""),"",TODAY()-$H5)'  # Days
        # row 54 — the last pre-wired row (blank but ready)
        assert last[0] == '=IF($B54="","","INV-"&TEXT(COUNTA($B$5:$B54),"000"))'
        assert last[18] == '=IF(OR($B54="",$E54="",$I54=""),"",$E54*$I54)'
        assert last[20] == '=IF(OR($B54="",$H54=""),"",TODAY()-$H54)'

    def test_seeded_item_values_and_blank_rows(self):
        rows = self.sheets["Inventory"]["tables"][0]["rows"]
        # the normalizer coerces ISO date strings → datetime.date
        seeded = rows[0]
        assert seeded[1] == "MacBook Pro 14"
        assert seeded[2] == "Electronics"
        assert seeded[3] == "Apple M3 Pro"
        assert seeded[4] == 1
        assert seeded[7] == date(2024, 6, 15)
        assert seeded[8] == 2400
        assert seeded[9] == 1800
        assert seeded[10] == "Excellent"
        assert seeded[11] == "C02X1234"
        assert seeded[12] == date(2027, 6, 15)
        assert seeded[13] == "In Stock"
        assert seeded[16] == "main laptop"
        assert seeded[17] == date.today()  # Last Updated = build date
        # seeded row 3 is the lent drill
        drill = rows[2]
        assert drill[13] == "Lent Out"
        assert drill[14] == "Sam"
        assert drill[15] == date(2025, 11, 1)
        # bare-string item lands with quantity 1 and status In Stock
        grinder = rows[4]
        assert grinder[1] == "Coffee grinder"
        assert grinder[4] == 1
        assert grinder[13] == "In Stock"
        # rows past the seeded items are blank (None) but keep formulas
        blank = rows[10]
        assert blank[1] is None
        assert blank[0].startswith("=IF($B15")

    def test_inventory_number_formats_and_alignments(self):
        table = self.sheets["Inventory"]["tables"][0]
        fmts = table["number_formats"]
        assert fmts["E"] == "#,##0"
        assert fmts["H"] == "yyyy-mm-dd"
        assert fmts["I"] == '"$"#,##0.00'
        assert fmts["S"] == '"$"#,##0.00'
        assert fmts["U"] == "0"
        aligns = table["alignments"]
        assert aligns["A"] == "center"
        assert aligns["S"] == "right"

    def test_inventory_dropdown_sources(self):
        dvs = self.sheets["Inventory"]["data_validation"]
        by_range = {dv["range"]: dv for dv in dvs}
        assert len(dvs) == 4
        cat = by_range["C5:C54"]
        assert cat["source_range"] == "Categories!$A$5:$A$24"
        assert cat["error_style"] == "warning"
        assert cat["prompt_title"] == "Category"
        loc = by_range["G5:G54"]
        assert loc["source_range"] == "Locations!$A$5:$A$19"
        cond = by_range["K5:K54"]
        assert cond["values"] == list(inv.CONDITIONS)
        status = by_range["N5:N54"]
        assert status["values"] == list(inv.STATUSES)
        assert status["error_style"] == "stop"

    def test_inventory_conditional_formats(self):
        cfs = {cf["range"]: cf for cf in self.sheets["Inventory"]["conditional_formats"]}
        # low stock — Qty at/below Min Stock with both present
        low = cfs["E5:E54"]["rules"][0]
        assert low["value"] == 'AND($E5<>"",$F5<>"",$E5<=$F5)'
        assert low["fill"] == inv.CF_RED_FILL
        assert low["bold"] is True
        # warranty: expiring inside 60 days (orange), expired (red text)
        wr = cfs["M5:M54"]["rules"]
        assert wr[0]["value"] == 'AND($M5<>"",$M5>=TODAY(),$M5<=TODAY()+60)'
        assert wr[0]["fill"] == inv.CF_ORANGE_FILL
        assert wr[1]["value"] == 'AND($M5<>"",$M5<TODAY())'
        assert wr[1]["font_color"] == inv.CF_RED_TEXT
        # status color coding — 5 cell_is rules
        status_rules = {r["value"]: r for r in cfs["N5:N54"]["rules"]}
        assert set(status_rules) == {"In Stock", "Lent Out", "Sold", "Lost",
                                     "Retired"}
        assert status_rules["In Stock"]["fill"] == inv.CF_GREEN_FILL
        assert status_rules["Lent Out"]["fill"] == inv.CF_GOLD_FILL
        assert status_rules["Retired"]["fill"] == inv.CF_GREY_FILL_2

    def test_inventory_protection(self):
        sheet = self.sheets["Inventory"]
        assert sheet["protect"]["unlocked"] == ["B5:R54"]
        assert sheet["protect"]["password"] is None
        # user notes replace the default hint
        spec = inv.build_inventory_spec(dict(INVENTORY_PARAMS, notes="my note"))
        norm = eg._normalize_spec(spec)
        inv_sheet = {s["name"]: s for s in norm["sheets"]}["Inventory"]
        assert inv_sheet["notes"] == "my note"
        # default hint when no notes given
        assert sheet["notes"].startswith("Type in the white cells")

    def test_plain_money_format_without_currency(self):
        spec = inv.build_inventory_spec({"items": ITEMS})
        norm = eg._normalize_spec(spec)
        fmts = {s["name"]: s for s in norm["sheets"]}["Inventory"]["tables"][0][
            "number_formats"
        ]
        assert fmts["I"] == "#,##0.00"  # MONEY fallback

    # ── Categories / Locations sheets ────────────────────────────────

    def test_categories_sheet(self):
        sheet = self.sheets["Categories"]
        assert sheet["tab_color"] == inv.STEEL
        table = sheet["tables"][0]
        assert table["headers"] == ["Category", "Items"]
        rows = table["rows"]
        assert len(rows) == inv.MAX_CATEGORY_ROWS
        assert [r[0] for r in rows[:5]] == ["Electronics", "Tools", "Kitchen",
                                            "Books", "Other"]
        assert rows[5][0] is None  # unused category rows stay blank
        assert rows[0][1] == (
            '=IF($A5="","",COUNTIF(Inventory!$C$5:$C$54,$A5))'
        )
        assert rows[7][1] == (
            '=IF($A12="","",COUNTIF(Inventory!$C$5:$C$54,$A12))'
        )
        assert sheet["protect"]["unlocked"] == ["A5:A24"]

    def test_locations_sheet(self):
        sheet = self.sheets["Locations"]
        table = sheet["tables"][0]
        assert table["headers"] == ["Location", "Items"]
        rows = table["rows"]
        assert len(rows) == inv.MAX_LOCATION_ROWS
        assert [r[0] for r in rows[:4]] == ["Home Office", "Living Room",
                                           "Garage", "Storage Unit"]
        assert rows[0][1] == (
            '=IF($A5="","",COUNTIF(Inventory!$G$5:$G$54,$A5))'
        )
        assert sheet["protect"]["unlocked"] == ["A5:A19"]

    def test_default_categories_when_none_given(self):
        norm = eg._normalize_spec(inv.build_inventory_spec({}))
        cats = {s["name"]: s for s in norm["sheets"]}["Categories"]
        names = [r[0] for r in cats["tables"][0]["rows"][
            : len(inv.DEFAULT_CATEGORIES)
        ]]
        assert names == list(inv.DEFAULT_CATEGORIES)

    # ── hidden Calc sheet ────────────────────────────────────────────

    def test_calc_sheet_hidden_and_helpers(self):
        sheet = self.sheets["Calc"]
        assert sheet["hidden"] is True
        assert sheet["protect"]["unlocked"] == []  # fully locked helper sheet
        table = sheet["tables"][0]
        assert table["start_cell"] == "C4"
        assert table["headers"] == [
            "Top Value", "Recency", "Lent Rank", "Lent To", "Lent Date",
            "Days Out", "Category", "Status",
        ]
        assert len(table["rows"]) == inv.MAX_ITEM_ROWS
        first = table["rows"][0]
        # C: top-value rank key — value (0 when missing) + ROW()/1e6
        assert first[0] == (
            '=IF(Inventory!$B5="","",IF(Inventory!$T5="",0,Inventory!$T5)'
            "+ROW()/1000000)"
        )
        # D: recency rank — named rows from here to the bottom
        assert first[1] == (
            '=IF(Inventory!$B5="","",COUNTA(Inventory!$B5:$B$54))'
        )
        # E: Lent Out rank; F/G/H: lending display values
        assert first[2] == (
            '=IF(Inventory!$N5="Lent Out",'
            'COUNTIF(Inventory!$N$5:$N5,"Lent Out"),"")'
        )
        assert first[3] == (
            '=IF(Inventory!$N5="Lent Out",'
            'IF(Inventory!$O5="","—",Inventory!$O5),"")'
        )
        assert first[4] == (
            '=IF(Inventory!$N5="Lent Out",'
            'IF(Inventory!$P5="","—",'
            'TEXT(Inventory!$P5,"yyyy-mm-dd")),"")'
        )
        assert first[5] == (
            '=IF(OR(Inventory!$N5<>"Lent Out",'
            'Inventory!$P5=""),"",TODAY()-Inventory!$P5)'
        )
        # I/J: category/status display sanitizers
        assert first[6] == (
            '=IF(Inventory!$B5="","",IF(Inventory!$C5="","—",Inventory!$C5))'
        )
        assert first[7] == (
            '=IF(Inventory!$B5="","",IF(Inventory!$N5="","—",Inventory!$N5))'
        )
        # last row 54 stays 1:1 with the Inventory grid
        assert table["rows"][-1][1] == (
            '=IF(Inventory!$B54="","",COUNTA(Inventory!$B54:$B$54))'
        )

    # ── Dashboard sheet ──────────────────────────────────────────────

    def test_dashboard_kpis(self):
        sheet = self.sheets["Dashboard"]
        assert sheet["no_freeze"] is True
        assert sheet["protect"]["unlocked"] == []  # never type on the dashboard
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["A1"]["text"] == "Home Inventory — Dashboard"
        assert blocks["B3"]["text"] == "=TODAY()"
        assert blocks["B6"]["text"] == "=COUNTA(Inventory!$B$5:$B$54)"
        assert blocks["B7"]["text"] == (
            '=SUMPRODUCT((Inventory!$B$5:$B$54<>"")/'
            'COUNTIF(Inventory!$B$5:$B$54,Inventory!$B$5:$B$54&""))'
        )
        assert blocks["B8"]["text"] == "=SUM(Inventory!$S$5:$S$54)"
        assert blocks["B9"]["text"] == "=SUM(Inventory!$T$5:$T$54)"
        assert blocks["B10"]["text"] == "=$B$9-$B$8"
        assert blocks["B11"]["text"] == '=COUNTIF(Inventory!$N$5:$N$54,"Lent Out")'
        assert blocks["B12"]["text"] == (
            '=SUMPRODUCT((Inventory!$E$5:$E$54<>"")*'
            '(Inventory!$F$5:$F$54<>"")*(Inventory!$E$5:$E$54<=Inventory!$F$5:$F$54))'
        )
        assert blocks["B13"]["text"] == (
            '=SUMPRODUCT((Inventory!$M$5:$M$54<>"")*(Inventory!$M$5:$M$54>=TODAY())*'
            "(Inventory!$M$5:$M$54<=TODAY()+60))"
        )
        assert blocks["B14"]["text"] == (
            '=SUMIF(Inventory!$N$5:$N$54,"In Stock",Inventory!$E$5:$E$54)'
        )
        # number formats: money on the value KPIs, counters elsewhere
        assert blocks["B8"]["number_format"] == '"$"#,##0.00'
        assert blocks["B6"]["number_format"] == "0"

    def test_dashboard_ranking_tables(self):
        sheet = self.sheets["Dashboard"]
        tables = {t["start_cell"]: t for t in sheet["tables"]}
        top = tables["A17"]
        assert top["headers"] == ["Rank", "Item", "Category", "Current Value"]
        assert len(top["rows"]) == inv.TOP_N
        first = top["rows"][0]
        pos = "MATCH(LARGE(Calc!$C$5:$C$54,1),Calc!$C$5:$C$54,0)"
        assert first == [
            1,
            f'=IFERROR(INDEX(Inventory!$B$5:$B$54,{pos}),"—")',
            f'=IFERROR(INDEX(Calc!$I$5:$I$54,{pos}),"—")',
            f'=IFERROR(INDEX(Inventory!$T$5:$T$54,{pos}),"—")',
        ]
        third = top["rows"][2]
        assert third[0] == 3
        assert "LARGE(Calc!$C$5:$C$54,3)" in third[1]

        recent = tables["A25"]
        assert recent["headers"] == ["Item", "Category", "Status"]
        assert len(recent["rows"]) == inv.RECENT_N
        assert recent["rows"][0][0] == (
            '=IFERROR(INDEX(Inventory!$B$5:$B$54,MATCH(1,Calc!$D$5:$D$54,0)),"—")'
        )
        assert recent["rows"][9][0] == (
            '=IFERROR(INDEX(Inventory!$B$5:$B$54,MATCH(10,Calc!$D$5:$D$54,0)),"—")'
        )

        lent = tables["A38"]
        assert lent["headers"] == ["Item", "Lent To", "Date Lent", "Days Out"]
        assert len(lent["rows"]) == inv.LENT_N
        assert lent["rows"][0][1] == (
            '=IFERROR(INDEX(Calc!$F$5:$F$54,MATCH(1,Calc!$E$5:$E$54,0)),"—")'
        )

    def test_dashboard_conditional_formats_and_charts(self):
        sheet = self.sheets["Dashboard"]
        # single-cell KPI ranges normalize to "B10:B10" style
        cfs = {cf["range"]: cf for cf in sheet["conditional_formats"]}
        gain = cfs["B10:B10"]["rules"]
        assert {r["operator"] for r in gain} == {"greater_than", "less_than"}
        assert gain[0]["font_color"] == inv.CF_GREEN_TEXT
        # attention counters light up when non-zero
        assert cfs["B11:B11"]["rules"][0]["operator"] == "greater_than"
        assert cfs["B12:B12"]["rules"][0]["font_color"] == inv.CF_RED_TEXT
        assert cfs["B13:B13"]["rules"][0]["font_color"] == inv.CF_ORANGE_TEXT
        charts = {c["title"]: c for c in sheet["charts"]}
        pie = charts["Items by Category"]
        assert pie["type"] == "pie"
        assert pie["categories_range"] == "Categories!$A$5:$A$24"
        assert pie["series"][0]["values_range"] == "Categories!$B$5:$B$24"
        bar = charts["Items by Location"]
        assert bar["type"] == "bar"
        assert bar["show_values"] is True
        assert bar["value_numfmt"] == "0"
        assert bar["categories_range"] == "Locations!$A$5:$A$19"

    # ── Instructions sheet ───────────────────────────────────────────

    def test_instructions_blocks(self):
        sheet = self.sheets["Instructions"]
        assert sheet["protect"]["unlocked"] == []  # documentation is locked
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["A1"]["text"] == "How to Use Your Inventory Tracker"
        assert blocks["A3"]["text"].startswith("1. ADDING A NEW ITEM")
        # the converter caps text blocks at 30 — the raw spec carries all 36
        raw = [s for s in self.spec["sheets"]
               if s["name"] == "Instructions"][0]["text_blocks"]
        raw_blocks = {b["cell"]: b for b in raw}
        assert raw_blocks["A35"]["text"] == (
            "Top 5 Most Valuable ranks by Total Current Value; the pie breaks "
            "items down by"
        )
        assert len(raw) == 36
        # the rank count is parameterized
        blocks3 = {b["cell"]: b for b in inv._instructions_blocks(3)}
        assert "Top 3" in blocks3["A35"]["text"]

    # ── independent math behind the formulas ─────────────────────────

    def test_python_math_matches_dashboard_formulas(self):
        """Re-compute the KPIs in Python for the seeded fixture."""
        items = inv.coerce_inventory_params(INVENTORY_PARAMS)["items"]
        purchase = [i["quantity"] * (i["purchase_price"] or 0) for i in items]
        current = [i["quantity"] * (i["current_value"] or 0) for i in items]
        assert sum(purchase) == pytest.approx(2400 + 40 + 150 + 300 + 0)
        assert sum(current) == pytest.approx(1800 + 30 + 90 + 450 + 0)
        # the fixture is deliberately mixed for the dashboard's counters
        low_stock = [
            i for i in items
            if i["min_stock"] is not None and i["quantity"] <= i["min_stock"]
        ]
        assert [i["name"] for i in low_stock] == ["Phone chargers"]
        lent = [i for i in items if i["status"] == "Lent Out"]
        assert [i["name"] for i in lent] == ["Cordless drill"]
        in_stock_qty = sum(
            i["quantity"] for i in items if i["status"] == "In Stock"
        )
        assert in_stock_qty == 5  # 1 + 2 + 1 + 1 (the grinder defaults to 1)
        # top-5 by current value: MacBook (1800) leads, camera (450) second
        ranked = sorted(
            zip(current, [i["name"] for i in items]), reverse=True
        )
        assert ranked[0][1] == "MacBook Pro 14"
        assert ranked[1][1] == "Vintage camera"

    def test_no_round_inside_formulas(self):
        for sheet in self.norm["sheets"]:
            for table in sheet.get("tables", []):
                for row in table["rows"]:
                    for cell in row:
                        if isinstance(cell, str):
                            assert "ROUND(" not in cell

    # ── built workbook round-trip ────────────────────────────────────

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "inventory.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == [
            "Inventory", "Dashboard", "Categories", "Locations",
            "Instructions", "Calc",
        ]
        assert wb["Calc"].sheet_state == "hidden"

        ws = wb["Inventory"]
        assert ws["A4"].value == "Item ID"
        assert ws["B5"].value == "MacBook Pro 14"
        assert ws["A5"].value == '=IF($B5="","","INV-"&TEXT(COUNTA($B$5:$B5),"000"))'
        assert ws["S5"].value == '=IF(OR($B5="",$E5="",$I5=""),"",$E5*$I5)'
        assert ws["U5"].value == '=IF(OR($B5="",$H5=""),"",TODAY()-$H5)'
        assert ws["B54"].value is None  # blank row beyond the seeded items
        assert ws["A54"].value.startswith('=IF($B54')
        assert ws["I5"].number_format == '"$"#,##0.00'
        assert ws["H5"].number_format == "yyyy-mm-dd"
        assert ws.freeze_panes == "C5"
        assert ws.auto_filter is not None
        # 4 dropdowns; the category one reads the live Categories range
        dvs = ws.data_validations.dataValidation
        assert len(dvs) == 4
        cat_dv = [d for d in dvs if d.sqref == "C5:C54"][0]
        assert cat_dv.formula1 == "Categories!$A$5:$A$24"
        status_dv = [d for d in dvs if d.sqref == "N5:N54"][0]
        assert status_dv.formula1 == '"In Stock,Lent Out,Sold,Lost,Retired"'
        # conditional formats present on all three ranges
        cf_ranges = {str(cf.sqref) for cf in ws.conditional_formatting}
        assert {"E5:E54", "M5:M54", "N5:N54"} <= cf_ranges
        # the sheet is protected with B5:R54 editable
        assert ws.protection.sheet is True
        assert ws["B5"].protection.locked is False
        assert ws["A5"].protection.locked is True  # ID formula stays locked

        dash = wb["Dashboard"]
        assert dash["B6"].value == "=COUNTA(Inventory!$B$5:$B$54)"
        assert dash["B10"].value == "=$B$9-$B$8"
        assert len(dash._charts) == 2
        assert dash.protection.sheet is True
        cats = wb["Categories"]
        assert cats["A5"].value == "Electronics"
        assert cats["B5"].value == '=IF($A5="","",COUNTIF(Inventory!$C$5:$C$54,$A5))'
        assert cats.protection.sheet is True
        assert cats["A5"].protection.locked is False

    def test_built_workbook_empty_inventory(self, tmp_path):
        """No items given → 50 blank pre-wired rows, workbook still builds."""
        norm = eg._normalize_spec(inv.build_inventory_spec({}))
        out = tmp_path / "blank.xlsx"
        eg._build_xlsx(norm, out)
        wb = load_workbook(out)
        ws = wb["Inventory"]
        assert ws["B5"].value is None
        assert ws["A5"].value.startswith('=IF($B5=')
        assert ws["R5"].value is None  # no Last Updated on blank rows


# ── routing ───────────────────────────────────────────────────────────


class TestInventoryRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg, "_call_llm",
            _classify_response({
                "pattern": "inventory",
                "params": {
                    "inventory_name": "Home Inventory",
                    "items": [
                        {"name": "Drill", "quantity": 1, "location": "Garage"},
                    ],
                },
            }),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "home inventory for my belongings with warranty dates"
        )
        assert result["pattern"] == "inventory"
        assert result["sheet_names"] == [
            "Inventory", "Dashboard", "Categories", "Locations",
            "Instructions", "Calc",
        ]
        assert result["filename"] == "inventory_tracker.xlsx"
        assert result["chart_count"] == 2
        # 50 rows × 4 formula columns + 20 COUNTIF + 15 COUNTIF + calc/…
        assert result["formula_count"] > 200
        assert "inventory template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
