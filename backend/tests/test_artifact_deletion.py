"""Explicit structural deletion, with versioned regeneration and no placeholders."""

import io
import json
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document as WordDocument
from fastapi import HTTPException
from openpyxl import load_workbook
from pptx import Presentation
from sqlalchemy import text

from app.agent.tools.artifacts import UpdateArtifactTool
from app.agent.validation import argument_errors
from app.api.artifacts import SectionEdit
from app.services import artifacts as service
from app.services.artifact_sources import sections_for
from test_artifacts import generated, store as artifact_store  # noqa: F401


@pytest.fixture
def store(request):
    return request.getfixturevalue("artifact_store")


async def conversation_report(store):
    conversation = uuid.uuid4()
    with store.engine.begin() as db:
        db.execute(
            text("INSERT INTO conversations (id) VALUES (:id)"),
            {"id": conversation.hex},
        )
    artifact = await generated(
        store, source=source_for("report"), conversation_id=conversation
    )
    return artifact, str(conversation)


def source_for(kind):
    if kind == "report":
        return (
            "# First\nKeep first.\n\n## Second\nRemove second.\n\n## Third\nKeep third."
        )
    if kind == "presentation":
        return "# First\nKeep first.\n---\n# Second\nRemove second.\n---\n# Third\nKeep third."
    return json.dumps(
        {
            "sheets": [
                {
                    "name": name,
                    "tables": [
                        {"start_cell": "A1", "headers": ["Value"], "rows": [["=1+2"]]}
                    ],
                }
                for name in ["First", "Second", "Third"]
            ]
        }
    )


@pytest.mark.parametrize(
    "kind,prefix",
    [("report", "section"), ("presentation", "slide"), ("excel", "sheet")],
)
def test_delete_and_replace_use_original_indices_in_any_order(kind, prefix):
    source = source_for(kind)
    version = SimpleNamespace(source=source, sections=sections_for(source, kind))
    replacement = (
        json.dumps({"name": "Updated", "tables": []})
        if kind == "excel"
        else "## Updated\nNew body."
    )
    changes = [
        {"section_id": f"{prefix}-1", "operation": "delete"},
        {"section_id": f"{prefix}-2", "operation": "delete"},
        {"section_id": f"{prefix}-3", "content": replacement},
    ]
    for edits in [changes, list(reversed(changes))]:
        result = service.edited_source(version, kind, edits)
        sections = sections_for(result, kind)
        assert len(sections) == 1
        assert sections[0]["title"] == "Updated"
        assert "Remove second." not in result
    assert version.source == source


@pytest.mark.parametrize(
    "kind,prefix",
    [("report", "section"), ("presentation", "slide"), ("excel", "sheet")],
)
def test_cannot_delete_all_sections(kind, prefix):
    source = source_for(kind)
    version = SimpleNamespace(source=source, sections=sections_for(source, kind))
    with pytest.raises(HTTPException, match="retain at least one"):
        service.edited_source(
            version,
            kind,
            [
                {"section_id": f"{prefix}-{i}", "operation": "delete"}
                for i in range(1, 4)
            ],
        )
    assert version.source == source


@pytest.mark.parametrize(
    "change",
    [
        {"section_id": "section-99", "operation": "delete"},
        {"section_id": "section-2", "operation": "delete", "content": "Removal notice"},
        {"section_id": "section-2", "operation": "unknown"},
        {"section_id": "section-2", "content": ""},
    ],
)
def test_invalid_deletion_is_rejected_without_modifying_source(change):
    source = source_for("report")
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    with pytest.raises(HTTPException):
        service.edited_source(version, "report", [change])
    assert version.source == source


def test_duplicate_delete_replace_target_rejected():
    source = source_for("report")
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    with pytest.raises(HTTPException, match="Duplicate"):
        service.edited_source(
            version,
            "report",
            [
                {"section_id": "section-2", "operation": "delete"},
                {"section_id": "section-2", "content": "## Second\nUpdated"},
            ],
        )


def test_delete_schema_and_api_validation_agree():
    tool = UpdateArtifactTool()
    schema = {
        "type": "object",
        "properties": tool.get_parameters(),
        "required": tool.get_required_params(),
    }
    base = {"artifact_id": "exact-id", "expected_version": 1}
    for change in [
        {"section_id": "section-2", "operation": "delete"},
        {"section_id": "section-2", "content": "## Second\nBody"},
        {
            "section_id": "section-2",
            "operation": "replace",
            "content": "## Second\nBody",
        },
    ]:
        assert not argument_errors({**base, "changes": [change]}, schema)
        SectionEdit(**change)
    for change in [
        {"section_id": "section-2"},
        {"section_id": "section-2", "operation": "replace"},
        {"section_id": "section-2", "operation": "delete", "content": "Removed"},
        {"section_id": "section-2", "operation": "unknown"},
    ]:
        assert argument_errors({**base, "changes": [change]}, schema)
        with pytest.raises(ValueError):
            SectionEdit(**change)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "kind,prefix,fmt",
    [
        ("report", "section", "docx"),
        ("presentation", "slide", "pptx"),
        ("excel", "sheet", "xlsx"),
    ],
)
async def test_api_delete_regenerates_file_and_preserves_history(
    store, kind, prefix, fmt
):
    artifact = await generated(store, kind, source_for(kind))
    url = f"/api/artifacts/{artifact['id']}"
    response = store.client.put(
        url,
        json={
            "expected_version": 1,
            "changes": [{"section_id": f"{prefix}-2", "operation": "delete"}],
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["selected_version"] == 2
    current = store.client.get(url + "/read").json()
    assert [part["title"] for part in current["sections"]] == ["First", "Third"]
    previous = store.client.get(url + "/read?version=1").json()
    assert [part["title"] for part in previous["sections"]] == [
        "First",
        "Second",
        "Third",
    ]
    output = store.client.get(url + f"/download/{fmt}?version=2")
    assert output.status_code == 200
    if kind == "report":
        text = "\n".join(
            paragraph.text
            for paragraph in WordDocument(io.BytesIO(output.content)).paragraphs
        )
        assert "Remove second." not in text and "Keep third." in text
    elif kind == "presentation":
        slides = Presentation(io.BytesIO(output.content)).slides
        assert len(slides) == 2
        assert "Third" in " ".join(
            shape.text for shape in slides[1].shapes if shape.has_text_frame
        )
    else:
        workbook = load_workbook(io.BytesIO(output.content))
        assert workbook.sheetnames == ["First", "Third"]
        assert workbook["Third"]["A2"].value == "=1+2"
    # IDs are resolved against the selected version, never a stale index.
    stale = store.client.put(
        url,
        json={
            "expected_version": 1,
            "changes": [{"section_id": f"{prefix}-2", "operation": "delete"}],
        },
    )
    assert stale.status_code == 409
    assert store.client.get(url).json()["selected_version"] == 2
    restore = store.client.post(
        url + "/restore", json={"expected_version": 2, "version": 1}
    )
    assert restore.status_code == 200
    assert len(store.client.get(url + "/read").json()["sections"]) == 3


@pytest.mark.asyncio
async def test_tool_delete_attaches_updated_file(store):
    artifact, conversation = await conversation_report(store)
    result = await UpdateArtifactTool().execute(
        artifact_id=artifact["id"],
        conversation_id=conversation,
        expected_version=1,
        changes=[{"section_id": "section-2", "operation": "delete"}],
    )
    assert result.success, result.output
    attachment = result.tool_call.gen_results[0]
    assert attachment["artifact_id"] == artifact["id"] and attachment["version"] == 2
    assert store.client.get(attachment["download_url"]).status_code == 200


@pytest.mark.asyncio
async def test_failed_deletion_render_keeps_previous_version(store, monkeypatch):
    artifact, conversation = await conversation_report(store)
    monkeypatch.setattr(
        service, "render", AsyncMock(side_effect=RuntimeError("Renderer unavailable"))
    )
    result = await UpdateArtifactTool().execute(
        artifact_id=artifact["id"],
        conversation_id=conversation,
        expected_version=1,
        changes=[{"section_id": "section-2", "operation": "delete"}],
    )
    assert not result.success
    url = f"/api/artifacts/{artifact['id']}"
    assert store.client.get(url).json()["selected_version"] == 1
    assert len(store.client.get(url + "/read").json()["sections"]) == 3
    assert store.client.get(url + "/download/docx?version=1").status_code == 200


def test_uploaded_artifact_deletion_is_read_only():
    version = SimpleNamespace(
        source="Uploaded", sections=[{"id": "part-1", "content": "Uploaded"}]
    )
    with pytest.raises(HTTPException) as error:
        service.edited_source(
            version, "upload", [{"section_id": "part-1", "operation": "delete"}]
        )
    assert error.value.status_code == 403
