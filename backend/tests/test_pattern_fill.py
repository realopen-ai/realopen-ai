"""
Tests for FILL MODE — guidance requests get POPULATED pattern sheets.

The scenario the user asked for, both directions:

  "Create me a meal planner sheet"
      → classifier routes with NO fill → blank dated template
        (covered by the per-pattern routing suites; here we pin that
        a missing/false/"false" fill NEVER triggers the populator).

  "Create me a meal planner sheet — my goal is weight loss, I want
   to eat more protein and vegetables"
      → classifier routes meal_planner with "fill": true → a SECOND
        small LLM call (prompts/pattern_populator.md, ONLY the routed
        pattern's stanza) drafts starter params → the deterministic
        builder emits a POPULATED week.

Safety contract under test:
  - fill is honored ONLY for the six fillable patterns (registry
    PATTERN_FILLABLE); a business-record pattern answering fill=true
    is ignored (never fabricated facts);
  - ANY populator failure (LLM down, unparseable JSON, params the
    builder rejects) falls back to the classifier's own extraction —
    i.e. the blank/partial template fill mode replaced, never worse;
  - the merge keeps extracted params the populator omits ({**params,
    **filled}), so drafting meals can never lose an extracted
    shopping list.
"""

import json
import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402

# ── helpers ──────────────────────────────────────────────────────────


def _llm_sequence(responses):
    """Fake _call_llm serving a fixed sequence of payloads.

    Responses are consumed in call order (classifier first, populator
    second). An Exception instance is raised instead of returned;
    a str is returned verbatim (raw unparseable output). ``.calls``
    collects the messages of every call for shape assertions.
    """

    async def fake_llm(messages, model=None):
        fake_llm.calls.append(messages)
        idx = len(fake_llm.calls) - 1
        item = responses[idx]
        if isinstance(item, Exception):
            raise item
        return item if isinstance(item, str) else json.dumps(item)

    fake_llm.calls = []
    return fake_llm


def _sheet_cells(spec: dict, sheet_name: str):
    """All non-None cell values of one sheet's tables (flattened)."""
    for sheet in spec["sheets"]:
        if sheet["name"] == sheet_name:
            cells = []
            for table in sheet.get("tables", []):
                for row in table.get("rows", []):
                    cells.extend(str(c) for c in row if c is not None)
            return cells
    raise AssertionError(f"sheet {sheet_name!r} not found in spec")


def _all_cells(spec: dict):
    out = []
    for sheet in spec["sheets"]:
        for table in sheet.get("tables", []):
            for row in table.get("rows", []):
                out.extend(str(c) for c in row if c is not None)
    return out


CLASSIFIER_MEAL = {
    "pattern": "meal_planner",
    "params": {"notes": "Goal: weight loss — more protein and vegetables, fewer carbs"},
    "fill": True,
}

POPULATOR_MEAL = {
    "params": {
        "meals": [
            {
                "day": "Monday",
                "meal": "Breakfast",
                "dish": "Greek yogurt with berries and walnuts",
            },
            {
                "day": "Monday",
                "meal": "Dinner",
                "dish": "Grilled chicken with broccoli and quinoa",
            },
            {
                "day": "Tuesday",
                "meal": "Lunch",
                "dish": "Tuna and white-bean salad",
            },
        ],
        "shopping_list": [
            {
                "item": "Chicken breast",
                "quantity": 1.5,
                "unit": "kg",
                "category": "Protein",
            },
            {
                "item": "Greek yogurt",
                "quantity": 1,
                "unit": "kg",
                "category": "Dairy",
            },
        ],
    }
}

MEAL_BRIEF = (
    "Create me a meal planner sheet — my goal is weight loss, I want "
    "to eat more protein and vegetables, less carbs"
)


# ── registry ↔ prompt consistency ────────────────────────────────────


class TestFillRegistry:
    def test_fillable_set_is_exactly_the_six_planners(self):
        assert set(ep.PATTERN_FILLABLE) == {
            "meal_planner",
            "workout_log",
            "habit_tracker",
            "content_calendar",
            "project_plan",
            "shift_schedule",
        }

    def test_no_business_record_pattern_is_fillable(self):
        for name in (
            "invoice",
            "amortization",
            "payroll",
            "receivables",
            "sales_tracker",
            "crm_pipeline",
            "price_list",
            "subscription_tracker",
            "timesheet",
            "expense_report",
        ):
            assert name not in ep.PATTERN_FILLABLE

    def test_populator_stanzas_match_fillable_registry(self):
        _, stanzas, _ = eg._parse_stanza_prompt("pattern_populator")
        assert set(stanzas) == set(ep.PATTERN_FILLABLE), (
            "populator stanzas and PATTERN_FILLABLE drifted apart: "
            f"{sorted(set(stanzas) ^ set(ep.PATTERN_FILLABLE))}"
        )

    def test_populator_prompt_is_single_stanza(self):
        # small-prompt discipline: one routed pattern only
        sp = eg._populator_system_prompt("meal_planner")
        assert "shopping_list" in sp  # meal stanza present
        assert "STARTER CONTENT" in sp  # header
        assert "Rules:" in sp  # footer
        for foreign in ("shift_codes", "training PLAN", "habits the stated"):
            assert foreign not in sp, f"stanza leak: {foreign!r}"

    def test_non_fillable_pattern_has_no_populator_prompt(self):
        assert eg._populator_system_prompt("invoice") == ""
        assert eg._populator_system_prompt("amortization") == ""

    def test_classifier_stanzas_document_fill_only_for_fillable(self):
        _, stanzas, _ = eg._parse_classifier_prompt()
        for name in ep.PATTERN_FILLABLE:
            assert 'Set \\"fill\\": true' in stanzas[name], name
        for name in ("invoice", "amortization", "sales_tracker", "payroll"):
            assert 'Set \\"fill\\"' not in stanzas[name], name

    def test_classifier_header_documents_the_fill_key(self):
        header, _, _ = eg._parse_classifier_prompt()
        assert '"fill"' in header
        assert "NEVER get fill" in header


# ── the user's exact scenario: guidance → populated sheet ────────────


class TestFillRouting:
    @pytest.mark.asyncio
    async def test_guidance_request_populates_meal_planner(self, monkeypatch):
        fake = _llm_sequence([CLASSIFIER_MEAL, POPULATOR_MEAL])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None
        spec, pattern = result
        assert pattern == "meal_planner"

        plan = _sheet_cells(spec, "Plan")
        assert "Greek yogurt with berries and walnuts" in plan
        assert "Grilled chicken with broccoli and quinoa" in plan
        assert "Tuna and white-bean salad" in plan
        shopping = _sheet_cells(spec, "Shopping List")
        assert "Chicken breast" in shopping
        assert "Greek yogurt" in shopping

    @pytest.mark.asyncio
    async def test_populator_gets_request_date_and_extracted_params(self, monkeypatch):
        fake = _llm_sequence([CLASSIFIER_MEAL, POPULATOR_MEAL])
        monkeypatch.setattr(eg, "_call_llm", fake)
        await eg._try_pattern_spec(MEAL_BRIEF, "keep dinners under 600 kcal")

        assert len(fake.calls) == 2
        system, user = fake.calls[1][0]["content"], fake.calls[1][1]["content"]
        # system: populator header + ONLY the meal_planner stanza
        assert "STARTER CONTENT" in system
        assert "shopping_list" in system
        assert "shift_codes" not in system
        # user: today's date + the request + requirements + extraction
        assert "Today's date:" in user
        assert date.today().isoformat() in user
        assert MEAL_BRIEF in user
        assert "keep dinners under 600 kcal" in user
        assert "Extracted parameters" in user
        assert (
            json.dumps(CLASSIFIER_MEAL["params"])
            in user.replace(" ", "").replace("  ", " ")
            or "weight loss" in user
        )

    @pytest.mark.asyncio
    async def test_merge_keeps_params_the_populator_omits(self, monkeypatch):
        classifier = {
            "pattern": "meal_planner",
            "params": {
                "shopping_list": [
                    {"item": "Chicken breast", "quantity": 1.5, "unit": "kg"}
                ]
            },
            "fill": True,
        }
        # populator drafts ONLY meals — the extracted shopping list
        # must survive the merge.
        populator = {
            "params": {
                "meals": [
                    {"day": "Monday", "meal": "Dinner", "dish": "Chicken stir-fry"}
                ]
            }
        }
        monkeypatch.setattr(eg, "_call_llm", _llm_sequence([classifier, populator]))

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None
        spec, _ = result
        assert "Chicken stir-fry" in _sheet_cells(spec, "Plan")
        assert "Chicken breast" in _sheet_cells(spec, "Shopping List")

    @pytest.mark.asyncio
    async def test_stringified_fill_true_still_triggers(self, monkeypatch):
        classifier = dict(CLASSIFIER_MEAL)
        classifier["fill"] = "true"  # small models stringify booleans
        fake = _llm_sequence([classifier, POPULATOR_MEAL])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None
        assert "Greek yogurt with berries and walnuts" in _sheet_cells(
            result[0], "Plan"
        )

    @pytest.mark.asyncio
    async def test_stringified_fill_false_does_not_trigger(self, monkeypatch):
        classifier = dict(CLASSIFIER_MEAL)
        classifier["fill"] = "false"  # must NOT read as truthy
        fake = _llm_sequence([classifier])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None  # blank template, no populate call
        assert len(fake.calls) == 1
        assert not any("Greek yogurt" in c for c in _all_cells(result[0]))

    @pytest.mark.asyncio
    async def test_fill_false_and_missing_do_not_trigger(self, monkeypatch):
        for fill_value in (False, None, "no"):
            classifier = dict(CLASSIFIER_MEAL)
            if fill_value is None:
                classifier.pop("fill")
            else:
                classifier["fill"] = fill_value
            fake = _llm_sequence([classifier])
            monkeypatch.setattr(eg, "_call_llm", fake)

            result = await eg._try_pattern_spec(MEAL_BRIEF, "")
            assert result is not None
            assert len(fake.calls) == 1

    @pytest.mark.asyncio
    async def test_fill_on_non_fillable_pattern_is_ignored(self, monkeypatch):
        classifier = {
            "pattern": "amortization",
            "params": {
                "loan_amount": 25000,
                "annual_rate": 0.065,
                "term_years": 3,
            },
            "fill": True,  # business-record pattern: fill must be ignored
        }
        fake = _llm_sequence([classifier])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(
            "loan schedule for $25,000 at 6.5% over 3 years", ""
        )
        assert result is not None
        assert result[1] == "amortization"
        assert len(fake.calls) == 1  # no populator call, ever


# ── fallbacks: fill must never be worse than the blank template ──────


class TestFillFallbacks:
    @pytest.mark.asyncio
    async def test_populator_llm_failure_builds_blank_template(self, monkeypatch):
        fake = _llm_sequence([CLASSIFIER_MEAL, RuntimeError("llm down")])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None  # pattern path retained!
        spec, pattern = result
        assert pattern == "meal_planner"
        assert "Greek yogurt" not in _all_cells(spec)

    @pytest.mark.asyncio
    async def test_populator_garbage_builds_blank_template(self, monkeypatch):
        fake = _llm_sequence([CLASSIFIER_MEAL, "not json at all"])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None
        assert "Greek yogurt" not in _all_cells(result[0])

    @pytest.mark.asyncio
    async def test_populator_params_without_params_key_falls_back(self, monkeypatch):
        fake = _llm_sequence(
            [CLASSIFIER_MEAL, {"drafted_meals": [{"dish": "Porridge"}]}]
        )
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None
        assert "Porridge" not in _all_cells(result[0])

    @pytest.mark.asyncio
    async def test_builder_rejecting_populated_params_retries_extraction(
        self, monkeypatch
    ):
        # structurally wrong populated params (meals as a string) →
        # builder ValueError → retry with the classifier's extraction.
        fake = _llm_sequence([CLASSIFIER_MEAL, {"params": {"meals": "not-a-list"}}])
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is not None
        spec, pattern = result
        assert pattern == "meal_planner"
        assert "Greek yogurt" not in _all_cells(spec)

    @pytest.mark.asyncio
    async def test_bad_params_without_fill_still_fall_to_ai_path(self, monkeypatch):
        # regression guard: the plain (no-fill) path keeps its old
        # semantics — structurally wrong params → AI path, not a retry.
        fake = _llm_sequence(
            [{"pattern": "meal_planner", "params": {"meals": "garbage"}}]
        )
        monkeypatch.setattr(eg, "_call_llm", fake)

        result = await eg._try_pattern_spec(MEAL_BRIEF, "")
        assert result is None


# ── every fillable pattern drafts ────────────────────────────────────


class TestPerPatternFill:
    @pytest.mark.asyncio
    async def test_workout_log_fill_drafts_sessions(self, monkeypatch):
        today = date.today().isoformat()
        classifier = {
            "pattern": "workout_log",
            "params": {"log_name": "Hypertrophy Block"},
            "fill": True,
        }
        populator = {
            "params": {
                "sessions": [
                    {
                        "date": today,
                        "exercise": "Bench Press",
                        "sets": 4,
                        "reps": 8,
                        "weight": 0,
                    },
                    {
                        "date": today,
                        "exercise": "Lat Pulldown",
                        "sets": 4,
                        "reps": 10,
                        "weight": 0,
                    },
                ]
            }
        }
        monkeypatch.setattr(eg, "_call_llm", _llm_sequence([classifier, populator]))

        result = await eg._try_pattern_spec(
            "make a workout plan for muscle gain, 4 days a week", ""
        )
        assert result is not None
        log = _sheet_cells(result[0], "Log")
        assert "Bench Press" in log
        assert "Lat Pulldown" in log

    @pytest.mark.asyncio
    async def test_habit_tracker_fill_drafts_habits(self, monkeypatch):
        classifier = {"pattern": "habit_tracker", "params": {}, "fill": True}
        populator = {
            "params": {
                "habits": [
                    {"name": "Drink 2 L water", "target": "daily", "active": True},
                    {"name": "Walk 10,000 steps", "target": "daily", "active": True},
                    {"name": "Gym before work", "target": "weekdays", "active": True},
                ]
            }
        }
        monkeypatch.setattr(eg, "_call_llm", _llm_sequence([classifier, populator]))

        result = await eg._try_pattern_spec(
            "set up a habit tracker — I want to drink more water, walk "
            "daily and hit the gym on weekdays",
            "",
        )
        assert result is not None
        habits = _sheet_cells(result[0], "Habits")
        assert "Drink 2 L water" in habits
        assert "Walk 10,000 steps" in habits
        assert "Gym before work" in habits

    @pytest.mark.asyncio
    async def test_content_calendar_fill_drafts_posts(self, monkeypatch):
        classifier = {
            "pattern": "content_calendar",
            "params": {"calendar_name": "Baking IG", "platforms": ["Instagram"]},
            "fill": True,
        }
        populator = {
            "params": {
                "platforms": ["Instagram"],
                "posts": [
                    {
                        "title": "Sourdough starter day 1 — what the bubbles mean",
                        "platform": "Instagram",
                        "date": "2026-03-02",
                        "topic": "How-to",
                        "status": "Idea",
                    },
                    {
                        "title": "5 mistakes that kill macarons",
                        "platform": "Instagram",
                        "date": "2026-03-06",
                        "topic": "Tips",
                        "status": "Idea",
                    },
                ],
            }
        }
        monkeypatch.setattr(eg, "_call_llm", _llm_sequence([classifier, populator]))

        result = await eg._try_pattern_spec(
            "content calendar for my baking Instagram, 3 posts a week", ""
        )
        assert result is not None
        posts = _sheet_cells(result[0], "Posts")
        assert "Sourdough starter day 1 — what the bubbles mean" in posts
        assert "5 mistakes that kill macarons" in posts

    @pytest.mark.asyncio
    async def test_project_plan_fill_drafts_tasks(self, monkeypatch):
        classifier = {
            "pattern": "project_plan",
            "params": {"project_name": "Office move", "deadline": "2026-06-30"},
            "fill": True,
        }
        populator = {
            "params": {
                "tasks": [
                    {
                        "name": "Measure the new office and draft a floor plan",
                        "status": "Not Started",
                        "progress": 0,
                        "milestone": False,
                    },
                    {
                        "name": "Sign the lease",
                        "status": "Not Started",
                        "progress": 0,
                        "milestone": True,
                    },
                    {
                        "name": "Order desks and chairs for 12 people",
                        "status": "Not Started",
                        "progress": 0,
                        "milestone": False,
                    },
                ]
            }
        }
        monkeypatch.setattr(eg, "_call_llm", _llm_sequence([classifier, populator]))

        result = await eg._try_pattern_spec(
            "project plan for our office move — deadline June 30", ""
        )
        assert result is not None
        plan = _sheet_cells(result[0], "Plan")
        assert "Measure the new office and draft a floor plan" in plan
        assert "Sign the lease" in plan
        assert "Order desks and chairs for 12 people" in plan

    @pytest.mark.asyncio
    async def test_shift_schedule_fill_drafts_rota(self, monkeypatch):
        classifier = {
            "pattern": "shift_schedule",
            "params": {
                "team_name": "Cafe Roast",
                "staff": [
                    {"name": "Sarah", "role": "Barista"},
                    {"name": "Tom", "role": "Barista"},
                ],
            },
            "fill": True,
        }
        populator = {
            "params": {
                "staff": [
                    {
                        "name": "Sarah",
                        "role": "Barista",
                        "mon": "M",
                        "tue": "M",
                        "wed": "E",
                        "fri": "E",
                    },
                    {
                        "name": "Tom",
                        "role": "Barista",
                        "mon": "E",
                        "wed": "M",
                        "thu": "E",
                        "sat": "M",
                    },
                ]
            }
        }
        monkeypatch.setattr(eg, "_call_llm", _llm_sequence([classifier, populator]))

        result = await eg._try_pattern_spec(
            "shift schedule for my cafe — Sarah and Tom, we open 9 to 5", ""
        )
        assert result is not None
        rota = _sheet_cells(result[0], "Rota")
        assert "Sarah" in rota
        assert "Tom" in rota
        assert "M" in rota and "E" in rota  # drafted codes landed


# ── end-to-end through the public entry point ────────────────────────


class TestFillEndToEnd:
    @pytest.mark.asyncio
    async def test_guidance_scenario_builds_populated_workbook(
        self, tmp_path, monkeypatch
    ):
        fake = _llm_sequence([CLASSIFIER_MEAL, POPULATOR_MEAL])
        monkeypatch.setattr(eg, "_call_llm", fake)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when fill mode works")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)

        result = await eg.generate_spreadsheet(MEAL_BRIEF)
        assert result["pattern"] == "meal_planner"
        assert "Plan" in result["sheet_names"]
        assert "Shopping List" in result["sheet_names"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

        # the workbook itself carries the drafted dishes
        from openpyxl import load_workbook

        wb = load_workbook(tmp_path / f"{result['report_id']}.xlsx")
        ws = wb["Plan"]
        dishes = {
            ws.cell(row=r, column=c).value for r in range(5, 12) for c in range(3, 7)
        }
        assert "Greek yogurt with berries and walnuts" in dishes
        assert "Grilled chicken with broccoli and quinoa" in dishes
        # live coverage formulas intact next to the drafted content
        assert str(ws.cell(row=5, column=7).value).startswith("=SUMPRODUCT")
