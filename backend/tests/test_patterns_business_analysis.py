"""
Tests for the business-analysis pattern batch (sales_tracker,
crm_pipeline, kpi_report) in app/services/patterns/ and their routing
integration in services/excel_gen.py.

Per pattern:
  • param coercion (quoted/currency/percent strings, alias keys, null
    semantics from the pattern_classifier.md stanzas, ValueErrors)
  • builder layout math — every formula references the row the
    converter will actually render (title rows included)
  • independent Python verification of the math the formulas encode
    (SUMIF totals, weighted values, direction-aware attainment)
  • a _build_xlsx round-trip (openpyxl load_workbook): formulas,
    dropdowns, conditional formats, charts, freeze panes
  • one routing test: classifier → template with the AI path skipped
"""

import re
import sys
from datetime import date, datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _classify_response(payload):
    import json

    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


def _sheet(norm, name):
    return next(s for s in norm["sheets"] if s["name"] == name)


def _as_date(value):
    return value.date() if isinstance(value, datetime) else value


# ═════════════════════════════════════════════════════════════════════
# Registry ↔ stanza contract
# ═════════════════════════════════════════════════════════════════════


class TestPatternRegistry:
    def test_pattern_names_match_classifier_stanzas(self):
        md = (
            Path(__file__).resolve().parent.parent
            / "app"
            / "prompts"
            / "pattern_classifier.md"
        ).read_text()
        stanzas = set(re.findall(r"<!--\s*stanza:\s*([A-Za-z0-9_]+)\s*-->", md))
        for name in ("sales_tracker", "crm_pipeline", "kpi_report"):
            assert name in stanzas  # stanza exists in the .md
            assert name in ep.PATTERN_BUILDERS  # registered under that key
            assert name in ep.PATTERN_KEYWORDS  # routable via keywords


# ═════════════════════════════════════════════════════════════════════
# sales_tracker
# ═════════════════════════════════════════════════════════════════════

SALES_PARAMS = {
    "period_label": "Q1 2026",
    "currency": "USD",
    "sales": [
        {
            "date": "2026-01-05",
            "product": "Widget",
            "channel": "Online",
            "customer": "Acme",
            "quantity": 10,
            "unit_price": 19.99,
            "unit_cost": 11.0,
        },
        {
            "date": "2026-01-12",
            "product": "Gadget",
            "channel": "Retail",
            "customer": "Beta",
            "quantity": 4,
            "unit_price": 49.5,
            "unit_cost": 30.0,
        },
        {
            "date": "2026-02-03",
            "product": "Widget",
            "channel": "Retail",
            "quantity": 6,
            "unit_price": 19.99,
        },
        {
            "date": "2026-02-10",
            "product": "Gizmo",
            "quantity": 2,
            "unit_price": 199.0,
            "unit_cost": 120.0,
        },
    ],
}


class TestSalesTrackerBuilder:
    def setup_method(self):
        self.spec = ep.build_sales_tracker_spec(SALES_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        # layout geometry, computed independently in the test
        self.first, self.last = 4, 7  # log data rows (header row 3)
        self.total = 8
        self.prod_title, self.prod_hdr = 9, 10
        self.prod_first, self.prod_last = 11, 13  # 3 products
        self.prod_total = 14
        self.ch_title, self.ch_hdr = 16, 17
        self.ch_first, self.ch_last = 18, 19  # 2 channels
        self.ch_total = 20
        self.mo_title, self.mo_hdr = 22, 23
        self.mo_first, self.mo_last = 24, 25  # 2 months
        self.mo_total = 26

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_sheet_names_and_tabs(self):
        assert [s["name"] for s in self.norm["sheets"]] == ["Sales", "Analysis"]
        assert self.norm["sheets"][0]["tab_color"] == "16304F"  # navy primary
        assert self.norm["sheets"][1]["tab_color"] == "1B3A5C"  # steel

    def test_log_geometry_and_formula_lattice(self):
        sales = _sheet(self.norm, "Sales")
        assert sales["freeze_panes"] == "A4"
        table = sales["tables"][0]
        assert table["start_cell"] == "A3"  # header row 3, data 4..7
        assert len(table["rows"]) == 4
        # every ref points at the row the converter actually renders
        assert table["rows"][0][7] == "=E4*F4"  # Revenue = qty x price
        assert table["rows"][2][7] == "=E6*F6"
        # Profit/Margin only when a cost is present (live IF guards)
        assert table["rows"][0][8] == '=IF(G4="","",H4-E4*G4)'
        assert table["rows"][2][8] == '=IF(G6="","",H6-E6*G6)'
        assert table["rows"][0][9] == '=IF(OR(G4="",H4=0),"",I4/H4)'
        # no cost on row 3 of the log → cost cell empty
        assert table["rows"][2][6] is None
        # totals SUM over the exact data rows; guarded overall margin
        assert table["total_row"][4] == "=SUM(E4:E7)"
        assert table["total_row"][7] == "=SUM(H4:H7)"
        assert table["total_row"][8] == "=SUM(I4:I7)"
        assert table["total_row"][9] == '=IF(OR(H8=0,I8=0),"",I8/H8)'
        assert table["auto_filter"] is True
        # money/percent/qty formats on every computed column
        assert table["number_formats"]["F"] == '"$"#,##0.00'
        assert table["number_formats"]["H"] == '"$"#,##0.00'
        assert table["number_formats"]["J"] == "0.00%"
        assert table["number_formats"]["E"] == "#,##0.##"
        # negative profit rows highlighted red
        cf = sales["conditional_formats"][0]
        assert cf["range"] == "I4:I7"
        assert cf["rules"][0]["operator"] == "less_than"
        assert cf["rules"][0]["value"] == 0
        assert cf["rules"][0]["fill"] == "FFC7CE"

    def test_analysis_topline_and_sumif_lattice(self):
        ana = _sheet(self.norm, "Analysis")
        blocks = {b["cell"]: b for b in ana["text_blocks"]}
        assert blocks["A1"]["text"] == "Sales Analysis"
        assert _sheet(self.norm, "Sales")["text_blocks"][0]["text"] == (
            "Sales Log — Q1 2026"
        )
        assert blocks["B3"]["text"] == "=COUNTA(Sales!$B$4:$B$7)"
        assert blocks["B4"]["text"] == "=SUM(Sales!$E$4:$E$7)"
        assert blocks["B5"]["text"] == "=SUM(Sales!$H$4:$H$7)"
        assert blocks["B6"]["text"] == "=SUM(Sales!$I$4:$I$7)"
        assert blocks["B7"]["text"] == "=IF($B$3=0,0,$B$5/$B$3)"  # guarded AOV

        prod, chan, month = ana["tables"]
        # product breakdown: title 9, header 10, data 11..13, total 14
        assert prod["start_cell"] == "A9"
        assert prod["title"] == "Revenue by Product"
        assert prod["rows"][0][0] == "Widget"
        assert prod["rows"][0][1] == ("=SUMIF(Sales!$B$4:$B$7,$A11,Sales!$E$4:$E$7)")
        assert prod["rows"][0][2] == ("=SUMIF(Sales!$B$4:$B$7,$A11,Sales!$H$4:$H$7)")
        assert prod["rows"][0][3] == ("=SUMIF(Sales!$B$4:$B$7,$A11,Sales!$I$4:$I$7)")
        assert prod["rows"][0][4] == "=IF($C$14=0,0,C11/$C$14)"
        assert prod["rows"][2][1] == ("=SUMIF(Sales!$B$4:$B$7,$A13,Sales!$E$4:$E$7)")
        assert prod["total_row"] == [
            "Total",
            "=SUM(B11:B13)",
            "=SUM(C11:C13)",
            "=SUM(D11:D13)",
            "=SUM(E11:E13)",
        ]
        # channel breakdown: criteria range pins to log column C
        assert chan["start_cell"] == "A16"
        assert chan["title"] == "Revenue by Channel"
        assert chan["rows"][0][1] == ("=SUMIF(Sales!$C$4:$C$7,$A18,Sales!$E$4:$E$7)")
        assert chan["rows"][1][2] == ("=SUMIF(Sales!$C$4:$C$7,$A19,Sales!$H$4:$H$7)")
        assert chan["total_row"][2] == "=SUM(C18:C19)"
        # month breakdown: SUMIF on date ranges [month start, next start)
        assert month["start_cell"] == "A22"
        assert month["title"] == "Revenue by Month"
        assert month["rows"][0][0] == "Jan 2026"
        assert month["rows"][0][1] == (
            '=SUMIF(Sales!$A$4:$A$7,">="&DATE(2026,1,1),Sales!$H$4:$H$7)'
            '-SUMIF(Sales!$A$4:$A$7,">="&DATE(2026,2,1),Sales!$H$4:$H$7)'
        )
        assert month["rows"][1][1] == (
            '=SUMIF(Sales!$A$4:$A$7,">="&DATE(2026,2,1),Sales!$H$4:$H$7)'
            '-SUMIF(Sales!$A$4:$A$7,">="&DATE(2026,3,1),Sales!$H$4:$H$7)'
        )
        assert month["rows"][1][3] == "=IF($B$26=0,0,B25/$B$26)"
        assert month["total_row"][1] == "=SUM(B24:B25)"

    def test_charts_pin_to_analysis_rows(self):
        ana = _sheet(self.norm, "Analysis")
        product_chart, month_chart = ana["charts"]
        assert product_chart["type"] == "bar"
        assert product_chart["title"] == "Revenue by Product"
        assert product_chart["categories_range"] == "Analysis!$A$11:$A$13"
        assert product_chart["series"][0]["values_range"] == ("Analysis!$C$11:$C$13")
        assert product_chart["value_numfmt"] == '"$"#,##0.00'
        assert month_chart["type"] == "bar"
        assert month_chart["categories_range"] == "Analysis!$A$24:$A$25"
        assert month_chart["series"][0]["values_range"] == "Analysis!$B$24:$B$25"

    def test_independent_math_sumif_simulation(self):
        """Simulate in Python exactly what the emitted SUMIF lattice
        computes, from the normalized spec's own cells."""
        log = _sheet(self.norm, "Sales")["tables"][0]["rows"]
        # the log's Revenue column is =E*F — simulate that product here
        revenue = [row[4] * row[5] for row in log]
        profit = [(row[4] * row[6]) if row[6] is not None else None for row in log]
        products = [row[1] for row in log]
        channels = [row[2] for row in log]
        dates = [row[0] for row in log]

        def sim_sumif(crit_values, label, values):
            return sum(
                v
                for c, v in zip(crit_values, values)
                if v is not None and str(c or "").lower() == str(label).lower()
            )

        ana = _sheet(self.norm, "Analysis")
        prod, chan, month = ana["tables"]

        # product table: labels are the distinct products, SUMIF totals
        # match the independently computed expectations
        expected_products = ["Widget", "Gadget", "Gizmo"]
        assert [r[0] for r in prod["rows"]] == expected_products
        for i, label in enumerate(expected_products):
            exp_rev = sim_sumif(products, label, revenue)
            exp_qty = sim_sumif(products, label, [row[4] for row in log])
            exp_profit = sum(  # noqa
                p
                for c, p in zip(products, profit)
                if p is not None and str(c).lower() == label.lower()
            )
            assert sim_sumif(products, prod["rows"][i][0], revenue) == exp_rev
            assert (
                sim_sumif(products, prod["rows"][i][0], [row[4] for row in log])
                == exp_qty
            )
            # the formulas' criteria ranges cover exactly the log rows
            assert "Sales!$B$4:$B$7" in prod["rows"][i][2]
            assert "Sales!$H$4:$H$7" in prod["rows"][i][2]
        # exact expected numbers (hand-computed from the params)
        assert sim_sumif(products, "Widget", revenue) == pytest.approx(199.9 + 119.94)
        assert sim_sumif(products, "Gadget", revenue) == pytest.approx(198.0)
        assert sim_sumif(products, "Gizmo", revenue) == pytest.approx(398.0)
        assert sum(revenue) == pytest.approx(915.84)

        # channel table (Gizmo has no channel → excluded from it)
        assert [r[0] for r in chan["rows"]] == ["Online", "Retail"]
        assert sim_sumif(channels, "Online", revenue) == pytest.approx(199.9)
        assert sim_sumif(channels, "Retail", revenue) == pytest.approx(198.0 + 119.94)

        # month table: [month start, next month start) buckets — the
        # two-SUMIF date-range formula computes exactly this
        for i, (label, row) in enumerate(zip(["Jan 2026", "Feb 2026"], month["rows"])):
            y, m = (2026, 1) if i == 0 else (2026, 2)
            lo = date(y, m, 1)
            hi = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
            in_bucket = sum(
                rev
                for d, rev in zip(dates, revenue)
                if d is not None and lo <= _as_date(d) < hi
            )
            two_sumif = sum(
                rev
                for d, rev in zip(dates, revenue)
                if d is not None and _as_date(d) >= lo
            ) - sum(
                rev
                for d, rev in zip(dates, revenue)
                if d is not None and _as_date(d) >= hi
            )
            assert two_sumif == pytest.approx(in_bucket)
            assert in_bucket == pytest.approx(397.9 if i == 0 else 517.94)
            # the formula's date criteria pin to the right month bounds
            assert f"DATE({y},{m},1)" in row[1]
            assert (
                f"DATE({2026 if m == 12 else y},{m + 1 if m < 12 else 1},1)" in row[1]
            )

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "sales_tracker.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Sales", "Analysis"]
        sales = wb["Sales"]
        assert sales.freeze_panes == "A4"
        assert sales.auto_filter.ref == "A3:J8"
        # live formulas at the exact rendered cells
        assert sales["H4"].value == "=E4*F4"
        assert sales["I4"].value == '=IF(G4="","",H4-E4*G4)'
        assert sales["J4"].value == '=IF(OR(G4="",H4=0),"",I4/H4)'
        assert sales["H8"].value == "=SUM(H4:H7)"
        assert sales["I8"].value == "=SUM(I4:I7)"
        # currency formats + real dates
        assert sales["F4"].number_format == '"$"#,##0.00'
        assert sales["J4"].number_format == "0.00%"
        assert _as_date(sales["A4"].value) == date(2026, 1, 5)
        assert sales["A4"].number_format == "yyyy-mm-dd"
        # negative-profit conditional format exists
        assert len(list(sales.conditional_formatting)) == 1
        ana = wb["Analysis"]
        assert ana["B5"].value == "=SUM(Sales!$H$4:$H$7)"
        assert ana["B11"].value == "=SUMIF(Sales!$B$4:$B$7,$A11,Sales!$E$4:$E$7)"
        assert ana["C11"].value == "=SUMIF(Sales!$B$4:$B$7,$A11,Sales!$H$4:$H$7)"
        assert ana["B24"].value == (
            '=SUMIF(Sales!$A$4:$A$7,">="&DATE(2026,1,1),Sales!$H$4:$H$7)'
            '-SUMIF(Sales!$A$4:$A$7,">="&DATE(2026,2,1),Sales!$H$4:$H$7)'
        )
        assert len(ana._charts) == 2

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after

    def test_minimal_analysis_without_channels_or_dates(self):
        spec = ep.build_sales_tracker_spec(
            {
                "sales": [
                    {"product": "Widget", "unit_price": 5},
                    {"product": "Widget", "unit_price": 5},
                ]
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        ana = _sheet(norm, "Analysis")
        # only the product table; no channel/month tables, one chart
        assert len(ana["tables"]) == 1
        assert ana["tables"][0]["title"] == "Revenue by Product"
        assert len(ana["charts"]) == 1
        # no costs anywhere → no Total Profit topline, profit SUMIF col gone
        blocks = {b["text"] for b in ana["text_blocks"]}
        assert "Total Profit" not in blocks
        assert "Average Order Value" in blocks
        assert ana["tables"][0]["headers"] == [
            "Product",
            "Units Sold",
            "Revenue",
            "Share of Revenue",
        ]


class TestSalesTrackerCoercion:
    def test_quoted_numbers_and_alias_keys(self):
        p = ep.coerce_sales_tracker_params(
            {
                "period_label": "Holiday Sales",
                "sales": [
                    {
                        "item": "Mug",
                        "qty": "3",
                        "price": "12.50",
                        "cost": "$7.00",
                        "channel": "Online",
                        "when": "2026-12-01",
                    }
                ],
            }
        )
        assert p["period_label"] == "Holiday Sales"
        assert p["sales"][0]["product"] == "Mug"
        assert p["sales"][0]["quantity"] == 3.0
        assert p["sales"][0]["unit_price"] == 12.5
        assert p["sales"][0]["unit_cost"] == 7.0
        assert p["sales"][0]["date"] == "2026-12-01"
        assert p["sales"][0]["channel"] == "Online"

    def test_null_semantics_defaults(self):
        p = ep.coerce_sales_tracker_params(
            {"sales": [{"product": "Mug", "unit_price": 12.5}]}
        )
        s = p["sales"][0]
        assert s["quantity"] == 1.0  # quantity 1 when not stated
        assert s["unit_cost"] is None  # only when given
        assert s["channel"] is None
        assert s["customer"] is None
        assert s["date"] is None
        assert p["period_label"] is None
        assert p["currency"] is None
        assert p["notes"] is None

    def test_sales_list_aliases(self):
        for key in ("sales", "transactions", "orders", "entries"):
            p = ep.coerce_sales_tracker_params(
                {key: [{"product": "Mug", "unit_price": 1}]}
            )
            assert len(p["sales"]) == 1

    def test_unusable_entries_dropped(self):
        p = ep.coerce_sales_tracker_params(
            {
                "sales": [
                    {"product": "Mug", "unit_price": 1},  # good
                    {"product": "NoPrice"},  # no price → dropped
                    {"unit_price": 5},  # no product → dropped
                    "junk",  # not a dict → dropped
                ]
            }
        )
        assert len(p["sales"]) == 1

    def test_missing_required_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_sales_tracker_params({})
        with pytest.raises(ValueError):
            ep.coerce_sales_tracker_params({"sales": []})
        with pytest.raises(ValueError):
            ep.coerce_sales_tracker_params({"sales": [{"product": "Mug"}]})
        with pytest.raises(ValueError):
            ep.coerce_sales_tracker_params("nope")

    def test_currency_lowercase_resolved(self):
        p = ep.coerce_sales_tracker_params(
            {"currency": "eur", "sales": [{"product": "Mug", "unit_price": 1}]}
        )
        assert p["currency"] == '#,##0.00" €"'


class TestSalesTrackerRouting:
    @pytest.mark.asyncio
    async def test_sales_tracker_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "sales_tracker", "params": SALES_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "sales tracker: log these sales and analyze revenue by product"
        )
        assert result["pattern"] == "sales_tracker"
        assert result["sheet_names"] == ["Sales", "Analysis"]
        assert result["table_count"] == 4  # log + product + channel + month
        assert result["chart_count"] == 2
        assert result["formula_count"] > 40
        assert "sales_tracker template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ═════════════════════════════════════════════════════════════════════
# crm_pipeline
# ═════════════════════════════════════════════════════════════════════

CRM_PARAMS = {
    "pipeline_name": "Q1 Pipeline",
    "currency": "EUR",
    "deals": [
        {
            "name": "Acme ERP",
            "company": "Acme",
            "stage": "Proposal",
            "value": 50000,
            "probability": 0.6,
            "owner": "Dana",
            "expected_close": "2026-03-31",
        },
        {
            "name": "Beta CRM",
            "company": "Beta",
            "stage": "Negotiation",
            "value": 20000,
            "probability": "80%",
            "expected_close": "2026-01-15",
        },
        {"name": "Gamma pilot", "value": 8000},  # no stage, no probability
        {"name": "Delta upsell", "stage": "Won", "value": 12000, "probability": 1},
    ],
}


class TestCrmPipelineBuilder:
    def setup_method(self):
        self.spec = ep.build_crm_pipeline_spec(CRM_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        self.first, self.last, self.total = 4, 7, 8  # log rows (header 3)
        self.stage_first, self.stage_last = 9, 15  # 6 defaults + Unassigned
        self.stage_total = 16

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_sheet_names_and_tabs(self):
        assert [s["name"] for s in self.norm["sheets"]] == ["Pipeline", "Stages"]
        assert self.norm["sheets"][0]["tab_color"] == "16304F"
        assert self.norm["sheets"][1]["tab_color"] == "1B3A5C"

    def test_log_geometry_and_formula_lattice(self):
        pipe = _sheet(self.norm, "Pipeline")
        assert pipe["freeze_panes"] == "A4"
        table = pipe["tables"][0]
        assert table["start_cell"] == "A3"
        assert len(table["rows"]) == 4
        # weighted value = value x probability when given, else value
        assert table["rows"][0][5] == '=IF(E4="",D4,D4*E4)'
        assert table["rows"][2][5] == '=IF(E6="",D6,D6*E6)'
        # days to close = expected close - TODAY(), blank without a date
        assert table["rows"][0][8] == '=IF(H4="","",H4-TODAY())'
        assert table["rows"][2][8] == '=IF(H6="","",H6-TODAY())'
        assert table["rows"][2][7] is None  # no expected close on Gamma
        # totals over the exact rows
        assert table["total_row"][3] == "=SUM(D4:D7)"
        assert table["total_row"][5] == "=SUM(F4:F7)"
        assert table["auto_filter"] is True
        assert table["number_formats"]["D"] == '#,##0.00" €"'
        assert table["number_formats"]["E"] == "0.00%"
        assert table["number_formats"]["I"] == "0"  # days as plain numbers

    def test_stage_dropdown_and_conditional_formats(self):
        pipe = _sheet(self.norm, "Pipeline")
        dv = pipe["data_validation"][0]
        assert dv["range"] == "C4:C7"
        assert dv["source_range"] == "Stages!$A$9:$A$15"
        by_range = {c["range"]: c["rules"] for c in pipe["conditional_formats"]}
        # Won green / Lost red on the stage column
        stage_rules = by_range["C4:C7"]
        assert stage_rules[0]["value"] == "Won"
        assert stage_rules[0]["fill"] == "C6EFCE"
        assert stage_rules[1]["value"] == "Lost"
        assert stage_rules[1]["fill"] == "FFC7CE"
        # overdue expected-close dates turn red (formula rule, no '=')
        overdue = by_range["H4:H7"][0]
        assert overdue["type"] == "formula"
        assert overdue["value"] == 'AND($H4<>"",$H4<TODAY())'
        assert overdue["fill"] == "FFC7CE"
        # negative days-to-close counts also flagged
        assert by_range["I4:I7"][0]["operator"] == "less_than"
        assert by_range["I4:I7"][0]["value"] == 0

    def test_stage_summary_lattice(self):
        stages = _sheet(self.norm, "Stages")
        blocks = {b["cell"]: b for b in stages["text_blocks"]}
        assert blocks["A1"]["text"] == "Pipeline by Stage"
        assert blocks["B3"]["text"] == "=SUM(Pipeline!$D$4:$D$7)"
        assert blocks["B4"]["text"] == "=SUM(Pipeline!$F$4:$F$7)"
        assert blocks["B5"]["text"] == "=COUNTA(Pipeline!$A$4:$A$7)"
        assert blocks["B6"]["text"] == '=COUNTIF(Pipeline!$I$4:$I$7,"<0")'

        table = stages["tables"][0]
        assert table["start_cell"] == "A8"  # header row 8, data 9..15
        labels = [r[0] for r in table["rows"]]
        assert labels == [
            "Lead",
            "Qualified",
            "Proposal",
            "Negotiation",
            "Won",
            "Lost",
            "Unassigned",
        ]
        proposal = table["rows"][2]
        assert proposal[1] == "=COUNTIF(Pipeline!$C$4:$C$7,$A11)"
        assert proposal[2] == "=SUMIF(Pipeline!$C$4:$C$7,$A11,Pipeline!$D$4:$D$7)"
        assert proposal[3] == "=SUMIF(Pipeline!$C$4:$C$7,$A11,Pipeline!$F$4:$F$7)"
        assert proposal[4] == "=IF($C$16=0,0,C11/$C$16)"
        # the Unassigned bucket is blank-aware so every deal is counted
        unassigned = table["rows"][6]
        assert (
            unassigned[1]
            == '=COUNTIF(Pipeline!$C$4:$C$7,"")+COUNTIF(Pipeline!$C$4:$C$7,$A15)'
        )
        assert unassigned[2] == (
            '=SUMIF(Pipeline!$C$4:$C$7,"",Pipeline!$D$4:$D$7)'
            "+SUMIF(Pipeline!$C$4:$C$7,$A15,Pipeline!$D$4:$D$7)"
        )
        assert table["total_row"] == [
            "Total",
            "=SUM(B9:B15)",
            "=SUM(C9:C15)",
            "=SUM(D9:D15)",
            "=SUM(E9:E15)",
        ]
        # chart: clustered pipeline vs weighted value by stage
        chart = stages["charts"][0]
        assert chart["type"] == "bar"
        assert chart["categories_range"] == "Stages!$A$9:$A$15"
        assert chart["series"][0]["values_range"] == "Stages!$C$9:$C$15"
        assert chart["series"][1]["values_range"] == "Stages!$D$9:$D$15"
        assert chart["value_numfmt"] == '#,##0.00" €"'

    def test_independent_math_weighted_and_stage_totals(self):
        """Verify in Python the math the emitted formulas encode."""
        deals = ep.coerce_crm_pipeline_params(CRM_PARAMS)["deals"]

        def weighted(d):
            return (
                d["value"] * d["probability"]
                if d["probability"] is not None
                else d["value"]
            )

        assert weighted(deals[0]) == pytest.approx(30000.0)
        assert weighted(deals[1]) == pytest.approx(16000.0)
        assert weighted(deals[2]) == pytest.approx(8000.0)  # no prob → value
        assert weighted(deals[3]) == pytest.approx(12000.0)
        assert sum(d["value"] for d in deals) == pytest.approx(90000.0)
        assert sum(weighted(d) for d in deals) == pytest.approx(66000.0)

        # simulate the stage summary (blank stage → Unassigned bucket)
        buckets = {}
        for d in deals:
            key = d["stage"] or "Unassigned"
            b = buckets.setdefault(key, {"count": 0, "value": 0.0, "weighted": 0.0})
            b["count"] += 1
            b["value"] += d["value"]
            b["weighted"] += weighted(d)
        assert buckets["Proposal"] == {
            "count": 1,
            "value": 50000.0,
            "weighted": 30000.0,
        }
        assert buckets["Negotiation"] == {
            "count": 1,
            "value": 20000.0,
            "weighted": 16000.0,
        }
        assert buckets["Won"] == {"count": 1, "value": 12000.0, "weighted": 12000.0}
        assert buckets["Unassigned"] == {
            "count": 1,
            "value": 8000.0,
            "weighted": 8000.0,
        }
        # stage totals cover EVERY deal (the Unassigned bucket guarantees it)
        assert sum(b["count"] for b in buckets.values()) == len(deals)
        assert sum(b["value"] for b in buckets.values()) == pytest.approx(90000.0)
        assert sum(b["weighted"] for b in buckets.values()) == pytest.approx(66000.0)

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "crm_pipeline.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Pipeline", "Stages"]
        pipe = wb["Pipeline"]
        assert pipe.freeze_panes == "A4"
        assert pipe.auto_filter.ref == "A3:I8"
        assert pipe["F4"].value == '=IF(E4="",D4,D4*E4)'
        assert pipe["F6"].value == '=IF(E6="",D6,D6*E6)'
        assert pipe["I4"].value == '=IF(H4="","",H4-TODAY())'
        assert pipe["D8"].value == "=SUM(D4:D7)"
        assert pipe["F8"].value == "=SUM(F4:F7)"
        assert pipe["E4"].number_format == "0.00%"
        assert pipe["I4"].number_format == "0"
        # stage dropdown fed live from the Stages sheet
        dvs = pipe.data_validations.dataValidation
        assert len(dvs) == 1
        assert dvs[0].formula1 == "Stages!$A$9:$A$15"
        assert dvs[0].sqref.__str__() == "C4:C7"
        # 3 conditional-format ranges: stage colors, overdue date, days<0
        assert len(list(pipe.conditional_formatting)) == 3
        stages = wb["Stages"]
        assert stages["B3"].value == "=SUM(Pipeline!$D$4:$D$7)"
        assert stages["B11"].value == "=COUNTIF(Pipeline!$C$4:$C$7,$A11)"
        assert (
            stages["C11"].value == "=SUMIF(Pipeline!$C$4:$C$7,$A11,Pipeline!$D$4:$D$7)"
        )
        assert stages["C16"].value == "=SUM(C9:C15)"
        assert len(stages._charts) == 1

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after

    def test_custom_stage_list_honored(self):
        spec = ep.build_crm_pipeline_spec(
            {
                "stages": ["Discovery", "Quote", "Commit"],
                "deals": [
                    {"name": "A", "stage": "Quote", "value": 100},
                    {"name": "B", "stage": "Commit", "value": 200},
                ],
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        labels = [r[0] for r in _sheet(norm, "Stages")["tables"][0]["rows"]]
        assert labels == ["Discovery", "Quote", "Commit"]
        # dropdown range shrinks with the shorter stage list
        pipe = _sheet(norm, "Pipeline")
        assert pipe["data_validation"][0]["source_range"] == "Stages!$A$9:$A$11"
        # no Won/Lost in a custom list is fine — the CF rules just idle

    def test_deal_only_stage_appended(self):
        spec = ep.build_crm_pipeline_spec(
            {
                "deals": [
                    {"name": "A", "stage": "Verbal OK", "value": 100},
                    {"name": "B", "stage": "Verbal OK", "value": 200},
                ]
            }
        )
        norm = eg._normalize_spec(spec)
        labels = [r[0] for r in _sheet(norm, "Stages")["tables"][0]["rows"]]
        assert labels == [
            "Lead",
            "Qualified",
            "Proposal",
            "Negotiation",
            "Won",
            "Lost",
            "Verbal OK",
        ]


class TestCrmPipelineCoercion:
    def test_probability_forms(self):
        p = ep.coerce_crm_pipeline_params(
            {
                "deals": [
                    {"name": "A", "value": 100, "probability": "30%"},
                    {"name": "B", "value": 100, "probability": 30},
                    {"name": "C", "value": 100, "probability": 0.3},
                    {"name": "D", "value": 100},
                ]
            }
        )
        probs = [d["probability"] for d in p["deals"]]
        assert probs == [0.3, 0.3, 0.3, None]  # null when not stated

    def test_quoted_value_and_aliases(self):
        p = ep.coerce_crm_pipeline_params(
            {
                "pipeline_name": "My Deals",
                "opportunities": [
                    {
                        "deal_name": "Acme",
                        "account": "Acme Corp",
                        "phase": "Proposal",
                        "amount": "$12,000",
                        "likelihood": "65%",
                        "sales_rep": "Dana",
                        "close_date": "2026-06-30",
                    }
                ],
            }
        )
        d = p["deals"][0]
        assert d["name"] == "Acme"
        assert d["company"] == "Acme Corp"
        assert d["stage"] == "Proposal"
        assert d["value"] == 12000.0
        assert d["probability"] == 0.65
        assert d["owner"] == "Dana"
        assert d["expected_close"] == "2026-06-30"

    def test_defaults(self):
        p = ep.coerce_crm_pipeline_params({"deals": [{"name": "A", "value": 1}]})
        assert p["pipeline_name"] == "Sales Pipeline"
        assert p["stages"] is None  # → standard six at build time
        assert p["currency"] is None
        d = p["deals"][0]
        assert d["stage"] is None and d["probability"] is None
        assert d["company"] is None and d["owner"] is None
        assert d["expected_close"] is None

    def test_stages_list_normalized(self):
        p = ep.coerce_crm_pipeline_params(
            {
                "stages": [" Lead ", "Qualified", "lead", "Proposal"],
                "deals": [{"name": "A", "value": 1}],
            }
        )
        assert p["stages"] == ["Lead", "Qualified", "Proposal"]  # deduped
        with pytest.raises(ValueError):
            ep.coerce_crm_pipeline_params({"stages": "not-a-list", "deals": []})

    def test_unusable_deals_dropped_and_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_crm_pipeline_params({})
        with pytest.raises(ValueError):
            ep.coerce_crm_pipeline_params({"deals": []})
        with pytest.raises(ValueError):
            ep.coerce_crm_pipeline_params({"deals": [{"name": "No Value"}]})
        with pytest.raises(ValueError):
            ep.coerce_crm_pipeline_params({"deals": [{"value": 100}]})
        with pytest.raises(ValueError):
            ep.coerce_crm_pipeline_params("nope")
        # one good deal among junk survives
        p = ep.coerce_crm_pipeline_params(
            {"deals": [{"name": "A", "value": 5}, {"name": "B"}, "junk"]}
        )
        assert len(p["deals"]) == 1


class TestCrmPipelineRouting:
    @pytest.mark.asyncio
    async def test_crm_pipeline_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "crm_pipeline", "params": CRM_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "track my sales pipeline: deals in negotiation with values and "
            "expected close dates"
        )
        assert result["pattern"] == "crm_pipeline"
        assert result["sheet_names"] == ["Pipeline", "Stages"]
        assert result["table_count"] == 2
        assert result["chart_count"] == 1
        assert result["formula_count"] > 40
        assert "crm_pipeline template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ═════════════════════════════════════════════════════════════════════
# kpi_report
# ═════════════════════════════════════════════════════════════════════

KPI_PARAMS = {
    "report_name": "Q3 KPIs",
    "period": "Q3 2025",
    "currency": "USD",
    "kpis": [
        {
            "name": "Monthly Revenue",
            "category": "Sales",
            "target": 500000,
            "actual": 480000,
            "unit": "currency",
        },
        {
            "name": "Customer Retention",
            "category": "Success",
            "target": 0.9,
            "actual": 0.94,
            "unit": "percent",
        },
        {
            "name": "Churn Rate",
            "category": "Success",
            "target": 0.05,
            "actual": 0.065,
            "unit": "percent",
            "higher_is_better": False,
        },
        {
            "name": "Avg Support Time",
            "category": "Ops",
            "target": 3,
            "actual": 2,
            "unit": "days",
            "higher_is_better": False,
        },
        {"name": "NPS", "target": 40, "actual": 44, "unit": "number"},
    ],
}


class TestKpiReportBuilder:
    def setup_method(self):
        self.spec = ep.build_kpi_report_spec(KPI_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        # unit-group order = first appearance: currency, percent, days, number
        self.cur_title, self.cur_hdr, self.cur_first, self.cur_last = 9, 10, 11, 11
        self.pct_title, self.pct_hdr, self.pct_first, self.pct_last = 14, 15, 16, 17
        self.day_title, self.day_hdr, self.day_first, self.day_last = 20, 21, 22, 22
        self.num_title, self.num_hdr, self.num_first, self.num_last = 25, 26, 27, 27

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry_and_formula_lattice(self):
        sc = _sheet(self.norm, "Scorecard")
        assert sc["freeze_panes"] == "A2"
        blocks = {b["cell"]: b for b in sc["text_blocks"]}
        assert blocks["A1"]["text"] == "Q3 KPIs"
        assert blocks["B3"]["text"] == "Q3 2025"
        # summary COUNTIFs span the REAL status-column ranges of all
        # four unit tables
        assert blocks["B4"]["text"] == (
            '=COUNTIF(H11:H11,"On Track")+COUNTIF(H16:H17,"On Track")'
            '+COUNTIF(H22:H22,"On Track")+COUNTIF(H27:H27,"On Track")'
        )
        assert blocks["B6"]["text"] == (
            '=COUNTIF(H11:H11,"Miss")+COUNTIF(H16:H17,"Miss")'
            '+COUNTIF(H22:H22,"Miss")+COUNTIF(H27:H27,"Miss")'
        )
        assert blocks["B7"]["text"] == (
            '=IFERROR(AVERAGE(G11:G11,G16:G17,G22:G22,G27:G27),"n/a")'
        )

        cur, pct, day, num = sc["tables"]
        assert cur["start_cell"] == "A9"
        assert cur["title"] == "Currency KPIs"
        assert cur["headers"] == [
            "KPI",
            "Category",
            "Target",
            "Actual",
            "Variance",
            "Variance %",
            "Attainment",
            "Status",
            "Better When",
        ]
        r = self.cur_first  # noqa
        row = cur["rows"][0]
        assert row[2] == 500000.0
        assert row[3] == 480000.0
        assert row[4] == "=D11-C11"  # Variance = Actual - Target
        assert row[5] == '=IF(C11=0,"n/a",E11/C11)'  # guarded Variance %
        assert row[6] == '=IF(C11=0,"n/a",D11/C11)'  # attainment higher-is-better
        assert row[7] == (
            '=IF(G11="n/a","n/a",IF(G11>=1,"On Track",' 'IF(G11>=0.9,"Watch","Miss")))'
        )
        assert row[8] == "Higher"

        # direction-aware attainment on the lower-is-better rows
        churn = pct["rows"][1]  # row 17: target 0.05, actual 0.065
        assert churn[6] == '=IF(D17=0,"n/a",C17/D17)'  # target / actual
        assert churn[8] == "Lower"
        support = day["rows"][0]  # row 22: target 3, actual 2
        assert support[6] == '=IF(D22=0,"n/a",C22/D22)'
        assert support[8] == "Lower"
        # higher-is-better NPS keeps actual / target
        nps = num["rows"][0]
        assert nps[6] == '=IF(C27=0,"n/a",D27/C27)'
        assert nps[8] == "Higher"

    def test_unit_aware_number_formats(self):
        sc = _sheet(self.norm, "Scorecard")
        cur, pct, day, num = sc["tables"]
        assert cur["number_formats"]["C"] == '"$"#,##0.00'
        assert cur["number_formats"]["D"] == '"$"#,##0.00'
        assert cur["number_formats"]["E"] == '"$"#,##0.00'
        assert cur["number_formats"]["F"] == "0.00%"
        assert pct["number_formats"]["C"] == "0.00%"
        assert day["number_formats"]["C"] == "#,##0"
        assert num["number_formats"]["C"] == "#,##0.##"

    def test_status_conditional_formats(self):
        sc = _sheet(self.norm, "Scorecard")
        by_range = {c["range"]: c["rules"] for c in sc["conditional_formats"]}
        assert set(by_range) == {"H11:H11", "H16:H17", "H22:H22", "H27:H27"}
        rules = by_range["H16:H17"]
        assert [r["value"] for r in rules] == ["On Track", "Watch", "Miss"]
        assert rules[0]["fill"] == "C6EFCE"  # green
        assert rules[0]["font_color"] == "1E4620"
        assert rules[1]["fill"] == "FFF2CC"  # amber
        assert rules[1]["font_color"] == "7F6000"
        assert rules[2]["fill"] == "FFC7CE"  # red
        assert rules[2]["font_color"] == "9C0006"

    def test_chart_pins_to_largest_group(self):
        sc = _sheet(self.norm, "Scorecard")
        chart = sc["charts"][0]
        assert chart["type"] == "bar"
        assert chart["anchor"] == "K3"  # clear of the 9-column tables
        # percent group has 2 KPIs — the largest group
        assert chart["categories_range"] == "Scorecard!$A$16:$A$17"
        assert chart["series"][0]["name"] == "Target"
        assert chart["series"][0]["values_range"] == "Scorecard!$C$16:$C$17"
        assert chart["series"][1]["name"] == "Actual"
        assert chart["series"][1]["values_range"] == "Scorecard!$D$16:$D$17"
        assert chart["value_numfmt"] == "0.00%"

    def test_independent_math_attainment_direction_aware(self):
        """Compute attainment + status in Python (direction-aware) and
        check the formulas the builder emitted encode exactly that."""
        kpis = ep.coerce_kpi_report_params(KPI_PARAMS)["kpis"]

        def attainment(k):
            if k["higher_is_better"]:
                return k["actual"] / k["target"] if k["target"] else None
            return k["target"] / k["actual"] if k["actual"] else None

        def status(a):
            if a is None:
                return "n/a"
            if a >= 1:
                return "On Track"
            if a >= 0.9:
                return "Watch"
            return "Miss"

        expected = [(k["name"], attainment(k), status(attainment(k))) for k in kpis]
        assert expected == [
            ("Monthly Revenue", pytest.approx(0.96), "Watch"),
            ("Customer Retention", pytest.approx(0.94 / 0.9), "On Track"),
            ("Churn Rate", pytest.approx(0.05 / 0.065), "Miss"),
            ("Avg Support Time", pytest.approx(3 / 2), "On Track"),
            ("NPS", pytest.approx(44 / 40), "On Track"),
        ]
        # the summary counts the same statuses the IF chains produce
        counts = {"On Track": 0, "Watch": 0, "Miss": 0}
        for _n, _a, st in expected:
            if st in counts:
                counts[st] += 1
        assert counts == {"On Track": 3, "Watch": 1, "Miss": 1}
        # variance: actual - target, verified per KPI
        for k in kpis:
            assert (k["actual"] - k["target"]) == pytest.approx(
                {
                    "Monthly Revenue": -20000.0,
                    "Customer Retention": 0.04,
                    "Churn Rate": 0.015,
                    "Avg Support Time": -1.0,
                    "NPS": 4.0,
                }[k["name"]]
            )

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "kpi_report.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Scorecard"]
        sc = wb["Scorecard"]
        assert sc.freeze_panes == "A2"
        assert sc["E11"].value == "=D11-C11"
        assert sc["F11"].value == '=IF(C11=0,"n/a",E11/C11)'
        assert sc["G11"].value == '=IF(C11=0,"n/a",D11/C11)'
        assert sc["H11"].value == (
            '=IF(G11="n/a","n/a",IF(G11>=1,"On Track",IF(G11>=0.9,"Watch","Miss")))'
        )
        assert sc["G17"].value == '=IF(D17=0,"n/a",C17/D17)'  # churn, lower better
        assert sc["B4"].value.startswith("=COUNTIF(H11:H11")
        # unit-aware number formats on the rendered cells
        assert sc["C11"].number_format == '"$"#,##0.00'
        assert sc["C16"].number_format == "0.00%"
        assert sc["C22"].number_format == "#,##0"
        assert sc["C27"].number_format == "#,##0.##"
        # one CF entry per unit table + the clustered target/actual chart
        assert len(list(sc.conditional_formatting)) == 4
        assert len(sc._charts) == 1

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after

    def test_zero_target_guards(self):
        spec = ep.build_kpi_report_spec(
            {
                "kpis": [
                    {"name": "Signups", "target": 0, "actual": 50, "unit": "number"},
                    {
                        "name": "Cost per Lead",
                        "target": 10,
                        "actual": 0,
                        "unit": "currency",
                        "higher_is_better": False,
                    },
                ]
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        sc = _sheet(norm, "Scorecard")
        # number group first (Signups), currency group second — no
        # period and no categories → tables at data rows 10 and 15
        tables = sc["tables"]
        assert len(tables) == 2
        assert [r[0] for r in tables[0]["rows"]] == ["Signups"]
        assert [r[0] for r in tables[1]["rows"]] == ["Cost per Lead"]
        zero_target_row = tables[0]["rows"][0]
        assert (
            zero_target_row[4] == '=IF(B10=0,"n/a",D10/B10)'
        )  # Variance % guarded on target
        assert (
            zero_target_row[5] == '=IF(B10=0,"n/a",C10/B10)'
        )  # attainment guarded on target
        zero_actual_row = tables[1]["rows"][0]
        # lower-is-better + zero ACTUAL → guard flips to the actual side
        assert zero_actual_row[5] == '=IF(C15=0,"n/a",B15/C15)'

    def test_no_category_column_when_absent(self):
        spec = ep.build_kpi_report_spec(
            {"kpis": [{"name": "NPS", "target": 40, "actual": 44}]}
        )
        norm = eg._normalize_spec(spec)
        sc = _sheet(norm, "Scorecard")
        table = sc["tables"][0]
        assert table["headers"] == [
            "KPI",
            "Target",
            "Actual",
            "Variance",
            "Variance %",
            "Attainment",
            "Status",
            "Better When",
        ]
        # no period → summary ends row 6 → data row 10; columns shift
        # left: Target B, Actual C, Variance D, attainment F
        assert table["rows"][0][3] == "=C10-B10"
        assert table["rows"][0][4] == '=IF(B10=0,"n/a",D10/B10)'
        assert table["rows"][0][5] == '=IF(B10=0,"n/a",C10/B10)'
        # chart anchor moves left with the narrower table
        assert sc["charts"][0]["anchor"] == "J3"
        assert sc["charts"][0]["series"][0]["values_range"] == "Scorecard!$B$10:$B$10"


class TestKpiReportCoercion:
    def test_quoted_numbers_and_percent_strings(self):
        p = ep.coerce_kpi_report_params(
            {
                "kpis": [
                    {
                        "name": "Revenue",
                        "target": "500,000",
                        "actual": "480000",
                        "unit": "currency",
                    },
                    {
                        "name": "Retention",
                        "target": "90%",
                        "actual": "94%",
                        "unit": "percent",
                    },
                    {
                        "name": "Retention2",
                        "target": 90,
                        "actual": 94,
                        "unit": "percent",
                    },
                ]
            }
        )
        assert p["kpis"][0]["target"] == 500000.0
        assert p["kpis"][0]["actual"] == 480000.0
        # percent values normalized to decimals whichever way given
        assert (p["kpis"][1]["target"], p["kpis"][1]["actual"]) == (0.9, 0.94)
        assert (p["kpis"][2]["target"], p["kpis"][2]["actual"]) == (0.9, 0.94)

    def test_kpi_missing_target_or_actual_omitted(self):
        p = ep.coerce_kpi_report_params(
            {
                "kpis": [
                    {"name": "A", "target": 1, "actual": 2},
                    {"name": "B", "target": 1},  # no actual → omitted
                    {"name": "C", "actual": 2},  # no target → omitted
                    {"target": 1, "actual": 2},  # no name → omitted
                ]
            }
        )
        assert [k["name"] for k in p["kpis"]] == ["A"]

    def test_all_omitted_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_kpi_report_params({"kpis": [{"name": "B", "target": 1}]})
        with pytest.raises(ValueError):
            ep.coerce_kpi_report_params({})
        with pytest.raises(ValueError):
            ep.coerce_kpi_report_params({"kpis": "not-a-list"})
        with pytest.raises(ValueError):
            ep.coerce_kpi_report_params("nope")

    def test_higher_is_better_semantics(self):
        p = ep.coerce_kpi_report_params(
            {
                "kpis": [
                    {"name": "A", "target": 1, "actual": 2},  # default True
                    {"name": "B", "target": 1, "actual": 2, "higher_is_better": False},
                    {
                        "name": "C",
                        "target": 1,
                        "actual": 2,
                        "higher_is_better": "false",
                    },
                    {"name": "D", "target": 1, "actual": 2, "direction": "lower"},
                    {"name": "E", "target": 1, "actual": 2, "higher_is_better": "true"},
                ]
            }
        )
        flags = [k["higher_is_better"] for k in p["kpis"]]
        assert flags == [True, False, False, False, True]

    def test_unit_aliases_and_defaults(self):
        p = ep.coerce_kpi_report_params(
            {
                "kpis": [
                    {"name": "A", "target": 1, "actual": 2, "unit": "money"},
                    {"name": "B", "target": 1, "actual": 2, "unit": "%"},
                    {"name": "C", "target": 1, "actual": 2, "unit": "hours"},
                    {"name": "D", "target": 1, "actual": 2},  # default number
                ]
            }
        )
        units = [k["unit"] for k in p["kpis"]]
        assert units == ["currency", "percent", "days", "number"]

    def test_aliases_and_defaults(self):
        p = ep.coerce_kpi_report_params(
            {
                "metrics": [
                    {"metric_name": "NPS", "goal": 40, "current": 44, "group": "CS"}
                ]
            }
        )
        assert p["report_name"] == "KPI Scorecard"  # default
        assert p["period"] is None
        assert p["kpis"][0]["name"] == "NPS"
        assert p["kpis"][0]["category"] == "CS"
        assert p["kpis"][0]["target"] == 40.0
        assert p["kpis"][0]["actual"] == 44.0


class TestKpiReportRouting:
    @pytest.mark.asyncio
    async def test_kpi_report_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "kpi_report", "params": KPI_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "Q3 KPI scorecard: target vs actual for our key metrics with "
            "attainment status"
        )
        assert result["pattern"] == "kpi_report"
        assert result["sheet_names"] == ["Scorecard"]
        assert result["table_count"] == 4  # one per unit group
        assert result["chart_count"] == 1
        assert result["formula_count"] > 20
        assert "kpi_report template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
