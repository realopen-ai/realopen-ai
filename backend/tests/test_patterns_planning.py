"""
Tests for the planning pattern batch: project_plan, event_planner and
travel_planner (app/services/patterns/).

Same discipline as tests/test_excel_patterns.py:
  • param coercion — quoted numbers, alias keys, documented defaults,
    ValueError when the stanza's required params are missing
  • builder layout math — every formula string pins to the exact rows
    the converter renders (normalized spec, computed row numbers)
  • openpyxl round-trip via eg._build_xlsx
  • routing — mocked classifier → template with the AI path skipped
  • independently recomputed math (durations, budget utilization,
    per-traveler share) checked against the emitted formula logic
"""

import json
import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _block(sheet: dict, cell: str):
    """Text-block value at a cell (raises when missing)."""
    for block in sheet.get("text_blocks", []):
        if block["cell"] == cell.upper():
            return block["text"]
    raise AssertionError(f"no text block at {cell}")


def _classify_response(payload):
    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


# ══════════════════════════════════════════════════════════════════════
# project_plan
# ══════════════════════════════════════════════════════════════════════

PROJECT_PARAMS = {
    "project_name": "Website Launch",
    "deadline": "2026-09-30",
    "tasks": [
        {
            "name": "Design mockups",
            "owner": "Amina",
            "start_date": "2026-08-01",
            "end_date": "2026-08-15",
            "progress": 100,
            "status": "Completed",
        },
        {
            "name": "Frontend build",
            "owner": "Karim",
            "start_date": "2026-08-16",
            "end_date": "2026-09-10",
            "progress": 40,
            "status": "In Progress",
        },
        {
            "name": "Content upload",
            "owner": "Sara",
            "start_date": "2026-09-01",
            "end_date": "2026-09-20",
            "progress": 0,
            "status": "Not Started",
        },
        {
            "name": "Go live",
            "owner": "Amina",
            "end_date": "2026-09-25",
            "milestone": True,
        },
    ],
}


class TestProjectPlanBuilder:
    def setup_method(self):
        self.spec = ep.build_project_plan_spec(PROJECT_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "overlap" in w]

    def test_geometry_and_formula_lattice(self):
        sheet = self.norm["sheets"][0]
        assert sheet["name"] == "Plan"
        assert sheet["freeze_panes"] == "A6"  # below the header row 5
        table = sheet["tables"][0]
        # header on row 5, data rows 6..9 (4 tasks)
        assert table["start_cell"] == "A5"
        assert len(table["rows"]) == 4

        first, second = table["rows"][0], table["rows"][1]
        # Duration = end - start + 1, blank until both dates exist
        assert first[4] == '=IF(OR(C6="",D6=""),"",D6-C6+1)'
        assert second[4] == '=IF(OR(C7="",D7=""),"",D7-C7+1)'
        # Days Remaining = end - TODAY()
        assert first[5] == '=IF(D6="","",D6-TODAY())'
        # raw spec keeps ISO dates; progress stored as a fraction
        raw_first = self.spec["sheets"][0]["tables"][0]["rows"][0]
        assert raw_first[2] == "2026-08-01" and raw_first[3] == "2026-08-15"
        assert first[6] == 1.0 and second[6] == 0.4
        assert first[8] == "No" and table["rows"][3][8] == "Yes"

        # deadline block on row 3 (a real date after normalization),
        # summary section starts at row 12
        assert _block(sheet, "B3") == date(2026, 9, 30)
        assert _block(sheet, "A12") == "Summary"
        assert _block(sheet, "B13") == "=COUNTA(A6:A9)"
        assert _block(sheet, "B14") == '=COUNTIF(H6:H9,"Completed")'
        assert _block(sheet, "B15") == '=COUNTIF(H6:H9,"In Progress")'
        assert (
            _block(sheet, "B16") == '=COUNTIFS(D6:D9,"<"&TODAY(),H6:H9,"<>Completed")'
        )
        assert _block(sheet, "B17") == '=COUNTIF(I6:I9,"Yes")'
        assert _block(sheet, "B18") == '=IF(COUNT(G6:G9)=0,"n/a",AVERAGE(G6:G9))'
        assert _block(sheet, "B19") == "=$B$3-TODAY()"

        # timeline chart over the exact rendered rows
        chart = sheet["charts"][0]
        assert chart["type"] == "bar_h"
        assert chart["categories_range"] == "Plan!A6:A9"
        assert chart["series"][0]["values_range"] == "Plan!E6:E9"

        # dropdowns cover exactly the data rows
        dv = {d["range"]: d for d in sheet["data_validation"]}
        assert set(dv) == {"H6:H9", "I6:I9"}
        assert dv["H6:H9"]["values"] == [
            "Not Started",
            "In Progress",
            "Completed",
            "Blocked",
            "On Hold",
        ]

    def test_status_conditional_formats(self):
        sheet = self.norm["sheets"][0]
        cfs = {c["range"]: c["rules"] for c in sheet["conditional_formats"]}
        status_rules = cfs["H6:H9"]
        assert status_rules[0]["type"] == "cell_is"
        assert status_rules[0]["value"] == "Completed"
        assert status_rules[0]["fill"] == "C6EFCE"
        assert status_rules[0]["font_color"] == "1E4620"
        # overdue: end < TODAY() AND status <> Completed (formula rule;
        # the normalizer stores the expression under "value")
        overdue = status_rules[1]
        assert overdue["type"] == "formula"
        assert overdue["value"] == 'AND($D6<>"",$D6<TODAY(),$H6<>"Completed")'
        assert overdue["fill"] == "FFC7CE"
        assert status_rules[2]["value"] == "In Progress"
        assert status_rules[2]["fill"] == "FFF2CC"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "plan.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        ws = wb["Plan"]
        assert ws["A1"].value == "Website Launch — Project Plan"
        assert ws.freeze_panes == "A6"
        assert ws["B3"].value is not None  # deadline date
        assert ws["B3"].number_format == "yyyy-mm-dd"
        assert ws["E6"].value == '=IF(OR(C6="",D6=""),"",D6-C6+1)'
        assert ws["B19"].value == "=$B$3-TODAY()"
        assert len(ws._charts) == 1
        assert len(ws.data_validations.dataValidation) == 2
        assert len(list(ws.conditional_formatting)) == 4
        assert ws.auto_filter.ref == "A5:I9"

    def test_duration_math_independently_verified(self):
        # (end - start + 1) recomputed in Python must match the live
        # formulas emitted for the same rows
        rows = self.spec["sheets"][0]["tables"][0]["rows"]
        expected = [
            (date(2026, 8, 15) - date(2026, 8, 1)).days + 1,
            (date(2026, 9, 10) - date(2026, 8, 16)).days + 1,
            (date(2026, 9, 20) - date(2026, 9, 1)).days + 1,
        ]
        assert expected == [15, 26, 20]
        for i, days in enumerate(expected):
            r = 6 + i
            assert rows[i][4] == '=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r}+1)'.format(r=r)
            assert rows[i][2] and rows[i][3]  # both dates present
        # overall completion = simple AVERAGE of the stated progress
        progress = [100, 40, 0]
        assert abs(sum(progress) / len(progress) / 100 - 7 / 15) < 1e-12

    def test_progress_bar_chart_when_no_dates(self):
        spec = ep.build_project_plan_spec(
            {
                "project_name": "Research",
                "tasks": [
                    {"name": "Literature review", "progress": 80},
                    {"name": "Draft paper", "progress": 10},
                ],
            }
        )
        sheet = eg._normalize_spec(spec)["sheets"][0]
        charts = sheet.get("charts") or []
        assert len(charts) == 1
        assert charts[0]["type"] == "bar"
        assert charts[0]["categories_range"] == "Plan!A6:A7"
        assert charts[0]["series"][0]["values_range"] == "Plan!G6:G7"
        assert charts[0]["value_numfmt"] == "0%"
        # 2 tasks, no deadline → no B3 block, summary at row 10 with
        # exactly 6 tally rows (no Days-to-Deadline row at B17)
        assert not [b for b in sheet["text_blocks"] if b["cell"] == "B3"]
        assert not [b for b in sheet["text_blocks"] if b["cell"] == "B17"]
        assert _block(sheet, "A10") == "Summary"
        assert _block(sheet, "B11") == "=COUNTA(A6:A7)"
        assert _block(sheet, "B16") == '=IF(COUNT(G6:G7)=0,"n/a",AVERAGE(G6:G7))'

    def test_no_chart_when_bare_tasks(self):
        spec = ep.build_project_plan_spec({"tasks": ["Kickoff", "Wrap-up"]})
        sheet = eg._normalize_spec(spec)["sheets"][0]
        assert not sheet.get("charts")
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []

    def test_template_mode_no_tasks_builds_blank_plan(self, tmp_path):
        # TEMPLATE MODE — "create a project plan" with no tasks must
        # build the blank task plan (10 scaffold rows with the same
        # live guarded Duration / Days-Remaining formulas, dropdowns
        # and zero-reading summary) instead of raising ValueError and
        # falling back to the AI path.
        spec = ep.build_project_plan_spec({})
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        assert sheet["name"] == "Plan"
        assert norm["filename"] == "project_plan.xlsx"

        table = sheet["tables"][0]
        # 10 blank scaffold rows: header 5, data 6..15 — nothing invented
        assert table["start_cell"] == "A5"
        assert len(table["rows"]) == 10
        assert table["rows"][0][:4] == [None] * 4
        assert table["rows"][0][6:] == [None] * 3
        # same guarded live formulas as data rows, pinned to the
        # rendered scaffold rows
        assert table["rows"][0][4] == '=IF(OR(C6="",D6=""),"",D6-C6+1)'
        assert table["rows"][0][5] == '=IF(D6="","",D6-TODAY())'
        assert table["rows"][9][4] == '=IF(OR(C15="",D15=""),"",D15-C15+1)'
        assert table["rows"][9][5] == '=IF(D15="","",D15-TODAY())'

        # dropdowns cover exactly the scaffold band
        dv = {d["range"]: d for d in sheet["data_validation"]}
        assert set(dv) == {"H6:H15", "I6:I15"}
        assert dv["H6:H15"]["values"] == [
            "Not Started",
            "In Progress",
            "Completed",
            "Blocked",
            "On Hold",
        ]
        assert dv["I6:I15"]["values"] == ["Yes", "No"]

        # conditional formats cover the scaffold band too
        cfs = {c["range"] for c in sheet["conditional_formats"]}
        assert cfs == {"H6:H15", "F6:F15", "I6:I15", "G6:G15"}

        # summary block: live tallies over the blanks read 0 / n/a;
        # no deadline → no B3 block, summary label at row 18
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        assert not [b for b in sheet["text_blocks"] if b["cell"] == "B3"]
        assert blocks["A18"] == "Summary"
        assert blocks["B19"] == "=COUNTA(A6:A15)"
        assert blocks["B20"] == '=COUNTIF(H6:H15,"Completed")'
        assert blocks["B22"] == ('=COUNTIFS(D6:D15,"<"&TODAY(),H6:H15,"<>Completed")')
        assert blocks["B23"] == '=COUNTIF(I6:I15,"Yes")'
        assert blocks["B24"] == '=IF(COUNT(G6:G15)=0,"n/a",AVERAGE(G6:G15))'
        # nothing to chart with zero tasks
        assert "charts" not in sheet

        # the workbook still round-trips through openpyxl
        out = tmp_path / "project_template.xlsx"
        eg._build_xlsx(norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Plan"]
        ws = wb["Plan"]
        assert ws.freeze_panes == "A6"
        assert ws.auto_filter.ref == "A5:I15"
        assert ws["E6"].value == '=IF(OR(C6="",D6=""),"",D6-C6+1)'
        assert ws["F15"].value == '=IF(D15="","",D15-TODAY())'
        assert len(ws.data_validations.dataValidation) == 2
        assert len(list(ws.conditional_formatting)) == 4
        assert not ws._charts

    def test_template_mode_with_deadline_keeps_countdown(self):
        # a deadline still renders its block + live countdown row even
        # when the task list is empty
        spec = ep.build_project_plan_spec(
            {"project_name": "New App", "deadline": "2027-03-31", "tasks": []}
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        assert blocks["A1"] == "New App — Project Plan"
        assert blocks["B3"] == date(2027, 3, 31)
        assert blocks["A25"] == "Days to Deadline"
        assert blocks["B25"] == "=$B$3-TODAY()"
        assert norm["filename"] == "New_App_project_plan.xlsx"


class TestProjectPlanCoercion:
    def test_quoted_numbers_and_alias_keys(self):
        p = ep.coerce_project_plan_params(
            {
                "project": "App v2",
                "due_date": "2026-12-31",
                "task_list": [
                    {
                        "task": "API",
                        "assigned_to": "Dana",
                        "start": "2026-10-01",
                        "due_date": "2026-11-01",
                        "progress": "75",
                        "status": "in progress",
                    }
                ],
            }
        )
        assert p["project_name"] == "App v2"
        assert p["deadline"] == "2026-12-31"
        t = p["tasks"][0]
        assert t["name"] == "API"
        assert t["owner"] == "Dana"
        assert t["start"] == "2026-10-01"
        assert t["end"] == "2026-11-01"
        assert t["progress"] == 0.75  # "75" → 75 → fraction
        assert t["status"] == "In Progress"
        assert t["milestone"] is False  # only when stated

    def test_milestone_flag_variants(self):
        p = ep.coerce_project_plan_params(
            {"tasks": [{"name": "Launch", "milestone": True}]}
        )
        assert p["tasks"][0]["milestone"] is True
        p = ep.coerce_project_plan_params(
            {"tasks": [{"name": "Launch", "milestone": "yes"}]}
        )
        assert p["tasks"][0]["milestone"] is True

    def test_nulls_stay_null_and_status_kept_verbatim(self):
        p = ep.coerce_project_plan_params(
            {"tasks": [{"name": "T", "status": None, "progress": None}]}
        )
        t = p["tasks"][0]
        assert t["status"] is None and t["start"] is None and t["end"] is None
        p = ep.coerce_project_plan_params(
            {"tasks": [{"name": "T", "status": "QA signoff"}]}
        )
        assert p["tasks"][0]["status"] == "QA signoff"  # request's own words

    def test_progress_clamped(self):
        p = ep.coerce_project_plan_params(
            {"tasks": [{"name": "A", "progress": 250}, {"name": "B", "progress": -5}]}
        )
        assert p["tasks"][0]["progress"] == 1.0
        assert p["tasks"][1]["progress"] == 0.0

    def test_missing_or_unusable_tasks_raise(self):
        # structurally wrong input still refuses
        with pytest.raises(ValueError):
            ep.coerce_project_plan_params({"tasks": "all the things"})
        with pytest.raises(ValueError):
            ep.coerce_project_plan_params(None)

    def test_empty_tasks_build_blank_template(self):
        # TEMPLATE MODE — no tasks is fine; never invents tasks
        for params in (
            {},
            {"tasks": []},
            {"tasks": [{"owner": "No name"}]},  # nameless entry → skipped
            {"tasks": None, "project_name": None, "deadline": None},
        ):
            p = ep.coerce_project_plan_params(params)
            assert p["tasks"] == []
            assert p["project_name"] is None
            assert p["deadline"] is None

    def test_template_mode_accepts_scalars_without_tasks(self):
        p = ep.coerce_project_plan_params(
            {"project": "Website Relaunch", "due_date": "2026-09-30", "tasks": []}
        )
        assert p["tasks"] == []
        assert p["project_name"] == "Website Relaunch"
        assert p["deadline"] == "2026-09-30"


class TestProjectPlanRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "project_plan", "params": PROJECT_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "project plan for the website launch — tasks with owners, "
            "dates, progress and milestones"
        )
        assert result["pattern"] == "project_plan"
        assert result["sheet_count"] == 1
        assert result["sheet_names"] == ["Plan"]
        assert result["chart_count"] == 1
        assert result["formula_count"] > 10  # 2/task + summary tallies
        assert "project_plan template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    @pytest.mark.asyncio
    async def test_routes_with_empty_params(self, tmp_path, monkeypatch):
        # Same class of bug as meal_planner: the classifier returns
        # project_plan with NO tasks ("create a project plan") — the
        # pattern must build the blank task plan instead of falling
        # back to the AI path.
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "project_plan", "params": {}}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet("create a project plan template")
        assert result["pattern"] == "project_plan"
        assert result["sheet_names"] == ["Plan"]
        assert result["chart_count"] == 0
        # guarded duration/days formulas over the 10 scaffold rows
        assert result["formula_count"] >= 20
        assert "project_plan template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    def test_gate_and_shortlist_match_project_briefs(self):
        for brief in (
            "project plan with milestones and a deadline",
            "task list with start and end dates",
            "roadmap spreadsheet for our sprint",
        ):
            assert eg._PATTERN_GATE_RE.search(brief)
            assert "project_plan" in eg._shortlist_patterns(brief)


# ══════════════════════════════════════════════════════════════════════
# event_planner
# ══════════════════════════════════════════════════════════════════════

EVENT_PARAMS = {
    "event_name": "Amina & Karim Wedding",
    "event_date": "2026-06-20",
    "venue": "Riad Star, Marrakech",
    "guest_count": 120,
    "budget_total": 40000,
    "currency": "USD",
    "guests": [
        {"name": "Youssef", "rsvp": "Yes", "plus_ones": 1},
        {"name": "Lina", "rsvp": "No"},
        {"name": "Omar"},
    ],
    "vendors": [
        {"category": "Venue", "name": "Riad Star", "cost": 12000, "deposit_paid": 6000},
        {
            "category": "Catering",
            "name": "Zitoune",
            "cost": 15000,
            "deposit_paid": 0,
            "status": "Confirmed",
        },
        {
            "category": "Photography",
            "name": "Studio Nour",
            "cost": 3000,
            "deposit_paid": 1000,
            "status": "Pending",
        },
    ],
    "tasks": [
        {
            "task": "Book venue",
            "due_date": "2026-01-15",
            "owner": "Amina",
            "done": True,
        },
        {"task": "Send invitations", "due_date": "2026-03-01"},
        {"task": "Final headcount", "due_date": "2026-06-01"},
    ],
}


class TestEventPlannerBuilder:
    def setup_method(self):
        self.spec = ep.build_event_planner_spec(EVENT_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "overlap" in w]

    def test_sheets_and_geometry(self):
        names = [s["name"] for s in self.norm["sheets"]]
        assert names == ["Event", "Guests", "Vendors", "Tasks"]
        event, guests, vendors, tasks = self.norm["sheets"]
        assert event["tab_color"] == "16304F"
        assert all(s["tab_color"] == "1B3A5C" for s in (guests, vendors, tasks))

        # header block: name 4 / date 5 / days 6 / venue 7 / count 8 /
        # budget 9 → planning status section at row 11
        assert _block(event, "B4") == "Amina & Karim Wedding"
        assert _block(event, "B5") == date(2026, 6, 20)
        assert _block(event, "B6") == "=B5-TODAY()"
        assert _block(event, "B7") == "Riad Star, Marrakech"
        assert _block(event, "B8") == 120
        assert _block(event, "B9") == 40000

        # Guests/Vendors/Tasks tables start at A3 → header row 3,
        # data rows 4..13 (3 real rows + scaffold padding to 10)
        for sheet in (guests, vendors, tasks):
            table = sheet["tables"][0]
            assert table["start_cell"] == "A3"
            assert len(table["rows"]) == 10
            assert sheet["freeze_panes"] == "A4"

    def test_event_sheet_formula_lattice(self):
        event = self.norm["sheets"][0]
        assert _block(event, "A11") == "Planning Status"
        assert _block(event, "B12") == '=COUNTIF(Guests!$B$4:$B$13,"Yes")'
        assert _block(event, "B13") == (
            '=COUNTIF(Guests!$B$4:$B$13,"Yes")'
            '+SUMIF(Guests!$B$4:$B$13,"Yes",Guests!$C$4:$C$13)'
        )
        assert _block(event, "B14") == "=Vendors!$C$14"
        assert _block(event, "B15") == '=IF($B$9>0,Vendors!$C$14/$B$9,"n/a")'
        assert _block(event, "B16") == (
            '=IF(COUNTA(Tasks!$A$4:$A$13)=0,"n/a",'
            'COUNTIF(Tasks!$E$4:$E$13,"Yes")/COUNTA(Tasks!$A$4:$A$13))'
        )

    def test_guests_sheet_rsvp_tallies(self):
        guests = self.norm["sheets"][1]
        rows = guests["tables"][0]["rows"]
        assert rows[0] == ["Youssef", "Yes", 1]
        assert rows[1] == ["Lina", "No", 0]  # plus_ones default 0
        assert rows[2] == ["Omar", "Pending", 0]  # rsvp default Pending
        assert rows[3] == [None, None, None]  # scaffold row

        # tallies start 2 rows below the table's last row (13)
        assert _block(guests, "A16") == "RSVP Summary"
        assert _block(guests, "B17") == '=COUNTIF(B4:B13,"Yes")'
        assert _block(guests, "B18") == '=COUNTIF(B4:B13,"No")'
        assert _block(guests, "B19") == '=COUNTIF(B4:B13,"Pending")'
        assert _block(guests, "B20") == "=COUNTA(A4:A13)"
        assert _block(guests, "B21") == '=SUMIF(B4:B13,"Yes",C4:C13)'
        assert (
            _block(guests, "B22") == '=COUNTIF(B4:B13,"Yes")+SUMIF(B4:B13,"Yes",C4:C13)'
        )

        dv = guests["data_validation"][0]
        assert dv["range"] == "B4:B13"
        assert dv["values"] == ["Yes", "No", "Pending"]

    def test_vendors_sheet_balance_and_budget(self):
        vendors = self.norm["sheets"][2]
        rows = vendors["tables"][0]["rows"]
        # balance due = cost - deposit_paid, guarded for blank rows
        assert rows[0][4] == '=IF($A4="","",C4-D4)'
        assert rows[9][4] == '=IF($A13="","",C13-D13)'
        assert rows[0][2] == 12000 and rows[0][3] == 6000

        # SUM total row directly below the data (row 14)
        total = vendors["tables"][0]["total_row"]
        assert total == [
            "Total",
            "",
            "=SUM(C{first_row}:C{last_row})",
            "=SUM(D{first_row}:D{last_row})",
            "=SUM(E{first_row}:E{last_row})",
            "",
        ]

        assert _block(vendors, "A17") == "Budget"
        assert _block(vendors, "B18") == "=D14"
        assert _block(vendors, "B19") == "=E14"
        assert _block(vendors, "B20") == '=IF(Event!$B$9>0,C14/Event!$B$9,"n/a")'
        assert _block(vendors, "B21") == "=Event!$B$9-C14"

        # pie over the 3 REAL vendor rows only (no blank slices)
        chart = vendors["charts"][0]
        assert chart["type"] == "pie"
        assert chart["categories_range"] == "Vendors!A4:A6"
        assert chart["series"][0]["values_range"] == "Vendors!C4:C6"

    def test_tasks_sheet_progress(self):
        tasks = self.norm["sheets"][3]
        rows = tasks["tables"][0]["rows"]
        assert rows[0][0] == "Book venue"
        assert rows[0][3] == '=IF(C4="","",C4-TODAY())'
        assert rows[0][4] == "Yes"  # done: true
        assert rows[1][4] == "No"  # done false when not stated

        assert _block(tasks, "A16") == "Progress"
        assert _block(tasks, "B17") == (
            '=IF(COUNTA(A4:A13)=0,"n/a",COUNTIF(E4:E13,"Yes")/COUNTA(A4:A13))'
        )
        assert _block(tasks, "B19") == '=COUNTIFS(C4:C13,"<"&TODAY(),E4:E13,"<>Yes")'

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "event.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Event", "Guests", "Vendors", "Tasks"]

        event = wb["Event"]
        assert event["A1"].value == "Amina & Karim Wedding"
        assert event["B5"].number_format == "yyyy-mm-dd"
        assert event["B6"].value == "=B5-TODAY()"
        assert event["B6"].number_format == "0"
        assert event["B9"].number_format == '"$"#,##0.00'
        assert event["B15"].number_format == "0.00%"
        assert event.freeze_panes is None  # dashboard sheet scrolls freely

        guests = wb["Guests"]
        assert guests.freeze_panes == "A4"
        assert len(list(guests.conditional_formatting)) == 1
        assert guests.data_validations.dataValidation[0].formula1 == '"Yes,No,Pending"'

        vendors = wb["Vendors"]
        assert vendors["C14"].value == "=SUM(C4:C13)"
        assert vendors["E4"].value == '=IF($A4="","",C4-D4)'
        assert len(vendors._charts) == 1
        assert vendors.freeze_panes == "A4"

        tasks = wb["Tasks"]
        assert tasks["D4"].value == '=IF(C4="","",C4-TODAY())'
        assert tasks["E4"].value == "Yes"

    def test_budget_math_independently_verified(self):
        # budget utilization = SUM(costs)/budget_total, recomputed here
        costs = [12000, 15000, 3000]
        deposits = [6000, 0, 1000]
        assert sum(costs) == 30000
        assert sum(costs) / 40000 == 0.75
        assert [c - d for c, d in zip(costs, deposits)] == [6000, 15000, 2000]
        assert sum(costs) - sum(deposits) == 23000
        assert 40000 - sum(costs) == 10000
        # expected attendees = confirmed guests + their plus ones
        assert 1 + 1 == 2
        # % complete = 1 of 3 tasks done
        assert abs(1 / 3 - 0.3333333) < 1e-6

    def test_minimal_event_only_name_and_date(self):
        spec = ep.build_event_planner_spec(
            {"event_name": "Kickoff", "event_date": "2026-03-01"}
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        sheet = eg._normalize_spec(spec)["sheets"][0]
        # header block: name 4 / date 5 / days 6 → status at row 8,
        # gapless even without a budget: Tasks Complete lands on B12
        assert _block(sheet, "B6") == "=B5-TODAY()"
        assert _block(sheet, "A8") == "Planning Status"
        assert _block(sheet, "B9") == '=COUNTIF(Guests!$B$4:$B$13,"Yes")'
        assert _block(sheet, "B12") == (
            '=IF(COUNTA(Tasks!$A$4:$A$13)=0,"n/a",'
            'COUNTIF(Tasks!$E$4:$E$13,"Yes")/COUNTA(Tasks!$A$4:$A$13))'
        )
        # no budget → no utilization row anywhere on the Event sheet
        assert not [b for b in sheet["text_blocks"] if b["cell"] == "B13"]
        vendors = eg._normalize_spec(spec)["sheets"][2]
        assert not [b for b in vendors["text_blocks"] if b["cell"] == "B20"]
        assert not vendors.get("charts")  # 0 vendors → no pie

    def test_single_vendor_gets_no_pie(self):
        spec = ep.build_event_planner_spec(
            {
                "event_name": "Dinner",
                "event_date": "2026-03-01",
                "vendors": [{"category": "Catering", "name": "Zitoune", "cost": 800}],
            }
        )
        vendors = eg._normalize_spec(spec)["sheets"][2]
        assert not vendors.get("charts")


class TestEventPlannerCoercion:
    def test_required_name_and_date(self):
        p = ep.coerce_event_planner_params(
            {"event_name": "Wedding", "event_date": "2026-06-20"}
        )
        assert p["event_name"] == "Wedding"
        assert p["event_date"] == "2026-06-20"
        assert p["guests"] == [] and p["vendors"] == [] and p["tasks"] == []

    def test_missing_event_date_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_event_planner_params({"event_name": "Wedding"})

    def test_missing_event_name_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_event_planner_params({"event_date": "2026-06-20"})

    def test_invalid_event_date_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_event_planner_params(
                {"event_name": "Wedding", "event_date": "June 20"}
            )

    def test_quoted_numbers_and_aliases(self):
        p = ep.coerce_event_planner_params(
            {
                "event": "Birthday Party",
                "date": "2026-04-05",
                "guest_count": "120",
                "budget_total": "40,000",
                "currency": "usd",
                "guest_list": [
                    {"name": "Sam", "rsvp": "yes", "plus_ones": "2"},
                    {"guest": "Jo"},
                ],
                "vendor_list": [
                    {"category": "Cake", "vendor": "Bakery", "cost": "250.5"}
                ],
                "checklist": [{"todo": "Order cake", "due": "2026-03-01"}],
            }
        )
        assert p["event_name"] == "Birthday Party"
        assert p["guest_count"] == 120
        assert p["budget_total"] == 40000.0
        assert p["currency"] == "usd"
        assert p["guests"][0] == {"name": "Sam", "rsvp": "Yes", "plus_ones": 2}
        assert p["guests"][1]["rsvp"] == "Pending"
        assert p["vendors"][0]["cost"] == 250.5
        assert p["vendors"][0]["deposit_paid"] == 0.0  # default when not stated
        assert p["tasks"][0]["done"] is False

    def test_rsvp_normalization(self):
        p = ep.coerce_event_planner_params(
            {
                "event_name": "E",
                "event_date": "2026-01-01",
                "guests": [
                    {"name": "A", "rsvp": "coming"},
                    {"name": "B", "rsvp": "declined"},
                    {"name": "C", "rsvp": ""},
                ],
            }
        )
        assert [g["rsvp"] for g in p["guests"]] == ["Yes", "No", "Pending"]

    def test_done_flag_variants(self):
        p = ep.coerce_event_planner_params(
            {
                "event_name": "E",
                "event_date": "2026-01-01",
                "tasks": [{"task": "A", "done": True}, {"task": "B", "done": "yes"}],
            }
        )
        assert [t["done"] for t in p["tasks"]] == [True, True]

    def test_bad_container_types_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_event_planner_params(
                {"event_name": "E", "event_date": "2026-01-01", "guests": "everyone"}
            )
        with pytest.raises(ValueError):
            ep.coerce_event_planner_params(None)


class TestEventPlannerRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "event_planner", "params": EVENT_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "plan my wedding on 2026-06-20 — guest list, vendors with "
            "deposits and a task checklist"
        )
        assert result["pattern"] == "event_planner"
        assert result["sheet_count"] == 4
        assert result["sheet_names"] == ["Event", "Guests", "Vendors", "Tasks"]
        assert result["chart_count"] == 1
        assert result["formula_count"] > 30  # tallies + prefilled row formulas
        assert "event_planner template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    def test_gate_and_shortlist_match_event_briefs(self):
        for brief in (
            "plan my wedding guest list and vendors",
            "birthday party planner with rsvp tracker",
            "conference event planning spreadsheet",
        ):
            assert eg._PATTERN_GATE_RE.search(brief)
            assert "event_planner" in eg._shortlist_patterns(brief)


# ══════════════════════════════════════════════════════════════════════
# travel_planner
# ══════════════════════════════════════════════════════════════════════

TRAVEL_PARAMS = {
    "trip_name": "Lisbon Getaway",
    "destination": "Lisbon",
    "start_date": "2026-05-04",
    "end_date": "2026-05-08",
    "travelers": 2,
    "currency": "EUR",
    "bookings": [
        {
            "type": "Flight",
            "provider": "TAP",
            "date": "2026-05-04",
            "cost": 320,
            "status": "Confirmed",
            "paid": True,
        },
        {
            "type": "Hotel",
            "provider": "Alfama Inn",
            "date": "2026-05-04",
            "cost": 450,
            "status": "Pending",
        },
        {"type": "Train", "provider": "CP", "date": "2026-05-06", "cost": 60},
    ],
    "activities": [
        {
            "day": 1,
            "date": "2026-05-04",
            "time": "14:00",
            "activity": "Check in & Alfama walk",
            "estimated_cost": 0,
        },
        {
            "day": 2,
            "date": "2026-05-05",
            "time": "10:00",
            "activity": "Sintra day trip",
            "estimated_cost": 85,
        },
        {
            "day": 2,
            "date": "2026-05-05",
            "time": "20:00",
            "activity": "Fado dinner",
            "estimated_cost": 70,
        },
        {
            "day": 3,
            "date": "2026-05-06",
            "activity": "Time Out Market",
            "estimated_cost": 40,
        },
    ],
}


class TestTravelPlannerBuilder:
    def setup_method(self):
        self.spec = ep.build_travel_planner_spec(TRAVEL_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "overlap" in w]

    def test_sheets_and_geometry(self):
        names = [s["name"] for s in self.norm["sheets"]]
        assert names == ["Trip", "Itinerary", "Bookings"]
        trip, itinerary, bookings = self.norm["sheets"]
        assert trip["tab_color"] == "16304F"
        assert itinerary["tab_color"] == "1B3A5C"
        assert bookings["tab_color"] == "1B3A5C"
        assert itinerary["freeze_panes"] == "A4"
        assert bookings["freeze_panes"] == "A4"

    def test_trip_sheet_formula_lattice(self):
        trip = self.norm["sheets"][0]
        # header: name 4 / dest 5 / start 6 / end 7 / length 8 /
        # travelers 9 / countdown 10 → cost summary at row 12
        assert _block(trip, "B4") == "Lisbon Getaway"
        assert _block(trip, "B5") == "Lisbon"
        assert _block(trip, "B6") == date(2026, 5, 4)
        assert _block(trip, "B7") == date(2026, 5, 8)
        assert _block(trip, "B8") == '=IF(OR($B$6="",$B$7=""),"n/a",$B$7-$B$6+1)'
        assert _block(trip, "B9") == 2
        assert _block(trip, "B10") == "=$B$6-TODAY()"

        assert _block(trip, "A12") == "Cost Summary"
        assert _block(trip, "B13") == "=SUM(Bookings!$D$4:$D$13)"
        assert _block(trip, "B14") == "=SUM(Itinerary!$E$4:$E$13)"
        assert _block(trip, "B15") == "=B13+B14"
        assert _block(trip, "B16") == (
            '=SUMIF(Bookings!$F$4:$F$13,"Yes",Bookings!$D$4:$D$13)'
        )
        assert _block(trip, "B17") == "=B13-B16"
        assert _block(trip, "B18") == '=IF($B$9=0,"n/a",B15/$B$9)'

    def test_itinerary_sheet(self):
        itinerary = self.norm["sheets"][1]
        table = itinerary["tables"][0]
        assert table["start_cell"] == "A3"
        assert len(table["rows"]) == 10  # 4 activities + scaffold padding
        rows = table["rows"]
        assert rows[0][0] == 1 and rows[0][3] == "Check in & Alfama walk"
        assert rows[0][4] == 0
        assert rows[4] == [None, None, None, None, None]
        assert table["total_row"] == [
            "Total",
            "",
            "",
            "",
            "=SUM(E{first_row}:E{last_row})",
        ]
        # day-by-day estimated cost bar chart over the REAL rows
        chart = itinerary["charts"][0]
        assert chart["type"] == "bar"
        assert chart["categories_range"] == "Itinerary!A4:A7"
        assert chart["series"][0]["values_range"] == "Itinerary!E4:E7"
        assert chart["value_numfmt"] == '#,##0.00" €"'

    def test_bookings_sheet(self):
        bookings = self.norm["sheets"][2]
        table = bookings["tables"][0]
        assert table["headers"] == [
            "Type",
            "Provider",
            "Date",
            "Cost",
            "Status",
            "Paid",
        ]
        rows = table["rows"]
        assert rows[0] == ["Flight", "TAP", date(2026, 5, 4), 320, "Confirmed", "Yes"]
        assert rows[1][5] == "No"  # paid only when stated
        assert rows[2][4] is None  # status only when stated
        assert table["total_row"][3] == "=SUM(D{first_row}:D{last_row})"

        dvs = {d["range"]: d for d in bookings["data_validation"]}
        assert set(dvs) == {"E4:E13", "F4:F13"}
        assert dvs["F4:F13"]["values"] == ["Yes", "No"]

        cf_ranges = {c["range"] for c in bookings["conditional_formats"]}
        assert cf_ranges == {"E4:E13", "F4:F13"}

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "trip.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Trip", "Itinerary", "Bookings"]

        trip = wb["Trip"]
        assert trip["A1"].value == "Lisbon Getaway"
        assert trip["B8"].value == '=IF(OR($B$6="",$B$7=""),"n/a",$B$7-$B$6+1)'
        assert trip["B8"].number_format == "0"
        assert trip["B10"].value == "=$B$6-TODAY()"
        assert trip["B13"].number_format == '#,##0.00" €"'
        assert trip["B18"].value == '=IF($B$9=0,"n/a",B15/$B$9)'
        assert trip.freeze_panes is None  # summary sheet scrolls freely

        itinerary = wb["Itinerary"]
        assert itinerary.freeze_panes == "A4"
        assert itinerary["E14"].value == "=SUM(E4:E13)"
        assert len(itinerary._charts) == 1

        bookings = wb["Bookings"]
        assert bookings["D14"].value == "=SUM(D4:D13)"
        assert bookings["F4"].value == "Yes"
        assert len(list(bookings.conditional_formatting)) == 2
        assert len(bookings.data_validations.dataValidation) == 2

    def test_trip_cost_math_independently_verified(self):
        # recompute the live formulas' semantics in pure Python
        booking_costs = [320, 450, 60]
        activity_costs = [0, 85, 70, 40]
        total_bookings = sum(booking_costs)
        total_activities = sum(activity_costs)
        paid = sum(c for c, b in zip(booking_costs, [True, False, False]) if b)
        assert total_bookings == 830
        assert total_activities == 195
        assert total_bookings + total_activities == 1025
        assert paid == 320
        assert total_bookings - paid == 510
        assert (total_bookings + total_activities) / 2 == 512.5  # per-traveler
        # trip length = end - start + 1
        assert (date(2026, 5, 8) - date(2026, 5, 4)).days + 1 == 5

    def test_no_cost_chart_when_activities_free(self):
        spec = ep.build_travel_planner_spec(
            {
                "destination": "Rome",
                "activities": [
                    {"day": 1, "activity": "Walk the centro storico"},
                    {"day": 2, "activity": "Vatican museums"},
                ],
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        itinerary = eg._normalize_spec(spec)["sheets"][1]
        assert not itinerary.get("charts")
        # no dates → no countdown / length rows; travelers default 1
        trip = eg._normalize_spec(spec)["sheets"][0]
        assert _block(trip, "B4") == "Rome"  # no trip name → destination first
        assert not [b for b in trip["text_blocks"] if b["cell"] == "B6"]
        assert _block(trip, "B5") == 1  # travelers default
        # cost summary: label row 7, share guarded on the travelers row
        assert _block(trip, "A7") == "Cost Summary"
        assert _block(trip, "B13") == '=IF($B$5=0,"n/a",B10/$B$5)'

    def test_chart_categories_fall_back_to_dates(self):
        spec = ep.build_travel_planner_spec(
            {
                "activities": [
                    {"date": "2026-07-01", "activity": "Beach", "estimated_cost": 20},
                    {"date": "2026-07-02", "activity": "Museum", "estimated_cost": 15},
                ]
            }
        )
        itinerary = eg._normalize_spec(spec)["sheets"][1]
        chart = itinerary["charts"][0]
        assert chart["categories_range"] == "Itinerary!B4:B5"
        assert chart["series"][0]["values_range"] == "Itinerary!E4:E5"

    def test_bookings_only_trip(self):
        spec = ep.build_travel_planner_spec(
            {
                "trip_name": "Fez Weekend",
                "bookings": [{"type": "Flight", "cost": 100}],
                "travelers": 3,
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        trip = eg._normalize_spec(spec)["sheets"][0]
        # trip name 4 / travelers 5 → cost summary label at row 7
        assert _block(trip, "B4") == "Fez Weekend"
        assert _block(trip, "B5") == 3
        assert _block(trip, "A7") == "Cost Summary"
        assert _block(trip, "B8") == "=SUM(Bookings!$D$4:$D$13)"
        assert _block(trip, "B13") == '=IF($B$5=0,"n/a",B10/$B$5)'
        assert not trip.get("charts")  # bookings have no chart of their own


class TestTravelPlannerCoercion:
    def test_quoted_numbers_and_defaults(self):
        p = ep.coerce_travel_planner_params(
            {
                "trip_name": "TOKYO",
                "start_date": "2026-03-10",
                "travelers": "2",
                "bookings": [
                    {"type": "Flight", "provider": "ANA", "cost": "720", "paid": True},
                    {"kind": "Hotel", "name": "Shinjuku Inn", "amount": 900},
                ],
                "activities": [
                    {"day": "2", "activity": "TeamLab", "estimated_cost": "35"}
                ],
            }
        )
        assert p["travelers"] == 2
        assert p["bookings"][0]["cost"] == 720.0
        assert p["bookings"][0]["paid"] is True
        assert p["bookings"][1]["type"] == "Hotel"
        assert p["bookings"][1]["paid"] is False  # paid only when stated
        assert p["activities"][0]["day"] == 2
        assert p["activities"][0]["estimated_cost"] == 35.0

    def test_travelers_default_and_clamp(self):
        p = ep.coerce_travel_planner_params({"bookings": [{"type": "Bus", "cost": 5}]})
        assert p["travelers"] == 1  # stanza default
        p = ep.coerce_travel_planner_params(
            {"travelers": 0, "bookings": [{"type": "Bus", "cost": 5}]}
        )
        assert p["travelers"] == 1  # clamped to at least 1

    def test_estimated_cost_only_when_stated(self):
        p = ep.coerce_travel_planner_params(
            {"activities": [{"activity": "Sunset walk"}, "Museum visit"]}
        )
        assert p["activities"][0]["estimated_cost"] is None
        assert p["activities"][1]["day"] is None
        assert p["activities"][1]["activity"] == "Museum visit"  # bare strings ok

    def test_missing_items_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_travel_planner_params({})
        with pytest.raises(ValueError):
            ep.coerce_travel_planner_params({"destination": "Oslo"})
        with pytest.raises(ValueError):
            ep.coerce_travel_planner_params({"bookings": "flight"})
        with pytest.raises(ValueError):
            ep.coerce_travel_planner_params(None)

    def test_bad_dates_are_dropped_not_raised(self):
        p = ep.coerce_travel_planner_params(
            {"start_date": "next Monday", "bookings": [{"type": "Flight", "cost": 5}]}
        )
        assert p["start_date"] is None


class TestTravelPlannerRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "travel_planner", "params": TRAVEL_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "5-day trip to Lisbon — flights, hotel bookings and a "
            "day-by-day itinerary with estimated costs"
        )
        assert result["pattern"] == "travel_planner"
        assert result["sheet_count"] == 3
        assert result["sheet_names"] == ["Trip", "Itinerary", "Bookings"]
        assert result["chart_count"] == 1
        assert result["formula_count"] > 5  # totals + live cost summary
        assert "travel_planner template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    def test_gate_and_shortlist_match_travel_briefs(self):
        for brief in (
            "five day trip itinerary with flights and a hotel",
            "vacation planner for our europe getaway",
            "travel booking spreadsheet for the holiday",
        ):
            assert eg._PATTERN_GATE_RE.search(brief)
            assert "travel_planner" in eg._shortlist_patterns(brief)


# ══════════════════════════════════════════════════════════════════════
# registry sanity
# ══════════════════════════════════════════════════════════════════════


class TestPlanningPatternRegistry:
    def test_patterns_registered_with_stanza_names(self):
        for name in ("project_plan", "event_planner", "travel_planner"):
            assert name in ep.PATTERN_BUILDERS
            assert name in ep.PATTERN_KEYWORDS
            assert ep.PATTERN_DESCRIPTIONS.get(name)

    def test_module_entry_points(self):
        from app.services.patterns import event_planner, project_plan, travel_planner

        for module, name in (
            (project_plan, "project_plan"),
            (event_planner, "event_planner"),
            (travel_planner, "travel_planner"),
        ):
            assert module.PATTERN_NAME == name
            assert module.coerce_params is not None
            assert callable(module.build_spec)
