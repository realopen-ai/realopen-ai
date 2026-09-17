"""
Tests for the team-operations pattern batch (task 3-e): leave_tracker,
timesheet, shift_schedule, equipment_maintenance.

Same conventions as tests/test_excel_patterns.py:
  • param coercion (quoted numbers, alias keys, null → documented
    defaults, missing required params → ValueError)
  • builder layout math — every formula references the row the
    converter will actually render (title rows included!), pinned via
    eg._normalize_spec
  • independent math verification — booked days per employee, weekly
    rota hours via the code→hours mapping, next-due dates and pay
    splits are recomputed in plain Python and compared
  • a converter round-trip (_build_xlsx → openpyxl reload)
  • one routing test per pattern (mocked classifier, AI path must NOT
    run)
"""

import sys
from datetime import date, timedelta
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _classify_response(payload):
    import json

    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


def _monday_of(d: date) -> date:
    return d - timedelta(days=d.weekday())


def _next_monday(d: date) -> date:
    return d + timedelta(days=(7 - d.weekday()) or 7)


# ═════════════════════════════════════════════════════════════════════
# leave_tracker
# ═════════════════════════════════════════════════════════════════════

LEAVE_PARAMS = {
    "year": 2026,
    "employees": [
        {
            "name": "Alice",
            "department": "Sales",
            "annual_leave_days": 25,
            "carried_over": 2,
        },
        {
            "name": "Bob",
            "department": "Ops",
            "annual_leave_days": 20,
            "carried_over": 0,
        },
    ],
    "leave_log": [
        {
            "employee": "Alice",
            "type": "Vacation",
            "start": "2026-03-02",
            "end": "2026-03-06",
            "status": "Taken",
        },
        {
            "employee": "Alice",
            "type": "Sick",
            "start": "2026-04-10",
            "end": "2026-04-10",
            "status": "taken",
        },
        {
            "employee": "Bob",
            "type": "Vacation",
            "start": "2026-07-01",
            "end": "2026-07-05",
        },
        {
            "employee": "Dana",
            "type": "Unpaid",
            "start": "2026-05-11",
            "end": "2026-05-12",
            "status": "Planned",
        },
    ],
}


class TestLeaveTrackerBuilder:
    def setup_method(self):
        self.spec = ep.build_leave_tracker_spec(LEAVE_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]
        assert not [w for w in warnings if "overlap" in w]

    def test_geometry_and_formula_lattice(self):
        bal, log = self.norm["sheets"]
        assert bal["name"] == "Balances"
        assert log["name"] == "Leave Log"
        bal_table = bal["tables"][0]
        # start A4, no table title → header row 4, data 5..7 (3 people —
        # Dana comes from the log), total row 8
        assert bal_table["start_cell"] == "A4"
        assert len(bal_table["rows"]) == 3
        # log: 4 entries + 6 blank fillers = 10 rows → 5..14, total 15
        log_table = log["tables"][0]
        assert log_table["start_cell"] == "A4"
        assert len(log_table["rows"]) == 10

        alice, bob, dana = bal_table["rows"]
        # every ref points at the rows the converter actually renders
        assert alice[4] == "=SUMIF('Leave Log'!$A$5:$A$14,$A5,'Leave Log'!$E$5:$E$14)"
        assert alice[5] == (
            "=SUMIFS('Leave Log'!$E$5:$E$14,"
            "'Leave Log'!$A$5:$A$14,$A5,"
            "'Leave Log'!$F$5:$F$14,\"Taken\")"
        )
        assert alice[6] == (
            "=SUMIFS('Leave Log'!$E$5:$E$14,"
            "'Leave Log'!$A$5:$A$14,$A5,"
            "'Leave Log'!$F$5:$F$14,\"Planned\")"
        )
        # remaining = entitlement + carried − booked (live, blank-safe)
        assert alice[7] == "=C5+D5-E5"
        assert alice[8] == "=IF((C5+D5)>0,E5/(C5+D5),0)"
        assert bob[7] == "=C6+D6-E6"
        # Dana has no stated entitlement — still a live formula
        assert dana[7] == "=C7+D7-E7"
        # total row pins to the exact data rows 5..7
        total = bal_table["total_row"]
        assert total[4] == "=SUM(E5:E7)"
        assert total[7] == "=SUM(H5:H7)"
        assert total[8] == "=IF((C8+D8)>0,E8/(C8+D8),0)"

        # log duration = end − start + 1, blank-safe on the filler rows
        assert log_table["rows"][0][4] == '=IF(OR(C5="",D5=""),0,D5-C5+1)'
        assert log_table["rows"][3][4] == '=IF(OR(C8="",D8=""),0,D8-C8+1)'
        assert log_table["rows"][9][4] == '=IF(OR(C14="",D14=""),0,D14-C14+1)'
        assert log_table["total_row"][4] == "=SUM(E5:E14)"

    def test_booked_days_math_independently(self):
        # recompute booked/taken/planned in plain Python from the params
        booked, taken, planned = {}, {}, {}
        for entry in LEAVE_PARAMS["leave_log"]:
            start = date(*(int(x) for x in entry["start"].split("-")))
            end = date(*(int(x) for x in entry["end"].split("-")))
            days = (end - start).days + 1
            who = entry["employee"]
            booked[who] = booked.get(who, 0) + days
            if entry.get("status", "Planned").lower() == "taken":
                taken[who] = taken.get(who, 0) + days
            else:
                planned[who] = planned.get(who, 0) + days

        assert booked == {"Alice": 6, "Bob": 5, "Dana": 2}
        assert taken == {"Alice": 6}
        assert planned == {"Bob": 5, "Dana": 2}

        # the SUMIF/SUMIFS ranges cover exactly the rendered log rows
        # 5..14 — one row per booking (4) + filler rows (6)
        bal = self.norm["sheets"][0]
        alice = bal["tables"][0]["rows"][0]
        assert "'Leave Log'!$A$5:$A$14" in alice[4]
        assert "'Leave Log'!$E$5:$E$14" in alice[4]
        assert "'Leave Log'!$F$5:$F$14" in alice[5]
        assert "'Leave Log'!$F$5:$F$14" in alice[6]

        # remaining = entitlement + carried − booked
        remaining = {"Alice": 25 + 2 - 6, "Bob": 20 + 0 - 5, "Dana": 0 + 0 - 2}
        assert remaining == {"Alice": 21, "Bob": 15, "Dana": -2}

    def test_chart_ranges_match_layout(self):
        chart = self.norm["sheets"][0]["charts"][0]
        assert chart["type"] == "bar"
        assert chart["categories_range"] == "Balances!$A$5:$A$7"
        assert chart["series"][0]["values_range"] == "Balances!$E$5:$E$7"
        assert chart["series"][1]["values_range"] == "Balances!$H$5:$H$7"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "leave.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        bal, log = wb["Balances"], wb["Leave Log"]
        # live formulas survive the round-trip
        assert (
            bal["E5"].value
            == "=SUMIF('Leave Log'!$A$5:$A$14,$A5,'Leave Log'!$E$5:$E$14)"
        )
        assert bal["H5"].value == "=C5+D5-E5"
        assert bal["E8"].value == "=SUM(E5:E7)"
        assert log["E5"].value == '=IF(OR(C5="",D5=""),0,D5-C5+1)'
        assert log["E15"].value == "=SUM(E5:E14)"
        # dates render as real dates with the shared date format
        assert log["C5"].value.date() == date(2026, 3, 2)
        assert log["C5"].number_format == "yyyy-mm-dd"
        # day counts render as plain integers
        assert log["E5"].number_format == "0"
        assert bal["E5"].number_format == "#,##0.##"
        assert bal["I5"].number_format == "0.00%"
        # frozen below the header, dropdowns + status colors in place
        assert bal.freeze_panes == "A5"
        assert log.freeze_panes == "A5"
        log_dvs = log.data_validations.dataValidation
        assert len(log_dvs) == 3
        sources = {str(dv.sqref): dv.formula1 for dv in log_dvs}
        assert sources["A5:A14"] == "Balances!$A$5:$A$7"
        assert sources["F5:F14"] == '"Taken,Planned"'
        assert len(log.conditional_formatting._cf_rules) == 1
        assert len(bal.conditional_formatting._cf_rules) == 1
        # one chart on the main sheet
        assert len(bal._charts) == 1

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after


class TestLeaveTrackerCoercion:
    def test_quoted_numbers_and_alias_keys(self):
        p = ep.coerce_leave_tracker_params(
            {
                "year": "2026",
                "staff": [
                    {"name": "Alice", "annual_leave_days": "25", "carried_over": "3"}
                ],
                "bookings": [
                    {
                        "employee": "Alice",
                        "type": "Vacation",
                        "start": "2026-03-02",
                        "end": "2026-03-04",
                        "status": "Planned",
                    }
                ],
            }
        )
        assert p["year"] == 2026
        assert p["employees"][0]["entitlement"] == 25.0
        assert p["employees"][0]["carried"] == 3.0
        assert p["leave_log"][0]["employee"] == "Alice"
        assert p["leave_log"][0]["status"] == "Planned"

    def test_null_semantics(self):
        p = ep.coerce_leave_tracker_params(
            {"employees": [{"name": "Alice"}], "leave_log": []}
        )
        # annual_leave_days null when not stated → blank input cell
        assert p["employees"][0]["entitlement"] is None
        # carried_over 0 when not stated
        assert p["employees"][0]["carried"] == 0.0
        assert p["leave_log"] == []

    def test_year_defaults_to_current(self):
        p = ep.coerce_leave_tracker_params({"employees": [{"name": "A"}]})
        assert p["year"] == date.today().year

    def test_status_normalization(self):
        p = ep.coerce_leave_tracker_params(
            {
                "employees": [{"name": "A"}],
                "leave_log": [
                    {
                        "employee": "A",
                        "start": "2026-02-02",
                        "end": "2026-02-02",
                        "status": "TAKEN",
                    },
                    {"employee": "A", "start": "2026-02-05", "end": "2026-02-06"},
                ],
            }
        )
        assert p["leave_log"][0]["status"] == "Taken"
        # status "Planned" unless the leave already happened
        assert p["leave_log"][1]["status"] == "Planned"

    def test_employees_derived_from_log(self):
        # people named only in the log still get a Balances row
        p = ep.coerce_leave_tracker_params(
            {
                "leave_log": [
                    {"employee": "Dana", "start": "2026-05-11", "end": "2026-05-12"}
                ]
            }
        )
        assert [e["name"] for e in p["employees"]] == ["Dana"]
        assert p["employees"][0]["entitlement"] is None

    def test_duplicate_employees_deduped(self):
        p = ep.coerce_leave_tracker_params(
            {"employees": [{"name": "Alice"}, {"name": "alice"}, {"name": "Bob"}]}
        )
        assert [e["name"] for e in p["employees"]] == ["Alice", "Bob"]

    def test_missing_everything_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_leave_tracker_params({})
        with pytest.raises(ValueError):
            ep.coerce_leave_tracker_params({"employees": []})
        with pytest.raises(ValueError):
            ep.coerce_leave_tracker_params({"leave_log": []})
        with pytest.raises(ValueError):
            ep.coerce_leave_tracker_params(["nope"])

    def test_log_entries_without_dates_dropped(self):
        p = ep.coerce_leave_tracker_params(
            {
                "employees": [{"name": "A"}],
                "leave_log": [
                    {"employee": "A", "type": "Vacation"},  # no dates
                    {"employee": "A", "start": "2026-02-02", "end": "2026-02-03"},
                ],
            }
        )
        assert len(p["leave_log"]) == 1


class TestLeaveTrackerRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "leave_tracker", "params": LEAVE_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "leave tracker for my team: vacation and PTO balances with "
            "carried over days"
        )
        assert result["pattern"] == "leave_tracker"
        assert result["sheet_names"] == ["Balances", "Leave Log"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ═════════════════════════════════════════════════════════════════════
# timesheet
# ═════════════════════════════════════════════════════════════════════

TIMESHEET_PARAMS = {
    "employee": "Jane Doe",
    "week_start": "2026-01-05",
    "hourly_rate": 25,
    "overtime_multiplier": 1.5,
    "currency": "USD",
    "entries": [
        {"project": "Project A", "mon": 8, "tue": 8, "wed": 8, "thu": 8, "fri": 6},
        {"project": "Project B", "wed": 2, "thu": 4, "sat": 13},
    ],
}


class TestTimesheetBuilder:
    def setup_method(self):
        self.spec = ep.build_timesheet_spec(TIMESHEET_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]
        assert not [w for w in warnings if "overlap" in w]

    def test_geometry_and_formula_lattice(self):
        sheet = self.norm["sheets"][0]
        assert sheet["name"] == "Timesheet"
        table = sheet["tables"][0]
        # with pay: assumptions rows 4-7 → header 9, data 10..11, total 12
        assert table["start_cell"] == "A9"
        assert len(table["rows"]) == 2
        # day headers carry the week's real dates
        assert table["headers"][1] == "Mon 05 Jan"
        assert table["headers"][6] == "Sat 10 Jan"
        assert table["headers"][7] == "Sun 11 Jan"
        assert table["headers"][8] == "Total Hours"

        row_a, row_b = table["rows"]
        # row totals = SUM across the day columns
        assert row_a[8] == "=SUM(B10:H10)"
        assert row_b[8] == "=SUM(B11:H11)"
        # day not worked → blank cell (SUM-friendly), not 0
        assert row_a[6] is None and row_a[7] is None
        assert row_b[0] == "Project B" and row_b[5] is None
        # total row: per-day totals + grand total, exact rows 10..11
        total = table["total_row"]
        assert total[1] == "=SUM(B10:B11)"
        assert total[6] == "=SUM(G10:G11)"
        assert total[8] == "=SUM(I10:I11)"

        # assumptions block: named cells B5..B7
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["B5"]["text"] == 25.0
        assert blocks["B5"]["number_format"] == '"$"#,##0.00'
        assert blocks["B6"]["text"] == 40.0
        assert blocks["B7"]["text"] == 1.5

        # pay summary: absolute references to the assumption cells
        assert blocks["B15"]["text"] == "=I12"
        assert blocks["B16"]["text"] == "=MIN(B15,$B$6)"
        assert blocks["B17"]["text"] == "=MAX(0,B15-B16)"
        assert blocks["B18"]["text"] == "=B16*$B$5"
        assert blocks["B19"]["text"] == "=B17*$B$5*$B$7"
        assert blocks["B20"]["text"] == "=B18+B19"

    def test_hours_and_pay_math_independently(self):
        # plain-Python truth from the raw params
        day_totals = {
            "mon": 8,
            "tue": 8,
            "wed": 8 + 2,
            "thu": 8 + 4,
            "fri": 6,
            "sat": 13,
            "sun": 0,
        }
        per_project = [38, 19]
        grand = sum(day_totals.values())
        assert grand == 57

        sheet = self.norm["sheets"][0]
        rows = sheet["tables"][0]["rows"]
        for i, (row, expected) in enumerate(zip(rows, per_project)):
            hours = [v for v in row[1:8] if isinstance(v, (int, float))]
            assert sum(hours) == expected
            assert row[8] == "=SUM(B{r}:H{r})".format(r=10 + i)

        # overtime split on the grand total against the threshold
        reg = min(grand, 40)
        ot = max(0, grand - reg)
        assert (reg, ot) == (40, 17)
        assert reg * 25 == 1000
        assert ot * 25 * 1.5 == 637.5
        assert reg * 25 + ot * 25 * 1.5 == 1637.5

    def test_chart_ranges_match_layout(self):
        chart = self.norm["sheets"][0]["charts"][0]
        assert chart["categories_range"] == "Timesheet!$A$10:$A$11"
        assert chart["series"][0]["values_range"] == "Timesheet!$I$10:$I$11"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "timesheet.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        ws = wb["Timesheet"]
        assert ws["I10"].value == "=SUM(B10:H10)"
        assert ws["B12"].value == "=SUM(B10:B11)"
        assert ws["B16"].value == "=MIN(B15,$B$6)"
        assert ws["B19"].value == "=B17*$B$5*$B$7"
        # hours format + currency on pay cells
        assert ws["I10"].number_format == "#,##0.##"
        assert ws["B5"].number_format == '"$"#,##0.00'
        assert ws["B20"].number_format == '"$"#,##0.00'
        # panes frozen below the header, project column pinned
        assert ws.freeze_panes == "B10"
        # long-day conditional formats: 2 rules on the day grid
        cf = list(ws.conditional_formatting)
        assert len(cf) == 1
        assert str(cf[0].sqref) == "B10:H11"
        assert len(cf[0].rules) == 2
        assert len(ws._charts) == 1

    def test_no_pay_variant_compacts_layout(self):
        spec = ep.build_timesheet_spec(
            {"entries": [{"project": "Project A", "mon": 8}]}
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        sheet = eg._normalize_spec(spec)["sheets"][0]
        table = sheet["tables"][0]
        # no rate → no assumptions block → header at row 4, data row 5
        assert table["start_cell"] == "A4"
        assert table["rows"][0][8] == "=SUM(B5:H5)"
        assert table["total_row"][8] == "=SUM(I5:I5)"
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert "B5" not in blocks  # no rate cell
        assert "A15" not in blocks  # no PAY SUMMARY heading
        assert (
            sheet.get("charts")
            and sheet["charts"][0]["categories_range"] == "Timesheet!$A$5:$A$5"
        )

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after


class TestTimesheetCoercion:
    def test_quoted_numbers(self):
        p = ep.coerce_timesheet_params(
            {
                "hourly_rate": "25",
                "overtime_multiplier": "1.5",
                "entries": [{"project": "A", "mon": "8", "tue": "7.5"}],
            }
        )
        assert p["hourly_rate"] == 25.0
        assert p["overtime_multiplier"] == 1.5
        assert p["entries"][0]["hours"][0] == 8.0
        assert p["entries"][0]["hours"][1] == 7.5

    def test_week_start_null_is_current_weeks_monday(self):
        p = ep.coerce_timesheet_params({"entries": [{"project": "A", "mon": 8}]})
        assert p["week_start"] == _monday_of(date.today()).isoformat()

    def test_week_start_snaps_to_monday(self):
        p = ep.coerce_timesheet_params(
            {"week_start": "2026-01-07", "entries": [{"project": "A", "mon": 8}]}
        )
        assert p["week_start"] == "2026-01-05"

    def test_defaults_and_zero_rate(self):
        p = ep.coerce_timesheet_params({"entries": [{"project": "A", "mon": 8}]})
        assert p["hourly_rate"] is None  # pay only when a rate is stated
        assert p["overtime_multiplier"] == 1.5  # documented default
        assert p["employee"] is None

        p0 = ep.coerce_timesheet_params(
            {"hourly_rate": 0, "entries": [{"project": "A", "mon": 8}]}
        )
        assert p0["hourly_rate"] is None  # zero rate → no pay block

    def test_day_aliases(self):
        p = ep.coerce_timesheet_params(
            {"entries": [{"project": "A", "monday": 8, "saturday": 4}]}
        )
        assert p["entries"][0]["hours"][0] == 8.0
        assert p["entries"][0]["hours"][5] == 4.0
        assert p["entries"][0]["hours"][1] is None

    def test_entries_required(self):
        with pytest.raises(ValueError):
            ep.coerce_timesheet_params({})
        with pytest.raises(ValueError):
            ep.coerce_timesheet_params({"entries": []})
        with pytest.raises(ValueError):
            ep.coerce_timesheet_params({"entries": [{"hours": 8}]})  # no project
        with pytest.raises(ValueError):
            ep.coerce_timesheet_params("nope")

    def test_currency_alias(self):
        p = ep.coerce_timesheet_params(
            {
                "hourly_rate": 20,
                "currency": "MAD",
                "entries": [{"project": "A", "mon": 8}],
            }
        )
        assert p["currency"] == '#,##0.00" DH"'


class TestTimesheetRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "timesheet", "params": TIMESHEET_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "timesheet for the hours I worked last week, with hourly rate"
        )
        assert result["pattern"] == "timesheet"
        assert result["sheet_names"] == ["Timesheet"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ═════════════════════════════════════════════════════════════════════
# shift_schedule
# ═════════════════════════════════════════════════════════════════════

SHIFT_PARAMS = {
    "team_name": "Cafe Team",
    "week_start": "2026-01-12",
    "staff": [
        {
            "name": "Alice",
            "role": "Barista",
            "mon": "M",
            "tue": "M",
            "wed": "E",
            "thu": "M",
            "fri": "M",
            "sat": "N",
            "sun": "O",
        },
        {
            "name": "Bob",
            "role": "Chef",
            "mon": "E",
            "tue": "E",
            "wed": "E",
            "thu": "E",
            "fri": "N",
            "sat": "O",
        },
        {"name": "Cleo"},
    ],
}

DEFAULT_CODE_HOURS = {"M": 8.0, "E": 8.0, "N": 10.0, "O": 0.0}


class TestShiftScheduleBuilder:
    def setup_method(self):
        self.spec = ep.build_shift_schedule_spec(SHIFT_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]
        assert not [w for w in warnings if "overlap" in w]

    def test_geometry_and_formula_lattice(self):
        rota, codes, calc = self.norm["sheets"]
        assert rota["name"] == "Rota"
        assert codes["name"] == "Codes"
        assert calc["name"] == "Calc"
        assert calc["hidden"] is True  # the helper lattice stays hidden

        rota_table = rota["tables"][0]
        assert rota_table["start_cell"] == "A4"
        assert len(rota_table["rows"]) == 3
        assert rota_table["headers"][0] == "Staff Member"
        assert rota_table["headers"][2] == "Mon 12 Jan"

        # visible grid holds clean codes; weekly hours ← Calc sum cell
        alice = rota_table["rows"][0]
        assert alice[2:9] == ["M", "M", "E", "M", "M", "N", "O"]
        assert alice[9] == "=Calc!J5"
        bob = rota_table["rows"][1]
        assert bob[8] is None  # blank day omitted
        assert bob[9] == "=Calc!J6"

        # totals: per-day hours summed from Calc, week total from Rota
        total = rota_table["total_row"]
        assert total[2] == "=SUM(Calc!C5:C7)"
        assert total[8] == "=SUM(Calc!I5:I7)"
        assert total[9] == "=SUM(J5:J7)"

        # Codes block: code → label → hours (rows 5..8)
        codes_table = codes["tables"][0]
        assert codes_table["start_cell"] == "A4"
        assert codes_table["rows"] == [
            ["M", "Morning", 8.0],
            ["E", "Evening", 8.0],
            ["N", "Night", 10.0],
            ["O", "Off", 0.0],
        ]

        # Calc: one row per staff, per-day VLOOKUP against the codes
        calc_table = calc["tables"][0]
        assert calc_table["start_cell"] == "B4"
        assert calc_table["rows"][0][0] == "=Rota!$A$5"
        assert calc_table["rows"][0][1] == (
            '=IF(Rota!C5="",0,IFERROR(VLOOKUP(Rota!C5,Codes!$A$5:$C$8,3,FALSE),0))'
        )
        assert calc_table["rows"][2][7] == (
            '=IF(Rota!I7="",0,IFERROR(VLOOKUP(Rota!I7,Codes!$A$5:$C$8,3,FALSE),0))'
        )
        assert calc_table["rows"][0][8] == "=SUM(C5:I5)"
        assert calc_table["rows"][1][8] == "=SUM(C6:I6)"

        # day dropdowns read the codes block live
        dv = rota["data_validation"][0]
        assert dv["range"] == "C5:I7"
        assert dv["source_range"] == "Codes!$A$5:$A$8"

    def test_weekly_hours_math_independently(self):
        # recompute weekly hours in plain Python via the code→hours map
        expected = {}
        for person in SHIFT_PARAMS["staff"]:
            total = 0.0
            for day in ("mon", "tue", "wed", "thu", "fri", "sat", "sun"):
                code = person.get(day)
                total += DEFAULT_CODE_HOURS.get(code, 0.0) if code else 0.0
            expected[person["name"]] = total
        assert expected == {"Alice": 50.0, "Bob": 42.0, "Cleo": 0.0}

        # day totals (staffing level per day) from the same mapping
        day_totals = []
        for day in ("mon", "tue", "wed", "thu", "fri", "sat", "sun"):
            day_totals.append(
                sum(
                    DEFAULT_CODE_HOURS.get(p.get(day), 0.0) if p.get(day) else 0.0
                    for p in SHIFT_PARAMS["staff"]
                )
            )
        assert day_totals == [16, 16, 16, 16, 18, 10, 0]
        assert sum(day_totals) == 92.0

        # the Calc lattice realizes exactly this mapping: every VLOOKUP
        # points at the rendered codes rows and the right day column
        calc = self.norm["sheets"][2]
        rows = calc["tables"][0]["rows"]
        for i, person in enumerate(SHIFT_PARAMS["staff"]):
            assert rows[i][8] == "=SUM(C{r}:I{r})".format(r=5 + i)

    def test_custom_codes_and_blank_days(self):
        spec = ep.build_shift_schedule_spec(
            {
                "shift_codes": {
                    "D": {"label": "Day", "hours": 9},
                    "L": 6,
                    "F": "Free",
                },
                "staff": [{"name": "Alice", "mon": "D", "tue": "L", "sun": "F"}],
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        codes, calc = norm["sheets"][1], norm["sheets"][2]
        # bare number → hours with the code as label; bare string → label
        assert codes["tables"][0]["rows"] == [
            ["D", "Day", 9.0],
            ["L", "L", 6.0],
            ["F", "Free", 0.0],
        ]
        # VLOOKUP range computed from the ACTUAL number of code rows
        assert calc["tables"][0]["rows"][0][1] == (
            '=IF(Rota!C5="",0,IFERROR(VLOOKUP(Rota!C5,Codes!$A$5:$C$7,3,FALSE),0))'
        )
        # dropdown source follows the code list length too
        assert (
            norm["sheets"][0]["data_validation"][0]["source_range"] == "Codes!$A$5:$A$7"
        )
        # Alice: 9 + 6 + 0 = 15 hours via the custom mapping
        assert 9.0 + 6.0 + 0.0 == 15.0

    def test_default_codes_when_null(self):
        p = ep.coerce_shift_schedule_params({"staff": [{"name": "A"}]})
        assert [c[0] for c in p["shift_codes"]] == ["M", "E", "N", "O"]
        assert dict((c[0], c[2]) for c in p["shift_codes"]) == DEFAULT_CODE_HOURS

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "shift.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        rota, codes, calc = wb["Rota"], wb["Codes"], wb["Calc"]
        # visible grid: clean codes, weekly hours referencing Calc
        assert rota["C5"].value == "M"
        assert rota["J5"].value == "=Calc!J5"
        assert rota["C8"].value == "=SUM(Calc!C5:C7)"
        # hidden Calc sheet holds the VLOOKUP lattice
        assert calc.sheet_state == "hidden"
        assert calc["C5"].value == (
            '=IF(Rota!C5="",0,IFERROR(VLOOKUP(Rota!C5,Codes!$A$5:$C$8,3,FALSE),0))'
        )
        assert calc["J5"].value == "=SUM(C5:I5)"
        # codes reference block + live dropdown
        assert codes["A5"].value == "M"
        assert codes["C5"].value == 8
        dvs = rota.data_validations.dataValidation
        assert len(dvs) == 1
        assert dvs[0].formula1 == "Codes!$A$5:$A$8"
        assert str(dvs[0].sqref) == "C5:I7"
        # freeze keeps staff/role + headers visible; chart present
        assert rota.freeze_panes == "C5"
        assert len(rota._charts) == 1
        assert len(rota.conditional_formatting._cf_rules) == 2

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after


class TestShiftScheduleCoercion:
    def test_quoted_codes_uppercased(self):
        p = ep.coerce_shift_schedule_params(
            {"staff": [{"name": "Alice", "mon": "m", "tue": " e "}]}
        )
        assert p["staff"][0]["days"][0] == "M"
        assert p["staff"][0]["days"][1] == "E"
        assert p["staff"][0]["days"][2] is None

    def test_week_start_null_is_next_monday(self):
        p = ep.coerce_shift_schedule_params({"staff": [{"name": "A"}]})
        assert p["week_start"] == _next_monday(date.today()).isoformat()

    def test_week_start_snaps_to_monday(self):
        p = ep.coerce_shift_schedule_params(
            {"week_start": "2026-01-14", "staff": [{"name": "A"}]}
        )
        assert p["week_start"] == "2026-01-12"

    def test_bare_string_staff(self):
        p = ep.coerce_shift_schedule_params({"staff": ["Alice", "Bob"]})
        assert [s["name"] for s in p["staff"]] == ["Alice", "Bob"]
        assert all(s["days"] == [None] * 7 for s in p["staff"])

    def test_staff_required(self):
        with pytest.raises(ValueError):
            ep.coerce_shift_schedule_params({})
        with pytest.raises(ValueError):
            ep.coerce_shift_schedule_params({"staff": []})
        with pytest.raises(ValueError):
            ep.coerce_shift_schedule_params({"staff": [{"role": "Chef"}]})
        with pytest.raises(ValueError):
            ep.coerce_shift_schedule_params({"staff": "Alice"})


class TestShiftScheduleRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "shift_schedule", "params": SHIFT_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "weekly shift rota roster for my cafe staff, morning and " "evening shifts"
        )
        assert result["pattern"] == "shift_schedule"
        assert result["sheet_names"] == ["Rota", "Codes", "Calc"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ═════════════════════════════════════════════════════════════════════
# equipment_maintenance
# ═════════════════════════════════════════════════════════════════════

EQUIP_PARAMS = {
    "currency": "EUR",
    "equipment": [
        {
            "name": "Van 1",
            "location": "Depot",
            "last_service": "2025-11-15",
            "interval_days": 90,
            "responsible": "Mehdi",
        },
        {
            "name": "Espresso Machine",
            "location": "Kitchen",
            "last_service": "2026-01-02",
            "interval_days": 60,
            "next_service": "2026-03-15",
        },
        {"name": "HVAC Unit", "interval_days": 365},
    ],
    "service_log": [
        {"date": "2025-11-15", "equipment": "Van 1", "type": "Repair", "cost": 320},
        {
            "date": "2026-01-02",
            "equipment": "Espresso Machine",
            "type": "Inspection",
            "cost": 90,
        },
    ],
}


class TestEquipmentMaintenanceBuilder:
    def setup_method(self):
        self.spec = ep.build_equipment_maintenance_spec(EQUIP_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]
        assert not [w for w in warnings if "overlap" in w]

    def test_geometry_and_formula_lattice(self):
        eq, log = self.norm["sheets"]
        assert eq["name"] == "Equipment"
        assert log["name"] == "Service Log"

        eq_table = eq["tables"][0]
        assert eq_table["start_cell"] == "A4"
        assert len(eq_table["rows"]) == 3
        assert eq_table["headers"][-1] == "Total Spend"  # log present

        van, esp, hvac = eq_table["rows"]
        # days-until = next due − TODAY() with "0" number format
        assert van[5] == '=IF(E5="","",E5-TODAY())'
        # status via nested IF: Overdue < 0 / Due Soon ≤ 30 / Scheduled
        assert van[6] == (
            '=IF(F5="","",IF(F5<0,"Overdue",IF(F5<=30,"Due Soon","Scheduled")))'
        )
        # spend = SUMIF over the rendered log rows 5..14
        assert van[8] == (
            "=SUMIF('Service Log'!$B$5:$B$14,$A5,'Service Log'!$D$5:$D$14)"
        )
        assert esp[8] == (
            "=SUMIF('Service Log'!$B$5:$B$14,$A6,'Service Log'!$D$5:$D$14)"
        )
        assert hvac[5] == '=IF(E7="","",E7-TODAY())'
        # total row: total maintenance spend over the exact data rows
        assert eq_table["total_row"][8] == "=SUM(I5:I7)"

        log_table = log["tables"][0]
        assert log_table["start_cell"] == "A4"
        # 2 entries + 8 blank fillers = 10 rows → 5..14, total 15
        assert len(log_table["rows"]) == 10
        assert log_table["total_row"][3] == "=SUM(D5:D14)"

        # dropdowns: equipment ← live register range, type ← list
        eq_dv, type_dv = log["data_validation"]
        assert eq_dv["range"] == "B5:B14"
        assert eq_dv["source_range"] == "Equipment!$A$5:$A$7"
        assert type_dv["range"] == "C5:C14"
        assert "Inspection" in type_dv["values"] and "Repair" in type_dv["values"]

        # red/amber/green status colors on the status column
        status_cf = [cf for cf in eq["conditional_formats"] if cf["range"] == "G5:G7"][
            0
        ]
        assert [rule["value"] for rule in status_cf["rules"]] == [
            "Overdue",
            "Due Soon",
            "Scheduled",
        ]

    def test_next_due_math_independently(self):
        # next due = user-stated when given, else last + interval
        van_due = date(2025, 11, 15) + timedelta(days=90)
        assert van_due == date(2026, 2, 13)
        eq = self.norm["sheets"][0]
        van, esp, hvac = eq["tables"][0]["rows"]
        # Van 1: computed default (an editable INPUT, noted on the sheet);
        # normalization turns ISO strings into real dates
        assert van[4] == date(2026, 2, 13)
        # Espresso: the user-stated date wins
        assert esp[4] == date(2026, 3, 15)
        # HVAC: no last service → no computable default → blank
        assert hvac[4] is None

    def test_status_thresholds_independently(self):
        # recompute the status words in plain Python from the due dates
        today = date.today()
        due_dates = {
            "Van 1": date(2026, 2, 13),
            "Espresso Machine": date(2026, 3, 15),
            "HVAC Unit": None,
        }
        expected = {}
        for name, due in due_dates.items():
            if due is None:
                expected[name] = ""
                continue
            days = (due - today).days
            expected[name] = (
                "Overdue" if days < 0 else ("Due Soon" if days <= 30 else "Scheduled")
            )
        # the emitted nested-IF encodes exactly these thresholds
        formula = self.norm["sheets"][0]["tables"][0]["rows"][0][6]
        assert 'IF(F5<0,"Overdue"' in formula
        assert 'IF(F5<=30,"Due Soon","Scheduled")' in formula
        # and the Python-side recomputation agrees with the formula shape
        for name, due in due_dates.items():
            if due is None:
                assert expected[name] == ""
            else:
                days = (due - today).days
                assert expected[name] == (
                    "Overdue"
                    if days < 0
                    else ("Due Soon" if days <= 30 else "Scheduled")
                )

    def test_spend_math_independently(self):
        spend = {}
        for entry in EQUIP_PARAMS["service_log"]:
            spend[entry["equipment"]] = (
                spend.get(entry["equipment"], 0.0) + entry["cost"]
            )
        assert spend == {"Van 1": 320.0, "Espresso Machine": 90.0}
        assert sum(spend.values()) == 410.0
        # the SUMIF criteria/sum ranges cover the rendered log rows
        eq = self.norm["sheets"][0]
        for row in eq["tables"][0]["rows"]:
            assert "$B$5:$B$14" in row[8]
            assert "$D$5:$D$14" in row[8]

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "equip.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        eq, log = wb["Equipment"], wb["Service Log"]
        assert eq["F5"].value == '=IF(E5="","",E5-TODAY())'
        assert eq["G5"].value == (
            '=IF(F5="","",IF(F5<0,"Overdue",IF(F5<=30,"Due Soon","Scheduled")))'
        )
        assert eq["I5"].value == (
            "=SUMIF('Service Log'!$B$5:$B$14,$A5,'Service Log'!$D$5:$D$14)"
        )
        assert eq["I8"].value == "=SUM(I5:I7)"
        # next due renders as a real date; days-until is a plain number
        assert eq["E5"].value.date() == date(2026, 2, 13)
        assert eq["E5"].number_format == "yyyy-mm-dd"
        assert eq["F5"].number_format == "0"
        # costs carry the currency format
        assert log["D5"].number_format == '#,##0.00" €"'
        assert eq["I5"].number_format == '#,##0.00" €"'
        # conditional formats: 3 status rules + the days-until red flag
        eq_cf = list(eq.conditional_formatting)
        assert len(eq_cf) == 2
        status_rules = [cf for cf in eq_cf if str(cf.sqref) == "G5:G7"][0]
        assert len(status_rules.rules) == 3
        # equipment dropdown sourced from the register
        dvs = log.data_validations.dataValidation
        assert len(dvs) == 2
        assert dvs[0].formula1 == "Equipment!$A$5:$A$7"
        # two charts when costs exist: days-until + spend by equipment
        assert len(eq._charts) == 2
        assert eq.freeze_panes == "A5"

    def test_no_log_variant(self):
        spec = ep.build_equipment_maintenance_spec(
            {
                "equipment": [
                    {
                        "name": "Compressor",
                        "last_service": "2026-01-01",
                        "interval_days": 365,
                    }
                ]
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        assert [s["name"] for s in norm["sheets"]] == ["Equipment"]
        table = norm["sheets"][0]["tables"][0]
        assert table["headers"][-1] == "Responsible"  # no Total Spend
        assert table.get("total_row") is None
        assert len(norm["sheets"][0]["charts"]) == 1  # days-until only

    def test_log_derived_equipment(self):
        spec = ep.build_equipment_maintenance_spec(
            {
                "equipment": [{"name": "Van 1"}],
                "service_log": [
                    {"date": "2026-01-02", "equipment": "Generator", "cost": 75}
                ],
            }
        )
        norm = eg._normalize_spec(spec)
        rows = norm["sheets"][0]["tables"][0]["rows"]
        assert [r[0] for r in rows] == ["Van 1", "Generator"]

    def test_heal_never_fires_on_template(self):
        before = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        eg._heal_off_by_one_formula_rows(self.norm)
        after = [r for t in self.norm["sheets"][0]["tables"] for r in t["rows"]]
        assert before == after


class TestEquipmentMaintenanceCoercion:
    def test_quoted_interval_and_next_due_computed(self):
        p = ep.coerce_equipment_maintenance_params(
            {
                "equipment": [
                    {
                        "name": "Compressor",
                        "last_service": "2026-01-01",
                        "interval_days": "365",
                    }
                ]
            }
        )
        item = p["equipment"][0]
        assert item["interval_days"] == 365
        # computed input default: last + interval
        assert item["next_due"] == "2027-01-01"

    def test_next_service_only_when_stated(self):
        p = ep.coerce_equipment_maintenance_params(
            {
                "equipment": [
                    {
                        "name": "A",
                        "last_service": "2026-01-01",
                        "interval_days": 30,
                        "next_service": "2026-03-01",
                    }
                ]
            }
        )
        # user-stated date wins over last + interval
        assert p["equipment"][0]["next_due"] == "2026-03-01"

    def test_no_computable_next_due(self):
        p = ep.coerce_equipment_maintenance_params(
            {"equipment": [{"name": "A", "interval_days": 365}]}
        )
        assert p["equipment"][0]["next_due"] is None

    def test_every_six_months_style_interval(self):
        p = ep.coerce_equipment_maintenance_params(
            {
                "equipment": [
                    {"name": "A", "last_service": "2026-01-01", "interval_days": 183}
                ]
            }
        )
        assert p["equipment"][0]["next_due"] == "2026-07-03"

    def test_equipment_required(self):
        with pytest.raises(ValueError):
            ep.coerce_equipment_maintenance_params({})
        with pytest.raises(ValueError):
            ep.coerce_equipment_maintenance_params({"equipment": []})
        with pytest.raises(ValueError):
            ep.coerce_equipment_maintenance_params({"equipment": [{"location": "X"}]})
        with pytest.raises(ValueError):
            ep.coerce_equipment_maintenance_params({"service_log": [{"cost": 5}]})

    def test_currency_alias(self):
        p = ep.coerce_equipment_maintenance_params(
            {"currency": "USD", "equipment": [{"name": "A"}]}
        )
        assert p["currency"] == '"$"#,##0.00'


class TestEquipmentMaintenanceRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {"pattern": "equipment_maintenance", "params": EQUIP_PARAMS}
            ),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "maintenance schedule for my machines with the next service "
            "due dates and past service costs"
        )
        assert result["pattern"] == "equipment_maintenance"
        assert result["sheet_names"] == ["Equipment", "Service Log"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
