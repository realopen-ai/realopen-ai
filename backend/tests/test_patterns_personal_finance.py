"""
Tests for the personal-finance pattern batch (savings_goal, net_worth,
break_even) in app/services/patterns/.

Same contract as tests/test_excel_patterns.py:
  • param coercion (quoted/currency/percent strings, alias keys,
    missing required params → ValueError)
  • builder layout math — every formula references the row the
    converter will actually render (title rows included)
  • _build_xlsx round-trip via openpyxl
  • independent Python verification of the math behind the formulas
    (months-to-goal simulation, SUMIF rollups, break-even units)
  • routing: classifier JSON → template with the AI path skipped
"""

import json
import math
import sys
from datetime import date, datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

# ── helpers ───────────────────────────────────────────────────────────


def _classify_response(payload):
    """Classifier LLM stub returning a fixed routing JSON object."""

    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


def _simulate_savings_chain(current, monthly, rate, months):
    """Independent re-implementation of the workbook's balance chain:
    opening = previous closing (first month = current savings),
    interest = opening x rate / 12, closing = opening + deposit + interest.
    Returns the list of closing balances, index 0 = month 1.
    """
    closings = []
    opening = current
    for _ in range(months):
        interest = opening * rate / 12.0
        closing = opening + monthly + interest
        closings.append(closing)
        opening = closing
    return closings


def _months_to_goal(closings, target):
    """First 1-based month whose closing balance reaches the target —
    the same semantics as the COUNTIF months-to-goal formula."""
    for m, closing in enumerate(closings, start=1):
        if closing >= target:
            return m
    return None


SAVINGS_PARAMS = dict(
    goal_name="New Car",
    target_amount=10000,
    current_saved=2000,
    monthly_contribution=250,
    annual_return=0.04,
    start_date="2026-01-01",
    currency="USD",
)

ASSETS = [
    {"category": "Cash", "item": "Checking account", "value": 4500},
    {"category": "Cash", "item": "Savings account", "value": 12000},
    {"category": "Investments", "item": "Brokerage", "value": 33000},
    {"category": "Real Estate", "item": "Apartment", "value": 210000},
]

LIABILITIES = [
    {"category": "Mortgage", "item": "Home loan", "value": 145000},
    {"category": "Credit Cards", "item": "Visa", "value": 2300},
    {"category": "Loans", "item": "Car loan", "value": 8200},
]

NET_WORTH_PARAMS = dict(
    as_of_date="2026-03-31",
    currency="USD",
    assets=ASSETS,
    liabilities=LIABILITIES,
)

BREAK_EVEN_PARAMS = dict(
    product_name="Soy Candles",
    fixed_costs=2000,
    price_per_unit=15,
    variable_cost_per_unit=6,
    target_profit=1500,
    currency="USD",
)


# ── registry / stanza contract ────────────────────────────────────────


class TestRegistry:
    def test_patterns_registered_with_stanza_names(self):
        for name in ("savings_goal", "net_worth", "break_even"):
            assert name in ep.PATTERN_BUILDERS
            assert callable(ep.PATTERN_BUILDERS[name])
            assert name in ep.PATTERN_DESCRIPTIONS

    def test_stanza_names_match_registry_keys(self):
        # PATTERN_NAME must equal the stanza marker in the classifier
        # prompt — the routing contract.
        prompt = (
            Path(__file__).resolve().parent.parent
            / "app"
            / "prompts"
            / "pattern_classifier.md"
        ).read_text(encoding="utf-8")
        for name in ("savings_goal", "net_worth", "break_even"):
            assert f"<!-- stanza: {name} -->" in prompt

    def test_keywords_registered(self):
        assert "savings goal" in ep.PATTERN_KEYWORDS["savings_goal"]
        assert "save for" in ep.PATTERN_KEYWORDS["savings_goal"]
        assert "net worth" in ep.PATTERN_KEYWORDS["net_worth"]
        assert "liabilit" in ep.PATTERN_KEYWORDS["net_worth"]
        assert "break-even" in ep.PATTERN_KEYWORDS["break_even"]
        assert "fixed cost" in ep.PATTERN_KEYWORDS["break_even"]

    def test_gate_and_shortlist(self):
        assert eg._PATTERN_GATE_RE.search("help me plan a savings goal")
        assert eg._PATTERN_GATE_RE.search("net worth for what I own and owe")
        assert eg._PATTERN_GATE_RE.search("break-even analysis for candles")
        assert "savings_goal" in eg._shortlist_patterns("savings goal for a new car")
        assert "net_worth" in eg._shortlist_patterns(
            "net worth: assets versus liabilities"
        )
        assert "break_even" in eg._shortlist_patterns(
            "break-even point with fixed costs and variable costs"
        )


# ── savings_goal: coercion ────────────────────────────────────────────


class TestSavingsGoalCoercion:
    def test_quoted_numbers_aliases_and_percent(self):
        p = ep.coerce_savings_goal_params(
            {
                "target_amount": "$10,000",
                "saved_so_far": "2,000",
                "monthly_deposit": "250",
                "annual_return": "4%",
                "goal": "New Car",
                "currency": "USD",
            }
        )
        assert p["target_amount"] == 10000.0
        assert p["current_saved"] == 2000.0
        assert p["monthly_contribution"] == 250.0
        assert p["annual_return"] == 0.04
        assert p["goal_name"] == "New Car"
        assert p["currency"] == '"$"#,##0.00'

    def test_decimal_return_string_stays_decimal(self):
        p = ep.coerce_savings_goal_params(
            {
                "target_amount": 5000,
                "monthly_contribution": 100,
                "annual_return": "0.04",
            }
        )
        assert p["annual_return"] == 0.04

    def test_documented_defaults(self):
        p = ep.coerce_savings_goal_params(
            {"target_amount": 5000, "monthly_contribution": 500}
        )
        assert p["current_saved"] == 0.0
        assert p["annual_return"] == 0.0
        assert p["goal_name"] == "Savings Goal"
        assert p["currency"] is None
        t = date.today()
        ny, nm = (t.year + 1, 1) if t.month == 12 else (t.year, t.month + 1)
        assert p["start_date"] == f"{ny:04d}-{nm:02d}-01"

    def test_missing_required_params_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_savings_goal_params({})
        with pytest.raises(ValueError):
            ep.coerce_savings_goal_params({"monthly_contribution": 100})
        with pytest.raises(ValueError):
            ep.coerce_savings_goal_params({"target_amount": 5000})
        with pytest.raises(ValueError):
            ep.coerce_savings_goal_params(
                {"target_amount": -5, "monthly_contribution": 10}
            )
        with pytest.raises(ValueError):
            ep.coerce_savings_goal_params(
                {"target_amount": 5000, "monthly_contribution": 0}
            )
        with pytest.raises(ValueError):
            ep.coerce_savings_goal_params("nope")

    def test_crazy_values_clamped_not_raised(self):
        p = ep.coerce_savings_goal_params(
            {
                "target_amount": 2e12,
                "monthly_contribution": 1e9,
                "annual_return": "150%",
            }
        )
        assert p["target_amount"] == 1e12
        assert p["monthly_contribution"] == 1e9
        assert p["annual_return"] == 1.0  # 150% clamps to 100%
        assert p["current_saved"] == 0.0  # negative → 0


# ── savings_goal: builder ─────────────────────────────────────────────


class TestSavingsGoalBuilder:
    def setup_method(self):
        self.spec = ep.build_savings_goal_spec(SAVINGS_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        # independent projection math for the fixture params
        self.months_to_goal = 30
        self.horizon = max(12, min(self.months_to_goal + 6, 360))

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry_and_formula_lattice(self):
        sheet = self.norm["sheets"][0]
        assert sheet["name"] == "Savings"
        assert sheet["tab_color"] == "16304F"
        table = sheet["tables"][0]
        # inputs rows 2-6, table at A8 (no title) → header 8, data 9..44
        assert table["start_cell"] == "A8"
        assert len(table["rows"]) == self.horizon
        first, second = table["rows"][0], table["rows"][1]
        # every ref points at the row the converter actually renders
        assert first[2] == "=$B$3"
        assert first[3] == "=$B$4"
        assert first[4] == "=C9*$B$5/12"
        assert first[5] == "=C9+D9+E9"
        assert first[6] == '=IF($B$2>0,F9/$B$2,"n/a")'
        assert first[7] == "=$B$2"
        # cumulative chain: opening = previous closing
        assert second[2] == "=F9"
        assert second[4] == "=C10*$B$5/12"
        assert second[5] == "=C10+D10+E10"
        last = table["rows"][-1]
        assert last[2] == "=F43"
        assert last[5] == "=C44+D44+E44"
        # totals: SUM over the exact data rows (placeholders expanded at
        # write time — pinned in the built-workbook test)
        assert table["total_row"][3] == "=SUM(D{first_row}:D{last_row})"
        assert table["total_row"][4] == "=SUM(E{first_row}:E{last_row})"
        # input blocks populate B2..B6 above the table
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["B2"]["text"] == 10000.0
        assert blocks["B3"]["text"] == 2000.0
        assert blocks["B4"]["text"] == 250.0
        assert blocks["B5"]["text"] == 0.04
        assert blocks["B6"]["text"] == date(2026, 1, 1)
        # goal readouts (row 45 = total row, readouts at 47-49)
        assert blocks["B47"]["text"] == (
            '=IF(COUNTIF(F9:F44,">="&$B$2)=0,"Not reached",'
            'COUNTIF(F9:F44,"<"&$B$2)+1)'
        )
        assert blocks["B48"]["text"] == ('=IF(ISNUMBER(B47),INDEX(B9:B44,B47),"—")')
        assert blocks["B49"]["text"] == '=IF($B$2>0,$B$3/$B$2,"n/a")'
        # summary table references the schedule's total-row cells
        summary = sheet["tables"][1]
        assert summary["start_cell"] == "A51"
        rows = {r[0]: r[1] for r in summary["rows"]}
        assert rows["Total Deposits"] == "=D45"
        assert rows["Total Interest Earned"] == "=E45"
        assert rows["Final Projected Balance"] == "=F44"

    def test_chart_ranges_match_layout(self):
        chart = self.norm["sheets"][0]["charts"][0]
        assert chart["type"] == "line"
        assert chart["categories_range"] == "Savings!B9:B44"
        assert chart["series"][0]["values_range"] == "Savings!F9:F44"
        assert chart["series"][1]["values_range"] == "Savings!H9:H44"

    def test_conditional_formats_and_freeze(self):
        sheet = self.norm["sheets"][0]
        cf = {c["range"]: c for c in sheet["conditional_formats"]}
        assert "F9:F44" in cf
        assert cf["F9:F44"]["rules"][0]["value"] == "F9>=$B$2"
        assert cf["F9:F44"]["rules"][0]["fill"] == "C6EFCE"
        assert cf["F9:F44"]["rules"][0]["font_color"] == "1E4620"
        assert "G9:G44" in cf
        assert cf["G9:G44"]["rules"][0]["value"] == "G9<0.5"
        assert cf["G9:G44"]["rules"][0]["fill"] == "FFF2CC"
        assert sheet["freeze_panes"] == "A9"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "savings.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        ws = wb["Savings"]
        assert ws["B2"].value == 10000
        assert ws["B5"].value == 0.04
        # openpyxl reads written dates back as datetime objects
        assert ws["B6"].value == datetime(2026, 1, 1)
        assert ws["C9"].value == "=$B$3"
        assert ws["C10"].value == "=F9"
        assert ws["E10"].value == "=C10*$B$5/12"
        assert ws["F44"].value == "=C44+D44+E44"
        # total row placeholders expanded to the real data rows
        assert ws["D45"].value == "=SUM(D9:D44)"
        assert ws["E45"].value == "=SUM(E9:E44)"
        assert ws["B47"].value.startswith("=IF(COUNTIF(F9:F44")
        assert ws["B53"].value == "=D45"
        # money / percent formats via currency_fmt / PCT_FMT
        assert ws["F9"].number_format == '"$"#,##0.00'
        assert ws["B5"].number_format == "0.00%"
        assert ws["G9"].number_format == "0.00%"
        assert len(ws._charts) == 1
        assert len(list(ws.conditional_formatting)) == 2
        assert ws.freeze_panes == "A9"

    def test_python_math_matches_formula_logic(self):
        # independent simulation of the emitted formula chain
        closings = _simulate_savings_chain(2000.0, 250.0, 0.04, 36)
        assert _months_to_goal(closings, 10000.0) == self.months_to_goal
        # the horizon must cover the goal month
        assert self.horizon >= self.months_to_goal
        # COUNTIF semantics: closings below target + 1 == first month
        below = sum(1 for c in closings if c < 10000.0)
        assert below + 1 == self.months_to_goal
        # closed-form annuity (deposit at month end) == the iteration
        i = 0.04 / 12.0
        growth = (1.0 + i) ** 36
        closed_form = 2000.0 * growth + 250.0 * (growth - 1.0) / i
        assert closings[-1] == pytest.approx(closed_form, rel=1e-12)
        # LibreOffice-verified recalc value for this fixture
        assert closings[-1] == pytest.approx(11799.9343378792, rel=1e-9)

    def test_zero_return_projection_is_straight_line(self):
        spec = ep.build_savings_goal_spec(dict(SAVINGS_PARAMS, annual_return=None))
        norm = eg._normalize_spec(spec)
        # 2000 + 250*m >= 10000 → month 32; horizon = 32 + 6
        assert len(norm["sheets"][0]["tables"][0]["rows"]) == 38
        blocks = {b["cell"]: b for b in norm["sheets"][0]["text_blocks"]}
        assert blocks["B5"]["text"] == 0.0
        # interest column still chained (computes to 0 with rate 0)
        assert norm["sheets"][0]["tables"][0]["rows"][3][4] == ("=C12*$B$5/12")

    def test_goal_already_within_reach(self):
        spec = ep.build_savings_goal_spec(dict(SAVINGS_PARAMS, current_saved=12000))
        norm = eg._normalize_spec(spec)
        # first projected month already clears the target → horizon 12
        assert len(norm["sheets"][0]["tables"][0]["rows"]) == 12


# ── savings_goal: routing ─────────────────────────────────────────────


class TestSavingsGoalRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {
                    "pattern": "savings_goal",
                    "params": {
                        "goal_name": "New Car",
                        "target_amount": 10000,
                        "current_saved": 2000,
                        "monthly_contribution": 250,
                        "annual_return": 0.04,
                    },
                }
            ),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "savings goal plan: save for a $10,000 car, already have "
            "2,000 saved, deposit 250 monthly at 4% annual return"
        )
        assert result["pattern"] == "savings_goal"
        assert result["sheet_names"] == ["Savings"]
        assert result["filename"] == "savings_goal.xlsx"
        assert result["chart_count"] == 1
        assert result["formula_count"] > 200  # 36 rows x 6 + totals + readouts
        assert "savings_goal template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ── net_worth: coercion ───────────────────────────────────────────────


class TestNetWorthCoercion:
    def test_entries_with_aliases_and_quoted_values(self):
        p = ep.coerce_net_worth_params(
            {
                "as_of_date": "2026-03-31",
                "assets": [
                    {"category": "Cash", "item": "Checking", "value": "$4,500"},
                    {"name": "Brokerage", "amount": "33000"},
                ],
                "debts": [
                    {"category": "Mortgage", "item": "Home loan", "value": 145000},
                    ["Car loan", 8200],
                ],
            }
        )
        assert p["as_of_date"] == "2026-03-31"
        assert p["assets"] == [
            ("Cash", "Checking", 4500.0),
            ("Other", "Brokerage", 33000.0),  # no category → Other
        ]
        assert p["liabilities"] == [
            ("Mortgage", "Home loan", 145000.0),
            ("Other", "Car loan", 8200.0),  # bare [item, value] list
        ]

    def test_entries_without_item_or_value(self):
        p = ep.coerce_net_worth_params(
            {
                "assets": [
                    {"category": "Cash", "value": 100},  # no item → skipped
                    {"item": "Wallet", "value": None},  # value → 0
                    {"item": "  ", "value": 5},  # blank item → skipped
                ]
            }
        )
        assert p["assets"] == [("Other", "Wallet", 0.0)]

    def test_one_side_only_is_valid(self):
        p = ep.coerce_net_worth_params({"assets": [{"item": "Cash", "value": 5}]})
        assert p["liabilities"] == []
        p2 = ep.coerce_net_worth_params({"debts": [{"item": "Loan", "value": 5}]})
        assert p2["assets"] == []

    def test_both_sides_empty_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_net_worth_params({})
        with pytest.raises(ValueError):
            ep.coerce_net_worth_params({"assets": [], "liabilities": []})
        with pytest.raises(ValueError):
            ep.coerce_net_worth_params("nope")

    def test_invalid_as_of_date_falls_back_to_none(self):
        p = ep.coerce_net_worth_params(
            {"assets": [{"item": "Cash", "value": 5}], "as_of_date": "March 2026"}
        )
        assert p["as_of_date"] is None  # → live =TODAY() in the workbook


# ── net_worth: builder ────────────────────────────────────────────────


class TestNetWorthBuilder:
    def setup_method(self):
        self.spec = ep.build_net_worth_spec(NET_WORTH_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_main_sheet_layout_and_totals(self):
        sheet = self.norm["sheets"][0]
        assert sheet["name"] == "NetWorth"
        assert sheet["tab_color"] == "16304F"
        assert sheet["no_freeze"] is True
        assets_t, liab_t, summary_t = sheet["tables"]
        # Assets anchored at row 4 with title → data 6-9, total 10
        assert assets_t["start_cell"] == "A4"
        assert assets_t["title"] == "Assets"
        assert len(assets_t["rows"]) == 4
        assert assets_t["total_row"] == [
            "Total Assets",
            "",
            "=SUM(C{first_row}:C{last_row})",
        ]
        # Liabilities anchored at row 12 → data 14-16, total 17
        assert liab_t["start_cell"] == "A12"
        assert len(liab_t["rows"]) == 3
        # Summary anchored at 19 → data 21-24 referencing the totals
        assert summary_t["start_cell"] == "A19"
        rows = {r[0]: r[1] for r in summary_t["rows"]}
        assert rows["Total Assets"] == "=C10"
        assert rows["Total Liabilities"] == "=C17"
        assert rows["Net Worth"] == "=C10-C17"
        assert rows["Debt-to-Asset Ratio"] == ('=IF(C10>0,TEXT(C17/C10,"0.0%"),"n/a")')
        # as-of date block
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["A1"]["text"] == "Net Worth Statement"
        assert blocks["B2"]["text"] == date(2026, 3, 31)
        assert blocks["B2"]["number_format"] == "yyyy-mm-dd"

    def test_missing_as_of_date_is_live_today_formula(self):
        spec = ep.build_net_worth_spec(dict(NET_WORTH_PARAMS, as_of_date=None))
        blocks = {
            b["cell"]: b for b in eg._normalize_spec(spec)["sheets"][0]["text_blocks"]
        }
        assert blocks["B2"]["text"] == "=TODAY()"
        assert blocks["B2"]["number_format"] == "yyyy-mm-dd"

    def test_breakdown_sumif_lattice(self):
        bd = self.norm["sheets"][1]
        assert bd["name"] == "Breakdown"
        assert bd["tab_color"] == "1B3A5C"
        assets_cat, liab_cat = bd["tables"]
        # asset categories anchored at 3 with title → data 5-7, total 8
        assert assets_cat["start_cell"] == "A3"
        assert assets_cat["rows"][0] == [
            "Cash",
            "=SUMIF(NetWorth!$A$6:$A$9,$A5,NetWorth!$C$6:$C$9)",
            '=IF($B$8>0,B5/$B$8,"n/a")',
        ]
        assert assets_cat["rows"][2] == [
            "Real Estate",
            "=SUMIF(NetWorth!$A$6:$A$9,$A7,NetWorth!$C$6:$C$9)",
            '=IF($B$8>0,B7/$B$8,"n/a")',
        ]
        assert assets_cat["total_row"][1] == "=SUM(B{first_row}:B{last_row})"
        # liability categories anchored at 10 → data 12-14, total 15
        assert liab_cat["start_cell"] == "A10"
        assert liab_cat["rows"][0] == [
            "Mortgage",
            "=SUMIF(NetWorth!$A$14:$A$16,$A12,NetWorth!$C$14:$C$16)",
            '=IF($B$15>0,B12/$B$15,"n/a")',
        ]
        # pie for the asset mix, horizontal bars for liabilities
        pie, bars = bd["charts"]
        assert pie["type"] == "pie"
        assert pie["categories_range"] == "Breakdown!A5:A7"
        assert pie["series"][0]["values_range"] == "Breakdown!B5:B7"
        assert bars["type"] == "bar_h"
        assert bars["categories_range"] == "Breakdown!A12:A14"
        assert bars["series"][0]["values_range"] == "Breakdown!B12:B14"

    def test_chart_conditional_format_and_dropdowns(self):
        sheet = self.norm["sheets"][0]
        chart = sheet["charts"][0]
        assert chart["type"] == "bar"
        assert chart["categories_range"] == "NetWorth!A21:A23"
        assert chart["series"][0]["values_range"] == "NetWorth!B21:B23"
        # net worth cell: red when negative, green when >= 0
        cf = {c["range"]: c for c in sheet["conditional_formats"]}
        assert "B23:B23" in cf
        rules = cf["B23:B23"]["rules"]
        assert rules[0]["operator"] == "less_than"
        assert rules[0]["value"] == 0
        assert rules[0]["fill"] == "FFC7CE"
        assert rules[1]["operator"] == "greater_than_or_equal"
        assert rules[1]["fill"] == "C6EFCE"
        # category dropdowns sourced from the Breakdown lists
        dvs = {dv["range"]: dv for dv in sheet["data_validation"]}
        assert dvs["A6:A9"]["source_range"] == "Breakdown!$A$5:$A$7"
        assert dvs["A14:A16"]["source_range"] == "Breakdown!$A$12:$A$14"
        assert dvs["A6:A9"]["error_style"] == "warning"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "networth.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["NetWorth", "Breakdown"]
        ws = wb["NetWorth"]
        assert ws["A6"].value == "Cash"
        assert ws["C6"].value == 4500
        assert ws["C10"].value == "=SUM(C6:C9)"
        assert ws["C17"].value == "=SUM(C14:C16)"
        assert ws["B21"].value == "=C10"
        assert ws["B23"].value == "=C10-C17"
        assert ws["C6"].number_format == '"$"#,##0.00'
        assert len(ws._charts) == 1
        assert len(ws.data_validations.dataValidation) == 2
        assert ws.data_validations.dataValidation[0].formula1 == "Breakdown!$A$5:$A$7"
        assert len(list(ws.conditional_formatting)) == 1
        bd = wb["Breakdown"]
        assert bd["B5"].value == ("=SUMIF(NetWorth!$A$6:$A$9,$A5,NetWorth!$C$6:$C$9)")
        assert bd["B8"].value == "=SUM(B5:B7)"
        assert len(bd._charts) == 2

    def test_python_math_matches_formulas(self):
        # totals the SUM/SUMIF formulas will compute
        total_assets = sum(a["value"] for a in ASSETS)
        total_liabilities = sum(ll["value"] for ll in LIABILITIES)
        assert total_assets == 259500
        assert total_liabilities == 155500
        assert total_assets - total_liabilities == 104000
        # category rollups (the SUMIF semantics)
        cash = sum(a["value"] for a in ASSETS if a["category"] == "Cash")
        assert cash == 16500
        mortgage = sum(ll["value"] for ll in LIABILITIES if ll["category"] == "Mortgage")
        assert mortgage == 145000
        # debt ratio guard semantics
        ratio = total_liabilities / total_assets
        assert ratio == pytest.approx(155500 / 259500, rel=1e-9)
        assert ratio == pytest.approx(0.599229, abs=1e-6)

    def test_assets_only_statement(self):
        spec = ep.build_net_worth_spec({"as_of_date": "2026-03-31", "assets": ASSETS})
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        # Assets + Summary tables only, italic placeholder for the rest
        assert [t.get("title") for t in sheet["tables"]] == [
            "Assets",
            "Net Worth Summary",
        ]
        rows = {r[0]: r[1] for r in sheet["tables"][1]["rows"]}
        assert rows["Total Liabilities"] == "=0"
        assert rows["Net Worth"] == "=C10-0"
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["A12"]["text"] == "No liabilities provided"
        # breakdown: one table + the asset-mix pie (3 categories), no
        # liability chart; the asset dropdown is still live
        bd = norm["sheets"][1]
        assert len(bd["tables"]) == 1
        assert [c["type"] for c in bd.get("charts", [])] == ["pie"]
        dvs = {dv["range"]: dv for dv in sheet["data_validation"]}
        assert "A6:A9" in dvs
        assert "A14:A16" not in dvs


# ── net_worth: routing ────────────────────────────────────────────────


class TestNetWorthRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {
                    "pattern": "net_worth",
                    "params": {
                        "as_of_date": "2026-03-31",
                        "assets": [
                            {"category": "Cash", "item": "Checking", "value": 4500},
                            {
                                "category": "Real Estate",
                                "item": "Apartment",
                                "value": 210000,
                            },
                        ],
                        "liabilities": [
                            {
                                "category": "Mortgage",
                                "item": "Home loan",
                                "value": 145000,
                            }
                        ],
                    },
                }
            ),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "net worth statement showing what I own and what I owe"
        )
        assert result["pattern"] == "net_worth"
        assert result["sheet_names"] == ["NetWorth", "Breakdown"]
        assert result["filename"] == "net_worth.xlsx"
        # headline bar + asset-mix pie (2 asset categories); the single
        # liability category gets no bar chart
        assert result["chart_count"] == 2
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()
        assert "net_worth template" in result["summary"]


# ── break_even: coercion ──────────────────────────────────────────────


class TestBreakEvenCoercion:
    def test_aliases_and_quoted_numbers(self):
        p = ep.coerce_break_even_params(
            {"fixed": "2,000", "price": 15, "unit_cost": "6", "product": "Candles"}
        )
        assert p["fixed_costs"] == 2000.0
        assert p["price_per_unit"] == 15.0
        assert p["variable_cost_per_unit"] == 6.0
        assert p["product_name"] == "Candles"
        assert p["target_profit"] is None

    def test_missing_core_numbers_raise(self):
        base = {"fixed_costs": 2000, "price_per_unit": 15, "variable_cost_per_unit": 6}
        for key in base:
            bad = dict(base)
            del bad[key]
            with pytest.raises(ValueError):
                ep.coerce_break_even_params(bad)
        with pytest.raises(ValueError):
            ep.coerce_break_even_params({})
        with pytest.raises(ValueError):
            ep.coerce_break_even_params("nope")

    def test_invalid_economics_raise(self):
        # price must exceed variable cost — no break-even exists otherwise
        with pytest.raises(ValueError):
            ep.coerce_break_even_params(
                {"fixed_costs": 100, "price_per_unit": 5, "variable_cost_per_unit": 5}
            )
        with pytest.raises(ValueError):
            ep.coerce_break_even_params(
                {"fixed_costs": 100, "price_per_unit": 4, "variable_cost_per_unit": 5}
            )
        with pytest.raises(ValueError):
            ep.coerce_break_even_params(
                {"fixed_costs": 100, "price_per_unit": 0, "variable_cost_per_unit": 5}
            )
        with pytest.raises(ValueError):
            ep.coerce_break_even_params(
                {"fixed_costs": -1, "price_per_unit": 15, "variable_cost_per_unit": 6}
            )
        with pytest.raises(ValueError):
            ep.coerce_break_even_params(
                {"fixed_costs": 100, "price_per_unit": 15, "variable_cost_per_unit": -6}
            )

    def test_target_profit_only_when_positive(self):
        base = {"fixed_costs": 2000, "price_per_unit": 15, "variable_cost_per_unit": 6}
        assert ep.coerce_break_even_params(base)["target_profit"] is None
        assert (
            ep.coerce_break_even_params(dict(base, target_profit=0))["target_profit"]
            is None
        )
        assert (
            ep.coerce_break_even_params(dict(base, target_profit=-50))["target_profit"]
            is None
        )
        assert (
            ep.coerce_break_even_params(dict(base, target_profit=1500))["target_profit"]
            == 1500.0
        )

    def test_product_name_default(self):
        p = ep.coerce_break_even_params(
            {"fixed_costs": 100, "price_per_unit": 5, "variable_cost_per_unit": 2}
        )
        assert p["product_name"] == "Product"


# ── break_even: builder ───────────────────────────────────────────────


class TestBreakEvenBuilder:
    def setup_method(self):
        self.spec = ep.build_break_even_spec(BREAK_EVEN_PARAMS)
        self.norm = eg._normalize_spec(self.spec)
        # break-even units for the fixture: 2000 / (15 - 6)
        self.be_units = 2000.0 / 9.0

    def _expected_units(self):
        lo = max(0, int(self.be_units * 0.5))
        hi = int(self.be_units * 1.5) + 1
        if hi - lo < 10:
            hi = lo + 10
        step = math.ceil((hi - lo + 1) / 20)
        return list(range(lo, hi + 1, step))

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert not [w for w in warnings if "above the table" in w]

    def test_geometry_and_formula_lattice(self):
        sheet = self.norm["sheets"][0]
        assert sheet["name"] == "BreakEven"
        assert sheet["tab_color"] == "16304F"
        assert sheet["no_freeze"] is True
        # inputs B2-B5 (target profit stated → row 5 present)
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["A1"]["text"] == "Break-Even Analysis — Soy Candles"
        assert blocks["B2"]["text"] == 2000.0
        assert blocks["B3"]["text"] == 15.0
        assert blocks["B4"]["text"] == 6.0
        assert blocks["B5"]["text"] == 1500.0
        # results section: header row 7, values rows 8-13 (blocks, so
        # each keeps its own number format)
        assert blocks["A7"]["text"] == "Break-Even Results"
        assert blocks["B8"]["text"] == "=B3-B4"
        assert blocks["B9"]["text"] == '=IF(B3>0,(B3-B4)/B3,"n/a")'
        assert blocks["B10"]["text"] == '=IF((B3-B4)>0,B2/(B3-B4),"n/a")'
        assert blocks["B11"]["text"] == '=IF(ISNUMBER(B10),B10*B3,"n/a")'
        assert blocks["B12"]["text"] == ('=IF((B3-B4)>0,(B2+B5)/(B3-B4),"n/a")')
        assert blocks["B13"]["text"] == '=IF(ISNUMBER(B12),B12*B3,"n/a")'
        assert blocks["B10"]["number_format"] == "#,##0.##"
        assert blocks["B9"]["number_format"] == "0.00%"
        # sensitivity table: title 15, header 16, data 17-35
        table = sheet["tables"][0]
        assert table["start_cell"] == "A15"
        assert table["title"] == "Sensitivity Analysis"
        units = self._expected_units()
        assert units[0] == 111 and units[-1] == 327
        assert len(table["rows"]) == len(units)
        first, second = table["rows"][0], table["rows"][1]
        assert first == [111, "=$B$2", "=A17*$B$4", "=B17+C17", "=A17*$B$3", "=E17-D17"]
        assert second[2] == "=A18*$B$4"
        assert second[5] == "=E18-D18"
        last = table["rows"][-1]
        assert last[0] == 327
        assert last[5] == "=E35-D35"

    def test_chart_and_conditional_formats(self):
        sheet = self.norm["sheets"][0]
        chart = sheet["charts"][0]
        assert chart["type"] == "line"
        assert chart["categories_range"] == "BreakEven!A17:A35"
        assert [s["values_range"] for s in chart["series"]] == [
            "BreakEven!E17:E35",
            "BreakEven!D17:D35",
            "BreakEven!F17:F35",
        ]
        cf = {c["range"]: c for c in sheet["conditional_formats"]}
        assert "F17:F35" in cf
        rules = cf["F17:F35"]["rules"]
        assert rules[0]["operator"] == "less_than"
        assert rules[0]["value"] == 0
        assert rules[0]["fill"] == "FFC7CE"
        assert rules[1]["operator"] == "greater_than_or_equal"
        assert rules[1]["value"] == 0
        assert rules[1]["fill"] == "C6EFCE"

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "breakeven.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        ws = wb["BreakEven"]
        assert ws["B2"].value == 2000
        assert ws["B5"].value == 1500
        assert ws["B10"].value == '=IF((B3-B4)>0,B2/(B3-B4),"n/a")'
        assert ws["A17"].value == 111
        assert ws["C17"].value == "=A17*$B$4"
        assert ws["F17"].value == "=E17-D17"
        assert ws["F35"].value == "=E35-D35"
        assert ws["C17"].number_format == '"$"#,##0.00'
        assert len(ws._charts) == 1
        assert len(list(ws.conditional_formatting)) == 1
        assert ws.freeze_panes is None  # single-screen analysis

    def test_without_target_profit_compact_layout(self):
        spec = ep.build_break_even_spec(
            {
                "fixed_costs": 2000,
                "price_per_unit": 15,
                "variable_cost_per_unit": 6,
            }
        )
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert "B5" not in blocks  # no target-profit input row
        # results header 6, values 7-10, sensitivity at 12 → data 14-32
        assert blocks["A6"]["text"] == "Break-Even Results"
        assert blocks["B9"]["text"] == '=IF((B3-B4)>0,B2/(B3-B4),"n/a")'
        assert blocks["B10"]["text"] == '=IF(ISNUMBER(B9),B9*B3,"n/a")'
        table = sheet["tables"][0]
        assert table["start_cell"] == "A12"
        assert len(table["rows"]) == 19
        assert table["rows"][0][2] == "=A14*$B$4"

    def test_python_math_matches_formulas(self):
        # break-even units = fixed / (price - variable)
        assert self.be_units == pytest.approx(222.2222, abs=1e-3)
        # target-profit units = (fixed + target) / contribution margin
        target_units = (2000.0 + 1500.0) / 9.0
        assert target_units == pytest.approx(388.8889, abs=1e-3)
        # the sensitivity rows bracket the break-even point
        units = self._expected_units()
        assert units[0] <= self.be_units <= units[-1]

        # profit(u) = u*(price - variable) - fixed → sign flips at BE
        def profit(u):
            return u * (15.0 - 6.0) - 2000.0

        assert profit(222) < 0 < profit(223)
        assert profit(units[0]) == 111 * 9 - 2000  # -1001, LibreOffice-verified
        # revenue and cost lines cross exactly at the break-even units
        assert profit(self.be_units) == pytest.approx(0.0, abs=1e-9)

    def test_tiny_numbers_edge(self):
        spec = ep.build_break_even_spec(
            {"fixed_costs": 10, "price_per_unit": 5, "variable_cost_per_unit": 3}
        )
        norm = eg._normalize_spec(spec)
        rows = norm["sheets"][0]["tables"][0]["rows"]
        # be = 5 → lo 2, hi forced to 12 → 11 unit rows starting at 2
        assert len(rows) == 11
        assert rows[0][0] == 2
        assert rows[-1][0] == 12
