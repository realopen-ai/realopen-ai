"""Artifact SQL/API, rendering, exact reads, OCR and scope regressions."""

import io
import json
import shutil
import uuid
from contextlib import asynccontextmanager
from datetime import date
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from app.db.models import Artifact, ArtifactVersion, Document
from app.services import artifacts as service, ocr, document_extraction
from app.services.artifact_sources import (
    sections_for,
    capture_source,
    workbook_json,
    workbook_spec,
)
from app.services.artifact_summary import summarize_sections
from app.services.artifact_refs import ArtifactReference, validate_reference
from app.api import artifacts as api
from test_flashcards import AsyncAdapter


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setenv("REALOPEN_DATA_DIR", str(tmp_path))
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE conversations (id UUID PRIMARY KEY)"))
        connection.execute(
            text(
                "CREATE TABLE messages (id UUID PRIMARY KEY, conversation_id UUID, deliverables JSON, created_at DATETIME)"
            )
        )
    for model in (Document, Artifact, ArtifactVersion):
        model.__table__.create(engine)

    @asynccontextmanager
    async def factory():
        with Session(engine, expire_on_commit=False) as db:
            yield AsyncAdapter(db)

    monkeypatch.setattr(api, "async_session_factory", factory)
    from app.agent.tools import artifacts as tool_module

    monkeypatch.setattr(tool_module, "async_session_factory", factory)
    app = FastAPI()
    app.include_router(api.router, prefix="/api")
    yield SimpleNamespace(engine=engine, root=tmp_path, client=TestClient(app), factory=factory)
    engine.dispose()


async def generated(
    store,
    kind="report",
    source="# Original\n\n## Detail\nKeep this.\n",
    conversation_id=None,
):
    identifier = uuid.uuid4()
    directory = store.root / "reports"
    directory.mkdir(exist_ok=True)
    fmt = {"report": "docx", "presentation": "pptx", "excel": "xlsx"}[kind]
    if kind == "report":
        from app.services.report_gen import _generate_docx

        _generate_docx(source, directory / f"{identifier}.docx", "Original")
    else:
        (directory / f"{identifier}.{fmt}").write_bytes(b"original")
    capture_source(
        directory,
        str(identifier),
        source,
        kind,
        {"format": fmt, "topic": "Original", "template": "modern"},
    )
    async with store.factory() as db:
        result = await service.register_generated(
            db,
            {
                "report_id": str(identifier),
                "filename": f"Original.{fmt}",
                "format": fmt,
                "file_path": f"reports/{identifier}.{fmt}",
            },
            conversation_id,
        )
        await db.commit()
    return result


@pytest.mark.asyncio
async def test_presentation_targeted_edit_and_historical_export(store):
    from pptx import Presentation

    artifact = await generated(store, "presentation", "# First\nOne.\n\n---\n\n# Second\nTwo.")
    url = f"/api/artifacts/{artifact['id']}"
    response = store.client.put(
        url,
        json={
            "expected_version": 1,
            "changes": [{"section_id": "slide-2", "content": "# Second\nUpdated."}],
        },
    )
    assert response.status_code == 200
    output = store.client.get(url + "/download/pptx?version=2")
    presentation = Presentation(io.BytesIO(output.content))
    assert len(presentation.slides) == 2
    text_content = " ".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Updated." in text_content and "First" in text_content
    assert (
        store.client.get(url + "/read?version=1&section_id=slide-2")
        .json()["content"]
        .endswith("Two.")
    )


@pytest.mark.asyncio
async def test_existing_outputs_without_source_are_read_only(store):
    from app.services.report_gen import _generate_docx

    identifier = uuid.uuid4()
    directory = store.root / "reports"
    directory.mkdir()
    _generate_docx(
        "# Historical report\nPreserved facts.",
        directory / f"{identifier}.docx",
        "Historical",
    )
    with store.engine.begin() as connection:
        connection.execute(
            text(
                "INSERT INTO messages (id, deliverables, created_at) VALUES (:id, :files, CURRENT_TIMESTAMP)"
            ),
            {
                "id": uuid.uuid4().hex,
                "files": json.dumps(
                    [
                        {
                            "report_id": str(identifier),
                            "format": "docx",
                            "filename": "Historical.docx",
                        }
                    ]
                ),
            },
        )
    response = store.client.get(f"/api/artifacts/{identifier}")
    assert response.status_code == 200
    assert response.json()["kind"] == "legacy"
    assert not response.json()["editable"]
    section = response.json()["outline"]["sections"][0]["id"]
    assert (
        "Preserved facts."
        in store.client.get(f"/api/artifacts/{identifier}/read?section_id={section}").json()[
            "content"
        ]
    )
    async with store.factory() as db:
        with pytest.raises(HTTPException) as error:
            await service.get_artifact(db, identifier, scoped=True)
        assert error.value.status_code == 403


@pytest.mark.asyncio
async def test_source_capture_and_targeted_edits_render_new_version(store):
    artifact = await generated(store)
    identifier = artifact["id"]
    detail = store.client.get(f"/api/artifacts/{identifier}").json()
    assert detail["editable"] and len(detail["outline"]["sections"]) == 2
    outline = store.client.get(f"/api/artifacts/{identifier}/read").json()
    assert "content" not in outline["sections"][0]
    response = store.client.put(
        f"/api/artifacts/{identifier}",
        json={
            "expected_version": 1,
            "changes": [{"section_id": "section-1", "content": "# Updated\n\n"}],
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["version"] == 2
    old = store.client.get(
        f"/api/artifacts/{identifier}/read?version=1&section_id=section-1"
    ).json()
    new = store.client.get(
        f"/api/artifacts/{identifier}/read?version=2&section_id=section-1"
    ).json()
    assert "Original" in old["content"] and "Updated" in new["content"]
    retained = store.client.get(
        f"/api/artifacts/{identifier}/read?version=2&section_id=section-2"
    ).json()
    assert "Keep this" in retained["content"]
    output = store.client.get(f"/api/artifacts/{identifier}/download/docx?version=2")
    assert output.status_code == 200
    from docx import Document as Docx

    assert "Updated" in "\n".join(p.text for p in Docx(io.BytesIO(output.content)).paragraphs)
    assert (
        store.client.put(
            f"/api/artifacts/{identifier}",
            json={
                "expected_version": 1,
                "changes": [{"section_id": "section-1", "content": "# Stale"}],
            },
        ).status_code
        == 409
    )
    restored = store.client.post(
        f"/api/artifacts/{identifier}/restore",
        json={"expected_version": 2, "version": 1},
    )
    assert restored.status_code == 200 and restored.json()["version"] == 3
    detail = store.client.get(f"/api/artifacts/{identifier}").json()
    assert [v["number"] for v in detail["versions"]] == [3, 2, 1]
    assert detail["versions"][0]["restored_from"] == 1


@pytest.mark.asyncio
async def test_failed_render_does_not_publish_or_modify_previous_file(store, monkeypatch):
    artifact = await generated(store)

    async def fail(*args, **kwargs):
        raise HTTPException(503, "Renderer unavailable")

    monkeypatch.setattr(service, "render", fail)
    path = f"/api/artifacts/{artifact['id']}"
    response = store.client.put(
        path,
        json={
            "expected_version": 1,
            "changes": [{"section_id": "section-1", "content": "# New"}],
        },
    )
    assert response.status_code == 503
    assert store.client.get(path).json()["version"] == 1
    assert len(store.client.get(path).json()["versions"]) == 1
    assert store.client.get(path + "/download/docx?version=1").status_code == 200


@pytest.mark.asyncio
async def test_publish_compare_and_swap_cleans_conflicting_candidate(store, monkeypatch):
    artifact = await generated(store)
    candidate = store.root / "reports/candidate.docx"

    async def render(*args, **kwargs):
        candidate.write_bytes(b"candidate")
        return "reports/candidate.docx"

    monkeypatch.setattr(service, "render", render)
    async with store.factory() as db:
        current = await service.get_artifact(db, uuid.UUID(artifact["id"]))
        await db.execute(
            __import__("sqlalchemy")
            .update(Artifact)
            .where(Artifact.id == current.id)
            .values(current_version=2)
            .execution_options(synchronize_session=False)
        )
        with pytest.raises(HTTPException) as error:
            await service.publish(db, current, 1, "# Changed")
        assert error.value.status_code == 409 and not candidate.exists()


def test_exact_reads_are_bounded_and_pagination_is_lossless():
    source = "# Big\n" + "word " * 6000
    version = SimpleNamespace(
        source=source,
        sections=sections_for(source, "report"),
        settings={},
        artifact_id=uuid.uuid4(),
        number=1,
    )
    first = service.read(version, "section-1")
    assert len(first["content"]) == 12000 and first["next_offset"] == 12000
    chunks = [first["content"]]
    while first["next_offset"] is not None:
        first = service.read(version, "section-1", first["next_offset"])
        chunks.append(first["content"])
    assert "".join(chunks) == source
    with pytest.raises(HTTPException):
        service.read(version, "invented")


def test_grounded_references_round_trip_display_metadata():
    reference = ArtifactReference.model_validate(
        {
            "artifact_id": str(uuid.uuid4()),
            "version": 2,
            "section_id": "part-4",
            "page": 4,
            "slide": None,
            "sheet": None,
        }
    )
    assert reference.page == 4 and reference.section_id == "part-4"


def test_chat_persistence_preserves_artifact_identity_and_revision():
    from app.services.blocks import BlockBuilder

    builder = BlockBuilder()
    builder.on_tool_call_update(
        "call",
        {
            "genResults": [
                {
                    "type": "report",
                    "report_id": "report",
                    "artifact_id": "artifact",
                    "version": 2,
                }
            ]
        },
    )
    assert builder.deliverables[0]["artifact_id"] == "artifact"
    assert builder.deliverables[0]["version"] == 2


def test_multi_section_edits_and_invalid_slide_changes():
    source = "# A\nSame\n## B\nSame\n## C\nKeep"
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    changes = [
        {"section_id": "section-2", "content": "## B\nNew\n"},
        {"section_id": "section-1", "content": "# A\nChanged\n"},
    ]
    assert (
        service.edited_source(version, "report", changes)
        == "# A\nChanged\n\n## B\nNew\n\n## C\nKeep"
    )
    with pytest.raises(HTTPException):
        service.edited_source(version, "report", [changes[0], changes[0]])
    version = SimpleNamespace(
        source="# A\n---\n# B", sections=sections_for("# A\n---\n# B", "presentation")
    )
    with pytest.raises(HTTPException):
        service.edited_source(
            version,
            "presentation",
            [{"section_id": "slide-1", "content": "# A\n---\n# Extra"}],
        )


@pytest.mark.asyncio
async def test_sheet_edit_preserves_formulas_and_typed_dates(store):
    spec = {
        "sheets": [
            {
                "name": "Overview",
                "tables": [
                    {
                        "start_cell": "A1",
                        "headers": ["Label", "Date", "Value"],
                        "rows": [["Old", date(2026, 10, 7), "=1+2"]],
                    }
                ],
            }
        ]
    }
    source = workbook_json(spec)
    assert workbook_spec(source) == spec
    artifact = await generated(store, "excel", source)
    sheet = json.loads(source)["sheets"][0]
    sheet["tables"][0]["rows"][0][0] = "New"
    response = store.client.put(
        f"/api/artifacts/{artifact['id']}",
        json={
            "expected_version": 1,
            "changes": [{"section_id": "sheet-1", "content": json.dumps(sheet)}],
        },
    )
    assert response.status_code == 200, response.text
    from openpyxl import load_workbook

    output = store.client.get(f"/api/artifacts/{artifact['id']}/download/xlsx?version=2")
    ws = load_workbook(io.BytesIO(output.content))["Overview"]
    assert ws["A2"].value == "New" and ws["C2"].value == "=1+2"
    assert ws["B2"].value.date() == date(2026, 10, 7)


@pytest.mark.asyncio
async def test_upload_reads_work_without_embeddings_and_are_read_only(store):
    identifier = uuid.uuid4()
    directory = store.root / f"documents/{identifier}"
    directory.mkdir(parents=True)
    (directory / "notes.md").write_text("# Upload\nExact uploaded text", encoding="utf-8")
    with Session(store.engine) as db:
        db.add(
            Document(
                id=identifier,
                filename="notes.md",
                original_filename="notes.md",
                file_path=f"documents/{identifier}/notes.md",
                scope="public",
                digestion_status="failed",
            )
        )
        db.commit()
    detail = store.client.get(f"/api/artifacts/{identifier}")
    assert detail.status_code == 200, detail.text
    assert detail.json()["editable"] is False
    result = store.client.get(f"/api/artifacts/{identifier}/read?section_id=part-1").json()
    assert "Exact uploaded text" in result["content"]
    assert (directory / "extraction.json").is_file()
    assert store.client.put(
        f"/api/artifacts/{identifier}",
        json={
            "expected_version": 1,
            "changes": [{"section_id": "part-1", "content": "changed"}],
        },
    ).status_code in (403, 422)
    assert store.client.post(f"/api/artifacts/{identifier}/export/pdf").status_code == 403


@pytest.mark.asyncio
async def test_scope_enforcement_and_reference_validation(store):
    conversation, other = uuid.uuid4(), uuid.uuid4()
    with store.engine.begin() as conn:
        conn.execute(
            text("INSERT INTO conversations (id) VALUES (:id)"),
            {"id": conversation.hex},
        )
    artifact = await generated(store, conversation_id=conversation)
    async with store.factory() as db:
        with pytest.raises(HTTPException) as error:
            await service.get_artifact(db, uuid.UUID(artifact["id"]), other, scoped=True)
        assert error.value.status_code == 403
        reference = ArtifactReference(artifact_id=artifact["id"], version=1, section_id="section-1")
        assert (await validate_reference(db, reference, conversation))["version"] == 1
        with pytest.raises(HTTPException):
            await validate_reference(db, reference, other)
        with pytest.raises(HTTPException):
            await validate_reference(
                db,
                reference.model_copy(update={"section_id": "invented"}),
                conversation,
            )
        assert not (await service.catalog(db, other, scoped=True))["items"]


@pytest.mark.asyncio
async def test_summary_is_explicit_hierarchical_and_reports_missing_sections():
    calls = []

    async def ask(prompt):
        calls.append(prompt)
        return "[part-1] Short summary."

    result = await summarize_sections(
        [
            {"id": "part-1", "title": "Long", "page": 1, "content": "a" * 25000},
            {"id": "part-2", "title": "Empty", "content": ""},
        ],
        "Summarize",
        ask,
    )
    assert result["missing_sections"] == ["part-2"] and result["sections_processed"] == 1
    assert len(calls) > 3 and all(len(prompt) < 13500 for prompt in calls)


@pytest.mark.asyncio
async def test_summary_batches_pages_and_reports_progress():
    prompts, updates = [], []

    async def ask(prompt):
        prompts.append(prompt)
        return "[part-1] Grounded summary"

    async def progress(update):
        updates.append(update)

    await summarize_sections(
        [{"id": f"part-{i}", "page": i, "content": "a" * 1000} for i in range(48)],
        "Summarize",
        ask,
        progress,
    )
    assert len(prompts) < 15
    assert all(len(prompt) < 13500 for prompt in prompts)
    assert updates[0]["stage"] == "batches"
    assert updates[-1]["stage"] == "complete"
    assert updates[-1]["completed"] == updates[-1]["total"]
    assert any(update["stage"] == "synthesis" for update in updates)


@pytest.mark.asyncio
async def test_update_tool_returns_versioned_downloadable_attachment(store):
    from app.agent.tools.artifacts import UpdateArtifactTool

    conversation = uuid.uuid4()
    with store.engine.begin() as db:
        db.execute(text("INSERT INTO conversations (id) VALUES (:id)"), {"id": conversation.hex})
    artifact = await generated(
        store, "report", "# Report\n\n## Plan: Alpha\nKeep body.", conversation
    )
    result = await UpdateArtifactTool().execute(
        artifact_id=artifact["id"],
        conversation_id=str(conversation),
        expected_version=1,
        changes=[{"section_id": "section-2", "content": "## Plan: Beta\nKeep body."}],
    )
    assert result.success
    attachment = result.tool_call.gen_results[0]
    assert attachment["type"] == "report"
    assert attachment["format"] == "docx"
    assert attachment["artifact_id"] == artifact["id"]
    assert attachment["version"] == 2
    assert attachment["filename"].endswith(".docx")
    assert attachment["file_path"]
    assert store.client.get(attachment["download_url"]).status_code == 200
    assert "?version=2" in attachment["download_url"]


@pytest.mark.asyncio
async def test_summary_tool_progress_uses_parent_id_and_persists(store, monkeypatch):
    import asyncio
    from app.agent.tools.artifacts import SummarizeArtifactTool
    from app.agent.service import _tool_call_to_update_dict

    artifact = SimpleNamespace(id=uuid.uuid4())
    version = SimpleNamespace(number=1)

    async def get_artifact(*args, **kwargs):
        return artifact

    async def get_version(*args, **kwargs):
        return version

    async def summarize(db, artifact, current, question, progress, model=None):
        assert model == "qwen3.5:0.8b-mlx"
        await progress({"stage": "batches", "completed": 1, "total": 2})
        await progress({"stage": "complete", "completed": 2, "total": 2})
        return {"summary": "Grounded summary"}

    monkeypatch.setattr(service, "get_artifact", get_artifact)
    monkeypatch.setattr(service, "get_version", get_version)
    monkeypatch.setattr(service, "summarize", summarize)
    from app.agent.tools import config_store
    from unittest.mock import AsyncMock

    monkeypatch.setattr(
        config_store, "tool_model_override", AsyncMock(return_value="qwen3.5:0.8b-mlx")
    )
    queue = asyncio.Queue()
    result = await SummarizeArtifactTool().execute(
        artifact_id=str(artifact.id), _event_queue=queue, _parent_tool_call_id="parent"
    )
    assert result.success
    assert queue.qsize() == 3
    assert all(queue.get_nowait()["id"] == "parent" for _ in range(3))
    assert _tool_call_to_update_dict(result.tool_call)["progress"]["stage"] == "complete"


@pytest.mark.asyncio
async def test_summary_retry_reuses_completed_batches(tmp_path, monkeypatch):
    import httpx
    from app.services import providers, model_prefs

    monkeypatch.setenv("REALOPEN_DATA_DIR", str(tmp_path))

    async def model(_):
        return "local-test"

    monkeypatch.setattr(model_prefs, "resolve_task_model", model)
    calls = []
    fail = True

    async def stream(model, messages, **options):
        calls.append(messages[-1]["content"])
        assert options["options"] == {"temperature": 0, "num_predict": 1536}
        assert options["think"] is False
        if fail and len(calls) == 2:
            raise httpx.ReadTimeout("test timeout")
        yield {"content": "[part-1] Summary"}

    monkeypatch.setattr(providers, "stream_chat", stream)
    artifact = SimpleNamespace(id=uuid.uuid4())
    version = SimpleNamespace(
        number=1,
        source="source",
        settings={},
        sections=[{"id": f"part-{i}", "content": str(i) * 8000} for i in range(3)],
    )
    with pytest.raises(httpx.ReadTimeout):
        await service.summarize(None, artifact, version)
    first_prompt = calls[0]
    fail = False
    result = await service.summarize(None, artifact, version)
    assert calls.count(first_prompt) == 1
    assert result["sections_processed"] == 3
    previous_calls = len(calls)
    assert await service.summarize(None, artifact, version) == result
    assert len(calls) == previous_calls


@pytest.mark.asyncio
async def test_scanned_pdf_uses_engine_neutral_ocr_only_on_empty_pages(monkeypatch):
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((50, 50), "Native text sufficient for this extraction page.")
    doc.new_page()
    calls = []

    class Engine:
        def recognize(self, image, language):
            calls.append(language)
            return ocr.OCRResult(
                "Scanned page text",
                "test",
                language,
                [{"text": "Scanned", "confidence": 90, "box": [0, 0, 10, 10]}],
            )

    monkeypatch.setattr(ocr, "default_engine", Engine())
    extracted, snapshot = await document_extraction.extract(doc.tobytes(), "scan.pdf", "fra")
    assert calls == ["fra"]
    assert snapshot["sections"][1]["method"] == "ocr" and snapshot["sections"][1]["page"] == 2
    assert extracted.pages[1].text == "Scanned page text"


@pytest.mark.asyncio
async def test_ocr_unavailable_does_not_drop_native_content(monkeypatch):
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((50, 50), "Small text")

    async def fail(*args, **kwargs):
        raise RuntimeError("unavailable")

    monkeypatch.setattr(ocr, "recognize", fail)
    result, snapshot = await document_extraction.extract(doc.tobytes(), "scan.pdf")
    assert "Small text" in result.pages[0].text and snapshot["warnings"]


@pytest.mark.skipif(
    not shutil.which("tesseract"), reason="Optional local Tesseract binary unavailable"
)
def test_real_tesseract_english_text():
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (900, 120), "white")
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    ImageDraw.Draw(image).text((20, 30), "Artifact OCR test 123", font=font, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    result = ocr.TesseractEngine().recognize(buffer.getvalue(), "eng")
    assert "123" in result.text and result.words
    assert all(len(word["box"]) == 4 for word in result.words)


def test_api_rejects_path_injection_bad_ids_and_unbounded_reads(store):
    assert store.client.get("/api/artifacts/not-a-uuid").status_code == 422
    assert store.client.get(f"/api/artifacts/{uuid.uuid4()}/read?offset=-1").status_code == 422
    assert store.client.get(f"/api/artifacts/{uuid.uuid4()}/read?limit=9999999").status_code == 422
    with pytest.raises(HTTPException):
        service.file_path("../escape")


@pytest.mark.asyncio
async def test_agent_read_schema_and_no_model_call(store):
    from app.agent.tools.artifacts import ReadArtifactTool, ListArtifactsTool

    conversation = uuid.uuid4()
    with store.engine.begin() as conn:
        conn.execute(
            text("INSERT INTO conversations (id) VALUES (:id)"),
            {"id": conversation.hex},
        )
    artifact = await generated(store, conversation_id=conversation)
    result = await ReadArtifactTool().execute(
        artifact_id=artifact["id"],
        conversation_id=str(conversation),
        section_id="section-1",
    )
    assert result.success and "Original" in json.loads(result.output)["content"]
    denied = await ReadArtifactTool().execute(
        artifact_id=artifact["id"], conversation_id=str(uuid.uuid4())
    )
    assert not denied.success
    assert "conversation_id" not in ReadArtifactTool().get_parameters()
    assert ListArtifactsTool().get_required_params() == []


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "kind,fmt,source",
    [
        ("presentation", "pdf", "# Cover\n\n---\n\n# Details\nBody."),
        (
            "excel",
            "pdf",
            '{"sheets":[{"name":"Data","tables":[{"start_cell":"A1","headers":["Value"],"rows":[[3]]}]}]}',
        ),
    ],
)
async def test_unavailable_office_export_keeps_history_and_cleans_candidates(
    store, monkeypatch, kind, fmt, source
):
    from unittest.mock import AsyncMock

    artifact = await generated(store, kind, source)
    monkeypatch.setattr(service, "convert_pptx_to_pdf", AsyncMock(return_value=None))
    monkeypatch.setattr(service, "_convert_office_to_pdf_in_cache", AsyncMock(return_value=None))
    before = set((store.root / "reports").iterdir())
    url = f"/api/artifacts/{artifact['id']}"
    response = store.client.post(url + f"/export/{fmt}?version=1")
    assert response.status_code == 503
    assert set((store.root / "reports").iterdir()) == before
    detail = store.client.get(url).json()
    assert detail["version"] == 1 and len(detail["versions"]) == 1
    assert fmt not in detail["outputs"]


@pytest.mark.asyncio
async def test_reference_rederives_source_metadata_instead_of_trusting_caller(store):
    conversation = uuid.uuid4()
    with store.engine.begin() as conn:
        conn.execute(text("INSERT INTO conversations (id) VALUES (:id)"), {"id": conversation.hex})
    artifact = await generated(
        store, "presentation", "# Cover\n\n---\n\n# Details\nFact.", conversation
    )
    reference = ArtifactReference(
        artifact_id=artifact["id"],
        version=1,
        section_id="slide-2",
        page=999,
        slide=99,
        sheet="Invented",
    )
    async with store.factory() as db:
        result = await validate_reference(db, reference, conversation)
    assert result == {
        "artifact_id": artifact["id"],
        "version": 1,
        "section_id": "slide-2",
        "slide": 2,
    }


@pytest.mark.asyncio
async def test_cancelled_summary_retry_retains_only_completed_batches(tmp_path, monkeypatch):
    import asyncio
    from app.services import providers

    monkeypatch.setenv("REALOPEN_DATA_DIR", str(tmp_path))
    artifact = SimpleNamespace(id=uuid.uuid4())
    version = SimpleNamespace(
        number=1,
        source="QA",
        settings={},
        sections=[
            {"id": "part-1", "content": "a" * 8000, "page": 1},
            {"id": "part-2", "content": "b" * 8000, "page": 2},
        ],
    )
    calls = []
    cancel = True

    async def stream(model, messages, **kwargs):
        calls.append(messages[-1]["content"])
        if cancel and len(calls) == 2:
            yield {"content": "Unfinished"}
            raise asyncio.CancelledError()
        yield {"content": "Fact [page 1]."}

    monkeypatch.setattr(providers, "stream_chat", stream)
    with pytest.raises(asyncio.CancelledError):
        await service.summarize(None, artifact, version, model="qa")
    first = calls[0]
    cancel = False
    result = await service.summarize(None, artifact, version, model="qa")
    assert calls.count(first) == 1
    assert "Unfinished" not in result["summary"]
    assert [ref["page"] for ref in result["source_references"]] == [1, 2]
