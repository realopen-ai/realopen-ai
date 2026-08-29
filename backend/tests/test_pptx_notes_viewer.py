"""
Tests for speaker-notes support in the PPTX slide viewer.

Three layers, mirroring test_pptx_slides_viewer.py:

1. Pure-filesystem / python-pptx units — notes extraction, manifest
   persistence, v1→v2 manifest invalidation (no LibreOffice needed).
2. API endpoint — /reports/{id}/slides returns per-slide `notes` (LibreOffice
   mocked out).
3. Real pipeline — PPTX → slides manifest carrying notes. Skips
   automatically when soffice / PyMuPDF are not installed.
"""

import asyncio
import json
import sys
import time
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI

BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.api import reports  # noqa: E402
from app.services.integrations import libreoffice  # noqa: E402

# ══════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════


def _make_pptx_with_notes(
    path: Path,
    slide_notes: list[str | None],
) -> Path:
    """Build a real PPTX; slide_notes entries map 1:1 to slides.

    None → slide without a notes slide; "" → notes slide with empty text;
    str → notes slide carrying that text.
    """
    from pptx import Presentation

    prs = Presentation()
    for i, notes in enumerate(slide_notes, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = f"Slide {i}"
        if notes is not None:
            tf = slide.notes_slide.notes_text_frame
            tf.text = notes
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _write_fake_cache(pptx: Path, count: int = 2, notes: list | None = None) -> Path:
    """Create a COMPLETE, fresh slide cache (manifest + all images)."""
    import io

    from PIL import Image

    cache_dir = libreoffice.slides_cache_dir(pptx)
    cache_dir.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    Image.new("RGB", (1600, 900), (60, 120, 200)).save(buf, format="JPEG")
    full_jpeg = buf.getvalue()
    buf2 = io.BytesIO()
    Image.new("RGB", (360, 202), (60, 120, 200)).save(buf2, format="JPEG")
    thumb_jpeg = buf2.getvalue()

    for i in range(1, count + 1):
        (cache_dir / libreoffice.slide_image_name(i)).write_bytes(full_jpeg)
        (cache_dir / libreoffice.slide_image_name(i, thumb=True)).write_bytes(
            thumb_jpeg
        )

    st = pptx.stat()
    manifest = {
        "version": libreoffice.SLIDES_MANIFEST_VERSION,
        "count": count,
        "width": 1600,
        "height": 900,
        "source_mtime": st.st_mtime,
        "source_size": st.st_size,
        "generated_at": time.time(),
    }
    if notes is not None:
        manifest["notes"] = notes
    (cache_dir / "manifest.json").write_text(json.dumps(manifest))
    return cache_dir


def _make_client(tmp_data: Path, monkeypatch) -> httpx.AsyncClient:
    """httpx client wired to a FastAPI app whose reports data dir is tmp_data."""
    app = FastAPI()
    app.include_router(reports.router, prefix="/api")
    monkeypatch.setattr(reports, "_get_data_dir", lambda: tmp_data)
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app),
        base_url="http://testserver",
    )


# ══════════════════════════════════════════════════════════════════════
# 1. Notes extraction units
# ══════════════════════════════════════════════════════════════════════


def test_extract_notes_basic(tmp_path):
    pptx = _make_pptx_with_notes(
        tmp_path / "a.pptx", ["First slide notes", None, "Third slide notes"]
    )
    assert libreoffice.extract_pptx_speaker_notes(pptx) == [
        "First slide notes",
        "",
        "Third slide notes",
    ]


def test_extract_notes_all_missing(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "b.pptx", [None, None])
    assert libreoffice.extract_pptx_speaker_notes(pptx) == ["", ""]


def test_extract_notes_strips_surrounding_whitespace(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "c.pptx", ["\n\n  hello there  \n\n"])
    assert libreoffice.extract_pptx_speaker_notes(pptx) == ["hello there"]


def test_extract_notes_preserves_internal_linebreaks(tmp_path):
    pptx = _make_pptx_with_notes(
        tmp_path / "d.pptx", ["Line one.\nLine two.\n\nLine three."]
    )
    notes = libreoffice.extract_pptx_speaker_notes(pptx)
    assert notes == ["Line one.\nLine two.\n\nLine three."]


def test_extract_notes_pads_to_expected_count(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "e.pptx", ["only", None])
    notes = libreoffice.extract_pptx_speaker_notes(pptx, expected_count=4)
    assert notes == ["only", "", "", ""]


def test_extract_notes_truncates_to_expected_count(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "f.pptx", ["one", "two", "three"])
    notes = libreoffice.extract_pptx_speaker_notes(pptx, expected_count=2)
    assert notes == ["one", "two"]


def test_extract_notes_missing_file(tmp_path):
    assert libreoffice.extract_pptx_speaker_notes(tmp_path / "nope.pptx") == []


def test_extract_notes_corrupt_file(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"this is not a real pptx zip archive")
    assert libreoffice.extract_pptx_speaker_notes(bad) == []


def test_extract_notes_never_raises_on_pptx_error(tmp_path, monkeypatch):
    """A python-pptx internal blow-up must degrade to [], not raise."""
    pptx = _make_pptx_with_notes(tmp_path / "g.pptx", ["hello"])

    from pptx import Presentation as _RealPresentation

    class _Boom:
        def __init__(self, *a, **kw):
            raise RuntimeError("simulated pptx parse failure")

    import pptx as _pptx_pkg

    monkeypatch.setattr(_pptx_pkg, "Presentation", _Boom)
    try:
        assert libreoffice.extract_pptx_speaker_notes(pptx) == []
    finally:
        monkeypatch.setattr(_pptx_pkg, "Presentation", _RealPresentation)


# ══════════════════════════════════════════════════════════════════════
# 2. Manifest persistence units
# ══════════════════════════════════════════════════════════════════════


def _pages(n: int) -> list[dict]:
    return [{"index": i, "width": 1600, "height": 900} for i in range(1, n + 1)]


def test_manifest_persists_notes(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "m.pptx", ["n1", "n2"])
    cache = tmp_path / "cache"
    cache.mkdir()

    manifest = libreoffice._write_slides_manifest(cache, pptx, _pages(2), ["n1", "n2"])
    assert manifest["notes"] == ["n1", "n2"]

    reread = libreoffice.read_slides_manifest(cache)
    assert reread is not None
    assert reread["notes"] == ["n1", "n2"]


def test_manifest_pads_short_notes(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "m2.pptx", ["n1"])
    cache = tmp_path / "cache"
    cache.mkdir()

    manifest = libreoffice._write_slides_manifest(cache, pptx, _pages(3), ["n1"])
    assert manifest["notes"] == ["n1", "", ""]


def test_manifest_truncates_long_notes(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "m3.pptx", ["n1", "n2", "n3", "n4"])
    cache = tmp_path / "cache"
    cache.mkdir()

    manifest = libreoffice._write_slides_manifest(
        cache, pptx, _pages(2), ["n1", "n2", "n3", "n4"]
    )
    assert manifest["notes"] == ["n1", "n2"]


def test_manifest_without_notes_defaults_empty(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "m4.pptx", ["n1"])
    cache = tmp_path / "cache"
    cache.mkdir()

    manifest = libreoffice._write_slides_manifest(cache, pptx, _pages(2))
    assert manifest["notes"] == ["", ""]


def test_v1_manifest_without_notes_is_rejected(tmp_path):
    """Old (v1) manifests must be invalidated so notes get backfilled."""
    pptx = _make_pptx_with_notes(tmp_path / "old.pptx", ["note"])
    _write_fake_cache(pptx, count=1)

    # Rewrite the manifest as if it came from the pre-notes (v1) era.
    cache = libreoffice.slides_cache_dir(pptx)
    v1 = json.loads((cache / "manifest.json").read_text())
    v1["version"] = 1
    v1.pop("notes", None)
    (cache / "manifest.json").write_text(json.dumps(v1))

    assert libreoffice.read_slides_manifest(cache) is None
    assert libreoffice.is_slide_cache_valid(pptx) is False


def test_manifest_version_is_2_with_notes(tmp_path):
    pptx = _make_pptx_with_notes(tmp_path / "v.pptx", ["x"])
    _write_fake_cache(pptx, count=1, notes=["x"])

    manifest = libreoffice.read_slides_manifest(libreoffice.slides_cache_dir(pptx))
    assert manifest is not None
    assert manifest["version"] == 2
    assert manifest["notes"] == ["x"]
    assert libreoffice.is_slide_cache_valid(pptx) is True


# ══════════════════════════════════════════════════════════════════════
# 3. API endpoint — notes in the /slides payload
# ══════════════════════════════════════════════════════════════════════


def test_slides_endpoint_returns_notes(tmp_path, monkeypatch):
    _make_pptx_with_notes(
        tmp_path / "reports" / "abc.pptx", ["alpha notes", None, "gamma notes"]
    )

    async def fake_convert(p, cache_dir=None, **kw):
        notes = libreoffice.extract_pptx_speaker_notes(p, expected_count=3)
        return json.loads(
            (_write_fake_cache(p, count=3, notes=notes) / "manifest.json").read_text()
        )

    monkeypatch.setattr(libreoffice, "is_available", lambda: True)
    monkeypatch.setattr(libreoffice, "convert_pptx_to_slide_images", fake_convert)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 200
            data = r.json()
            assert data["count"] == 3
            slides = data["slides"]
            assert slides[0]["notes"] == "alpha notes"
            assert slides[1]["notes"] == ""
            assert slides[2]["notes"] == "gamma notes"
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_defaults_notes_when_manifest_lacks_them(tmp_path, monkeypatch):
    """Defensive: a manifest without a notes key still yields notes="\"."""
    pptx = _make_pptx_with_notes(tmp_path / "reports" / "abc.pptx", ["x"])
    _write_fake_cache(pptx, count=2, notes=None)  # no "notes" key written

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 200
            for slide in r.json()["slides"]:
                assert slide["notes"] == ""
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_short_notes_list_padded_with_empty(tmp_path, monkeypatch):
    pptx = _make_pptx_with_notes(tmp_path / "reports" / "abc.pptx", ["x"])
    _write_fake_cache(pptx, count=3, notes=["only-first"])

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 200
            slides = r.json()["slides"]
            assert slides[0]["notes"] == "only-first"
            assert slides[1]["notes"] == ""
            assert slides[2]["notes"] == ""
        finally:
            await client.aclose()

    asyncio.run(run())


def test_slides_endpoint_serves_cached_manifest_with_notes(tmp_path, monkeypatch):
    """Fresh cache → no conversion, notes come straight from the manifest."""
    pptx = _make_pptx_with_notes(
        tmp_path / "reports" / "abc.pptx", ["cached note", "other note"]
    )
    _write_fake_cache(pptx, count=2, notes=["cached note", "other note"])

    async def boom(*a, **kw):  # pragma: no cover - must not run
        raise AssertionError("convert must not be called for a fresh cache")

    monkeypatch.setattr(libreoffice, "convert_pptx_to_slide_images", boom)

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/abc/slides")
            assert r.status_code == 200
            assert r.json()["slides"][1]["notes"] == "other note"
        finally:
            await client.aclose()

    asyncio.run(run())


# ══════════════════════════════════════════════════════════════════════
# 4. Real pipeline (skips when tools are absent)
# ══════════════════════════════════════════════════════════════════════


def _fitz_available() -> bool:
    try:
        import fitz  # noqa: F401

        return True
    except ImportError:
        return False


@pytest.mark.skipif(not libreoffice.is_available(), reason="LibreOffice not installed")
@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_end_to_end_notes_in_manifest_and_api(tmp_path, monkeypatch):
    """Real pipeline: pptx with notes → slide renders + notes in manifest/API."""
    pptx = _make_pptx_with_notes(
        tmp_path / "reports" / "e2e-notes.pptx",
        ["Opening remarks", None, "Closing remarks"],
    )

    manifest = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))
    assert manifest is not None
    assert manifest["count"] == 3
    assert manifest["notes"] == ["Opening remarks", "", "Closing remarks"]

    async def run():
        client = _make_client(tmp_path, monkeypatch)
        try:
            r = await client.get("/api/reports/e2e-notes/slides")
            assert r.status_code == 200
            slides = r.json()["slides"]
            assert slides[0]["notes"] == "Opening remarks"
            assert slides[1]["notes"] == ""
            assert slides[2]["notes"] == "Closing remarks"
        finally:
            await client.aclose()

    asyncio.run(run())


@pytest.mark.skipif(not libreoffice.is_available(), reason="LibreOffice not installed")
@pytest.mark.skipif(not _fitz_available(), reason="PyMuPDF not installed")
def test_stale_v1_cache_regenerates_with_notes(tmp_path):
    """A v1-era cache (no notes) is invalid → re-render backfills notes."""
    pptx = _make_pptx_with_notes(tmp_path / "reports" / "old.pptx", ["kept note"])
    _write_fake_cache(pptx, count=1)

    # Downgrade the manifest to v1 (pre-notes).
    cache = libreoffice.slides_cache_dir(pptx)
    v1 = json.loads((cache / "manifest.json").read_text())
    v1["version"] = 1
    v1.pop("notes", None)
    (cache / "manifest.json").write_text(json.dumps(v1))

    manifest = asyncio.run(libreoffice.convert_pptx_to_slide_images(pptx))
    assert manifest is not None
    assert manifest["version"] == 2
    assert manifest["notes"] == ["kept note"]
