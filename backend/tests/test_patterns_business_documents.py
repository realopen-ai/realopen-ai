"""
Tests for the business-document pattern batch: price_list, receivables
and payroll (app/services/patterns/{price_list,receivables,payroll}.py).

Same conventions as tests/test_excel_patterns.py:
  • param coercion — quoted/currency/percent strings, alias keys, null
    semantics from the pattern stanzas, ValueError on missing required
    params
  • builder layout math — every formula string pins to the EXACT row
    the converter renders (specs normalized via eg._normalize_spec)
  • independent math — expected margins / outstanding sums / net pay
    recomputed in plain Python and cross-checked against the emitted
    formulas (arithmetic formulas evaluated over the emitted grid)
  • converter round-trip via eg._build_xlsx + openpyxl load_workbook
  • routing — mocked classifier → pattern, AI path never runs
"""

import json
import re
import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import excel_gen as eg  # noqa: E402
from app.services import patterns as ep  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

# ── Shared helpers ────────────────────────────────────────────────────


def _classify_response(payload):
    async def fake_llm(messages, model=None):
        assert "route spreadsheet requests" in messages[0]["content"]
        return json.dumps(payload)

    return fake_llm


class _NotEvaluable(Exception):
    """Formula carries functions (IF/MAX/…) the mini-evaluator skips."""


_SUM_RE = re.compile(r"SUM\((\$?[A-Z]{1,2})\$?(\d+):(\$?[A-Z]{1,2})\$?(\d+)\)")
_REF_RE = re.compile(r"\$?([A-Z]{1,2})\$?(\d+)")


def _eval_formula(formula, values):
    """Evaluate an arithmetic formula (SUM + cell refs) over a grid.

    Raises KeyError while a referenced cell is not resolved yet (used by
    _resolve_formulas for dependency ordering) and _NotEvaluable for
    formulas with functions the evaluator does not model.
    """
    expr = formula.lstrip("=")

    def _sum(m):
        c1, r1, c2, r2 = (
            m.group(1).replace("$", ""),
            int(m.group(2)),
            m.group(3).replace("$", ""),
            int(m.group(4)),
        )
        if c1 != c2:
            raise _NotEvaluable()
        total = 0.0
        for r in range(r1, r2 + 1):
            v = values.get(f"{c1}{r}")  # empty cells sum as 0, like Excel
            if v is not None:
                total += float(v)
        return repr(total)

    expr = _SUM_RE.sub(_sum, expr)
    if "(" in expr or re.search(r"[A-Z]{2,}", expr):
        raise _NotEvaluable()

    def _ref(m):
        key = f"{m.group(1)}{m.group(2)}"
        if key not in values:
            raise KeyError(key)
        v = values[key]
        return "0.0" if v is None else repr(float(v))

    expr = _REF_RE.sub(_ref, expr)
    return eval(expr, {"__builtins__": {}}, {})  # noqa: S307 — test helper


def _resolve_formulas(cells):
    """{"D7": 18.0, "G7": "=D7*E7+…"} → values with arithmetic formulas
    resolved in dependency order (guarded IF/MAX formulas skipped)."""
    values = {
        k: v for k, v in cells.items() if not (isinstance(v, str) and v.startswith("="))
    }
    pending = {
        k: v for k, v in cells.items() if isinstance(v, str) and v.startswith("=")
    }
    for _ in range(30):
        if not pending:
            break
        progressed = False
        for ref in list(pending):
            try:
                values[ref] = _eval_formula(pending[ref], values)
            except KeyError:
                continue
            except _NotEvaluable:
                pass  # guarded logic — verified via string pins instead
            del pending[ref]
            progressed = True
        if not progressed:
            break
    return values


def _cell_date(value):
    """openpyxl returns datetime.datetime for date cells → date."""
    return value.date() if hasattr(value, "date") else value


# ══════════════════════════════════════════════════════════════════════
# price_list
# ══════════════════════════════════════════════════════════════════════

PRICE_PARAMS = {
    "list_name": "Bakery Price List",
    "currency": "USD",
    "vat_rate": "20%",
    "products": [
        {
            "sku": "BR-01",
            "name": "Baguette",
            "category": "Bread",
            "cost": 0.45,
            "price": 1.20,
        },
        {
            "sku": "CR-02",
            "name": "Croissant",
            "category": "Pastry",
            "cost": 0.38,
            "price": 1.05,
        },
        {"name": "Brioche", "category": "Bread", "cost": 0.62},
    ],
}


class TestPriceListBuilder:
    def setup_method(self):
        self.spec = ep.build_price_list_spec(PRICE_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert warnings == []

    def test_geometry_and_formula_lattice(self):
        # title 1, VAT assumption 2, blank 3, headers 4, data 5-7, total 8
        assert self.spec["filename"] == "bakery_price_list.xlsx"
        assert [s["name"] for s in self.norm["sheets"]] == ["Pricing", "Categories"]
        sheet = self.norm["sheets"][0]
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        assert blocks["A1"] == "Bakery Price List"
        assert blocks["A2"] == "VAT Rate"
        assert blocks["B2"] == 0.2  # "20%" → 0.2

        table = sheet["tables"][0]
        assert table["start_cell"] == "A4"
        assert table["headers"] == [
            "SKU",
            "Product",
            "Category",
            "Cost",
            "Net Price",
            "Markup %",
            "Margin %",
            "Gross Price",
        ]
        # every per-product formula pins to its OWN rendered row
        assert table["rows"][0][5] == '=IF(AND(D5>0,E5<>""),(E5-D5)/D5,"n/a")'
        assert table["rows"][0][6] == '=IF(E5>0,(E5-D5)/E5,"n/a")'
        assert table["rows"][0][7] == '=IF(E5>0,E5*(1+$B$2),"n/a")'
        assert table["rows"][1][5] == '=IF(AND(D6>0,E6<>""),(E6-D6)/D6,"n/a")'
        assert table["rows"][2][4] is None  # price unknown → n/a guards
        assert table["rows"][2][7] == '=IF(E7>0,E7*(1+$B$2),"n/a")'
        # total row: SUMs over the exact data rows + guarded averages
        assert table["total_row"][3] == "=SUM(D5:D7)"
        assert table["total_row"][4] == "=SUM(E5:E7)"
        assert table["total_row"][5] == '=IF(COUNT(F5:F7)>0,AVERAGE(F5:F7),"n/a")'
        assert table["total_row"][6] == '=IF(COUNT(G5:G7)>0,AVERAGE(G5:G7),"n/a")'
        assert table["total_row"][7] == '=IF(COUNT(H5:H7)>0,SUM(H5:H7),"n/a")'

        # charts / dropdown / conditional format pin to the same rows
        chart = sheet["charts"][0]
        assert chart["categories_range"] == "Pricing!B5:B7"
        assert chart["series"][0]["values_range"] == "Pricing!G5:G7"
        assert chart["anchor"] == "J4"
        dv = sheet["data_validation"][0]
        assert dv["range"] == "C5:C7"
        assert dv["source_range"] == "Categories!$A$4:$A$5"
        cf = sheet["conditional_formats"][0]
        assert cf["range"] == "G5:G7"
        assert cf["rules"][0]["value"] == "AND(ISNUMBER($G5),$G5<0)"

        # Categories sheet: headers 3, data 4-5, total 6
        cat_sheet = self.norm["sheets"][1]
        cat_table = cat_sheet["tables"][0]
        assert cat_table["start_cell"] == "A3"
        assert cat_table["rows"][0][0] == "Bread"
        assert cat_table["rows"][0][1] == "=COUNTIF(Pricing!$C$5:$C$7,$A4)"
        assert cat_table["rows"][0][2] == (
            "=SUMIF(Pricing!$C$5:$C$7,$A4,Pricing!$D$5:$D$7)"
        )
        assert cat_table["rows"][0][4] == (
            "=IFERROR(AVERAGEIF(Pricing!$C$5:$C$7,$A4," 'Pricing!$G$5:$G$7),"n/a")'
        )
        assert cat_table["total_row"][1] == "=SUM(B4:B5)"
        assert cat_table["total_row"][4] == ('=IF(COUNT(E4:E5)>0,AVERAGE(E4:E5),"n/a")')
        assert cat_sheet["charts"][0]["categories_range"] == "Categories!A4:A5"

    def test_margin_math_independent(self):
        """Recompute markup/margin in Python from the emitted data and
        check they match the guard semantics the formulas encode."""
        table = self.norm["sheets"][0]["tables"][0]
        expected_margins, expected_markups = [], []
        for row in table["rows"]:
            cost, price = row[3], row[4]
            expected_markups.append(
                (price - cost) / cost
                if (cost or 0) > 0 and price is not None
                else "n/a"
            )
            expected_margins.append(
                (price - cost) / price if (price or 0) > 0 else "n/a"
            )
        # Baguette / Croissant priced, Brioche unpriced → n/a
        assert expected_markups[0] == pytest.approx(0.75 / 0.45)
        assert expected_margins[0] == pytest.approx(0.75 / 1.20)
        assert expected_markups[1] == pytest.approx(0.67 / 0.38)
        assert expected_margins[1] == pytest.approx(0.67 / 1.05)
        assert expected_markups[2] == "n/a"
        assert expected_margins[2] == "n/a"
        # average margin over the priced rows only (as AVERAGE does)
        avg_margin = sum(m for m in expected_margins if m != "n/a") / sum(
            1 for m in expected_margins if m != "n/a"
        )
        assert avg_margin == pytest.approx((0.75 / 1.20 + 0.67 / 1.05) / 2)

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "price_list.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Pricing", "Categories"]
        ws = wb["Pricing"]
        assert ws["A1"].value == "Bakery Price List"
        assert ws["B2"].value == 0.2
        assert ws["B2"].number_format == "0.00%"
        assert ws["A4"].value == "SKU"
        assert ws["B5"].value == "Baguette"
        assert ws["D5"].value == 0.45
        assert ws["D5"].number_format == '"$"#,##0.00'
        assert ws["F5"].value == '=IF(AND(D5>0,E5<>""),(E5-D5)/D5,"n/a")'
        assert ws["G5"].number_format == "0.00%"
        assert ws["H5"].value == '=IF(E5>0,E5*(1+$B$2),"n/a")'
        assert ws["H5"].number_format == '"$"#,##0.00'
        assert ws["E7"].value is None  # unknown price stays empty
        assert ws["D8"].value == "=SUM(D5:D7)"
        assert ws.freeze_panes == "A5"
        assert len(ws._charts) == 1
        (dv,) = ws.data_validations.dataValidation
        assert dv.formula1 == "Categories!$A$4:$A$5"
        cf_ranges = {str(c.sqref) for c in ws.conditional_formatting}
        assert cf_ranges == {"G5:G7"}
        cats = wb["Categories"]
        assert cats["A1"].value == "Category Summary"
        assert cats["A4"].value == "Bread"
        assert cats["B4"].value == "=COUNTIF(Pricing!$C$5:$C$7,$A4)"
        assert cats["B6"].value == "=SUM(B4:B5)"
        assert len(cats._charts) == 1

    def test_no_vat_layout(self):
        spec = ep.build_price_list_spec(
            {
                "products": [
                    {"name": "A", "cost": 1, "price": 2},
                    {"name": "B", "cost": 2, "price": 5},
                ]
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        table = sheet["tables"][0]
        assert table["headers"][-1] == "Margin %"  # no Gross Price column
        assert len(table["headers"]) == 7
        assert len(table["rows"][0]) == 7
        blocks = {b["cell"] for b in sheet["text_blocks"]}
        assert "B2" not in blocks  # no VAT assumption cell
        assert sheet["charts"][0]["anchor"] == "I4"
        assert norm["filename"] == "price_list.xlsx"  # default list name

    def test_single_product_single_category(self):
        spec = ep.build_price_list_spec(
            {"products": [{"name": "Widget", "cost": 4, "price": 10}]}
        )
        errors, warnings = eg.validate_workbook_spec(spec)
        assert errors == []
        assert warnings == []
        norm = eg._normalize_spec(spec)
        cat_sheet = norm["sheets"][1]
        assert cat_sheet["tables"][0]["rows"][0][0] == "Uncategorized"
        assert "charts" not in cat_sheet  # one category → no chart
        dv = norm["sheets"][0]["data_validation"][0]
        assert dv["source_range"] == "Categories!$A$4:$A$4"

    def test_template_mode_no_products_builds_blank_catalog(self, tmp_path):
        # "create a price list" with no products must build the blank
        # pricing template instead of raising → AI path → invalid spec.
        spec = ep.build_price_list_spec({})
        errors, warnings = eg.validate_workbook_spec(spec)
        assert errors == []
        assert warnings == []
        norm = eg._normalize_spec(spec)
        assert norm["filename"] == "price_list.xlsx"

        sheet = norm["sheets"][0]
        table = sheet["tables"][0]
        assert table["start_cell"] == "A4"
        assert len(table["rows"]) == 8  # blank scaffold rows 5..12
        # nothing invented: sku/name/category/cost/price all blank
        assert all(r[:5] == [None] * 5 for r in table["rows"])
        # guarded markup/margin formulas on every scaffold row
        assert table["rows"][0][5] == (
            '=IF(OR($D5="",$D5=0,$E5=""),"n/a",($E5-$D5)/$D5)'
        )
        assert table["rows"][0][6] == '=IF(OR($E5="",$E5=0),"n/a",($E5-$D5)/$E5)'
        assert table["rows"][7][5] == (
            '=IF(OR($D12="",$D12=0,$E12=""),"n/a",($E12-$D12)/$D12)'
        )
        # guarded average total row over the scaffold band
        assert table["total_row"][3] == "=SUM(D5:D12)"
        assert table["total_row"][5] == ('=IF(COUNT(F5:F12)>0,AVERAGE(F5:F12),"n/a")')
        assert table["total_row"][6] == ('=IF(COUNT(G5:G12)>0,AVERAGE(G5:G12),"n/a")')
        # no VAT given → no Gross Price column, no B2 assumption
        assert table["headers"] == [
            "SKU",
            "Product",
            "Category",
            "Cost",
            "Net Price",
            "Markup %",
            "Margin %",
        ]
        blocks = {b["cell"] for b in sheet["text_blocks"]}
        assert "B2" not in blocks
        # no charts when there are no products
        assert "charts" not in sheet
        # dropdown covers the scaffold rows, fed by the category scaffold
        dv = sheet["data_validation"][0]
        assert dv["range"] == "C5:C12"
        assert dv["source_range"] == "Categories!$A$4:$A$11"
        # margin traffic lights still cover the blank band
        cf = sheet["conditional_formats"][0]
        assert cf["range"] == "G5:G12"

        # Categories sheet: default-free scaffold (blank labels +
        # guarded live subtotals) — no category is invented
        cats = norm["sheets"][1]
        cat_table = cats["tables"][0]
        assert len(cat_table["rows"]) == 8
        assert all(r[0] is None for r in cat_table["rows"])
        assert cat_table["rows"][0][1] == (
            '=IF($A4="","",COUNTIF(Pricing!$C$5:$C$12,$A4))'
        )
        assert cat_table["rows"][3][4] == (
            '=IF($A7="","",IFERROR(AVERAGEIF(Pricing!$C$5:$C$12,$A7,'
            'Pricing!$G$5:$G$12),"n/a"))'
        )
        assert cat_table["total_row"][1] == "=SUM(B4:B11)"
        assert "charts" not in cats

        # the workbook still round-trips through openpyxl
        out = tmp_path / "price_blank.xlsx"
        eg._build_xlsx(norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Pricing", "Categories"]
        ws = wb["Pricing"]
        assert ws["B5"].value is None
        assert ws["F5"].value == '=IF(OR($D5="",$D5=0,$E5=""),"n/a",($E5-$D5)/$D5)'
        assert ws["G5"].number_format == "0.00%"
        assert ws["D13"].value == "=SUM(D5:D12)"
        assert len(ws._charts) == 0
        (dv,) = ws.data_validations.dataValidation
        assert dv.formula1 == "Categories!$A$4:$A$11"
        assert str(dv.sqref) == "C5:C12"
        cats_ws = wb["Categories"]
        assert cats_ws["A4"].value is None
        assert cats_ws["B4"].value == '=IF($A4="","",COUNTIF(Pricing!$C$5:$C$12,$A4))'


class TestPriceListCoercion:
    def test_quoted_numbers_and_percent_strings(self):
        p = ep.coerce_price_list_params(
            {
                "vat_rate": "20%",
                "products": [
                    {"name": "Rye", "cost": "$0.90", "price": "2.50"},
                    {"name": "Oat", "cost": "0.55", "price": 1.4},
                ],
            }
        )
        assert p["vat_rate"] == 0.2
        assert p["products"][0]["cost"] == 0.9
        assert p["products"][0]["price"] == 2.5
        assert p["products"][1]["cost"] == 0.55

    def test_alias_keys(self):
        p = ep.coerce_price_list_params(
            {
                "items": [
                    {
                        "product": "Sourdough",
                        "unit_cost": 1.1,
                        "selling_price": 3.4,
                        "cat": "Bread",
                    },
                ],
            }
        )
        assert p["products"][0]["name"] == "Sourdough"
        assert p["products"][0]["cost"] == 1.1
        assert p["products"][0]["price"] == 3.4
        assert p["products"][0]["category"] == "Bread"

    def test_null_cost_and_price_semantics(self):
        p = ep.coerce_price_list_params(
            {
                "products": [
                    {"name": "OnlyPrice", "price": 9},
                    {"name": "OnlyCost", "cost": 2},
                ],
            }
        )
        assert p["products"][0]["cost"] is None
        assert p["products"][1]["price"] is None
        assert p["vat_rate"] is None  # not stated → no VAT column

    def test_product_without_name_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_price_list_params(
                {"products": [{"sku": "X-1", "cost": 1, "price": 2}]}
            )

    def test_products_missing_or_empty_is_template_mode(self):
        # "create a price list template" — no products given: blank
        # template mode, NOT a refusal (products are never invented).
        for params in ({}, {"products": []}, {"list_name": "Bakery"}):
            p = ep.coerce_price_list_params(params)
            assert p["products"] == []
            assert p["list_name"] in ("Price List", "Bakery")

    def test_products_not_array_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_price_list_params({"products": "Widget"})
        with pytest.raises(ValueError):
            ep.coerce_price_list_params("nope")
        # entries given but none usable — structural garbage
        with pytest.raises(ValueError):
            ep.coerce_price_list_params({"products": [None, 42]})

    def test_empty_product_rows_raise(self):
        with pytest.raises(ValueError):
            ep.coerce_price_list_params({"products": [{}, {"sku": "A"}]})

    def test_default_list_name_and_filename(self):
        p = ep.coerce_price_list_params({"products": [{"name": "A", "price": 1}]})
        assert p["list_name"] == "Price List"
        spec = ep.build_price_list_spec({"products": [{"name": "A", "price": 1}]})
        assert spec["filename"] == "price_list.xlsx"

    def test_string_product_entry_treated_as_name(self):
        p = ep.coerce_price_list_params({"products": ["Rye Bread"]})
        assert p["products"][0]["name"] == "Rye Bread"

    def test_user_notes_rendered_on_sheet(self):
        spec = ep.build_price_list_spec(
            {
                "products": [{"name": "A", "cost": 1, "price": 2}],
                "notes": "Wholesale prices, VAT excluded",
            }
        )
        norm = eg._normalize_spec(spec)
        assert norm["sheets"][0]["notes"].startswith("Wholesale prices, VAT excluded")


class TestPriceListRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response(
                {
                    "pattern": "price_list",
                    "params": {
                        "list_name": "Bakery Price List",
                        "currency": "USD",
                        "vat_rate": 0.2,
                        "products": PRICE_PARAMS["products"],
                    },
                }
            ),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "product price list for my bakery with cost and margin"
        )
        assert result["pattern"] == "price_list"
        assert result["sheet_count"] == 2
        assert result["sheet_names"] == ["Pricing", "Categories"]
        assert result["table_count"] == 2
        assert result["chart_count"] == 2
        assert result["formula_count"] > 20
        assert result["filename"] == "bakery_price_list.xlsx"
        assert "price_list template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    @pytest.mark.asyncio
    async def test_routes_with_empty_params_to_blank_template(
        self, tmp_path, monkeypatch
    ):
        # Same class of bug as meal_planner: the classifier returns
        # price_list with NO products ("create a price list template")
        # — the pattern must build the blank template instead of
        # falling back to the AI path.
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "price_list", "params": {}}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "create a price list template for my products"
        )
        assert result["pattern"] == "price_list"
        assert result["sheet_names"] == ["Pricing", "Categories"]
        assert result["chart_count"] == 0  # nothing to chart yet
        assert result["filename"] == "price_list.xlsx"
        assert "price_list template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ══════════════════════════════════════════════════════════════════════
# receivables
# ══════════════════════════════════════════════════════════════════════

REC_PARAMS = {
    "currency": "EUR",
    "as_of_date": "2026-08-15",
    "invoices": [
        {
            "customer": "Acme",
            "invoice_number": "INV-100",
            "date": "2026-06-01",
            "due_date": "2026-07-01",
            "amount": 1500,
            "paid_amount": 1500,
        },
        {
            "customer": "Beta",
            "invoice_number": "INV-101",
            "date": "2026-07-01",
            "due_date": "2026-08-01",
            "amount": 2000,
            "paid_amount": 500,
        },
        {
            "customer": "Gamma",
            "invoice_number": "INV-103",
            "date": "2026-04-01",
            "due_date": "2026-05-01",
            "amount": 600,
            "paid_amount": 0,
        },
        {
            "customer": "Acme",
            "invoice_number": "INV-102",
            "date": "2026-07-15",
            "due_date": "2026-08-15",
            "amount": 800,
            "paid_amount": 0,
        },
    ],
}

_AS_OF = date(2026, 8, 15)


def _rec_expected():
    """Independent Python model of the tracker semantics."""
    rows = []
    for inv in REC_PARAMS["invoices"]:
        due = date.fromisoformat(inv["due_date"]) if inv["due_date"] else None
        days = 0 if due is None else max(0, (_AS_OF - due).days)
        outstanding = inv["amount"] - inv["paid_amount"]
        if inv["paid_amount"] >= inv["amount"]:
            status = "Paid"
        elif inv["paid_amount"] > 0:
            status = "Partial"
        elif due is not None and due < _AS_OF:
            status = "Overdue"
        else:
            status = "Current"
        rows.append((inv, days, outstanding, status))
    return rows


class TestReceivablesBuilder:
    def setup_method(self):
        self.spec = ep.build_receivables_spec(REC_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert warnings == []

    def test_geometry_and_formula_lattice(self):
        assert self.spec["filename"] == "receivables.xlsx"
        assert [s["name"] for s in self.norm["sheets"]] == [
            "Receivables",
            "Summary",
        ]
        sheet = self.norm["sheets"][0]
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        assert blocks["A1"] == "Accounts Receivable"
        assert blocks["A2"] == "As Of"
        assert blocks["B2"] == date(2026, 8, 15)

        table = sheet["tables"][0]
        assert table["start_cell"] == "A4"
        assert table["headers"] == [
            "Customer",
            "Invoice #",
            "Date",
            "Due Date",
            "Amount",
            "Paid",
            "Outstanding",
            "Days Overdue",
            "Status",
        ]
        # per-invoice formulas pin to their own rendered rows (5..8)
        assert table["rows"][0][6] == "=E5-F5"
        assert table["rows"][0][7] == '=IF(D5="",0,MAX(0,$B$2-D5))'
        assert table["rows"][0][8] == (
            '=IF(F5>=E5,"Paid",IF(F5>0,"Partial",'
            'IF(AND(D5<>"",D5<$B$2),"Overdue","Current")))'
        )
        assert table["rows"][3][6] == "=E8-F8"
        assert table["rows"][3][7] == '=IF(D8="",0,MAX(0,$B$2-D8))'
        assert table["total_row"][4] == "=SUM(E5:E8)"
        assert table["total_row"][5] == "=SUM(F5:F8)"
        assert table["total_row"][6] == "=SUM(G5:G8)"

        # conditional formats on the status column
        cf = sheet["conditional_formats"][0]
        assert cf["range"] == "I5:I8"
        assert cf["rules"][0]["value"] == "Overdue"
        assert cf["rules"][1]["value"] == "Partial"
        assert cf["rules"][2]["value"] == "Paid"

        # Summary: aging headers 3, data 4-8, total 9; customer headers
        # 12, data 13-15, total 16
        summary = self.norm["sheets"][1]
        aging = summary["tables"][0]
        assert aging["start_cell"] == "A3"
        assert aging["rows"][0][0] == "Current (not due)"
        assert aging["rows"][0][1] == "=COUNTIF(Receivables!$H$5:$H$8,0)"
        assert aging["rows"][0][2] == (
            "=SUMIF(Receivables!$H$5:$H$8,0,Receivables!$G$5:$G$8)"
        )
        assert aging["rows"][1][1] == (
            '=COUNTIFS(Receivables!$H$5:$H$8,">=1",' 'Receivables!$H$5:$H$8,"<=30")'
        )
        assert aging["rows"][1][2] == (
            '=SUMIFS(Receivables!$G$5:$G$8,Receivables!$H$5:$H$8,">=1",'
            'Receivables!$H$5:$H$8,"<=30")'
        )
        assert aging["rows"][4][1] == '=COUNTIF(Receivables!$H$5:$H$8,">90")'
        assert aging["rows"][4][2] == (
            '=SUMIF(Receivables!$H$5:$H$8,">90",Receivables!$G$5:$G$8)'
        )
        assert aging["total_row"][1] == "=SUM(B4:B8)"
        assert aging["total_row"][2] == "=SUM(C4:C8)"

        cust = summary["tables"][1]
        assert cust["start_cell"] == "A12"
        assert cust["rows"][0][0] == "Acme"
        assert cust["rows"][0][1] == "=COUNTIF(Receivables!$A$5:$A$8,$A13)"
        assert cust["rows"][0][4] == (
            "=SUMIF(Receivables!$A$5:$A$8,$A13,Receivables!$G$5:$G$8)"
        )
        assert cust["rows"][2][4] == (
            "=SUMIF(Receivables!$A$5:$A$8,$A15,Receivables!$G$5:$G$8)"
        )
        assert cust["total_row"][4] == "=SUM(E13:E15)"

        # charts pin to the rendered bucket/customer rows
        aging_chart, cust_chart = summary["charts"]
        assert aging_chart["categories_range"] == "Summary!A4:A8"
        assert aging_chart["series"][0]["values_range"] == "Summary!C4:C8"
        assert cust_chart["categories_range"] == "Summary!A13:A15"
        assert cust_chart["series"][0]["values_range"] == "Summary!E13:E15"
        assert cust_chart["type"] == "bar_h"

    def test_outstanding_math_independent(self):
        """Python-computed outstanding sums vs the emitted formulas."""
        expected = _rec_expected()
        # days overdue / statuses replicate the emitted guard semantics
        assert [d for _inv, d, _out, _st in expected] == [45, 14, 106, 0]
        assert [st for _inv, _d, _out, st in expected] == [
            "Paid",
            "Partial",
            "Overdue",
            "Current",
        ]
        assert [out for _inv, _d, out, _st in expected] == [0, 1500, 600, 800]

        # evaluate the emitted Outstanding/total formulas over the grid
        table = self.norm["sheets"][0]["tables"][0]
        cells = {"B2": 0}  # placeholder — Outstanding rows don't use B2
        for i, row in enumerate(table["rows"]):
            r = 5 + i
            cells[f"E{r}"] = row[4]
            cells[f"F{r}"] = row[5]
            cells[f"G{r}"] = row[6]
        cells["G9"] = table["total_row"][6]
        resolved = _resolve_formulas(cells)
        assert resolved["G5"] == pytest.approx(0)
        assert resolved["G6"] == pytest.approx(1500)
        assert resolved["G7"] == pytest.approx(600)
        assert resolved["G8"] == pytest.approx(800)
        assert resolved["G9"] == pytest.approx(2900)

        # aging buckets recomputed in Python (the SUMIFS lattice above
        # pins to exactly these H/G rows)
        buckets = {
            "Current (not due)": [0, 0.0],
            "1-30 days": [0, 0.0],
            "31-60 days": [0, 0.0],
            "61-90 days": [0, 0.0],
            "Over 90 days": [0, 0.0],
        }
        for _inv, days, out, _st in expected:
            if days == 0:
                buckets["Current (not due)"][0] += 1
                buckets["Current (not due)"][1] += out
            elif days <= 30:
                buckets["1-30 days"][0] += 1
                buckets["1-30 days"][1] += out
            elif days <= 60:
                buckets["31-60 days"][0] += 1
                buckets["31-60 days"][1] += out
            elif days <= 90:
                buckets["61-90 days"][0] += 1
                buckets["61-90 days"][1] += out
            else:
                buckets["Over 90 days"][0] += 1
                buckets["Over 90 days"][1] += out
        assert buckets["Current (not due)"] == [1, 800]
        assert buckets["1-30 days"] == [1, 1500]
        assert buckets["31-60 days"] == [1, 0]
        assert buckets["61-90 days"] == [0, 0]
        assert buckets["Over 90 days"] == [1, 600]

        # per-customer rollup recomputed in Python
        per_customer = {}
        for inv, _d, out, _st in expected:
            agg = per_customer.setdefault(inv["customer"], [0, 0.0, 0.0, 0.0])
            agg[0] += 1
            agg[1] += inv["amount"]
            agg[2] += inv["paid_amount"]
            agg[3] += out
        assert per_customer["Acme"] == [2, 2300, 1500, 800]
        assert per_customer["Beta"] == [1, 2000, 500, 1500]
        assert per_customer["Gamma"] == [1, 600, 0, 600]

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "receivables.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Receivables", "Summary"]
        ws = wb["Receivables"]
        assert ws["A1"].value == "Accounts Receivable"
        as_of = ws["B2"].value
        assert (as_of.year, as_of.month, as_of.day) == (2026, 8, 15)
        assert ws["B2"].number_format == "yyyy-mm-dd"
        assert ws["A5"].value == "Acme"
        assert _cell_date(ws["D5"].value) == date(2026, 7, 1)
        assert ws["D5"].number_format == "yyyy-mm-dd"
        assert ws["E5"].value == 1500
        assert ws["E5"].number_format == '#,##0.00" €"'
        assert ws["G6"].value == "=E6-F6"
        assert ws["H6"].value == '=IF(D6="",0,MAX(0,$B$2-D6))'
        assert ws["H6"].number_format == "0"
        assert ws["I7"].value == (
            '=IF(F7>=E7,"Paid",IF(F7>0,"Partial",'
            'IF(AND(D7<>"",D7<$B$2),"Overdue","Current")))'
        )
        assert ws["G9"].value == "=SUM(G5:G8)"
        assert ws.freeze_panes == "A5"
        cf_ranges = {str(c.sqref) for c in ws.conditional_formatting}
        assert cf_ranges == {"I5:I8"}
        summary = wb["Summary"]
        assert summary["A1"].value == "AR Summary"
        assert summary["A4"].value == "Current (not due)"
        assert summary["C5"].value == (
            '=SUMIFS(Receivables!$G$5:$G$8,Receivables!$H$5:$H$8,">=1",'
            'Receivables!$H$5:$H$8,"<=30")'
        )
        assert summary["A13"].value == "Acme"
        assert summary["E16"].value == "=SUM(E13:E15)"
        assert len(summary._charts) == 2

    def test_missing_as_of_uses_live_today_anchor(self):
        spec = ep.build_receivables_spec({"invoices": REC_PARAMS["invoices"]})
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        blocks = {
            b["cell"]: b for b in eg._normalize_spec(spec)["sheets"][0]["text_blocks"]
        }
        assert blocks["B2"]["text"] == "=TODAY()"
        assert blocks["B2"]["number_format"] == "yyyy-mm-dd"

    def test_missing_due_date_is_current_not_overdue(self):
        spec = ep.build_receivables_spec(
            {
                "as_of_date": "2026-08-15",
                "invoices": [
                    {"customer": "Delta", "amount": 100, "paid_amount": 0},
                ],
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        table = eg._normalize_spec(spec)["sheets"][0]["tables"][0]
        row = table["rows"][0]
        assert row[3] is None  # no due date → empty cell
        assert row[7] == '=IF(D5="",0,MAX(0,$B$2-D5))'
        assert row[8] == (
            '=IF(F5>=E5,"Paid",IF(F5>0,"Partial",'
            'IF(AND(D5<>"",D5<$B$2),"Overdue","Current")))'
        )

    def test_single_customer_one_chart(self):
        spec = ep.build_receivables_spec(
            {
                "invoices": [
                    {
                        "customer": "Solo",
                        "due_date": "2026-01-01",
                        "amount": 100,
                        "paid_amount": 0,
                    },
                ],
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        summary = eg._normalize_spec(spec)["sheets"][1]
        assert len(summary["charts"]) == 1  # aging chart only

    def test_template_mode_no_invoices_builds_blank_tracker(self, tmp_path):
        # "create a receivables tracker" with no invoices must build
        # the blank tracker instead of raising → AI path → invalid spec.
        spec = ep.build_receivables_spec({})
        errors, warnings = eg.validate_workbook_spec(spec)
        assert errors == []
        assert warnings == []
        norm = eg._normalize_spec(spec)
        assert norm["filename"] == "receivables.xlsx"

        sheet = norm["sheets"][0]
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        # as-of anchor kept and live (=TODAY() when as_of_date null)
        assert blocks["A2"] == "As Of"
        assert blocks["B2"] == "=TODAY()"
        table = sheet["tables"][0]
        assert table["start_cell"] == "A4"
        assert len(table["rows"]) == 8  # blank scaffold rows 5..12
        # nothing invented: customer..paid all blank
        assert all(r[:6] == [None] * 6 for r in table["rows"])
        # guarded live formulas on every scaffold row
        assert table["rows"][0][6] == '=IF($E5="","",E5-F5)'
        assert table["rows"][0][7] == '=IF($D5="","",MAX(0,$B$2-$D5))'
        assert table["rows"][0][8] == (
            '=IF(OR($A5="",$E5=""),"",'
            'IF(F5>=E5,"Paid",'
            'IF(F5>0,"Partial",'
            'IF(AND(D5<>"",D5<$B$2),"Overdue","Current"))))'
        )
        assert table["rows"][7][7] == '=IF($D12="","",MAX(0,$B$2-$D12))'
        # totals SUM over the scaffold band (blank → 0)
        assert table["total_row"][4] == "=SUM(E5:E12)"
        assert table["total_row"][6] == "=SUM(G5:G12)"
        # status conditional format covers the scaffold rows
        cf = sheet["conditional_formats"][0]
        assert cf["range"] == "I5:I12"

        # Summary: aging buckets over the blank rows read 0; no
        # customer rollup (none invented), no charts
        summary = norm["sheets"][1]
        assert len(summary["tables"]) == 1
        aging = summary["tables"][0]
        assert aging["start_cell"] == "A3"
        assert [r[0] for r in aging["rows"]] == [
            "Current (not due)",
            "1-30 days",
            "31-60 days",
            "61-90 days",
            "Over 90 days",
        ]
        assert aging["rows"][0][1] == "=COUNTIF(Receivables!$H$5:$H$12,0)"
        assert aging["rows"][0][2] == (
            "=SUMIF(Receivables!$H$5:$H$12,0,Receivables!$G$5:$G$12)"
        )
        assert aging["rows"][1][1] == (
            '=COUNTIFS(Receivables!$H$5:$H$12,">=1",' 'Receivables!$H$5:$H$12,"<=30")'
        )
        assert aging["total_row"][1] == "=SUM(B4:B8)"
        assert "charts" not in summary

        # the workbook still round-trips through openpyxl
        out = tmp_path / "receivables_blank.xlsx"
        eg._build_xlsx(norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Receivables", "Summary"]
        ws = wb["Receivables"]
        assert ws["A5"].value is None
        assert ws["G5"].value == '=IF($E5="","",E5-F5)'
        assert ws["H5"].value == '=IF($D5="","",MAX(0,$B$2-$D5))'
        assert ws["I5"].value == (
            '=IF(OR($A5="",$E5=""),"",'
            'IF(F5>=E5,"Paid",'
            'IF(F5>0,"Partial",'
            'IF(AND(D5<>"",D5<$B$2),"Overdue","Current"))))'
        )
        assert ws["B2"].value == "=TODAY()"
        assert ws["E13"].value == "=SUM(E5:E12)"
        assert len(ws._charts) == 0
        cf_ranges = {str(c.sqref) for c in ws.conditional_formatting}
        assert cf_ranges == {"I5:I12"}
        summary_ws = wb["Summary"]
        assert summary_ws["A4"].value == "Current (not due)"
        assert summary_ws["B4"].value == "=COUNTIF(Receivables!$H$5:$H$12,0)"


class TestReceivablesCoercion:
    def test_quoted_numbers(self):
        p = ep.coerce_receivables_params(
            {
                "invoices": [
                    {"customer": "Acme", "amount": "$2,000", "paid_amount": "500"},
                ],
            }
        )
        assert p["invoices"][0]["amount"] == 2000.0
        assert p["invoices"][0]["paid_amount"] == 500.0

    def test_alias_keys(self):
        p = ep.coerce_receivables_params(
            {
                "invoice_list": [
                    {
                        "client": "Beta",
                        "invoice_no": "B-1",
                        "invoice_date": "2026-05-01",
                        "due": "2026-06-01",
                        "total": 750,
                        "paid": 250,
                    },
                ],
            }
        )
        inv = p["invoices"][0]
        assert inv["customer"] == "Beta"
        assert inv["invoice_number"] == "B-1"
        assert inv["date"] == "2026-05-01"
        assert inv["due_date"] == "2026-06-01"
        assert inv["amount"] == 750.0
        assert inv["paid_amount"] == 250.0

    def test_paid_defaults_to_zero_and_null_as_of(self):
        p = ep.coerce_receivables_params(
            {
                "invoices": [{"customer": "Acme", "amount": 100}],
                "as_of_date": None,
            }
        )
        assert p["invoices"][0]["paid_amount"] == 0.0
        assert p["as_of_date"] is None  # → live =TODAY() anchor

    def test_invalid_as_of_date_falls_back_to_today(self):
        p = ep.coerce_receivables_params(
            {
                "as_of_date": "August 15",
                "invoices": [{"customer": "Acme", "amount": 100}],
            }
        )
        assert p["as_of_date"] is None

    def test_invoices_missing_or_empty_is_template_mode(self):
        # "create a receivables tracker" — no invoices given: blank
        # template mode, NOT a refusal (invoices are never invented).
        for params in ({}, {"invoices": []}, {"currency": "USD"}):
            p = ep.coerce_receivables_params(params)
            assert p["invoices"] == []
            assert p["as_of_date"] is None  # → live =TODAY() anchor

    def test_invoices_not_array_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_receivables_params({"invoices": "INV-1"})
        with pytest.raises(ValueError):
            ep.coerce_receivables_params("nope")
        # entries given but none usable — structural garbage
        with pytest.raises(ValueError):
            ep.coerce_receivables_params({"invoices": ["junk", 42]})

    def test_invoice_without_customer_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_receivables_params(
                {"invoices": [{"invoice_number": "X", "amount": 10}]}
            )

    def test_invoice_without_amount_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_receivables_params(
                {"invoices": [{"customer": "Acme", "paid_amount": 5}]}
            )

    def test_user_notes_rendered_on_sheet(self):
        spec = ep.build_receivables_spec(
            {
                "invoices": [
                    {"customer": "Acme", "amount": 100, "paid_amount": 0},
                ],
                "notes": "Chase anything over 60 days",
            }
        )
        norm = eg._normalize_spec(spec)
        assert norm["sheets"][0]["notes"].startswith("Chase anything over 60 days")


class TestReceivablesRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "receivables", "params": REC_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "aged receivables tracker for my unpaid customer invoices"
        )
        assert result["pattern"] == "receivables"
        assert result["sheet_count"] == 2
        assert result["sheet_names"] == ["Receivables", "Summary"]
        assert result["table_count"] == 3
        assert result["chart_count"] == 2
        assert result["formula_count"] > 30
        assert result["filename"] == "receivables.xlsx"
        assert "receivables template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    @pytest.mark.asyncio
    async def test_routes_with_empty_params_to_blank_template(
        self, tmp_path, monkeypatch
    ):
        # Same class of bug as meal_planner: the classifier returns
        # receivables with NO invoices ("create a receivables tracker")
        # — the pattern must build the blank tracker instead of
        # falling back to the AI path.
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "receivables", "params": {}}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "create a receivables tracker for my unpaid customer invoices"
        )
        assert result["pattern"] == "receivables"
        assert result["sheet_names"] == ["Receivables", "Summary"]
        assert result["chart_count"] == 0  # nothing to chart yet
        assert result["filename"] == "receivables.xlsx"
        assert "receivables template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ══════════════════════════════════════════════════════════════════════
# payroll
# ══════════════════════════════════════════════════════════════════════

PAY_PARAMS = {
    "pay_period": "March 2026",
    "pay_date": "2026-03-31",
    "currency": "USD",
    "employees": [
        {
            "name": "Alice",
            "role": "Baker",
            "pay_basis": "hourly",
            "rate": 18,
            "hours": 160,
            "overtime_hours": 5,
            "tax_rate": "20%",
            "deductions": 120,
        },
        {
            "name": "Bob",
            "role": "Manager",
            "pay_basis": "monthly",
            "rate": 3200,
            "tax_rate": 0.25,
            "deductions": 200,
        },
        {
            "name": "Cara",
            "role": "Cashier",
            "pay_basis": "hourly",
            "rate": 15,
            "hours": 120,
        },
    ],
}

# Python-side expectations (multiplier 1.5):
#   Alice gross 18*160 + 5*18*1.5 = 3015, tax 603, net 3015-603-120 = 2292
#   Bob   gross 3200, tax 800, net 2200
#   Cara  gross 15*120 = 1800, tax 0 (rate blank), net 1800
PAY_EXPECTED = {
    "gross": [3015.0, 3200.0, 1800.0],
    "tax": [603.0, 800.0, 0.0],
    "net": [2292.0, 2200.0, 1800.0],
}


class TestPayrollBuilder:
    def setup_method(self):
        self.spec = ep.build_payroll_spec(PAY_PARAMS)
        self.norm = eg._normalize_spec(self.spec)

    def test_spec_validates_clean(self):
        errors, warnings = eg.validate_workbook_spec(self.spec)
        assert errors == []
        assert warnings == []

    def test_geometry_and_formula_lattice(self):
        assert self.spec["filename"] == "payroll_march_2026.xlsx"
        sheet = self.norm["sheets"][0]
        assert sheet["name"] == "Payroll"

        # assumptions block: period 2, pay date 3, multiplier 4
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["A1"]["text"] == "Payroll — March 2026"
        assert blocks["A2"]["text"] == "Pay Period"
        assert blocks["B2"]["text"] == "March 2026"
        assert blocks["A3"]["text"] == "Pay Date"
        assert blocks["B3"]["text"] == date(2026, 3, 31)
        assert blocks["A4"]["text"] == "Overtime Multiplier"
        assert blocks["B4"]["text"] == 1.5
        assert blocks["B4"]["number_format"] == "0.0"

        table = sheet["tables"][0]
        # headers row 6, data 7-9, total 10
        assert table["start_cell"] == "A6"
        assert table["headers"] == [
            "Employee",
            "Role",
            "Basis",
            "Rate",
            "Hours",
            "OT Hours",
            "Gross Pay",
            "Tax Rate",
            "Tax",
            "Deductions",
            "Net Pay",
        ]
        # hourly gross references the multiplier assumption cell $B$4
        assert table["rows"][0][6] == "=D7*E7+F7*D7*$B$4"
        assert table["rows"][0][8] == "=G7*H7"
        assert table["rows"][0][10] == "=G7-I7-J7"
        # monthly gross is just the salary
        assert table["rows"][1][2] == "Monthly"
        assert table["rows"][1][6] == "=D8"
        assert table["rows"][1][10] == "=G8-I8-J8"
        # Cara: no tax rate stated → blank cell, tax formula still live
        assert table["rows"][2][7] is None
        assert table["rows"][2][8] == "=G9*H9"
        # totals over the exact data rows
        assert table["total_row"][4] == "=SUM(E7:E9)"
        assert table["total_row"][6] == "=SUM(G7:G9)"
        assert table["total_row"][8] == "=SUM(I7:I9)"
        assert table["total_row"][9] == "=SUM(J7:J9)"
        assert table["total_row"][10] == "=SUM(K7:K9)"

        # chart / dropdown / conditional format pin to rows 7-9
        chart = sheet["charts"][0]
        assert chart["categories_range"] == "Payroll!A7:A9"
        assert chart["series"][0]["values_range"] == "Payroll!G7:G9"
        assert chart["anchor"] == "M6"
        dv = sheet["data_validation"][0]
        assert dv["range"] == "C7:C9"
        assert dv["values"] == ["Hourly", "Monthly"]
        cf = sheet["conditional_formats"][0]
        assert cf["range"] == "K7:K9"
        assert cf["rules"][0]["operator"] == "less_than"
        assert cf["rules"][0]["value"] == 0

    def test_payroll_math_independent(self):
        """Evaluate the emitted gross/tax/net formulas over the emitted
        grid and compare with plain-Python payroll math."""
        table = self.norm["sheets"][0]["tables"][0]
        cells = {"B4": 1.5}
        for i, row in enumerate(table["rows"]):
            r = 7 + i
            cells[f"D{r}"] = row[3]
            cells[f"E{r}"] = row[4]
            cells[f"F{r}"] = row[5]
            cells[f"G{r}"] = row[6]
            cells[f"H{r}"] = row[7]
            cells[f"I{r}"] = row[8]
            cells[f"J{r}"] = row[9]
            cells[f"K{r}"] = row[10]
        for col_idx, key in ((6, "G"), (8, "I"), (9, "J"), (10, "K")):
            cells[f"{key}10"] = table["total_row"][col_idx]
        resolved = _resolve_formulas(cells)

        for i in range(3):
            r = 7 + i
            assert resolved[f"G{r}"] == pytest.approx(PAY_EXPECTED["gross"][i])
            assert resolved[f"I{r}"] == pytest.approx(PAY_EXPECTED["tax"][i])
            assert resolved[f"K{r}"] == pytest.approx(PAY_EXPECTED["net"][i])
        assert resolved["G10"] == pytest.approx(8015.0)
        assert resolved["I10"] == pytest.approx(1403.0)
        assert resolved["J10"] == pytest.approx(320.0)
        assert resolved["K10"] == pytest.approx(6292.0)

    def test_built_workbook(self, tmp_path):
        out = tmp_path / "payroll.xlsx"
        eg._build_xlsx(self.norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Payroll"]
        ws = wb["Payroll"]
        assert ws["A1"].value == "Payroll — March 2026"
        assert ws["B2"].value == "March 2026"
        assert _cell_date(ws["B3"].value) == date(2026, 3, 31)
        assert ws["B3"].number_format == "yyyy-mm-dd"
        assert ws["B4"].value == 1.5
        assert ws["B4"].number_format == "0.0"
        assert ws["A6"].value == "Employee"
        assert ws["C7"].value == "Hourly"
        assert ws["C8"].value == "Monthly"
        assert ws["D7"].value == 18
        assert ws["D7"].number_format == '"$"#,##0.00'
        assert ws["G7"].value == "=D7*E7+F7*D7*$B$4"
        assert ws["G8"].value == "=D8"
        assert ws["H7"].value == 0.2
        assert ws["H7"].number_format == "0.00%"
        assert ws["I7"].value == "=G7*H7"
        assert ws["K7"].value == "=G7-I7-J7"
        assert ws["K7"].number_format == '"$"#,##0.00'
        assert ws["E8"].value is None  # monthly staff carry no hours
        assert ws["G10"].value == "=SUM(G7:G9)"
        assert ws.freeze_panes == "A7"
        assert len(ws._charts) == 1
        (dv,) = ws.data_validations.dataValidation
        assert dv.formula1 == '"Hourly,Monthly"'
        cf_ranges = {str(c.sqref) for c in ws.conditional_formatting}
        assert cf_ranges == {"K7:K9"}

    def test_minimal_params_block(self):
        spec = ep.build_payroll_spec({"employees": [{"name": "Solo", "rate": 3000}]})
        errors, warnings = eg.validate_workbook_spec(spec)
        assert errors == []
        assert warnings == []
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        assert blocks["A1"] == "Payroll Register"
        assert blocks["A2"] == "Overtime Multiplier"
        assert blocks["B2"] == 1.5
        table = sheet["tables"][0]
        assert table["start_cell"] == "A4"
        # no hours stated → monthly basis, gross = rate
        assert table["rows"][0][2] == "Monthly"
        assert table["rows"][0][6] == "=D5"
        assert sheet["charts"][0]["anchor"] == "M4"
        assert norm["filename"] == "payroll.xlsx"

    def test_custom_overtime_multiplier(self):
        spec = ep.build_payroll_spec(
            {
                "overtime_multiplier": 2,
                "employees": [
                    {
                        "name": "Alice",
                        "pay_basis": "hourly",
                        "rate": 18,
                        "hours": 160,
                        "overtime_hours": 5,
                    },
                ],
            }
        )
        errors, _ = eg.validate_workbook_spec(spec)
        assert errors == []
        norm = eg._normalize_spec(spec)
        sheet = norm["sheets"][0]
        blocks = {b["cell"]: b for b in sheet["text_blocks"]}
        assert blocks["B2"]["text"] == 2.0
        table = sheet["tables"][0]
        # headers row 4, data row 5 → gross references $B$2
        assert table["rows"][0][6] == "=D5*E5+F5*D5*$B$2"
        cells = {
            "B2": 2.0,
            "D5": 18.0,
            "E5": 160.0,
            "F5": 5.0,
            "G5": table["rows"][0][6],
        }
        resolved = _resolve_formulas(cells)
        assert resolved["G5"] == pytest.approx(18 * 160 + 5 * 18 * 2)

    def test_template_mode_no_employees_builds_blank_register(self, tmp_path):
        # "create a payroll register" with no employees must build the
        # blank register instead of raising → AI path → invalid spec.
        spec = ep.build_payroll_spec({})
        errors, warnings = eg.validate_workbook_spec(spec)
        assert errors == []
        assert warnings == []
        norm = eg._normalize_spec(spec)
        assert norm["filename"] == "payroll.xlsx"

        sheet = norm["sheets"][0]
        assert sheet["name"] == "Payroll"
        # assumptions block keeps the overtime multiplier (row 2)
        blocks = {b["cell"]: b["text"] for b in sheet["text_blocks"]}
        assert blocks["A1"] == "Payroll Register"
        assert blocks["A2"] == "Overtime Multiplier"
        assert blocks["B2"] == 1.5

        table = sheet["tables"][0]
        # headers row 4, blank scaffold rows 5..12, total 13
        assert table["start_cell"] == "A4"
        assert len(table["rows"]) == 8
        # nothing invented: employee..deductions all blank
        assert all(
            r[c] is None for r in table["rows"] for c in (0, 1, 2, 3, 4, 5, 7, 9)
        )
        # guarded gross/tax/net formulas on every scaffold row —
        # gross reads the overtime-multiplier assumption cell $B$2
        assert table["rows"][0][6] == (
            '=IF(OR($A5="",$D5=""),"",'
            'IF($C5="Hourly",'
            "$D5*$E5+$F5*$D5*$B$2,"
            "$D5))"
        )
        assert table["rows"][0][8] == '=IF($G5="","",IF($H5="",0,$G5*$H5))'
        assert table["rows"][0][10] == '=IF($G5="","",$G5-$I5-$J5)'
        assert table["rows"][7][6] == (
            '=IF(OR($A12="",$D12=""),"",'
            'IF($C12="Hourly",'
            "$D12*$E12+$F12*$D12*$B$2,"
            "$D12))"
        )
        # totals SUM over the scaffold band (blank → 0)
        assert table["total_row"][4] == "=SUM(E5:E12)"
        assert table["total_row"][6] == "=SUM(G5:G12)"
        assert table["total_row"][10] == "=SUM(K5:K12)"

        # Hourly/Monthly dropdown + negative-net CF cover the scaffold
        dv = sheet["data_validation"][0]
        assert dv["range"] == "C5:C12"
        assert dv["values"] == ["Hourly", "Monthly"]
        cf = sheet["conditional_formats"][0]
        assert cf["range"] == "K5:K12"
        # no chart when there are no employees
        assert "charts" not in sheet

        # the workbook still round-trips through openpyxl
        out = tmp_path / "payroll_blank.xlsx"
        eg._build_xlsx(norm, out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["Payroll"]
        ws = wb["Payroll"]
        assert ws["A5"].value is None
        assert ws["B2"].value == 1.5
        assert ws["G5"].value == (
            '=IF(OR($A5="",$D5=""),"",'
            'IF($C5="Hourly",'
            "$D5*$E5+$F5*$D5*$B$2,"
            "$D5))"
        )
        assert ws["I5"].value == '=IF($G5="","",IF($H5="",0,$G5*$H5))'
        assert ws["K5"].value == '=IF($G5="","",$G5-$I5-$J5)'
        assert ws["G13"].value == "=SUM(G5:G12)"
        assert ws.freeze_panes == "A5"
        assert len(ws._charts) == 0
        (dv,) = ws.data_validations.dataValidation
        assert dv.formula1 == '"Hourly,Monthly"'
        assert str(dv.sqref) == "C5:C12"
        cf_ranges = {str(c.sqref) for c in ws.conditional_formatting}
        assert cf_ranges == {"K5:K12"}


class TestPayrollCoercion:
    def test_quoted_numbers_and_percent_strings(self):
        p = ep.coerce_payroll_params(
            {
                "employees": [
                    {
                        "name": "Ann",
                        "pay_basis": "hourly",
                        "rate": "18",
                        "hours": "40",
                        "tax_rate": "20%",
                        "deductions": "50",
                    },
                ],
            }
        )
        emp = p["employees"][0]
        assert emp["rate"] == 18.0
        assert emp["hours"] == 40.0
        assert emp["tax_rate"] == 0.2
        assert emp["deductions"] == 50.0

    def test_alias_keys(self):
        p = ep.coerce_payroll_params(
            {
                "staff": [
                    {
                        "employee": "Bill",
                        "job": "Cook",
                        "hourly_rate": "$15",
                        "hours_worked": 100,
                        "ot_hours": 4,
                        "tax_pct": 15,
                        "other_deductions": 30,
                    },
                ],
            }
        )
        emp = p["employees"][0]
        assert emp["name"] == "Bill"
        assert emp["role"] == "Cook"
        assert emp["rate"] == 15.0
        assert emp["hours"] == 100.0
        assert emp["overtime_hours"] == 4.0
        assert emp["tax_rate"] == 0.15
        assert emp["deductions"] == 30.0

    def test_null_semantics_and_defaults(self):
        p = ep.coerce_payroll_params(
            {
                "employees": [
                    {"name": "Cara", "pay_basis": "hourly", "rate": 15, "hours": 120},
                    {"name": "Dave", "monthly_salary": 4000},
                ],
            }
        )
        cara, dave = p["employees"]
        assert cara["overtime_hours"] == 0.0  # not stated → 0
        assert cara["tax_rate"] is None  # not stated → blank → 0 tax
        assert cara["deductions"] == 0.0
        assert dave["pay_basis"] == "monthly"  # no hours → salary
        assert dave["hours"] is None
        assert p["overtime_multiplier"] == 1.5  # documented default

    def test_basis_inference_from_hours(self):
        p = ep.coerce_payroll_params(
            {
                "employees": [
                    {"name": "Eve", "rate": 20, "hours": 80},
                    {"name": "Faye", "rate": 5000},
                ],
            }
        )
        assert p["employees"][0]["pay_basis"] == "hourly"
        assert p["employees"][1]["pay_basis"] == "monthly"

    def test_employees_missing_or_empty_is_template_mode(self):
        # "create a payroll register" — no employees given: blank
        # template mode, NOT a refusal (employees are never invented).
        for params in ({}, {"employees": []}, {"pay_period": "March 2026"}):
            p = ep.coerce_payroll_params(params)
            assert p["employees"] == []
            assert p["overtime_multiplier"] == 1.5  # documented default

    def test_employees_not_array_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_payroll_params({"employees": "Alice"})
        with pytest.raises(ValueError):
            ep.coerce_payroll_params("nope")
        # entries given but none usable — structural garbage
        with pytest.raises(ValueError):
            ep.coerce_payroll_params({"employees": ["junk", 7]})

    def test_employee_without_name_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_payroll_params({"employees": [{"rate": 20, "hours": 10}]})

    def test_employee_without_rate_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_payroll_params({"employees": [{"name": "Gus", "hours": 10}]})

    def test_hourly_employee_without_hours_raises(self):
        with pytest.raises(ValueError):
            ep.coerce_payroll_params(
                {
                    "employees": [
                        {"name": "Hana", "pay_basis": "hourly", "rate": 20},
                    ],
                }
            )

    def test_user_notes_rendered_on_sheet(self):
        spec = ep.build_payroll_spec(
            {
                "employees": [{"name": "Ivan", "rate": 3000}],
                "notes": "Paid on the last working day",
            }
        )
        norm = eg._normalize_spec(spec)
        assert norm["sheets"][0]["notes"].startswith("Paid on the last working day")


class TestPayrollRouting:
    @pytest.mark.asyncio
    async def test_routes_to_template(self, tmp_path, monkeypatch):
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "payroll", "params": PAY_PARAMS}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "payroll register for March with employee salaries " "and hourly wages"
        )
        assert result["pattern"] == "payroll"
        assert result["sheet_count"] == 1
        assert result["sheet_names"] == ["Payroll"]
        assert result["table_count"] == 1
        assert result["chart_count"] == 1
        assert result["formula_count"] > 10
        assert result["filename"] == "payroll_march_2026.xlsx"
        assert "payroll template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()

    @pytest.mark.asyncio
    async def test_routes_with_empty_params_to_blank_template(
        self, tmp_path, monkeypatch
    ):
        # Same class of bug as meal_planner: the classifier returns
        # payroll with NO employees ("create a payroll register") — the
        # pattern must build the blank register instead of falling back
        # to the AI path.
        monkeypatch.setattr(
            eg,
            "_call_llm",
            _classify_response({"pattern": "payroll", "params": {}}),
        )

        async def must_not_run(brief, requirements, model=None):
            raise AssertionError("AI path must not run when pattern matches")

        monkeypatch.setattr(eg, "_generate_workbook_json", must_not_run)
        monkeypatch.setattr(eg, "_get_reports_dir", lambda: tmp_path)
        result = await eg.generate_spreadsheet(
            "create a payroll register spreadsheet for my staff"
        )
        assert result["pattern"] == "payroll"
        assert result["sheet_names"] == ["Payroll"]
        assert result["chart_count"] == 0  # nothing to chart yet
        assert result["filename"] == "payroll.xlsx"
        assert "payroll template" in result["summary"]
        assert (tmp_path / f"{result['report_id']}.xlsx").exists()


# ── Registry wiring ───────────────────────────────────────────────────


class TestRegistryWiring:
    def test_patterns_registered_with_matching_names(self):
        for name in ("price_list", "receivables", "payroll"):
            assert name in ep.PATTERN_BUILDERS
            assert name in ep.PATTERN_KEYWORDS
            assert ep.PATTERN_DESCRIPTIONS.get(name)

    def test_keywords_route_the_gate(self):
        for brief in (
            "product price list with cost and margin",
            "pricing catalog for my bakery",
            "track my unpaid invoices and aged receivables",
            "who owes me money — outstanding customer balances",
            "payroll register with salaries and wages",
            "payslips for my staff this pay period",
        ):
            assert eg._PATTERN_GATE_RE.search(brief), brief

    def test_gate_still_skips_non_documents(self):
        assert not eg._PATTERN_GATE_RE.search("grocery list")
