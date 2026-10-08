from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from app.services import artifacts, report_gen
from app.services.artifact_sources import sections_for
from app.services.artifact_summary import summarize_sections
from app.agent.tools import report_gen as report_tool


def test_section_replacement_without_final_newline_preserves_next_heading():
    source = "# Title\nIntro\n\n## Summary\nFacts.\n"
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    edited = artifacts.edited_source(
        version, "report", [{"section_id": "section-1", "content": "# New\nBeta."}]
    )
    assert "Beta.\n\n## Summary" in edited
    assert len(sections_for(edited, "report")) == 2


@pytest.mark.parametrize("replacement", ["Plan: Beta", "Changed body without a heading."])
def test_report_edit_rejects_lost_section_boundary(replacement):
    source = (
        "# Report\n\n## Overview\nIntro\n\n## Plan: Alpha\nKeep this body.\n\n## Results\nPending"
    )
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    with pytest.raises(artifacts.HTTPException) as error:
        artifacts.edited_source(
            version, "report", [{"section_id": "section-3", "content": replacement}]
        )
    assert error.value.status_code == 422
    assert version.source == source


def test_report_heading_rename_retains_body_and_neighbors():
    source = (
        "# Report\n\n## Overview\nIntro\n\n## Plan: Alpha\nKeep this body.\n\n## Results\nPending"
    )
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    edited = artifacts.edited_source(
        version,
        "report",
        [{"section_id": "section-3", "content": "## Plan: Beta\nKeep this body."}],
    )
    assert [part["title"] for part in sections_for(edited, "report")] == [
        "Report",
        "Overview",
        "Plan: Beta",
        "Results",
    ]
    assert "Keep this body." in edited


@pytest.mark.parametrize(
    ("replacement", "titles", "body"),
    [
        ("## Topic\nNew body.", ["Report", "Topic", "Next"], "New body."),
        ("## Renamed\nNew body.", ["Report", "Renamed", "Next"], "New body."),
        ("# Topic\nNew body.", ["Report", "Topic", "Next"], "New body."),
        (
            "## Part one\nFirst body.\n\n## Part two\nSecond body.",
            ["Report", "Part one", "Part two", "Next"],
            "Second body.",
        ),
    ],
    ids=["body-only", "heading-and-body", "heading-level", "split-section"],
)
def test_report_edits_support_requested_content_and_structure_changes(replacement, titles, body):
    source = "# Report\n\n## Topic\nOld body.\n\n## Next\nUntouched body."
    version = SimpleNamespace(source=source, sections=sections_for(source, "report"))
    edited = artifacts.edited_source(
        version, "report", [{"section_id": "section-2", "content": replacement}]
    )
    sections = sections_for(edited, "report")
    assert [section["title"] for section in sections] == titles
    assert body in edited
    assert "Old body." not in edited
    assert sections[-1]["content"] == "## Next\nUntouched body."


@pytest.mark.asyncio
async def test_source_provenance_survives_model_dropping_markers():
    result = await summarize_sections(
        [
            {"id": "page-4", "page": 4, "content": "A grounded definition."},
        ],
        "One bullet",
        AsyncMock(return_value="- A definition."),
    )
    assert result["source_references"] == [{"id": "page-4", "page": 4}]
    assert "page 4" in result["summary"]


@pytest.mark.asyncio
async def test_report_constraints_are_required_not_suggested(monkeypatch):
    chat = AsyncMock(return_value={"message": {"content": "# QA\n## Overview\nTest."}})
    monkeypatch.setattr(report_gen.providers, "chat_once", chat)
    await report_gen._generate_markdown("QA", "Exactly three short sections", "test")
    prompt = chat.call_args.args[1][1]["content"]
    assert "Required structure" in prompt
    assert "Exactly three short sections" in prompt
    assert "you may adapt" not in prompt
    assert "Do not add an executive summary" in chat.call_args.args[1][0]["content"]


@pytest.mark.asyncio
async def test_report_keeps_original_user_constraints(monkeypatch):
    from app.agent.tools import config_store

    generate = AsyncMock(
        return_value={
            "report_id": "qa",
            "filename": "qa.docx",
            "format": "docx",
            "file_path": "reports/qa.docx",
            "download_url": "/api/reports/qa/download",
            "created_at": "2026-10-08T00:00:00Z",
        }
    )
    monkeypatch.setattr(report_tool, "generate_report", generate)
    monkeypatch.setattr(config_store, "tool_model_override", AsyncMock(return_value=None))
    result = await report_tool.ReportGenTool().execute(
        topic="QA",
        outline="Overview | Plan | Results",
        format="docx",
        _user_request="Under 100 words, no executive summary.",
    )
    assert result.success
    assert "Under 100 words" in generate.call_args.kwargs["outline"]


@pytest.mark.asyncio
async def test_truncated_summary_is_not_cached(tmp_path, monkeypatch):
    monkeypatch.setenv("REALOPEN_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(artifacts.model_prefs, "resolve_task_model", AsyncMock(return_value="test"))

    async def stream(*args, **kwargs):
        yield {"content": "An incomplete", "done": True, "finish_reason": "length"}

    monkeypatch.setattr(artifacts.providers, "stream_chat", stream)
    with pytest.raises(RuntimeError, match="output limit"):
        await artifacts.summarize(
            None,
            SimpleNamespace(id="test"),
            SimpleNamespace(
                number=1,
                source="Facts",
                sections=[{"id": "one", "content": "Facts"}],
                settings={},
            ),
        )
    assert not list(tmp_path.rglob("*.json"))


@pytest.mark.asyncio
async def test_cutoff_omits_unfinished_sentence_and_marks_abridged(tmp_path, monkeypatch):
    monkeypatch.setenv("REALOPEN_DATA_DIR", str(tmp_path))

    async def stream(*args, **kwargs):
        assert args[0] == "override"
        yield {
            "content": "- Complete fact. Unfinished ending",
            "done": True,
            "finish_reason": "length",
        }

    monkeypatch.setattr(artifacts.providers, "stream_chat", stream)
    result = await artifacts.summarize(
        None,
        SimpleNamespace(id="test"),
        SimpleNamespace(
            number=1,
            source="Facts",
            sections=[{"id": "one", "content": "Facts"}],
            settings={},
        ),
        model="override",
    )
    assert "Unfinished ending" not in result["summary"]
    assert "abbreviated" in result["summary"]
    assert result["warnings"]
@pytest.mark.parametrize("template", ["modern", "corporate", "elegant"])
def test_presentation_layouts_preserve_authored_body(tmp_path, template):
    from pptx import Presentation
    from app.services.pptx_gen import _parse_slides, _resolve_theme, _build_pptx

    source = "# Cover\nFirst body.\n\n---\n\n## Section\nSection body.\n\n---\n\n# Conclusion\nClosing body."
    _, theme = _resolve_theme(template)
    output = tmp_path / "preserved.pptx"
    _build_pptx(_parse_slides(source), theme, "QA", output)
    deck = Presentation(output)
    assert len(deck.slides) == 3
    for slide, expected in zip(deck.slides, ["First body.", "Section body.", "Closing body."]):
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert expected in text
